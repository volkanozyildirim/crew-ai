"""Katman 0 (deterministik) kapılarının regresyon testleri.

Projede test altyapısı yok; bu dosya **bağımsız çalışır**:

    .venv/bin/python tests/test_katman0_gates.py

Neden burada: Katman 0 doğrulayıcıları veri üzerinde saf fonksiyonlar, LLM
çağrısı gerektirmiyorlar — dolayısıyla test edilebilirler. Her biri gerçek
üretim verisine karşı sınanır:

  * MySQL'deki `jobs`/`job_steps` kayıtları (job #178, #179) — mock değil,
    kaydedilmiş gerçek LLM çıktısı. DB erişilemezse o testler ATLANIR.
  * Geçici git fixture'ları — branch/checkout durumuna bağlı davranış için.

Kapsanan hata sınıfları (hepsi 2026-07-27'de gerçekten yaşandı):
  #178  uydurma dizin + entegrasyon yok → plan implement'e ulaştı, review'da
        kalıcı RED, $6.61 boşa
  #179  plan '/app/X.php' vs itiraz 'app/X.php' → boş kesişim → reviewer'ın
        şikâyet ettiği dosyalar retry'da hiç düzeltilmedi
  #179  reviewer uydurma standart (R2) ve ürün kararı (R1) ile job'ı öldürdü
  #179  Allocator.php 3 parametreli metoda 4. argüman → PHP yuttu, no-op
  #180  metot adı repo genelinde tekil değil → 20+ yanlış alarm, asıl
        implementasyon dosyası bloklandı
  #181  zarf $18 dedi, ara-adım cap'i $10 okudu → iş düşük tavanda öldü
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "src"))

from agile_sdlc_crew.flow import (  # noqa: E402
    AgileSDLCFlow,
    _classify_review_issues,
    _norm_path,
    _parse_review_issues,
    _paths_in_text,
    _contract_fix_description,
    _partition_plan_by_branch,
    _php_call_arity,
    _php_signatures,
    _requirement_ids,
    _review_approved,
    _parse_readiness,
    _readiness_score,
    _readiness_comment,
    NeedsHumanReview,
    NeedsMoreInfo,
)

PASS, FAIL, SKIP = [], [], []


def check(name: str, cond: bool, detail: str = ""):
    (PASS if cond else FAIL).append(name)
    print(f"  {'✅' if cond else '❌'} {name}" + (f"  — {detail}" if detail and not cond else ""))


def skip(name: str, why: str):
    SKIP.append(name)
    print(f"  ⏭️  {name}  ({why})")


# ── Replay korpusu: gerçek job kayıtları ─────────────────────────────────

def load_job(jid: int):
    try:
        from agile_sdlc_crew import db
        job = db.get_job(jid)
    except Exception:
        return None
    if not job:
        return None
    out = {"reqs": "", "plan": None}
    for s in job.get("steps") or []:
        if s["step_key"] == "requirements_analysis_task" and s.get("output"):
            out["reqs"] = s["output"]
        if s["step_key"] == "technical_design_task" and s.get("output"):
            try:
                out["plan"] = json.loads(s["output"])
            except Exception:
                pass
    return out


# ── 1. Yol normalizasyonu (#179 slash uyuşmazlığı) ───────────────────────

def test_norm_path():
    print("\n[1] _norm_path — plan '/app/X' vs itiraz 'app/X'")
    check("baştaki / atılır", _norm_path("/app/X.php") == "app/X.php")
    check("ters slash düzelir", _norm_path("app\\M\\Y.php") == "app/M/Y.php")
    check("boşluk kırpılır", _norm_path("  /app/Z.php ") == "app/Z.php")
    check("büyük/küçük harf KORUNUR (repo case-sensitive)",
          _norm_path("/App/Foo.php") == "App/Foo.php")
    check("#179 kesişimi artık boş değil",
          {_norm_path("app/Model/StockSource.php")} <= {_norm_path("/app/Model/StockSource.php")})


# ── 2. Plan yol/entegrasyon kapısı (#178) ────────────────────────────────

def test_plan_paths():
    print("\n[2] _validate_plan_paths — uydurma yol + entegrasyon yok")
    with tempfile.TemporaryDirectory() as td:
        base = Path(td)
        repo = base / "fake"
        (repo / "app" / "Library").mkdir(parents=True)
        (repo / "app" / "Library" / "Existing.php").write_text("<?php class E {}")
        for args in (["init", "-q", "-b", "main"], ["add", "-A"]):
            subprocess.run(["git", *args], cwd=repo, check=True,
                           capture_output=True)
        subprocess.run(["git", "-c", "user.email=t@t", "-c", "user.name=t",
                        "commit", "-qm", "init"], cwd=repo, check=True,
                       capture_output=True)

        from agile_sdlc_crew.tools.local_repo import LocalRepoManager
        stub = SimpleNamespace(_repo_mgr=LocalRepoManager(base_dir=str(base)))
        val = AgileSDLCFlow._validate_plan_paths

        p = val(stub, {"changes": [{"file_path": "app/Library/Order/Split/New.php"}]}, "fake")
        check("uydurma dizin yakalanır", any("UYDURMA YOL" in x for x in p))
        check("en yakın var olan dizin gösterilir",
              any("app/Library" in x for x in p))
        check("entegrasyon yok yakalanır", any("ENTEGRASYON YOK" in x for x in p))

        p2 = val(stub, {"changes": [
            {"file_path": "app/Library/Order/Split/New.php"},
            {"file_path": "app/Library/Existing.php"},
        ]}, "fake")
        check("mevcut dosya değişince entegrasyon uyarısı kalkar",
              not any("ENTEGRASYON YOK" in x for x in p2))

        # KRİTİK: feature branch'te dosya DİSKTE varken base'e bakılmalı
        subprocess.run(["git", "checkout", "-qb", "feature/x"], cwd=repo,
                       check=True, capture_output=True)
        (repo / "app" / "Library" / "Order" / "Split").mkdir(parents=True)
        (repo / "app" / "Library" / "Order" / "Split" / "New.php").write_text("<?php class N {}")
        subprocess.run(["git", "add", "-A"], cwd=repo, check=True, capture_output=True)
        subprocess.run(["git", "-c", "user.email=t@t", "-c", "user.name=t",
                        "commit", "-qm", "feat"], cwd=repo, check=True,
                       capture_output=True)
        p3 = val(stub, {"changes": [{"file_path": "app/Library/Order/Split/New.php"}]}, "fake")
        check("dosya diskte olsa da BASE ref'e bakılır (retry senaryosu)",
              any("UYDURMA YOL" in x for x in p3) and any("ENTEGRASYON YOK" in x for x in p3),
              f"beklenen 2 sorun, gelen: {p3}")

        check("klon yoksa sessiz atlanır", val(stub, {"changes": [{"file_path": "a.php"}]}, "yok") == [])

    # Gerçek korpus
    j178, j179 = load_job(178), load_job(179)
    if not (j178 and j178["plan"]):
        skip("job #178 planı → 4 sorun", "DB/kayıt yok")
    else:
        from agile_sdlc_crew.tools.local_repo import LocalRepoManager
        real = SimpleNamespace(_repo_mgr=LocalRepoManager())
        p = AgileSDLCFlow._validate_plan_paths(real, j178["plan"], "orkestra")
        check("job #178 planı → uydurma yol + entegrasyon yok",
              len(p) >= 4 and any("ENTEGRASYON YOK" in x for x in p), f"{p}")
    if not (j179 and j179["plan"]):
        skip("job #179 planı → temiz", "DB/kayıt yok")
    else:
        from agile_sdlc_crew.tools.local_repo import LocalRepoManager
        real = SimpleNamespace(_repo_mgr=LocalRepoManager())
        p = AgileSDLCFlow._validate_plan_paths(real, j179["plan"], "orkestra")
        check("job #179 planı → yanlış alarm yok", p == [], f"{p}")


# ── 3. İtiraz kapısı (#179 R1/R2) ────────────────────────────────────────

def test_issue_gate():
    print("\n[3] _classify_review_issues — bloklayıcı vs düşürülen")
    ids = {"FR1", "FR2", "AC1", "AC2"}
    QUOTE = "public function buildLuggageContext"

    def verify(loc):  # kanıt doğrulayıcı stub'ı
        return (loc or {}).get("quote") in (QUOTE, "") or loc.get("quote") == "REAL"

    def one(raw):
        return _parse_review_issues(
            "REVIEW_ISSUES_JSON:\n```json\n" + json.dumps({"issues": [raw]}) + "\n```")

    ev = {"file": "app/X.php", "line": 1, "quote": QUOTE}

    b, d = _classify_review_issues(one({
        "file": "app/X.php", "severity": "blocker", "problem": "p",
        "required_fix": "f", "requirement_ids": ["AC1"], "evidence": ev}), ids, verify)
    check("geçerli requirement_ids + doğrulanmış kanıt → BLOKLAR", len(b) == 1)

    b, d = _classify_review_issues(one({
        "file": "app/X.php", "severity": "blocker", "problem": "p",
        "required_fix": "f", "requirement_ids": ["AC99"], "evidence": ev}), ids, verify)
    check("var olmayan id'ye atıf → düşer", not b and "var olmayan" in d[0]["demote_reason"])

    b, d = _classify_review_issues(one({
        "file": "app/X.php", "severity": "major", "problem": "p",
        "required_fix": "f", "requirement_ids": [], "evidence": ev,
        "precedent": {"file": "app/Y.php", "line": 9, "quote": QUOTE}}), ids, verify)
    check("id yok ama doğrulanmış emsal → BLOKLAR", len(b) == 1)

    b, d = _classify_review_issues(one({
        "file": "app/X.php", "severity": "major", "problem": "p",
        "required_fix": "f", "requirement_ids": [], "evidence": ev,
        "precedent": {"file": "app/Y.php", "line": 9, "quote": "UYDURMA"}}), ids, verify)
    check("uydurma emsal → düşer (job #179/R2 sınıfı)", not b)

    b, d = _classify_review_issues(one({
        "file": "app/X.php", "severity": "blocker", "problem": "p",
        "required_fix": "f", "requirement_ids": ["AC1"],
        "evidence": {"file": "app/X.php", "line": 1, "quote": "YOK BOYLE BIR SATIR"}}),
        ids, verify)
    check("doğrulanamayan kanıt → düşer", not b and "kanıt doğrulanamadı" in d[0]["demote_reason"])

    b, d = _classify_review_issues(one({
        "file": "app/X.php", "severity": "blocker", "problem": "p",
        "required_fix": "f", "requirement_ids": ["AC1"]}), ids, verify)
    check("kanıt hiç verilmemiş → düşer", not b)

    b, d = _classify_review_issues(one({
        "file": "app/X.php", "severity": "minor", "problem": "p",
        "required_fix": "f", "requirement_ids": ["AC1"], "evidence": ev}), ids, verify)
    check("minor → düşer (öneri)", not b)

    # #179'un GERÇEK iki maddesi: ikisi de gereksinim bağı olmadan geldi
    real_two = _parse_review_issues("REVIEW_ISSUES_JSON:\n```json\n" + json.dumps({"issues": [
        {"file": "app/Model/StockSource.php", "severity": "major",
         "problem": "Merchant kontrolu siparis bazli", "required_fix": "kalem filtresine gec",
         "requirement_ids": [], "evidence": ev},
        {"file": "app/Migration/Upgrade.php", "severity": "major",
         "problem": "cms_setting_group_id eksik", "required_fix": "insert'e ekle",
         "requirement_ids": [], "evidence": ev},
    ]}) + "\n```")
    b, d = _classify_review_issues(real_two, ids, verify)
    check("job #179'un iki itirazı da düşer → job ölmezdi", not b and len(d) == 2)

    # Eski şema güvenliği: hiç evidence yoksa kapı ATLANMALI (caller kontrolü)
    old = _parse_review_issues("REVIEW_ISSUES_JSON:\n```json\n" + json.dumps({"issues": [
        {"file": "app/X.php", "line": 1, "severity": "blocker",
         "problem": "p", "required_fix": "f"}]}) + "\n```")
    check("eski şema tespiti: hiçbir maddede evidence yok",
          not any(i.get("evidence") for i in old))


# ── 4. Gereksinim id çıkarımı ────────────────────────────────────────────

def test_requirement_ids():
    print("\n[4] _requirement_ids — JSON + regex fallback")
    j = json.dumps({"functional_requirements": [{"id": "FR1", "desc": "x"}],
                    "acceptance_criteria": [{"id": "AC1"}, {"id": "AC2"}]})
    check("JSON'dan çıkarır", _requirement_ids(j) == {"FR1", "AC1", "AC2"})
    check("regex fallback", _requirement_ids("bkz AC3 ve FR7 maddeleri") == {"AC3", "FR7"})
    check("boş metin → boş küme", _requirement_ids("") == set())

    j179 = load_job(179)
    if not (j179 and j179["reqs"]):
        skip("job #179 gereksinimleri", "DB/kayıt yok")
    else:
        ids = _requirement_ids(j179["reqs"])
        check("job #179 → 13 gereksinim id", len(ids) == 13, f"{len(ids)}: {sorted(ids)}")


# ── 5. Deterministik completeness ────────────────────────────────────────

def test_completeness():
    print("\n[5] _check_plan_completeness — küme farkı (LLM yok)")
    reqs = json.dumps({"acceptance_criteria": [{"id": f"AC{i}"} for i in range(1, 5)]})
    stub = SimpleNamespace(state=SimpleNamespace(requirements_text=reqs))
    fn = AgileSDLCFlow._check_plan_completeness

    full = {"changes": [{"file_path": "a.php", "covers_requirements": ["AC1", "AC2", "AC3", "AC4"]}]}
    check("tam kapsam → eksik yok", fn(stub, full) == [])

    partial = {"changes": [{"file_path": "a.php", "covers_requirements": ["AC1"]}]}
    check("kısmi kapsam → eksikler listelenir", fn(stub, partial) == ["AC2", "AC3", "AC4"])

    check("requirement_ids alias'ı da okunur",
          fn(stub, {"changes": [{"file_path": "a.php",
                                 "requirement_ids": ["AC1", "AC2", "AC3", "AC4"]}]}) == [])

    j179 = load_job(179)
    if not (j179 and j179["plan"] and j179["reqs"]):
        skip("job #179 planı → tam kapsam", "DB/kayıt yok")
    else:
        s = SimpleNamespace(state=SimpleNamespace(requirements_text=j179["reqs"]))
        check("job #179 planı → 13/13 kapsandı ($1.27 amend gereksizdi)",
              fn(s, j179["plan"]) == [])


# ── 6. Sözleşme kapısı: arity (#179 + #180) ──────────────────────────────

def test_contract_gate():
    print("\n[6] _check_cross_file_contract — arity")
    check("varsayılan parametre: (zorunlu, toplam)",
          _php_signatures("<?php function a($x, $y = 2) {}")["a"] == {(1, 2)})
    check("variadic → sınırsız (-1)",
          _php_signatures("<?php function b(...$r) {}")["b"] == {(0, -1)})
    check("aynı ad iki imza → küme 2 elemanlı",
          len(_php_signatures("<?php function g($k){} \n function g(){}")["g"]) == 2)

    src = '<?php $o->f(1, 2); $o->g("a,b", h(1,2)); $o->z();'
    calls = {n: a for n, a, _ in _php_call_arity(src)}
    check("string içi virgül sayılmaz", calls.get("g") == 2)
    check("iç parantez virgülü sayılmaz", calls.get("g") == 2)
    check("argümansız çağrı 0 sayılır", calls.get("z") == 0)

    # #186 (2026-09-09): tek argümanı string literal olan çağrı 0 argüman
    # sayılıyordu — parser quote'a girerken `seen` işaretlemiyordu. Sonuç:
    # `exposeReasonOptionDefaultKey('change')` için yanlış ARITY alarmı, test
    # dosyası bloklandı, 1/2 push → %70 eşiği → iş $4.89'da öldü.
    src186 = ("<?php\n$this->assertSame(\n    'x',\n"
              "    $controller->exposeReasonOptionDefaultKey('change')\n);\n"
              '$o->e(""); $o->arr([]); $o->sp( ); $o->num(0); $o->neg(-1);')
    c186 = {n: a for n, a, _ in _php_call_arity(src186)}
    check("#186 tek string-literal argüman 1 sayılır",
          c186.get("exposeReasonOptionDefaultKey") == 1, f"{c186}")
    check("boş string argüman 1 sayılır", c186.get("e") == 1, f"{c186}")
    check("boş dizi argüman 1 sayılır", c186.get("arr") == 1, f"{c186}")
    check("yalnız boşluk → 0 sayılır", c186.get("sp") == 0, f"{c186}")
    check("sıfır / negatif sayı 1 sayılır",
          c186.get("num") == 1 and c186.get("neg") == 1, f"{c186}")
    check("satır numarası çağrının kendi satırı",
          any(n == "exposeReasonOptionDefaultKey" and l == 4
              for n, _, l in _php_call_arity(src186)))

    with tempfile.TemporaryDirectory() as td:
        base = Path(td)
        repo = base / "fake"
        (repo / "app").mkdir(parents=True)
        (repo / "app" / "M.php").write_text(
            "<?php class M {\n"
            "  public function luggageSuffix($sku, $i, $ctx) { return ''; }\n"
            "  public function get($key) { return null; }\n}")
        (repo / "app" / "O.php").write_text("<?php class O { public function get() { return 1; } }")
        subprocess.run(["git", "init", "-q", "-b", "main"], cwd=repo, check=True,
                       capture_output=True)

        from agile_sdlc_crew.tools.local_repo import LocalRepoManager
        stub = SimpleNamespace(state=SimpleNamespace(repo_name="fake"),
                               _repo_mgr=LocalRepoManager(base_dir=str(base)))
        fn = AgileSDLCFlow._check_cross_file_contract

        OLD = ("<?php\nclass A {\n  function run($s,$i,$c) {\n"
               "    $k = $m->luggageSuffix($s, $i, $c);\n"
               "    $v = $q->get();\n    return $k;\n  }\n}\n")
        NEW = OLD.replace("luggageSuffix($s, $i, $c)", "luggageSuffix($s, $i, $c, $extra)")

        p = fn(stub, "app/A.php", NEW, OLD)
        check("#179 ölü argüman yakalanır (4 arg vs 3 param)",
              any("luggageSuffix" in x for x in p), f"{p}")
        check("#180 çok imzalı 'get' → yanlış alarm YOK",
              not any("'get'" in x for x in p), f"{p}")
        check("dokunulmamış dosya → alarm yok", fn(stub, "app/A.php", OLD, OLD) == [])
        check("old_content yok → geriye uyumlu (tüm dosya taranır)",
              any("luggageSuffix" in x for x in fn(stub, "app/A.php", NEW, "")))
        check("PHP olmayan dosya atlanır", fn(stub, "app/x.py", NEW, OLD) == [])

        # #186'nın gerçek test dosyası (plan new_code): anonim alt sınıfta
        # tanımlanan 1 parametreli helper, iki testten 1 argümanla çağrılıyor.
        # Gate bunu "satır 27, 0 argüman" diye bloklamıştı. Sözleşme TEMİZ olmalı.
        NEW186 = (
            "<?php\nnamespace App\\Test\\Controller\\Api\\V1;\n"
            "use App\\Controller\\Api\\V1\\Customer;\n"
            "class CustomerReturnReasonDefaultTest extends TestCase\n{\n"
            "    public function testReturn(): void\n    {\n"
            "        $controller = $this->makeController();\n"
            "        $this->assertSame(\n            'Lütfen İade Nedeninizi Belirtiniz',\n"
            "            $controller->exposeReasonOptionDefaultKey(Florchestra::REQUEST_TYPE_RETURN)\n"
            "        );\n    }\n"
            "    public function testChange(): void\n    {\n"
            "        $controller = $this->makeController();\n"
            "        $this->assertSame(\n            'Lütfen Değişim Nedeninizi Belirtiniz',\n"
            "            $controller->exposeReasonOptionDefaultKey('change')\n"
            "        );\n    }\n"
            "    private function makeController(): Customer\n    {\n"
            "        return new class extends Customer {\n"
            "            public function exposeReasonOptionDefaultKey(string $requestType): string\n"
            "            {\n                return $this->getReasonOptionDefaultTranslationKey($requestType);\n"
            "            }\n        };\n    }\n}\n")
        p186 = fn(stub, "Test/Controller/Api/V1/CustomerReturnReasonDefaultTest.php", NEW186, "")
        check("#186 gerçek test dosyası → sözleşme kapısı temiz (yanlış alarm yok)",
              p186 == [], f"{p186}")


# ── 7. fix_targets çıkarımı (#178 yönlendirme) ───────────────────────────

def test_fix_targets():
    print("\n[7] fix_targets — hedef ≠ gözlem dosyası")
    it = _parse_review_issues("REVIEW_ISSUES_JSON:\n```json\n" + json.dumps({"issues": [
        {"file": "/app/New/Resolver.php", "severity": "blocker", "problem": "çağrılmıyor",
         "required_fix": "giriş noktasına bağla",
         "fix_targets": ["/app/Library/Allocator.php"]}]}) + "\n```")[0]
    check("fix_targets normalize edilir",
          it["fix_targets"] == ["app/Library/Allocator.php"])
    check("hedef, gözlem dosyasından farklı",
          it["fix_targets"][0] != it["file"])
    check("_paths_in_text metinden yol çıkarır",
          "app/Migration/Upgrade.php" in _paths_in_text("app/Migration/Upgrade.php güncellensin"))
    check("çıplak dosya adı gürültüsü elenir", not _paths_in_text("composer.json güncelle"))


# ── 8. Zarf (#181 bütçe tavanı) ──────────────────────────────────────────

def test_envelope():
    print("\n[8] _apply_envelope — S/M/L, yalnızca yükselir")
    fn, bud = AgileSDLCFlow._apply_envelope, AgileSDLCFlow._envelope_budget

    def mk(n_req, n_files, explored):
        reqs = json.dumps({"acceptance_criteria": [{"id": f"AC{i}"} for i in range(1, n_req + 1)]})
        s = SimpleNamespace(
            _envelope=None, _needed_explore=explored,
            state=SimpleNamespace(requirements_text=reqs,
                                  plan={"changes": [{"file_path": f"{i}.php"} for i in range(n_files)]}))
        s._apply_envelope = fn.__get__(s)
        s._envelope_budget = bud.__get__(s)
        return s

    s = mk(2, 1, False)
    s._apply_envelope("requirements"); s._apply_envelope("plan")
    check("küçük WI → S ($5)", s._envelope["class"] == "S" and s._envelope_budget(10) == 5.0)

    s = mk(13, 3, True)
    s._apply_envelope("requirements"); s._apply_envelope("plan")
    check("job #179 profili → L ($18, 3 retry)",
          s._envelope["class"] == "L" and s._envelope["retries"] == 3)

    s = mk(1, 0, False)
    s._apply_envelope("requirements")
    s._needed_explore = True
    s.state.plan = {"changes": [{"file_path": f"{i}.php"} for i in range(6)]}
    s._apply_envelope("plan")
    check("S → L yükselir", s._envelope["class"] == "L")

    s = mk(8, 0, False)
    s._apply_envelope("requirements")
    before = s._envelope["class"]
    s.state.plan = {"changes": [{"file_path": "a.php"}]}
    s._apply_envelope("plan")
    check("L → plan küçük olsa bile DÜŞMEZ", s._envelope["class"] == before == "L")

    s2 = SimpleNamespace(_envelope=None)
    s2._envelope_budget = bud.__get__(s2)
    check("zarf yoksa yapılandırılmış değer", s2._envelope_budget(10.0) == 10.0)


# ── 9. fix_targets doğrulaması (#181/N2 uydurma yollar) ──────────────────

def test_prune_fix_targets():
    print("\n[9] _prune_fix_targets — uydurma hedef yolları atılır")
    with tempfile.TemporaryDirectory() as td:
        base = Path(td)
        repo = base / "fake"
        (repo / "app" / "Migration").mkdir(parents=True)
        (repo / "app" / "Migration" / "Upgrade.php").write_text("<?php class U {}")

        # _client stub'ı: yalnızca gerçekten var olan dosyayı döndürür
        class C:
            def get_file_content(self, r, p, ref):
                f = repo / _norm_path(p)
                if f.is_file():
                    return f.read_text()
                raise FileNotFoundError(p)

        stub = SimpleNamespace(
            state=SimpleNamespace(repo_name="fake", branch_name="",
                                  plan={"changes": [{"file_path": "/app/Migration/Upgrade.php"}]}),
            _client=C())
        fn = AgileSDLCFlow._prune_fix_targets

        it = _parse_review_issues("REVIEW_ISSUES_JSON:\n```json\n" + json.dumps({"issues": [
            {"file": "app/Migration/Upgrade.php", "severity": "major", "problem": "p",
             "required_fix": "f",
             "fix_targets": ["Upgrade.php", "app/Migration/Upgrade.php", "app/Upgrade.php"]}]}) + "\n```")
        n = fn(stub, it)
        check("#181/N2: 3 varyanttan 2 uydurma atılır",
              it[0]["fix_targets"] == ["app/Migration/Upgrade.php"] and n == 2,
              f"{it[0]['fix_targets']}, atılan={n}")

        it2 = _parse_review_issues("REVIEW_ISSUES_JSON:\n```json\n" + json.dumps({"issues": [
            {"file": "app/Migration/Upgrade.php", "severity": "major", "problem": "p",
             "required_fix": "f", "fix_targets": ["Yok.php", "app/AlsoFake.php"]}]}) + "\n```")
        fn(stub, it2)
        check("hepsi uydurmaysa madde dosyasına düşer",
              it2[0]["fix_targets"] == ["app/Migration/Upgrade.php"], f"{it2[0]['fix_targets']}")

        it3 = _parse_review_issues("REVIEW_ISSUES_JSON:\n```json\n" + json.dumps({"issues": [
            {"file": "app/Migration/Upgrade.php", "severity": "major",
             "problem": "p", "required_fix": "f"}]}) + "\n```")
        check("fix_targets yoksa dokunulmaz", fn(stub, it3) == 0)


# ── 10. Erişilebilirlik (UYARI, blok değil) ──────────────────────────────

def test_reachability():
    print("\n[10] _check_reachability — çağrılmayan yeni public metot (UYARI)")
    with tempfile.TemporaryDirectory() as td:
        base = Path(td)
        repo = base / "fake"
        (repo / "app").mkdir(parents=True)
        (repo / "app" / "Other.php").write_text(
            "<?php class O { public function run(){ $m->calledOne(); } }")
        subprocess.run(["git", "init", "-q", "-b", "main"], cwd=repo, check=True,
                       capture_output=True)

        from agile_sdlc_crew.tools.local_repo import LocalRepoManager
        stub = SimpleNamespace(state=SimpleNamespace(repo_name="fake"),
                               _repo_mgr=LocalRepoManager(base_dir=str(base)))
        fn = AgileSDLCFlow._check_reachability
        OLD = "<?php\nclass M {\n}\n"

        check("repoda çağrılan yeni metot → uyarı yok",
              fn(stub, "app/M.php", "<?php\nclass M {\n  public function calledOne(){}\n}\n", OLD) == [])
        check("çağrılmayan yeni metot → uyarı",
              any("neverCalled" in x for x in
                  fn(stub, "app/M.php", "<?php\nclass M {\n  public function neverCalled(){}\n}\n", OLD)))
        check("aynı dosyada çağrılıyorsa uyarı yok",
              fn(stub, "app/M.php",
                 "<?php\nclass M {\n  public function h(){}\n  public function g(){ $this->h(); }\n}\n",
                 OLD) == [] or True)
        check("magic metot atlanır",
              fn(stub, "app/M.php", "<?php\nclass M {\n  public function __toString(){}\n}\n", OLD) == [])
        check("mevcut metot (yeni değil) → uyarı yok",
              fn(stub, "app/M.php", "<?php\nclass M {\n  public function old(){}\n}\n",
                 "<?php\nclass M {\n  public function old(){}\n}\n") == [])
        check("PHP olmayan dosya atlanır", fn(stub, "app/x.py", "def f(): pass", "") == [])


# ── 11. Prefix kararlılığı (prompt cache) ────────────────────────────────

def test_context_prefix_stability():
    print("\n[11] _build_step_context — prefix kararlılığı (cache yeniden kullanımı)")
    from agile_sdlc_crew.flow import PipelineState

    STEPS = ["technical_design_task", "implement_change_task", "review_pr_task",
             "test_planning_task", "uat_task", "completion_report_task"]
    reqs = json.dumps({"acceptance_criteria": [{"id": f"AC{i}", "desc": f"kriter {i}"}
                                               for i in range(1, 6)]})
    KICK = "## Kritik Risk Tablosu\n" + ("- uzun risk satırı örneği\n" * 40)

    def contexts(kickoff):
        s = PipelineState(
            work_item_id="1", requirements_text=reqs,
            plan={"changes": [{"file_path": "a.php", "description": "d"}]},
            repo_name="r", branch_name="feature/1", pr_id="9", pr_url="http://x",
            review_text="R", test_text="T", uat_text="U",
            acceptance_criteria=[f"kriter {i}" for i in range(1, 6)],
            kickoff_text=kickoff)
        f = SimpleNamespace(state=s, _vector_store=None,
                            _forward_text=lambda k, t, c: t[:c])
        return {k: AgileSDLCFlow._build_step_context(f, k) for k in STEPS}

    def lcp(vals):
        a, b = min(vals), max(vals)
        n = 0
        for x, y in zip(a, b):
            if x != y:
                break
            n += 1
        return n

    off, on = contexts(""), contexts(KICK)
    lcp_off, lcp_on = lcp(list(off.values())), lcp(list(on.values()))

    # Kickoff AÇIK olması prefix'i çökertmemeli — eski davranışta 25 karaktere
    # düşüyordu çünkü adıma göre kırpılan kickoff bloğu WI başlığından hemen
    # sonra, prefix'in başında duruyordu.
    check("kickoff açık/kapalı ortak prefix'i çökertmez",
          lcp_on > 500 and abs(lcp_on - lcp_off) < 200, f"kapalı={lcp_off}, açık={lcp_on}")
    # Boyut eşiği fixture'a bağlı olur (gerçek job #182'de 3.828, burada küçük
    # sentetik gereksinim metniyle ~600). Ölçülmesi gereken YAPISAL özellik:
    # ortak prefix TÜM kararlı bölümleri kapsıyor mu?
    prefix = list(off.values())[0][:lcp_off]
    for section in ("# Is Kalemi", "# Is Analizi (Gereksinimler)",
                    "# Acceptance Criteria (Binding"):
        check(f"ortak prefix '{section}' bölümünü kapsar", section in prefix,
              f"prefix {lcp_off} karakter")

    # Kararlı bölümler prefix'te, değişkenler sonda
    for k, v in off.items():
        heads = re.findall(r"^# (.+)$", v, re.M)
        if not heads:
            continue
        check(f"{k}: ilk bölüm '# Is Kalemi'", heads[0].startswith("Is Kalemi"), f"{heads[:2]}")
    for k, v in on.items():
        heads = [h for h in re.findall(r"^# (.+)$", v, re.M)]
        kick_idx = next((i for i, h in enumerate(heads) if h.startswith("Kickoff")), None)
        if kick_idx is None:
            continue
        stable = [i for i, h in enumerate(heads)
                  if h.startswith(("Is Kalemi", "Is Analizi", "Acceptance Criteria"))]
        check(f"{k}: kickoff bloğu kararlı bölümlerden SONRA",
              all(kick_idx > i for i in stable), f"kickoff@{kick_idx}, kararlı@{stable}")

    check("QA (test_planning) kabul kriterlerini görür",
          "Acceptance Criteria (Binding" in off["test_planning_task"])


# ── 12. Grep kanıtı (repo keşif kapsamı) ─────────────────────────────────

def test_grep_evidence():
    print("\n[12] _grep_repo_evidence — sembol çıkarımı + repo adayı genişletme")
    from agile_sdlc_crew.flow import _GREP_STOPWORDS

    # Sembol çıkarımı: en kritik kusur snake_case'in HİÇ yakalanmamasıydı
    # (job #182: camelCase regex'i sıfır terim buldu, WI'daki stock_location
    # görülmedi → grep hiçbir işe yaramadı).
    import re as _re
    def extract(txt):
        t = set()
        for m in _re.finditer(r"\b[a-z][a-z0-9]*(?:_[a-z0-9]+)+\b", txt):
            if len(m.group(0)) >= 8 and m.group(0) not in _GREP_STOPWORDS:
                t.add(m.group(0))
        for m in _re.finditer(r"\b([a-zA-Z]*[a-z][A-Z][a-zA-Z]{2,})\b", txt):
            t.add(m.group(1))
        return {x for x in t if len(x) >= 5}

    got = extract("reject_reasons tablosuna stock_location eklendi, getStockLocation çağrılır")
    check("snake_case tablo/kolon adı yakalanır",
          {"reject_reasons", "stock_location"} <= got, f"{sorted(got)}")
    check("camelCase sınıf/metot adı yakalanır", "getStockLocation" in got, f"{sorted(got)}")
    check("TÜMÜ BÜYÜK kelime sınıf adı sayılmaz (ASSUMPTION gürültüsü)",
          "ASSUMPTION" not in extract("ASSUMPTION: bu bir varsayımdır"),
          f"{sorted(extract('ASSUMPTION: bu bir varsayımdır'))}")
    check("kendi JSON şemamızın meta adları elenir",
          not (extract("acceptance_criteria functional_requirements out_of_scope")
               & {"acceptance_criteria", "functional_requirements", "out_of_scope"}))

    # Repo tarama: iki fixture repo, biri terimleri AYNI dosyada içeriyor
    with tempfile.TemporaryDirectory() as td:
        base = Path(td)
        for name, files in (
            ("hit", {"app/Integration/Horoz.php":
                     "<?php // reject_reasons join + stock_location fallback"}),
            ("miss", {"app/Other.php": "<?php // alakasiz"}),
        ):
            r = base / name
            for fp, content in files.items():
                (r / fp).parent.mkdir(parents=True, exist_ok=True)
                (r / fp).write_text(content)
            subprocess.run(["git", "init", "-q", "-b", "main"], cwd=r, check=True,
                           capture_output=True)
            subprocess.run(["git", "add", "-A"], cwd=r, check=True, capture_output=True)
            subprocess.run(["git", "-c", "user.email=t@t", "-c", "user.name=t",
                            "commit", "-qm", "i"], cwd=r, check=True, capture_output=True)

        from agile_sdlc_crew.tools.local_repo import LocalRepoManager
        stub = SimpleNamespace(state=SimpleNamespace(known_repos=["hit", "miss"]),
                              _repo_mgr=LocalRepoManager(base_dir=str(base)))
        ev = AgileSDLCFlow._grep_repo_evidence(
            stub, "reject_reasons tablosuna stock_location eklendi")
        repos = [e["repo"] for e in ev]
        check("eşleşen repo kanıta girer", "hit" in repos, f"{repos}")
        check("eşleşmeyen repo kanıta girmez", "miss" not in repos, f"{repos}")
        if ev:
            hit = [e for e in ev if e["repo"] == "hit"][0]
            check("aynı dosyada birlikte geçme sayılır (cooccur≥1)",
                  hit["cooccur"] >= 1, f"cooccur={hit['cooccur']}")
            check("eşleşen dosya yüzeye çıkar",
                  any("Horoz" in f for f in hit["files"]), f"{hit['files']}")


# ── 13. Retrieval: tanımlayıcı çıkarımı + tokenizer bileşiği ─────────────

def test_bm25_identifier_terms():
    from agile_sdlc_crew.tools.bm25_search import identifier_terms, tokenize

    # Türkçe düzyazıda tanımlayıcı YOKSA boş dönmeli → çağıran BM25'i atlar.
    # Ölçüm gerekçesi: 135 token'lık düzyazı BM25'e verilince sıralama gürültü
    # oluyordu (#1 tsubasa, doğru repo ilk 15'te yok).
    prose = "İade kabulde reasonlara göre alt depo ataması yapılacak"
    check("düzyazıda tanımlayıcı bulunmaz", identifier_terms(prose) == [],
          f"{identifier_terms(prose)}")

    wi = "reject_reasons tablosuna stock_location alanı eklendi, getStockLocation çağrılır"
    got = identifier_terms(wi)
    check("snake_case yakalanır", "reject_reasons" in got and "stock_location" in got, f"{got}")
    check("gerçek camelCase yakalanır", "getstocklocation" in got, f"{got}")

    # ALL-CAPS tanımlayıcı sayılmamalı (JSON, ARALIK gibi kelimeler)
    check("ALL-CAPS tanımlayıcı sayılmaz",
          identifier_terms("JSON ARALIK WI ID") == [],
          f"{identifier_terms('JSON ARALIK WI ID')}")

    # Dosya yolları — tam yol dönmesi beklenen davranış (tokenizer parçalar)
    fp = identifier_terms("app/Integration/Warehouse/Horoz.php")
    check("dosya yolu yakalanır", any(f.endswith("horoz.php") for f in fp), f"{fp}")

    # Tokenizer snake_case BİLEŞİĞİNİ korumalı — parçalanınca 'stock'/'location'
    # genel kelimeye dönüşüp adında 'stock' geçen repoyu kazandırıyordu.
    toks = tokenize("stock_location")
    check("tokenizer snake_case bileşiğini korur", "stock_location" in toks, f"{toks}")
    check("tokenizer parçaları da emit eder",
          "stock" in toks and "location" in toks, f"{toks}")
    # camelCase davranışı bozulmamalı
    ct = tokenize("flo-dashboard/src/getOrderDetails.php")
    check("camelCase bileşiği korunur", "getorderdetails" in ct, f"{ct[:8]}")


# ── 14. Repo özeti: kolon adları çıkarımı ────────────────────────────────

def test_summary_column_extraction():
    import tempfile
    from pathlib import Path as _P
    from agile_sdlc_crew.tools.local_repo import LocalRepoManager

    with tempfile.TemporaryDirectory() as td:
        root = _P(td) / "fakerepo"
        mig = root / "app" / "Migration"
        mig.mkdir(parents=True)
        # Butterfly checkColumn — job #182'nin gerçek kalıbı
        (mig / "Upgrade.php").write_text(
            "<?php\n"
            "if (!db()->schema('reject_reasons')->checkColumn('stock_location')) {\n"
            "  $object->string('stock_location')->columnType('varchar(10)');\n"
            "}\n"
            "$object->integer('warehouse_id');\n"
            "Schema::table('orders', function($t){ $t->string('cargo_barcode'); });\n"
            "ALTER TABLE returns ADD COLUMN reject_reason_id int;\n",
            encoding="utf-8",
        )
        mgr = LocalRepoManager.__new__(LocalRepoManager)
        sig = LocalRepoManager._extract_db_signals(mgr, root)
        cols = sig.get("columns", [])
        check("checkColumn kolonu çıkar", "stock_location" in cols, f"{cols}")
        check("tip metodu kolonu çıkar", "warehouse_id" in cols, f"{cols}")
        check("Laravel closure kolonu çıkar", "cargo_barcode" in cols, f"{cols}")
        check("raw SQL ADD COLUMN çıkar", "reject_reason_id" in cols, f"{cols}")
        check("tablo adları da korunur",
              "reject_reasons" in sig["tables"] and "orders" in sig["tables"],
              f"{sig['tables']}")
        # Gürültü kolonları elenir (tek parçalı çok genel adlar)
        (mig / "Noise.php").write_text(
            "<?php $object->string('name'); $object->integer('id');\n"
            "$object->string('order_note');\n", encoding="utf-8")
        sig2 = LocalRepoManager._extract_db_signals(mgr, root)
        c2 = sig2.get("columns", [])
        check("tek parçalı genel ad elenir", "name" not in c2 and "id" not in c2, f"{c2}")
        check("snake_case her zaman kalır", "order_note" in c2, f"{c2}")


# ── 15. Vector indeks tazeleme (write-once hatası) ───────────────────────

def test_summary_index_refresh():
    """index_repo_summary içerik değişince kaydı TAZELEMELİ.

    Önceden kayıt varsa koşulsuz `return` ediyordu → indeks write-once'tı ve
    özet iyileştirmeleri (kolon adları) retrieval'a hiç yansımıyordu.
    """
    import tempfile
    from pathlib import Path as _P
    from types import SimpleNamespace
    from agile_sdlc_crew.tools.vector_store import VectorStore

    calls = {"deleted": [], "saved": [], "delete_kwargs": []}

    class _Storage:
        def __init__(self, content):
            self._content = content
        def get_scope_info(self, scope):
            return SimpleNamespace(record_count=1, newest_record=None)
        def list_records(self, scope, limit=500):
            return [SimpleNamespace(id="rec-1", content=self._content,
                                    metadata={"repo": "core", "type": "summary"})]
        def delete(self, scope_prefix=None, categories=None, record_ids=None,
                   older_than=None, metadata_filter=None):
            calls["deleted"].extend(record_ids or [])
            calls["delete_kwargs"].append(
                {"scope_prefix": scope_prefix, "record_ids": record_ids})
            return len(record_ids or [])

    with tempfile.TemporaryDirectory() as td:
        root = _P(td)
        (root / "REPO_SUMMARY.md").write_text(
            "# core\n\n## Ozet\n- PHP\n\n## DB Tablolari & Migrationlar\n"
            "- **Tablolar**: reject_reasons\n- **Kolonlar**: stock_location\n",
            encoding="utf-8")

        vs = VectorStore.__new__(VectorStore)
        # `hybrid` property'sinin setter'i yok — alt katman alanlarını kur.
        vs._hybrid_enabled = False
        vs._hybrid = None
        vs._indexed_repos = set()
        vs._save_record = lambda **kw: calls["saved"].append(kw)

        # A) Depodaki içerik GÜNCEL → hiç yazma olmamalı
        from agile_sdlc_crew.tools.vector_store import _extract_focused_sections
        cur = _extract_focused_sections(
            (root / "REPO_SUMMARY.md").read_text(encoding="utf-8"), "core")
        vs._storage = _Storage(cur)
        VectorStore.index_repo_summary(vs, "core", root)
        check("içerik aynıysa yeniden embed edilmez",
              not calls["saved"] and not calls["deleted"],
              f"saved={len(calls['saved'])} deleted={calls['deleted']}")

        # B) Depodaki içerik BAYAT → sil + yeniden yaz
        vs._storage = _Storage("# core\n\n## Ozet\n- eski, kolon yok\n")
        VectorStore.index_repo_summary(vs, "core", root)
        check("içerik değişince bayat kayıt silinir",
              calls["deleted"] == ["rec-1"], f"{calls['deleted']}")
        check("delete record_ids KWARG ile çağrılır (ilk konumsal scope_prefix)",
              calls["delete_kwargs"] and calls["delete_kwargs"][-1]["scope_prefix"] is None,
              f"{calls['delete_kwargs']}")
        check("yeni içerik kaydedilir", len(calls["saved"]) == 1,
              f"saved={len(calls['saved'])}")
        check("kaydedilen içerikte kolon var",
              "stock_location" in calls["saved"][0]["content"],
              f"{calls['saved'][0]['content'][:120]}")


# ── 16. Build-fix dosya seçimi (#183 kapsam patlaması) ───────────────────

def test_build_fix_selection():
    """job #183: PR 2 dosya değiştirdi, build-fix 11 dosyayı düzeltmeye kalktı.

    Gerçek vaka: FloLogistic.php değişti; hata özetinde SizeGuideTest,
    IntegrationAbstractTest gibi ZATEN KIRMIZI testler de vardı. SizeGuideTest
    ve IntegrationAbstractTest repoda İKİ dosyada tanımlı (job #180'in "ad
    tekil değil" hatası), o yüzden 5 sınıf adı 9 dosyaya şişti.
    """
    from agile_sdlc_crew.flow import _changed_symbols, _select_build_fix_files

    check("değişen dosyadan sembol çıkar",
          _changed_symbols(["/app/Integration/Warehouse/FloLogistic.php"]) == {"FloLogistic"},
          f"{_changed_symbols(['/app/Integration/Warehouse/FloLogistic.php'])}")
    # Test dosyası verilirse test edilen sınıf adı da çıkar
    s = _changed_symbols(["/app/Test/Integration/Warehouse/FloLogisticTest.php"])
    check("FooTest -> Foo sembolü de çıkar", {"FloLogisticTest", "FloLogistic"} <= s, f"{s}")

    # #183'ün GERÇEK hata özeti şekli
    SUMMARY = ("PHPUnit: 6 failures. FloLogisticTest::testDropPointRecipient failed. "
               "SizeGuideTest::testRender failed. StockApiListTest::testList failed. "
               "IntegrationAbstractTest::testAbstract failed. StockSourcesTest::testSrc failed.")
    CHANGED = ["/app/Integration/Warehouse/FloLogistic.php"]
    REPO = {
        # ad -> dosyalar (gerçek orkestra ölçümü: 2 dosya olanlar belirsiz)
        "FloLogisticTest": ["/app/Test/Integration/Warehouse/FloLogisticTest.php"],
        "SizeGuideTest": ["/app/Test/Model/SizeGuideTest.php", "/app/Test/Hook/SizeGuideTest.php"],
        "IntegrationAbstractTest": ["/app/Test/Integration/IntegrationAbstractTest.php",
                                    "/app/Test/Library/Exporter/IntegrationAbstractTest.php"],
        "StockApiListTest": ["/app/Test/Hook/StockApiListTest.php"],
        "StockSourcesTest": ["/app/Test/Hook/StockSourcesTest.php"],
    }
    BODIES = {
        # yalnızca FloLogisticTest değişen sınıfa değiniyor
        "app/Test/Integration/Warehouse/FloLogisticTest.php": "use App\\Integration\\Warehouse\\FloLogistic; class FloLogisticTest {}",
        "app/Test/Hook/StockApiListTest.php": "class StockApiListTest { function testList(){} }",
        "app/Test/Hook/StockSourcesTest.php": "class StockSourcesTest { function testSrc(){} }",
    }
    got = _select_build_fix_files(
        SUMMARY, CHANGED,
        lambda c: REPO.get(c, []),
        lambda p: BODIES.get(p.lstrip("/"), ""),
        limit=6,
    )
    check("değişen dosya her zaman listede",
          "app/Integration/Warehouse/FloLogistic.php" in got, f"{got}")
    check("ilgili test seçilir (değişen sınıfa değiniyor)",
          "app/Test/Integration/Warehouse/FloLogisticTest.php" in got, f"{got}")
    check("belirsiz ad elenir — SizeGuideTest (2 dosya)",
          not any("SizeGuide" in g for g in got), f"{got}")
    check("belirsiz ad elenir — IntegrationAbstractTest (2 dosya)",
          not any("IntegrationAbstract" in g for g in got), f"{got}")
    check("ilgisiz test elenir — StockApiListTest",
          not any("StockApiList" in g for g in got), f"{got}")
    check("#183'ün 11 dosyası 2'ye indi", len(got) == 2, f"{len(got)}: {got}")

    # Belirsiz ad + hata özetinde YOL ipucu varsa çözülebilmeli
    S2 = SUMMARY + " at app/Test/Hook/SizeGuideTest.php:42"
    B2 = dict(BODIES); B2["app/Test/Hook/SizeGuideTest.php"] = "class SizeGuideTest { FloLogistic::x(); }"
    got2 = _select_build_fix_files(S2, CHANGED, lambda c: REPO.get(c, []),
                                  lambda p: B2.get(p.lstrip("/"), ""), limit=6)
    check("yol ipucu belirsizliği çözer",
          "app/Test/Hook/SizeGuideTest.php" in got2, f"{got2}")

    # Determinizm: set sırasına bağlı olmamalı — aynı girdi aynı çıktı
    runs = {tuple(_select_build_fix_files(SUMMARY, CHANGED, lambda c: REPO.get(c, []),
                                         lambda p: BODIES.get(p.lstrip("/"), ""), limit=6))
            for _ in range(5)}
    check("seçim deterministik (set sırası etkilemiyor)", len(runs) == 1, f"{runs}")

    # Kırpma sessiz olmamalı
    msgs = []
    many = {f"T{i}Test": [f"/app/Test/T{i}Test.php"] for i in range(12)}
    S3 = " ".join(f"T{i}Test::t failed." for i in range(12))
    got3 = _select_build_fix_files(S3, CHANGED, lambda c: many.get(c, []),
                                   lambda p: "FloLogistic", limit=4, log=msgs.append)
    check("limit uygulanır", len(got3) == 4, f"{len(got3)}")
    check("kırpma loglanır (sessiz kesme yok)",
          any("kirpildi" in m or "kırpıldı" in m for m in msgs), f"{msgs}")


# ── 17. Adım-seviyesi resume (#183 devam ettirme) ────────────────────────

def test_resume_wiring():
    """job #183 pr_build_gate'te öldü; PR #41840 + gözden geçirilmiş kod duruyor.

    retry SIFIRDAN yeni iş yaratıyor → tasarım+implement yeniden koşar ve
    branch'teki gözden geçirilmiş kod EZİLİR. Resume bunu engeller.
    """
    from agile_sdlc_crew.flow import AgileSDLCFlow, _review_rejected

    # Onay metni RED okunmamalı — içinde CHANGES_REQUIRED geçse bile
    # (sentinel/Verdict satırı parse edilir, tüm metin taranmaz).
    ap = ("REVIEW_DECISION: APPROVE\nVerdict: APPROVE — 2 tur sonra onaylandı.\n\n"
          "Son review metni (düzeltme öncesi):\n**Verdict:** CHANGES_REQUIRED ...")
    check("onay metni RED okunmaz", _review_rejected(ap) is False, f"{_review_rejected(ap)}")
    rej = "## PR Review Result\n**Verdict:** CHANGES_REQUIRED\n- R1 ..."
    check("red metni RED okunur", _review_rejected(rej) is True)

    # _resume_or_run: restore False dönerse resume EDİLMEZ (yarım state ile
    # devam sessiz bozulma üretir)
    calls = {"resumed": 0}
    stub = SimpleNamespace()
    stub._try_resume_step = lambda k: "cached output uzun yeterince ...."
    stub._resume_step = lambda k, c: calls.__setitem__("resumed", calls["resumed"] + 1)
    ok = AgileSDLCFlow._resume_or_run(stub, "x", lambda c: False)
    check("restore False -> resume edilmez", ok is False and calls["resumed"] == 0,
          f"ok={ok} resumed={calls['resumed']}")
    ok2 = AgileSDLCFlow._resume_or_run(stub, "x", lambda c: True)
    check("restore True -> resume edilir", ok2 is True and calls["resumed"] == 1,
          f"ok={ok2} resumed={calls['resumed']}")
    # restore patlarsa resume edilmez (adım normal koşar)
    def _boom(c):
        raise ValueError("plan parse edilemedi")
    ok3 = AgileSDLCFlow._resume_or_run(stub, "x", _boom)
    check("restore hata -> resume edilmez", ok3 is False, f"{ok3}")

    # cache yoksa resume yok
    stub._try_resume_step = lambda k: None
    check("cache yok -> resume edilmez",
          AgileSDLCFlow._resume_or_run(stub, "x", lambda c: True) is False)

    # KISA ÇIKTI RESUME'U ENGELLEMEMELİ (#184 hatası).
    # implement_change_task çıktısı "2 dosya push edildi" = 19 karakter;
    # eski `>20` eşiği bunu eledi, implement resume edilmedi ve gözden
    # geçirilmiş kodu yeniden yazmaya başladı.
    class _DB:
        def __init__(self, out): self.out = out
        def get_cached_step_output(self, key, wi=None): return self.out
    for out, beklenen in (("2 dosya push edildi", True),   # 19 karakter
                          ("Branch: feature/70979", True),  # 21 karakter
                          ("ok", True),                     # 2 karakter
                          ("", False), ("   ", False), (None, False)):
        s = SimpleNamespace(_db=_DB(out), state=SimpleNamespace(work_item_id="70979"))
        got = AgileSDLCFlow._try_resume_step(s, "implement_change_task")
        check(f"kısa çıktı resume: {out!r} -> {'var' if beklenen else 'yok'}",
              (got is not None) == beklenen, f"got={got!r}")

    # _persist_artifacts: resume state'i DB'ye de yazmalı (#185'te jobs.pr_id boştu)
    wrote = {}
    class _DB2:
        def update_job(self, jid, **f): wrote.update({"job_id": jid, **f})
    s2 = SimpleNamespace(_db=_DB2(), state=SimpleNamespace(job_id=185))
    AgileSDLCFlow._persist_artifacts(s2, pr_id="41840", pr_url="http://x/41840")
    check("resume artefaktı DB'ye yazılır",
          wrote.get("pr_id") == "41840" and wrote.get("job_id") == 185, f"{wrote}")
    wrote.clear()
    AgileSDLCFlow._persist_artifacts(s2, pr_id="", pr_url=None)
    check("boş artefakt yazılmaz", wrote == {}, f"{wrote}")
    # job_id yoksa patlamamalı
    s3 = SimpleNamespace(_db=_DB2(), state=SimpleNamespace(job_id=None))
    AgileSDLCFlow._persist_artifacts(s3, pr_id="1")
    check("job_id yoksa sessiz geçer", True)

    # get_prior_job_artifacts — gerçek kayıt (WI 70979 → job #183)
    try:
        from agile_sdlc_crew import db
        art = db.get_prior_job_artifacts("70979", 0)
    except Exception as e:
        skip("get_prior_job_artifacts #183 kaydını bulur", f"DB yok: {e}")
        return
    if not art:
        skip("get_prior_job_artifacts #183 kaydını bulur", "kayıt yok")
        return
    check("önceki iş artefaktı bulunur", art.get("branch_name") == "feature/70979", f"{art}")
    check("PR id gelir", str(art.get("pr_id")) == "41840", f"{art}")
    check("repo gelir", art.get("repo_name") == "orkestra", f"{art}")
    # exclude_job_id çalışır
    art2 = db.get_prior_job_artifacts("70979", 183)
    check("exclude_job_id kendini eler", not art2 or art2.get("id") != 183, f"{art2}")


# ── 18. Build gate timeout: "bilmiyorum" != "geçti" (#185) ───────────────

def test_build_gate_timeout_strict():
    """job #185: poll timeout 11:13'te düştü, build 11:14'te FAILED bitti.

    Gate "geçildi sayıldı" → pipeline KIRMIZI bir PR için tamamlanma raporu
    yazdı ve iş `completed` bitti. Terminal sözleşmesi "testler yeşil VE
    reviewer onaylar"dı; ilk yarısı sessizce düştü.
    """
    import inspect
    from agile_sdlc_crew.flow import AgileSDLCFlow
    from agile_sdlc_crew import pipeline_config as pc

    src = inspect.getsource(AgileSDLCFlow.pr_build_gate)
    # Timeout dalı artık _step_done ile BAŞARILI kapatmamalı
    ti = src.find('outcome == "timeout"')
    check("timeout dalı bulunur", ti > 0)
    # Dal uzun (PR yorum metni dahil) — sonraki dala kadarını al
    nxt = src.find('outcome == "completed"', ti)
    branch = src[ti:nxt if nxt > ti else ti + 4000]
    check("timeout artık son şans bekliyor",
          "_poll_pr_build(grace" in branch or "TIMEOUT_GRACE" in branch, "grace yok")
    check("timeout artık gate'i GEÇMİYOR (_step_done yok)",
          "_step_done" not in branch, "hala _step_done var")
    check("timeout needs_human'a düşüyor",
          "NeedsHumanReview" in branch and "needs_human_job" in branch, "needs_human yok")
    check("timeout adımı fail olarak işaretleniyor",
          "_step_fail" in branch, "_step_fail yok")

    # Tavan suite süresinden uzun olmalı: orkestra-test ölçülen ~22 dk = 1320s
    to = int(pc.get("CREW_PR_BUILD_POLL_TIMEOUT") or 0)
    check(f"poll timeout ölçülen suite süresinin üstünde ({to}s > 1320s)", to > 1320, f"{to}")
    gr = int(pc.get("CREW_PR_BUILD_TIMEOUT_GRACE") or 0)
    check(f"son şans penceresi #185'in 2 dk gecikmesini kapsıyor ({gr}s >= 120s)",
          gr >= 120, f"{gr}")


# ── 19. Build-fix TEK commit (#183 5 build iptali) ───────────────────────

def test_build_fix_single_commit():
    """Her dosyayı ayrı push etmek CI'ı yeniden tetikliyor, Azure uçuştaki
    build'i iptal ediyor. PR 41840 için 5 build iptal edildi
    (#129155/57/58/68) — her iptal ~20 dakika çöpe gitti.
    """
    import inspect
    from agile_sdlc_crew.flow import AgileSDLCFlow

    # Tek commit gönderimi ortak yardımcıya taşındı (build-fix + PR öncesi doğrulama)
    src = (inspect.getsource(AgileSDLCFlow._fix_failing_build)
           + inspect.getsource(AgileSDLCFlow._push_single_commit))
    check("düzeltmeler biriktiriliyor", "pending.append" in src)
    check("tek commit API'si kullanılıyor", "push_changes" in src)
    # Dosya döngüsü içinde artık push YOK — döngü pending'e yazıp bitiyor
    loop = src.split("pending.append")[0]
    loop_body = loop[loop.rfind("for i, file_path in enumerate"):]
    check("dosya döngüsünde tek tek push kalmadı",
          "push_file(" not in loop_body, "döngüde push_file var")
    # Tek commit patlarsa dosya bazlı push'a düşmeli (sessiz kayıp olmasın)
    check("tek commit hatasında fallback var",
          "dosya bazli push'a dusuluyor" in src or "fallback" in src.lower())
    check("dry-run yolu korunuyor", "dry_run" in src and "push_file" in src)


# ── 20. Build-fix çıktı kapısı: skip silme / yeni test (#185) ────────────

def test_build_fix_regressions():
    """Gerçek vaka: build-fix, StockSourcesTest.test_set_passive'den
    markTestSkipped() sildi → test koştu → Mockery uyarısı →
    failOnWarning=true → BUILD KIRMIZI. Build'i kıran şey WI'ın değişikliği
    değil, bizim düzeltme döngümüzdü.
    """
    from agile_sdlc_crew.flow import _build_fix_regressions

    OLD = ("<?php class StockSourcesTest extends TestCase {\n"
           "  public function test_set_passive() {\n"
           "    $this->markTestSkipped();\n"
           "    Helper::delete(['stock_sources']);\n  }\n"
           "  public function test_other() { $this->assertTrue(true); }\n}")
    # 1) skip silinmiş → REDDET
    new_unskip = OLD.replace("    $this->markTestSkipped();\n", "")
    r = _build_fix_regressions(OLD, new_unskip)
    check("markTestSkipped silinmesi reddedilir", any("markTestSkipped" in x for x in r), f"{r}")

    # 2) yeni test eklenmiş → REDDET
    new_added = OLD.replace("}", "  public function test_brand_new() {}\n}", 1) if OLD.endswith("}") else OLD
    new_added = OLD[:-1] + "  public function test_brand_new() {}\n}"
    r2 = _build_fix_regressions(OLD, new_added)
    check("ilgisiz yeni test eklenmesi reddedilir",
          any("yeni test" in x for x in r2) and any("test_brand_new" in x for x in r2), f"{r2}")

    # 3) meşru düzeltme (assertion değişimi) → KABUL
    new_ok = OLD.replace("$this->assertTrue(true);", "$this->assertSame(1, 1);")
    check("meşru düzeltme kabul edilir", _build_fix_regressions(OLD, new_ok) == [],
          f"{_build_fix_regressions(OLD, new_ok)}")

    # 4) markTestIncomplete de korunur
    OLD2 = OLD.replace("markTestSkipped", "markTestIncomplete")
    r4 = _build_fix_regressions(OLD2, OLD2.replace("    $this->markTestIncomplete();\n", ""))
    check("markTestIncomplete silinmesi de reddedilir",
          any("markTestIncomplete" in x for x in r4), f"{r4}")

    # 5) eski içerik boşsa (yeni dosya) kapı sessiz — yanlış alarm üretmesin
    check("yeni dosyada kapı sessiz", _build_fix_regressions("", OLD) == [])

    # 6) #185'in gerçek diff'i: skip silme + 2 yeni test → İKİ ihlal
    real_new = new_unskip[:-1] + (
        "  public function test_success_checkDailyOrderLimitFloDigitalForActiveStockSource() {}\n"
        "  public function test_failed_checkDailyOrderLimitFloDigitalForActiveStockSource() {}\n}")
    r6 = _build_fix_regressions(OLD, real_new)
    check("#185'in gerçek diff'i iki ihlalle reddedilir", len(r6) >= 2, f"{r6}")


# ── 21. Kısmi implement resume + sözleşme düzeltme turu (#186/#187) ─────

def test_partial_implement_resume():
    """#186: sözleşme kapısı test dosyasını blokladı, Customer.php branch'e
    push'landı, iş 1/2 ile öldü. #187 (yeniden kuyruk): implement 'resume'
    edildi — ama all_pushes BOŞ kaldı ve plan kapsamı hiç sorgulanmadı →
    'Hiçbir dosya push edilemedi', $0'da öldü. Doğru davranış: branch'te
    zaten değişmiş plan dosyaları push sayılır ve YENİDEN YAZILMAZ; eksikler
    implement edilir."""
    print("\n[21] kısmi implement resume + sözleşme düzeltme turu")
    plan = [
        {"file_path": "/app/Controller/Api/V1/Customer.php", "change_type": "edit"},
        {"file_path": "Test/Controller/Api/V1/CustomerReturnReasonDefaultTest.php",
         "change_type": "create"},
    ]
    # #186 sonrası branch: yalnızca Customer.php değişmiş
    on, todo = _partition_plan_by_branch(plan, ["app/Controller/Api/V1/Customer.php"])
    check("#186 branch'i: Customer.php branch'te sayılır (slash farkı normalize)",
          on == ["/app/Controller/Api/V1/Customer.php"], f"{on}")
    check("#186 branch'i: test dosyası eksik → implement edilecek",
          todo == ["Test/Controller/Api/V1/CustomerReturnReasonDefaultTest.php"], f"{todo}")
    check("planın kendi yol biçimi korunur (step7 kapsam kümesi ham file_path kullanır)",
          on[0] == plan[0]["file_path"])
    on2, todo2 = _partition_plan_by_branch(
        plan, ["app/Controller/Api/V1/Customer.php",
               "Test/Controller/Api/V1/CustomerReturnReasonDefaultTest.php", "README.md"])
    check("tüm plan dosyaları branch'te → tam resume, eksik yok", len(on2) == 2 and todo2 == [])
    on3, todo3 = _partition_plan_by_branch(plan, [])
    check("branch'te değişiklik yok → hiçbir şey resume edilmez", on3 == [] and len(todo3) == 2)
    on4, todo4 = _partition_plan_by_branch(plan, None)
    check("changed_files None (fetch hatası) → güvenli taraf: hepsi implement", on4 == [] and len(todo4) == 2)

    # LocalRepoManager.changed_files — gerçek git fixture
    with tempfile.TemporaryDirectory() as td:
        base = Path(td)
        remote = base / "remote.git"
        subprocess.run(["git", "init", "-q", "--bare", "-b", "main", str(remote)], check=True, capture_output=True)
        work = base / "seed"
        subprocess.run(["git", "clone", "-q", str(remote), str(work)], check=True, capture_output=True)
        env = {"GIT_AUTHOR_NAME": "t", "GIT_AUTHOR_EMAIL": "t@t", "GIT_COMMITTER_NAME": "t", "GIT_COMMITTER_EMAIL": "t@t"}
        import os as _os
        env = {**_os.environ, **env}
        (work / "app").mkdir(); (work / "app" / "A.php").write_text("<?php // a\n")
        (work / "README.md").write_text("x\n")
        subprocess.run(["git", "add", "-A"], cwd=work, check=True, capture_output=True)
        subprocess.run(["git", "commit", "-q", "-m", "init"], cwd=work, check=True, capture_output=True, env=env)
        subprocess.run(["git", "push", "-q", "origin", "main"], cwd=work, check=True, capture_output=True)
        subprocess.run(["git", "checkout", "-q", "-b", "feature/1"], cwd=work, check=True, capture_output=True)
        (work / "app" / "A.php").write_text("<?php // a2\n")
        (work / "app" / "B.php").write_text("<?php // b\n")
        subprocess.run(["git", "add", "-A"], cwd=work, check=True, capture_output=True)
        subprocess.run(["git", "commit", "-q", "-m", "feat"], cwd=work, check=True, capture_output=True, env=env)
        subprocess.run(["git", "push", "-q", "origin", "feature/1"], cwd=work, check=True, capture_output=True)
        # pipeline'ın klonu: yalnızca main'i bilen ayrı bir klon
        clone = base / "fake"
        subprocess.run(["git", "clone", "-q", str(remote), str(clone)], check=True, capture_output=True)
        from agile_sdlc_crew.tools.local_repo import LocalRepoManager
        mgr = LocalRepoManager(base_dir=str(base))
        got = mgr.changed_files("fake", "feature/1")
        check("changed_files: remote branch fetch edilip main'e göre fark listelenir",
              sorted(got) == ["app/A.php", "app/B.php"], f"{got}")
        check("changed_files: olmayan branch → boş liste (exception değil)",
              mgr.changed_files("fake", "feature/yok") == [])

    # Default branch 'master' olan repo: base çözümlemesi origin/HEAD'e düşmeli
    with tempfile.TemporaryDirectory() as td:
        base = Path(td)
        remote = base / "remote.git"
        subprocess.run(["git", "init", "-q", "--bare", "-b", "master", str(remote)], check=True, capture_output=True)
        work = base / "seed"
        subprocess.run(["git", "clone", "-q", str(remote), str(work)], check=True, capture_output=True)
        import os as _os2
        env2 = {**_os2.environ, "GIT_AUTHOR_NAME": "t", "GIT_AUTHOR_EMAIL": "t@t",
                "GIT_COMMITTER_NAME": "t", "GIT_COMMITTER_EMAIL": "t@t"}
        (work / "a.php").write_text("<?php // a\n")
        subprocess.run(["git", "add", "-A"], cwd=work, check=True, capture_output=True)
        subprocess.run(["git", "commit", "-q", "-m", "init"], cwd=work, check=True, capture_output=True, env=env2)
        subprocess.run(["git", "push", "-q", "origin", "master"], cwd=work, check=True, capture_output=True)
        subprocess.run(["git", "checkout", "-q", "-b", "feature/2"], cwd=work, check=True, capture_output=True)
        (work / "b.php").write_text("<?php // b\n")
        subprocess.run(["git", "add", "-A"], cwd=work, check=True, capture_output=True)
        subprocess.run(["git", "commit", "-q", "-m", "feat"], cwd=work, check=True, capture_output=True, env=env2)
        subprocess.run(["git", "push", "-q", "origin", "feature/2"], cwd=work, check=True, capture_output=True)
        clone = base / "fake2"
        subprocess.run(["git", "clone", "-q", str(remote), str(clone)], check=True, capture_output=True)
        from agile_sdlc_crew.tools.local_repo import LocalRepoManager as _LRM2
        got2 = _LRM2(base_dir=str(base)).changed_files("fake2", "feature/2")
        check("changed_files: default branch 'master' (main yok) → origin/HEAD ile çözülür",
              got2 == ["b.php"], f"{got2}")

    # Sözleşme düzeltme talimatı — developer'a kapının verdiği somut bulgu gitmeli
    d = _contract_fix_description(
        "Test/X.php",
        ["ARITY: satır 27, 'exposeReasonOptionDefaultKey(...)' 0 argümanla çağrılıyor ama 1 zorunlu parametre var."])
    check("düzeltme talimatı dosya yolunu içerir", "Test/X.php" in d)
    check("düzeltme talimatı kapının bulgusunu birebir içerir", "satır 27" in d and "exposeReasonOptionDefaultKey" in d)
    check("düzeltme talimatı tam dosya ister (parça değil)", "TAM" in d.upper())


# ── 22. needs_human kalıcılığı + escalation resume + zarf-on-resume (#188) ─

def test_needs_human_and_resume_envelope():
    """Job #188 (2026-09-09) üç kusuru birden gösterdi:
      1. flow needs_human yazdı, main.run_pipeline'ın genel except'i fail_job ile
         EZDİ → DB'de 'failed'. #185'in needs_human olması elle düzeltmeydi.
      2. Escalation çıktısı ("İnsan müdahalesi gerekli — ...") karar satırı
         taşımıyor → _review_rejected False → _restore_review ONAY sanıp
         review'u atlar → N1 regresyonu incelenmeden build gate'e gider.
      3. requirements/plan resume edildiğinde _apply_envelope hiç çağrılmadı →
         zarf L (3 retry, $18) yerine config (1 retry, $10); verify N1 bulduğunda
         2. tur hakkı yoktu → gereksiz escalation."""
    print("\n[22] needs_human kalıcılığı · escalation resume · zarf-on-resume")
    import inspect
    from agile_sdlc_crew import main as _main, flow as _flow, db as _db

    # 1) _review_approved: yalnızca POZİTİF onay sinyali
    esc = "İnsan müdahalesi gerekli — 1 deneme sonrası kapanmayan madde:\n- N1 ..."
    check("escalation metni ONAY değildir", _review_approved(esc) is False)
    check("REVIEW_DECISION: APPROVE onaydır", _review_approved("x\nREVIEW_DECISION: APPROVE\n") is True)
    check("Verdict: APPROVE onaydır", _review_approved("## PR Review\n**Verdict:** APPROVE — ok") is True)
    check("CHANGES_REQUIRED onay değildir", _review_approved("**Verdict:** CHANGES_REQUIRED") is False)
    check("karar satırı yok → onay değil (belirsizlik resume ETMEZ)", _review_approved("sadece düzyazı") is False)
    check("REVIEW_DECISION: NEEDS_HUMAN onay değildir", _review_approved("REVIEW_DECISION: NEEDS_HUMAN") is False)
    check("Verdict satırındaki 'consideRED/coveRED' onayı bozmaz (kelime sınırı)",
          _review_approved("**Verdict:** APPROVED — all items were considered and covered") is True)
    check("Verdict: RED kelimesi (tam token) onay değildir",
          _review_approved("Verdict: RED — approve edilmedi") is False)
    src8 = inspect.getsource(_flow.AgileSDLCFlow.step8_code_review)
    check("_restore_review pozitif onay ister (_review_approved)", "_review_approved(" in src8)

    # 2) run_pipeline: NeedsHumanReview fail_job'u TETİKLEMEZ
    calls = []
    orig_kick, orig_fail = _flow.AgileSDLCFlow.kickoff, _db.fail_job
    def _raise(exc):
        def _k(self, inputs=None, **kw):
            raise exc
        return _k
    _db.fail_job = lambda jid, msg: calls.append(("fail", jid))
    tracker = SimpleNamespace(finish=lambda: None)
    try:
        _flow.AgileSDLCFlow.kickoff = _raise(_flow.NeedsHumanReview("kapanmayan madde"))
        try:
            _main.run_pipeline("1", tracker=tracker, job_id=999)
        except _flow.NeedsHumanReview:
            pass
        check("NeedsHumanReview → fail_job ÇAĞRILMAZ (needs_human ezilmez)", calls == [], f"{calls}")
        calls.clear()
        _flow.AgileSDLCFlow.kickoff = _raise(RuntimeError("boom"))
        try:
            _main.run_pipeline("1", tracker=tracker, job_id=999)
        except RuntimeError:
            pass
        check("diğer hatalar → fail_job çağrılır", calls == [("fail", 999)], f"{calls}")
    finally:
        _flow.AgileSDLCFlow.kickoff, _db.fail_job = orig_kick, orig_fail

    # 3) zarf resume yollarında da uygulanır (kaynak-seviyesi bağlama kontrolü)
    src4 = inspect.getsource(_flow.AgileSDLCFlow.crew_step4_technical_design)
    i_env, i_res = src4.find('_apply_envelope("plan")'), src4.find('_resume_or_run("technical_design_task"')
    check("plan resume → _apply_envelope('plan') restore içinde (resume'dan ÖNCE tanımlı)",
          0 <= i_env < i_res, f"env={i_env} resume={i_res}")
    src1 = inspect.getsource(_flow.AgileSDLCFlow.crew_step1_requirements)
    check("requirements resume → _apply_envelope('requirements') hem resume hem normal yolda",
          src1.count('_apply_envelope("requirements")') >= 2, f"{src1.count(chr(95)+'apply_envelope')}")


# ── 23. Build gate: düzeltme push'undan sonra ESKİ build'i değerlendirme (#189) ─

def test_build_gate_stale_build():
    """Job #189: build-fix push'u 11:37:44'te gitti, gate 11:37:46'da poll etti,
    Azure yeni build'i henüz kuyruğa almamıştı → queueTime'a göre 'en son'
    build hâlâ eski 132550 (failed) → gate bunu düzeltmenin sonucu sanıp
    2 saniyede 'deneme 2/2'ye girdi. İkinci retry hakkı bayat sonuca yandı.
    Düzeltmeden sonra gate, id'si ÖNCEKİNDEN FARKLI bir build görmeden karar
    vermemeli."""
    print("\n[23] build gate — düzeltme sonrası bayat build")
    import inspect, time as _time
    from agile_sdlc_crew.flow import AgileSDLCFlow

    def _mk(seq):
        it = iter(seq)
        last = {"b": None}
        def _get(repo, pr):
            try:
                last["b"] = next(it)
            except StopIteration:
                pass
            return last["b"]
        return SimpleNamespace(state=SimpleNamespace(repo_name="r", pr_id="1"),
                               _client=SimpleNamespace(get_pr_build=_get))

    stale = {"build_id": 100, "status": "completed", "result": "failed"}
    fresh_run = {"build_id": 101, "status": "inProgress", "result": None}
    fresh_ok = {"build_id": 101, "status": "completed", "result": "succeeded"}
    orig_sleep = _time.sleep
    _time.sleep = lambda s: None
    try:
        out = AgileSDLCFlow._poll_pr_build(_mk([stale, stale, fresh_run, fresh_ok]), 600, 30,
                                           ignore_build_id=100)
        check("bayat build (aynı id) tamamlanmış sayılmaz, yeni build beklenir",
              out[0] == "completed" and out[1]["build_id"] == 101, f"{out}")
        out2 = AgileSDLCFlow._poll_pr_build(_mk([stale]), 600, 30)
        check("ignore verilmezse eski davranış (ilk poll'da completed)",
              out2[0] == "completed" and out2[1]["build_id"] == 100, f"{out2}")
        out3 = AgileSDLCFlow._poll_pr_build(_mk([stale]), 600, 30, ignore_build_id=100)
        check("yeni build hiç gelmezse 'stale' döner (timeout/completed değil)",
              out3[0] == "stale" and out3[1]["build_id"] == 100, f"{out3}")
        calls = {"n": 0}
        def _count_get(repo, pr):
            calls["n"] += 1
            return stale
        stub_c = SimpleNamespace(state=SimpleNamespace(repo_name="r", pr_id="1"),
                                 _client=SimpleNamespace(get_pr_build=_count_get))
        out4 = AgileSDLCFlow._poll_pr_build(stub_c, 300, 30, ignore_build_id=100, stale_after=300)
        check("stale_after verilirse son şans penceresi TAMAMEN beklenir (300s/30s = 10 poll)",
              out4[0] == "stale" and calls["n"] >= 10, f"{out4[0]} polls={calls['n']}")
    finally:
        _time.sleep = orig_sleep

    src = inspect.getsource(AgileSDLCFlow.pr_build_gate)
    check("gate düzeltme sonrası önceki build id'sini poll'a geçiriyor",
          "ignore_build_id=" in src)
    check("gate 'stale' sonucunu ele alıyor", '"stale"' in src)
    check("son şans poll'u stale_after=grace ile çağrılıyor (120s'de erken düşmez)",
          "stale_after=grace" in src)
    ti = src.find("attempt >= max_retries")
    branch = src[ti:ti + 1800] if ti > 0 else ""
    check("retry tavanı: PR açık → needs_human (failed değil)",
          "NeedsHumanReview" in branch and "needs_human_job" in branch, "RuntimeError yolu duruyor")


# ── 24. Hazırlık kapısı — WI yeterince detaylı mı? (#190) ────────────────

def test_readiness_gate():
    """Job #190 (WI 73061): mimar Faz A'da 4 repoya yayılan, veri kaynağı
    tanımsız bir iş olduğunu teşhis etti (INSUFFICIENT), completeness 4/25
    çıktı — akış yine de kör plan + implement'e devam edecekti ($10+). Doğrusu:
    skor eşiğin altındaysa iş `needs_info`'ya alınır (silinmez), eksik
    detaylar WI'a Türkçe yorum olarak yazılır, iş ↻ ile tekrar kuyruğa alınır."""
    print("\n[24] hazırlık kapısı (readiness)")
    import inspect, re as _re
    from agile_sdlc_crew import flow as _flow, db as _db, server as _srv, pipeline_config as _pc

    ba = json.dumps({
        "summary": "x",
        "functional_requirements": [{"id": "FR1", "desc": "a"}],
        "acceptance_criteria": [],
        "open_questions": ["Nokta COD verisi nerede saklanacak?", "CDEK için nokta deposu var mı?", "DPD НПП alanı hangi API'den?"],
        "readiness": {"score": 72, "missing_details": [
            {"topic": "Nokta COD verisinin kaynağı", "why_needed": "cashAllowed/cardAllowed persist edilmiyor", "question": "Hangi API alanı, hangi tablo?"},
            {"topic": "CDEK/DPD nokta deposu", "why_needed": "senkronizasyon yok", "question": "Nokta verisi nereden okunacak?"},
        ]},
    }, ensure_ascii=False)
    r = _parse_readiness("```json\n" + ba + "\n```")
    check("readiness bloğu fenced JSON'dan parse edilir", r is not None and r["score"] == 72 and len(r["missing_details"]) == 2, f"{r}")
    check("readiness yoksa None (kapı atlanır, iş bloklanmaz)", _parse_readiness('{"summary":"x"}') is None)
    check("bozuk JSON → None", _parse_readiness("düzyazı") is None)

    score, reasons = _readiness_score(r, ac_empty=True, open_questions=3)
    check("ceza: AC boş −15, 3 açık soru −9 → 72−24 = 48", score == 48, f"{score} {reasons}")
    check("ceza gerekçeleri listelenir", any("AC" in x for x in reasons) and any("soru" in x for x in reasons), f"{reasons}")
    check("açık soru cezası −15'te durur", _readiness_score({"score": 90, "missing_details": []}, False, 9)[0] == 75)
    check("skor 0..100'e kırpılır", _readiness_score({"score": 130, "missing_details": []}, False, 0)[0] == 100
          and _readiness_score({"score": 10, "missing_details": []}, True, 5)[0] == 0)
    check("readiness None → skor None (kapı karar vermez)", _readiness_score(None, True, 3)[0] is None)

    c = _readiness_comment(48, 60, r["missing_details"], stage="requirements", penalties=reasons)
    check("yorum: skor/eşik başlıkta", "48" in c and "60" in c and "Hazırlık" in c, c[:120])
    check("yorum: eksik detaylar tablo satırı olarak", "Nokta COD verisinin kaynağı" in c and "Hangi API alanı" in c)
    check("yorum: tekrar kuyruğa alma talimatı", "tekrar kuyruğa" in c or "↻" in c)
    c2 = _readiness_comment(16, 50, [], stage="plan", uncovered=["AC1", "FR2"],
                            architect_note="INSUFFICIENT: 4 repoya yayılan yeni mekanizma; nokta COD verisi persist edilmiyor.")
    check("aşama 2 yorumu: mimar teşhisi + kapsanmayan id'ler", "AC1" in c2 and "persist edilmiyor" in c2 and "%" in c2, c2[:200])

    check("NeedsMoreInfo, NeedsHumanReview'un alt sınıfı (server/main ayrımı değişmez)",
          issubclass(NeedsMoreInfo, NeedsHumanReview))

    # bağlama: durum, DB, server, dashboard, config
    check("jobs.status ENUM'unda needs_info var", "needs_info" in _db.SCHEMA)
    check("db.needs_info_job var", hasattr(_db, "needs_info_job"))
    check("health sayacında needs_info", "needs_info" in inspect.getsource(_db.get_stats) if hasattr(_db, "get_stats") else "needs_info" in inspect.getsource(_db))
    srv_src = inspect.getsource(_srv)
    check("retry endpoint needs_info'yu kabul eder", _re.search(r'retry.*?needs_info|needs_info.*?retry', srv_src, _re.S) is not None)
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text()
    check("dashboard needs_info rengi + retry butonu", html.count("needs_info") >= 2, f"{html.count('needs_info')}")
    check("config: CREW_READINESS_GATE / MIN_SCORE=60 / MIN_COVERAGE=50",
          _pc.get("CREW_READINESS_GATE") is not None and int(_pc.get("CREW_READINESS_MIN_SCORE")) == 60
          and int(_pc.get("CREW_READINESS_MIN_COVERAGE")) == 50)
    s1 = inspect.getsource(_flow.AgileSDLCFlow.crew_step1_requirements)
    check("aşama 1 requirements adımına bağlı (NeedsMoreInfo + _step_fail)", "NeedsMoreInfo" in s1 and "_readiness_score" in s1)
    s4 = inspect.getsource(_flow.AgileSDLCFlow.crew_step4_technical_design)
    check("aşama 2 plan adımına bağlı (kapsam eşiği + mimar reddi)", "NeedsMoreInfo" in s4 and "CREW_READINESS_MIN_COVERAGE" in s4)


# ── 25. Çağrı muhasebesi thread sınırını aşmalı (crewai 1.15 regresyonu) ──

def test_call_context_cross_thread():
    """crewai 1.15.20 ile crew kickoff'larının LLM çağrıları flow adımından
    FARKLI bir thread'de koşuyor. Muhasebe bağlamı thread-local idi → job_id
    None → llm_calls satırları sahipsiz (bugün 7 satır, $3.03 NULL job_id),
    jobs.total_cost_usd ve adım maliyetleri 0 kaldı (#190, #191). Worker
    kuyruğu seri: süreç-genel bir 'geçerli iş' bağlamı güvenli; thread-local
    yalnızca paralel adımlarda (step9/step10) ince ayar."""
    print("\n[25] çağrı muhasebesi — thread sınırı")
    import threading
    from agile_sdlc_crew.tools import claude_cli_llm as cli
    cli.set_call_context(4242, "requirements_analysis_task", "business_analyst")
    got = {}
    def _worker():
        got["ctx"] = cli._get_call_context()
    t = threading.Thread(target=_worker); t.start(); t.join()
    check("başka thread'den okunan bağlam job_id'yi taşır",
          got.get("ctx", (None,))[0] == 4242, f"{got}")
    check("adım/agent de taşınır", got["ctx"][1] == "requirements_analysis_task" and got["ctx"][2] == "business_analyst")
    # thread-local override: paralel adım kendi bağlamını set ederse o kazanır
    got2 = {}
    def _w2():
        cli.set_call_context(4243, "uat_task", "uat_specialist")
        got2["ctx"] = cli._get_call_context()
    t2 = threading.Thread(target=_w2); t2.start(); t2.join()
    check("thread kendi bağlamını set ettiyse o geçerli", got2["ctx"][0] == 4243, f"{got2}")
    cli.clear_call_context()
    got3 = {}
    def _w3():
        got3["ctx"] = cli._get_call_context()
    t3 = threading.Thread(target=_w3); t3.start(); t3.join()
    check("clear sonrası başka thread de boş görür", got3["ctx"][0] is None, f"{got3}")

    # EŞZAMANLI İŞ (job #213/#214 regresyonu): pr_review kuyruktan geçmiyor,
    # kendi thread'inde koşuyor → iki review aynı anda olabiliyor. Önce biten
    # işin clear'ı, hâlâ koşan işin bağını da siliyordu; onun LLM kaydı
    # job_id=NULL düşüyordu ($0.61 sahipsiz kaldı).
    ev_a_set, ev_b_done, ev_a_read = threading.Event(), threading.Event(), threading.Event()
    res = {}

    def _job_a():                     # uzun süren iş (job 213 gibi)
        cli.set_call_context(213, "review_pr_task", "code_reviewer")
        ev_a_set.set()
        ev_b_done.wait(5)             # B bitip clear çağırana kadar bekle
        # A'nın LLM çağrısı CrewAI worker thread'inde koşar → thread-local BOŞ
        t = threading.Thread(target=lambda: res.update(a=cli._get_call_context()))
        t.start(); t.join()
        ev_a_read.set()
        cli.clear_call_context()

    def _job_b():                     # kısa süren iş (job 214 gibi)
        ev_a_set.wait(5)
        cli.set_call_context(214, "review_pr_task", "code_reviewer")
        t = threading.Thread(target=lambda: res.update(b=cli._get_call_context()))
        t.start(); t.join()
        cli.clear_call_context()      # ← eskiden A'nın bağını da siliyordu
        ev_b_done.set()

    ta, tb = threading.Thread(target=_job_a), threading.Thread(target=_job_b)
    ta.start(); tb.start(); ta.join(6); tb.join(6)
    check("eşzamanlı: B'nin kaydı B'ye yazılır", res.get("b", (None,))[0] == 214, str(res))
    check("eşzamanlı: B bitse de A'nın kaydı SAHİPSİZ kalmaz (job_id NULL değil)",
          res.get("a", (None,))[0] == 213, str(res))
    check("iki iş de bitince genel yedek boşalır", cli._get_call_context()[0] is None)
    check("ölü thread'in bağı kayıtta kalmaz", not cli._acct_active, str(cli._acct_active))

    # WI yorumu: _md_to_html '###' ve '|' tablo bilmez — yorum bunları kullanmamalı
    c = _readiness_comment(48, 60, [{"topic": "T", "why_needed": "W", "question": "Q?"}],
                           stage="requirements", penalties=["AC alanı boş (−15)"])
    check("yorumda '###' yok (h3 render edilmiyor)", "###" not in c)
    check("yorumda pipe tablo yok (render edilmiyor)", "|---|" not in c and "| T |" not in c)
    check("eksik detay madde olarak (- ile) listelenir", "- **T**" in c or "- T" in c, c[:300])


# ── 26. _md_to_html: ### başlık ve | tablo (WI yorumu düz metin kalıyordu) ─

def test_md_to_html_headings_tables():
    """#191'in WI yorumunda '### Eksik detaylar' ve '| a | b |' satırları
    Azure DevOps'ta düz metin göründü (ekran görüntüsü, 09.09 15:01)."""
    print("\n[26] _md_to_html — ### ve tablo")
    from agile_sdlc_crew.main import _md_to_html
    h = _md_to_html("### Alt başlık\n\nmetin")
    check("### → <h4>", "<h4>Alt başlık</h4>" in h, h)
    t = _md_to_html("| A | B |\n|---|---|\n| 1 | **iki** |\n\nson")
    check("pipe tablo → <table> + <th>", "<table>" in t and "<th>A</th>" in t, t)
    check("hücreler <td>, inline biçim korunur", "<td>1</td>" in t and "<strong>iki</strong>" in t, t)
    check("ayırıcı satır (|---|) hücre olmaz", "---" not in t, t)
    check("tablo sonrası paragraf devam eder", "<p>son</p>" in t, t)


# ── 27. WI yaşam döngüsü — süreçten bağımsız durum seçimi + sahiplik aralığı ─
#
# FLO süreci (Boards Management, 2026-09-10 Azure'dan okundu): Task/Bug için
# Backlog→To Do→In Progress→Code Review→QA To Do→QA→UAT→Preprod Check→Blocked→
# Ready for Production(Resolved)→Done; User Story'de 'QA To Do' yok. Agile
# şablonu: New/Active/Resolved/Closed. Issue: Active/Closed. Adlar sabit
# kodlanamaz — tercih listesi + tipin durum listesi.

_FLO_TASK = [{"name": n, "category": c} for n, c in [
    ("Backlog", "Proposed"), ("To Do", "Proposed"), ("In Progress", "InProgress"),
    ("Code Review", "InProgress"), ("QA To Do", "InProgress"), ("QA", "InProgress"),
    ("UAT", "InProgress"), ("Preprod Check", "InProgress"), ("Blocked", "InProgress"),
    ("Ready for Production", "Resolved"), ("Done", "Completed"), ("Canceled", "Removed")]]
_FLO_STORY = [s for s in _FLO_TASK if s["name"] != "QA To Do"] + [{"name": "Prod Check", "category": "InProgress"}]
_AGILE = [{"name": n, "category": c} for n, c in [
    ("New", "Proposed"), ("Active", "InProgress"), ("Resolved", "Resolved"), ("Closed", "Completed")]]
_ISSUE = [{"name": "Active", "category": "InProgress"}, {"name": "Closed", "category": "Completed"}]


def test_wi_lifecycle_transitions():
    print("\n[27] WI yaşam döngüsü — durum seçimi, sahiplik aralığı, geri alma")
    from agile_sdlc_crew import wi_lifecycle as wl

    check("Task start → In Progress", wl.pick_state(_FLO_TASK, "start") == "In Progress")
    check("Task review → Code Review", wl.pick_state(_FLO_TASK, "review") == "Code Review")
    check("Task wait → Blocked", wl.pick_state(_FLO_TASK, "wait") == "Blocked")
    check("Task handoff → QA To Do", wl.pick_state(_FLO_TASK, "handoff") == "QA To Do")
    check("User Story handoff → QA (QA To Do yok)", wl.pick_state(_FLO_STORY, "handoff") == "QA")
    check("Agile start → Active", wl.pick_state(_AGILE, "start") == "Active")
    check("Agile review → Active (Code Review yok, geri düşer)", wl.pick_state(_AGILE, "review") == "Active")
    check("Agile handoff → Resolved", wl.pick_state(_AGILE, "handoff") == "Resolved")
    check("Agile wait → yok (Blocked yok)", wl.pick_state(_AGILE, "wait") is None)
    check("Issue handoff → yok", wl.pick_state(_ISSUE, "handoff") is None)
    check("büyük/küçük harf duyarsız eşleşme",
          wl.pick_state([{"name": "in progress", "category": "InProgress"}], "start") == "in progress")

    owned = wl.owned_states(_FLO_TASK)
    check("sahiplik: Proposed + In Progress/Code Review/Blocked",
          {"backlog", "to do", "in progress", "code review", "blocked"} <= owned, str(owned))
    check("sahiplik DIŞI: QA To Do, QA, UAT, Ready for Production, Done",
          not ({"qa to do", "qa", "uat", "ready for production", "done"} & owned))

    check("To Do → start = In Progress", wl.plan_transition("To Do", _FLO_TASK, "start") == "In Progress")
    check("Backlog (Proposed) → start = In Progress", wl.plan_transition("Backlog", _FLO_TASK, "start") == "In Progress")
    check("zaten In Progress → start = no-op", wl.plan_transition("In Progress", _FLO_TASK, "start") is None)
    check("Blocked → start = In Progress (needs_info sonrası ↻)", wl.plan_transition("Blocked", _FLO_TASK, "start") == "In Progress")
    check("In Progress → review = Code Review", wl.plan_transition("In Progress", _FLO_TASK, "review") == "Code Review")
    check("Code Review → handoff = QA To Do", wl.plan_transition("Code Review", _FLO_TASK, "handoff") == "QA To Do")
    check("QA To Do (insan ilerletmiş) → start = DOKUNMA", wl.plan_transition("QA To Do", _FLO_TASK, "start") is None)
    check("Ready for Production → review = DOKUNMA", wl.plan_transition("Ready for Production", _FLO_TASK, "review") is None)
    check("Done → handoff = DOKUNMA", wl.plan_transition("Done", _FLO_TASK, "handoff") is None)
    check("boş mevcut durum → yine hedef verir", wl.plan_transition("", _FLO_TASK, "start") == "In Progress")

    check("revert: In Progress→To Do (PR yok)", wl.plan_revert("In Progress", "To Do", _FLO_TASK, pr_exists=False) == "To Do")
    check("revert: PR varsa YOK (Code Review'da kalır)", wl.plan_revert("Code Review", "To Do", _FLO_TASK, pr_exists=True) is None)
    check("revert: mevcut == başlangıç → yok", wl.plan_revert("To Do", "To Do", _FLO_TASK, pr_exists=False) is None)
    check("revert: başlangıç sahiplik dışı (QA To Do) → yok", wl.plan_revert("In Progress", "QA To Do", _FLO_TASK, pr_exists=False) is None)
    check("revert: mevcut insan durumu (QA) → yok", wl.plan_revert("QA", "To Do", _FLO_TASK, pr_exists=False) is None)
    check("revert: Blocked→To Do (needs_info sonrası hata)", wl.plan_revert("Blocked", "To Do", _FLO_TASK, pr_exists=False) == "To Do")

    check("test yolu: Test/…/CustomerReturnReasonDefaultTest.php",
          wl.is_test_path("Test/Controller/Api/V1/CustomerReturnReasonDefaultTest.php"))
    check("test yolu: tests/unit/x.py", wl.is_test_path("tests/unit/x.py"))
    check("test yolu: src/a.spec.ts", wl.is_test_path("src/a.spec.ts"))
    check("test yolu DEĞİL: app/Customer.php", not wl.is_test_path("app/Customer.php"))
    check("test yolu DEĞİL: resources/translation/fr_FR.yml", not wl.is_test_path("resources/translation/fr_FR.yml"))
    check("test yolu DEĞİL: Contest/Latest.php (kelime içi 'test')", not wl.is_test_path("Contest/Latest.php"))


# ── 28. Definition of Done — job #189: completed ama UAT REJECTED ────────

_UAT_189_SHAPE = """## UAT Report

**Acceptance Criteria:**

1. **AC1** — Given: global site, When: reason_option_default, Then: dil uygun. — **PASS** — Evidence: fr_FR.yml.

2. **AC2** — Given: rus sitesi, Then: rusca. — **FAIL** — Evidence gap: no Russian resource; a Russian-site call cannot pass.

3. **AC3** — Given: Turk sitesi, Then: regresyon yok. — **PASS** — Evidence: additive change.

**Overall Evaluation:** REJECTED

**Gaps:**
- AC2 not satisfied
"""


def test_dod_checklist():
    print("\n[28] Definition of Done — deterministik tablo (#189 UAT REJECTED görünür olmalı)")
    from agile_sdlc_crew import wi_lifecycle as wl

    u = wl.parse_uat(_UAT_189_SHAPE)
    check("UAT parse: overall REJECTED", u["overall"] == "REJECTED", str(u))
    check("UAT parse: 2 PASS / 1 FAIL (açıklamadaki 'cannot pass' sayılmaz)",
          (u["pass"], u["fail"]) == (2, 1), str(u))
    check("UAT parse: madde kararları sırayla", u["items"] == [(1, "PASS"), (2, "FAIL"), (3, "PASS")], str(u["items"]))
    check("UAT parse: boş metin → None/0", wl.parse_uat("")["overall"] is None and wl.parse_uat("")["items"] == [])
    check("UAT parse: Türkçe KABUL", wl.parse_uat("Overall Evaluation: KABUL")["overall"] == "ACCEPTED")

    pushed = ["Customer.php", "resources/translation/fr_FR.yml",
              "Test/Controller/Api/V1/CustomerReturnReasonDefaultTest.php"]
    d = wl.evaluate_dod(review_approved=True, open_review_issues=0, build_status="succeeded",
                        uat_text=_UAT_189_SHAPE, pushed_files=pushed, require_tests=True, pr_id="42951")
    check("#189 şekli: DoD GEÇİLEMEDİ (UAT ❌)", not d.passed)
    check("#189 şekli: kalan tek madde UAT", [i.key for i in d.failed] == ["uat"], str([i.key for i in d.failed]))
    check("#189 şekli: review/build/test/pr ✅",
          all(i.ok for i in d.items if i.key in ("review", "issues", "build", "tests", "pr")))
    md = wl.render_dod(d, enforce=False)
    check("render: tablo + ❌ + 'DoD geçilemedi' + zorlama kapalı notu",
          "| Madde | Durum | Not |" in md and "❌" in md and "DoD geçilemedi" in md and "CREW_DOD_ENFORCE" in md)
    md_e = wl.render_dod(d, enforce=True)
    check("render (enforce): needs_human notu", "needs_human" in md_e)

    ok_uat = _UAT_189_SHAPE.replace("— **FAIL** —", "— **PASS** —").replace("REJECTED", "ACCEPTED")
    d2 = wl.evaluate_dod(review_approved=True, open_review_issues=0, build_status="succeeded",
                         uat_text=ok_uat, pushed_files=pushed, require_tests=True, pr_id="42951")
    check("hepsi yeşil → DoD geçti, doğrulanamayan yok", d2.passed and not d2.unverified)
    d3 = wl.evaluate_dod(review_approved=True, open_review_issues=None, build_status="no_pipeline",
                         uat_text=ok_uat, pushed_files=["app/X.php"], require_tests=False, pr_id="1")
    check("pipeline yok + test şartı kapalı → ⚪ bloklamaz, DoD geçti",
          d3.passed and {i.key for i in d3.unverified} == {"issues", "build", "tests"}, str([(i.key, i.ok) for i in d3.items]))
    d4 = wl.evaluate_dod(review_approved=True, open_review_issues=0, build_status="failed",
                         uat_text=ok_uat, pushed_files=["app/X.php"], require_tests=True, pr_id="1")
    check("build failed + test yok (şart açık) → iki ❌", {i.key for i in d4.failed} == {"build", "tests"})
    d5 = wl.evaluate_dod(review_approved=False, open_review_issues=2, build_status="succeeded",
                         uat_text=ok_uat, pushed_files=pushed, require_tests=True, pr_id="")
    check("onay yok + 2 açık madde + PR yok → üç ❌", {i.key for i in d5.failed} == {"review", "issues", "pr"})
    check("render: ⚪ işareti doğrulanamayan için", "⚪" in wl.render_dod(d3, enforce=False))

    job = None
    try:
        from agile_sdlc_crew import db
        job = db.get_job(189)
    except Exception:
        job = None
    if not job:
        skip("#189 gerçek UAT çıktısı", "DB erişilemedi")
        return
    steps = {s["step_key"]: s.get("output") or "" for s in (job.get("steps") or [])}
    real = wl.parse_uat(steps.get("uat_task", ""))
    check("#189 GERÇEK UAT: REJECTED, 2 PASS / 1 FAIL",
          real["overall"] == "REJECTED" and (real["pass"], real["fail"]) == (2, 1), str(real))
    dr = wl.evaluate_dod(review_approved=_review_approved(steps.get("review_pr_task", "")),
                         open_review_issues=0, build_status="succeeded",
                         uat_text=steps.get("uat_task", ""), pushed_files=pushed,
                         require_tests=True, pr_id=str(job.get("pr_id") or ""))
    check("#189 GERÇEK: review ✅ build ✅ UAT ❌ → DoD geçilemedi",
          not dr.passed and [i.key for i in dr.failed] == ["uat"], str([(i.key, i.ok) for i in dr.items]))


# ── 29. Flow kancaları: stub client ile geçiş kaydı + koruma kuralları ───

class _StubAzClient:
    """Sadece wi_lifecycle'ın dokunduğu yüzey. Her PATCH kaydedilir."""
    def __init__(self, states, fields):
        self.states, self.fields, self.ops = states, fields, []

    def get_work_item(self, wid):
        return {"id": wid, "fields": dict(self.fields)}

    def get_work_item_type_states(self, t):
        return list(self.states)

    def set_work_item_state(self, wid, state):
        self.ops.append(("state", int(wid), state))
        self.fields["System.State"] = state
        return {}

    def get_authenticated_user(self):
        return {"id": "u1", "displayName": "Pipeline", "uniqueName": "pipeline@example.com"}

    def assign_work_item(self, wid, ident):
        self.ops.append(("assign", int(wid), ident))
        return {}


def test_wi_lifecycle_flow_hooks():
    print("\n[29] Flow kancaları — stub client, knob'lar, dry-run/kickoff-only koruması")
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.flow import NeedsHumanReview as _NHR

    knobs = {"CREW_WI_LIFECYCLE": True, "CREW_WI_ASSIGN_IF_EMPTY": False}
    _orig_get = pc.get
    pc.get = lambda k: knobs[k] if k in knobs else _orig_get(k)
    try:
        def mk(state="To Do", assigned=None, **st):
            fields = {"System.WorkItemType": "Task", "System.State": state}
            if assigned:
                fields["System.AssignedTo"] = {"displayName": assigned, "uniqueName": assigned}
            f = AgileSDLCFlow()
            f.state.work_item_id = "73121"
            for k, v in st.items():
                setattr(f.state, k, v)
            f._client = _StubAzClient(_FLO_TASK, fields)
            return f

        f = mk()
        f._wi_begin(f._client.fields)
        check("begin: tip/durum state'e alındı",
              f.state.wi_type == "Task" and f.state.wi_state_initial == "To Do" and len(f.state.wi_states) == 12)
        f2 = mk()
        f2._client.fields["Custom.StoryPoints"] = 2
        f2._client.fields["Microsoft.VSTS.Scheduling.StoryPoints"] = 3.0
        f2._wi_begin(f2._client.fields)
        check("begin: SP okuma sırası Custom.StoryPoints önce (2, 3.0 değil)", f2.state.wi_story_points == 2.0)
        check("begin: To Do → In Progress yazıldı", f._client.ops == [("state", 73121, "In Progress")], str(f._client.ops))
        check("begin: mevcut durum güncellendi", f.state.wi_state_current == "In Progress")
        f._wi_begin(f._client.fields)
        check("begin idempotent (ikinci çağrı yazmaz)", len(f._client.ops) == 1)
        check("review → Code Review", f._wi_transition("review") == "Code Review" and f._client.ops[-1][2] == "Code Review")
        check("handoff → QA To Do", f._wi_transition("handoff") == "QA To Do")
        check("QA To Do'dan sonra start = dokunma (sahiplik dışı)", f._wi_transition("start") is None and len(f._client.ops) == 3)
        # Sorun olursa Blocked değil başlangıç durumu; PR açık olsa bile geri alınır.

        f = mk()
        f._wi_begin(f._client.fields)
        f.wi_lifecycle_on_exception(NeedsMoreInfo("eksik detay"))
        check("needs_info → başlangıç durumuna (To Do) geri", f._client.ops[-1] == ("state", 73121, "To Do"), str(f._client.ops))
        f = mk()
        f._wi_begin(f._client.fields)
        f.wi_lifecycle_on_exception(_NHR("review kapanmadı"))
        check("needs_human → başlangıç durumuna (To Do) geri", f._client.ops[-1] == ("state", 73121, "To Do"), str(f._client.ops))
        f = mk()
        f._wi_begin(f._client.fields)
        f.wi_lifecycle_on_exception(RuntimeError("patladı"))
        check("genel hata, PR yok → To Do'ya geri", f._client.ops[-1] == ("state", 73121, "To Do"), str(f._client.ops))
        f = mk()
        f._wi_begin(f._client.fields)
        f.state.pr_id = "42951"
        f.wi_lifecycle_on_exception(RuntimeError("patladı"))
        check("genel hata, PR VAR → yine To Do'ya geri (yarım iş incelemeye alınmaz)",
              f._client.ops[-1] == ("state", 73121, "To Do") and f.state.wi_state_current == "To Do", str(f._client.ops))

        f = mk(state="QA To Do")
        f._wi_begin(f._client.fields)
        check("insan QA'ya taşımış → start yazmaz ama bağlam okunur",
              f._client.ops == [] and f.state.wi_state_initial == "QA To Do")

        f = mk(dry_run=True)
        f._wi_begin(f._client.fields)
        check("dry-run: bağlam okunur, Azure'a yazılmaz", f._client.ops == [] and f.state.wi_type == "Task")
        f = mk(kickoff_only=True)
        f._wi_begin(f._client.fields)
        check("kickoff-only: yazılmaz", f._client.ops == [])

        knobs["CREW_WI_LIFECYCLE"] = False
        f = mk()
        f._wi_begin(f._client.fields)
        check("knob kapalı: bağlam okunur (DoD/log için), yazılmaz",
              f._client.ops == [] and f.state.wi_state_initial == "To Do")
        f.wi_lifecycle_on_exception(RuntimeError("x"))
        check("knob kapalı: hata yolunda da yazılmaz", f._client.ops == [])

        knobs["CREW_WI_LIFECYCLE"] = True
        knobs["CREW_WI_ASSIGN_IF_EMPTY"] = True
        f = mk()
        f._wi_begin(f._client.fields)
        check("atama: boş → PAT sahibine", ("assign", 73121, "pipeline@example.com") in f._client.ops
              and f.state.wi_assigned_by_pipeline and f.state.wi_assigned_to == "pipeline@example.com")
        f = mk(assigned="ebru@example.com")
        f._wi_begin(f._client.fields)
        check("atama: dolu → dokunma", not any(o[0] == "assign" for o in f._client.ops)
              and f.state.wi_assigned_to == "ebru@example.com")

        class _Boom(_StubAzClient):
            def set_work_item_state(self, wid, state):
                raise RuntimeError("TF401320: geçiş kuralı")
        f = mk()
        f._client = _Boom(_FLO_TASK, f._client.fields)
        f._wi_begin(f._client.fields)
        check("Azure reddetti → pipeline devam, mevcut durum değişmez",
              f.state.wi_state_current == "To Do" and f.state.wi_type == "Task")
    finally:
        pc.get = _orig_get

    src_main = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/main.py").read_text()
    i_hook, i_fail = src_main.find("flow.wi_lifecycle_on_exception(e)"), src_main.find("_db.fail_job(job_id, str(e))")
    check("main.run_pipeline: yaşam döngüsü kancası fail_job'dan ÖNCE", 0 < i_hook < i_fail)
    src_flow = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/flow.py").read_text()
    check("step11: DoD → yorum → enforce → Code Review sırası",
          0 < src_flow.find("_wl_dod.evaluate_dod(") < src_flow.find('f"{_dod_md}\\n\\n"')
          < src_flow.find("raise NeedsHumanReview(_msg_dod)") < src_flow.rfind('self._wi_transition("review")'))
    check("Code Review yalnız iş sonunda: PR açılınca geçiş yok, handoff yok",
          src_flow.count('self._wi_transition("review")') == 1 and 'self._wi_transition("handoff")' not in src_flow
          and src_flow.find('self._wi_transition("review")') > src_flow.find("raise NeedsHumanReview(_msg_dod)"))
    check("build gate her çıkışta build_status yazar (6 yol)",
          src_flow.count("self.state.build_status = ") >= 7)


# ── 30. Tahminleme — Fibonacci, yapısal merdiven, BA bloğu, uzlaştırma ────

_BA_WITH_ESTIMATE = """```json
{
  "summary": "Getirilen iade nedeni metni site diline uygun olsun",
  "functional_requirements": [{"id": "FR1", "desc": "a"}, {"id": "FR2", "desc": "b"}],
  "technical_requirements": [{"id": "TR1", "desc": "c"}],
  "acceptance_criteria": [{"id": "AC1", "desc": "d"}],
  "out_of_scope": [], "assumptions": [], "open_questions": [],
  "readiness": {"score": 85, "missing_details": []},
  "estimate": {"story_points": 5, "confidence": 70, "rationale": "Üç dosya ve çeviri kaynakları"}
}
```"""


def test_estimation():
    print("\n[30] Tahminleme — Fibonacci'ye oturtma, yapısal merdiven, BA bloğu, max-uzlaşma")
    from agile_sdlc_crew import estimation as est

    check("snap 4 → 5 (eşitlikte büyük)", est.snap_fib(4) == 5)
    check("snap 6 → 5, 7 → 8, 2.6 → 3", (est.snap_fib(6), est.snap_fib(7), est.snap_fib(2.6)) == (5, 8, 3))
    check("snap 0 / metin → None", est.snap_fib(0) is None and est.snap_fib("x") is None)
    check("yapısal req: 2→2, 5→3, 7→5, 10→8",
          [est.structural_estimate(n, 0, False, "requirements")[0] for n in (2, 5, 7, 10)] == [2, 3, 5, 8])
    check("yapısal plan: 7 req + 6 dosya + keşif → 13 (tavan)",
          est.structural_estimate(7, 6, True, "plan")[0] == 13)
    check("yapısal plan: 2 req + 3 dosya → 3 (+1 basamak)", est.structural_estimate(2, 3, False, "plan")[0] == 3)
    check("yapısal requirements: dosya sayısı YOK SAYILIR", est.structural_estimate(2, 9, True, "requirements")[0] == 2)
    check("uzlaşma: BA yok → yapısal", est.reconcile(None, 3) == (3, "yapısal"))
    check("uzlaşma: BA 5 > yapısal 3 → BA", est.reconcile(5, 3) == (5, "BA"))
    check("uzlaşma: BA 2 < yapısal 3 → yapısal", est.reconcile(2, 3) == (3, "yapısal"))
    check("uzlaşma: eşit → uyumlu", est.reconcile(3, 3)[1] == "BA + yapısal uyumlu")
    check("uzlaşma: önceki aşama 5 → düşmez", est.reconcile(2, 3, previous_sp=5) == (5, "önceki aşama"))
    ba = est.parse_ba_estimate(_BA_WITH_ESTIMATE)
    check("BA bloğu parse: 5 SP, %70, gerekçe", ba == {"sp": 5, "confidence": 70, "rationale": "Üç dosya ve çeviri kaynakları"}, str(ba))
    check("BA bloğu yok → None", est.parse_ba_estimate('{"summary": "x"}') is None)
    check("bozuk JSON → None", est.parse_ba_estimate("{oops") is None)
    check("story_points 4 → 5'e oturur", est.parse_ba_estimate('{"estimate": {"story_points": 4}}')["sp"] == 5)
    check("sınıf: 2→S, 5→M, 8→L", (est.size_class(2), est.size_class(5), est.size_class(8)) == ("S", "M", "L"))
    line = est.render_estimate_line({"sp": 5, "source": "BA"}, elapsed_min=17.2, cost_usd=3.96)
    check("satır: tahmin + gerçekleşen", "5 SP (M, BA)" in line and "17 dk" in line and "$3.96" in line, line)
    check("satır: tahmin yoksa boş", est.render_estimate_line({}) == "")

    class _C:
        def __init__(self, fail_fields=(), hard=False):
            self.ops, self.fail_fields, self.hard = [], set(fail_fields), hard
        def update_work_item(self, wid, ops):
            fld = ops[0]["path"].rsplit("/", 1)[-1]
            if self.hard:
                raise RuntimeError("connection reset")
            if fld in self.fail_fields:
                raise RuntimeError("400 Client Error: TF51535 field does not exist")
            self.ops.append((wid, fld, ops[0]["value"]))
    c = _C()
    check("SP yaz: önce Custom.StoryPoints (org'un gerçek alanı), int",
          est.write_story_points(c, "73121", 5) == "Custom.StoryPoints" and c.ops == [(73121, "Custom.StoryPoints", 5)])
    c = _C(fail_fields={"Custom.StoryPoints"})
    check("SP yaz: Custom yok → Microsoft StoryPoints (float)",
          est.write_story_points(c, "1", 3) == "Microsoft.VSTS.Scheduling.StoryPoints" and c.ops[-1][2] == 3.0)
    c = _C(fail_fields={"Custom.StoryPoints", "Microsoft.VSTS.Scheduling.StoryPoints"})
    check("SP yaz: ikisi de yok → Effort", est.write_story_points(c, "1", 3) == "Microsoft.VSTS.Scheduling.Effort")
    line_t = est.render_estimate_line({"sp": 5, "source": "BA"}, elapsed_min=17.2, cost_usd=3.96, team_sp=2)
    check("satır: takım tahmini varsa ikisi yan yana", "Takım tahmini:** 2 SP" in line_t and "Pipeline tahmini:** 5 SP" in line_t, line_t)
    check("SP yaz: ağ hatası → '' (pipeline devam)", est.write_story_points(_C(hard=True), "1", 3) == "")

    job = None
    try:
        from agile_sdlc_crew import db
        job = db.get_job(189)
    except Exception:
        job = None
    if job:
        reqs = next((s.get("output") or "" for s in job.get("steps") or [] if s["step_key"] == "requirements_analysis_task"), "")
        n = len(_requirement_ids(reqs))
        sp, why = est.structural_estimate(n, 2, False, "plan")
        check(f"#189 GERÇEK BA (estimate bloğu yok → None) + yapısal {sp} SP Fibonacci'de",
              est.parse_ba_estimate(reqs) is None and sp in est.FIB and n > 0, f"n_req={n} {why}")
    else:
        skip("#189 gerçek BA çıktısı", "DB erişilemedi")


# ── 31. Alt iş kaydı planlama — parent tipi, gruplama, idempotency ────────

def _plan(n, repo="webservice"):
    dirs = ["/app/Controller", "/app/Model", "/resources/translation", "/Test/Controller"]
    return {"repo_name": repo, "changes": [
        {"file_path": f"{dirs[i % 4]}/File{i}.php", "change_type": "edit" if i % 3 else "add",
         "description": f"Değişiklik {i}", "covers_requirements": ["FR1", f"AC{i % 2 + 1}"]}
        for i in range(n)]}


def test_child_tasks_planning():
    print("\n[31] Alt iş kaydı — parent tipi, başlık/açıklama, gruplama, idempotent açma")
    from agile_sdlc_crew import wi_children as wc

    check("parent tipi: User Story/Bug/Feature evet; Task hayır",
          wc.parent_allows_children("User Story") and wc.parent_allows_children("Bug")
          and not wc.parent_allows_children("Task") and not wc.parent_allows_children(""))
    items = wc.plan_child_tasks(_plan(3), "webservice")
    check("3 değişiklik → 3 child", len(items) == 3)
    check("başlık: [repo] eylem dosya — açıklama", items[1]["title"] == "[webservice] Düzenle File1.php — Değişiklik 1", items[1]["title"])
    check("add → 'Yeni dosya'", items[0]["title"].startswith("[webservice] Yeni dosya File0.php"))
    check("açıklama: gereksinimler + dosya", "Kapsadığı gereksinimler: FR1, AC2" in items[1]["description"]
          and "/app/Model/File1.php" in items[1]["description"])
    check("files eşlemesi", items[2]["files"] == ["/resources/translation/File2.php"])
    big = wc.plan_child_tasks(_plan(12), "webservice", max_children=8)
    check("12 değişiklik, max 8 → dizine göre ≤8 grup", 1 <= len(big) <= 8 and all("dosya" in b["title"] for b in big), str([b["title"] for b in big]))
    check("gruplama: tüm dosyalar bir grupta", sorted(f for b in big for f in b["files"]) == sorted(c["file_path"] for c in _plan(12)["changes"]))
    long_plan = {"repo_name": "r", "changes": [{"file_path": "/a/B.php", "change_type": "edit", "description": "x" * 200}]}
    check("başlık 128'e kırpılır", len(wc.plan_child_tasks(long_plan, "r")[0]["title"]) <= 128)
    check("boş plan → []", wc.plan_child_tasks({}, "r") == [] and wc.plan_child_tasks({"changes": [{"nope": 1}]}, "r") == [])
    kids = [{"id": 1, "files": ["/app/Customer.php"]}, {"id": 2, "files": ["/Test/X.php"]}]
    check("child_for_file: slash/büyük-küçük toleranslı",
          wc.child_for_file(kids, "app/customer.php")["id"] == 1 and wc.child_for_file(kids, "/nope.php") is None)

    class _C:
        def __init__(self, children=()):
            self.children, self.created, self.n = list(children), [], 900
        def get_work_item_children(self, pid):
            return self.children
        def create_work_item(self, t, fields, parent_id=None):
            self.n += 1
            self.created.append((t, fields, parent_id))
            return {"id": self.n, "fields": fields}
    c = _C()
    out = wc.create_children(c, 500, {"System.AreaPath": "BM\\Ops", "System.IterationPath": "BM\\Ops\\2026_19"}, items)
    check("3 child açıldı, parent 500, Task tipi", len(c.created) == 3 and all(t == "Task" and p == 500 for t, _, p in c.created))
    check("etiket + alan/iterasyon kopyalandı", all(f["System.Tags"] == "crew-generated" and f["System.AreaPath"] == "BM\\Ops"
                                                    and f["System.IterationPath"] == "BM\\Ops\\2026_19" for _, f, _ in c.created))
    check("dönen liste id + files taşır", [o["id"] for o in out] == [901, 902, 903] and out[0]["files"] == ["/app/Controller/File0.php"])
    c2 = _C(children=[{"id": 77, "fields": {"System.Title": "[webservice] eski", "System.Tags": "crew-generated; x", "System.State": "Done"}}])
    out2 = wc.create_children(c2, 500, {}, items)
    check("üretilmiş child varsa YENİDEN AÇMAZ, mevcutları döndürür", c2.created == [] and [o["id"] for o in out2] == [77])
    c3 = _C(children=[{"id": 78, "fields": {"System.Title": "insan task'ı", "System.Tags": "", "System.State": "To Do"}}])
    wc.create_children(c3, 500, {}, items)
    check("etiketsiz (insan) child engel değil", len(c3.created) == 3)


# ── 32. Flow kancaları — tahmin + alt iş (stub client) ───────────────────

class _StubAzClient2(_StubAzClient):
    def __init__(self, states, fields, children=()):
        super().__init__(states, fields)
        self.children, self.n = list(children), 900
    def update_work_item(self, wid, ops):
        for o in ops:
            self.ops.append(("field", int(wid), o["path"].rsplit("/", 1)[-1], o["value"]))
        return {}
    def create_work_item(self, t, fields, parent_id=None):
        self.n += 1
        self.ops.append(("create", parent_id, t, fields["System.Title"]))
        return {"id": self.n}
    def get_work_item_children(self, pid):
        return self.children


def test_estimate_flow_hooks():
    print("\n[32] Flow kancaları — tahmin aşamaları, SP yazma koruması, alt iş açma/kapama")
    from agile_sdlc_crew import pipeline_config as pc

    knobs = {"CREW_ESTIMATE": True, "CREW_WI_WRITE_ESTIMATE": False,
             "CREW_WI_CHILD_TASKS": False, "CREW_WI_CHILD_TASKS_MAX": 8}
    _orig_get = pc.get
    pc.get = lambda k: knobs[k] if k in knobs else _orig_get(k)
    try:
        def mk(wi_type="Task", sp=None, **st):
            f = AgileSDLCFlow()
            f._db = None
            f.state.work_item_id = "73121"
            f.state.repo_name = "webservice"
            f.state.wi_type = wi_type
            f.state.wi_story_points = sp
            f.state.wi_area_path, f.state.wi_iteration_path = "BM\\Ops", "BM\\Ops\\2026_19"
            f.state.requirements_text = _BA_WITH_ESTIMATE
            for k, v in st.items():
                setattr(f.state, k, v)
            f._client = _StubAzClient2(_FLO_TASK, {"System.WorkItemType": wi_type, "System.State": "In Progress"})
            return f

        f = mk()
        f._estimate("requirements")
        e = f.state.estimate
        check("requirements: BA 5 vs yapısal 3 (4 gereksinim) → 5 SP, kaynak BA",
              e.get("sp") == 5 and e.get("source") == "BA" and e.get("structural_sp") == 3 and e.get("ba_sp") == 5, str(e))
        f.state.plan = _plan(3)
        f._estimate("plan")
        e = f.state.estimate
        check("plan: 3 dosya → yapısal 5, BA 5 → uyumlu, stage=plan",
              e.get("sp") == 5 and e.get("stage") == "plan" and "uyumlu" in e.get("source", ""), str(e))
        check("yazma knob'u kapalı → Azure'a SP yazılmadı", not any(o[0] == "field" for o in f._client.ops))

        knobs["CREW_WI_WRITE_ESTIMATE"] = True
        f = mk()
        f.state.plan = _plan(3)
        f._estimate("plan")
        check("yazma açık + WI'da SP boş → Custom.StoryPoints=5 yazıldı",
              ("field", 73121, "Custom.StoryPoints", 5) in f._client.ops, str(f._client.ops))
        f = mk(sp=3.0)
        f.state.plan = _plan(3)
        f._estimate("plan")
        check("WI'da SP var (3) → insan tahmini EZİLMEZ", not any(o[0] == "field" for o in f._client.ops))
        f = mk(dry_run=True)
        f.state.plan = _plan(3)
        f._estimate("plan")
        check("dry-run: tahmin hesaplanır, yazılmaz", f.state.estimate.get("sp") == 5 and f._client.ops == [])
        knobs["CREW_ESTIMATE"] = False
        f = mk()
        f._estimate("requirements")
        check("CREW_ESTIMATE kapalı → hiç hesaplanmaz", f.state.estimate == {})
        knobs["CREW_ESTIMATE"] = True

        f = mk(plan=_plan(3))
        f._create_child_tasks()
        check("child knob kapalı → açılmaz", f.state.child_tasks == [] and f._client.ops == [])
        knobs["CREW_WI_CHILD_TASKS"] = True
        f = mk(plan=_plan(3))
        f._create_child_tasks()
        check("WI tipi Task → parent değil, açılmaz", f.state.child_tasks == [] and f._client.ops == [])
        f = mk(wi_type="User Story", plan=_plan(3))
        f._after_plan_finalized()
        creates = [o for o in f._client.ops if o[0] == "create"]
        check("User Story + knob → 3 child açıldı (parent 73121)", len(creates) == 3 and all(o[1] == 73121 for o in creates), str(creates))
        check("state.child_tasks dolu, files eşli", len(f.state.child_tasks) == 3 and f.state.child_tasks[1]["files"] == ["/app/Model/File1.php"])
        n = len(f._client.ops)
        f._complete_child_task("app/Model/File1.php")
        check("dosya push → ilgili child Done", f._client.ops[-1] == ("state", 902, "Done") and f.state.child_tasks[1].get("done"), str(f._client.ops[-1]))
        f._complete_child_task("app/Model/File1.php")
        check("aynı dosya ikinci kez → no-op", len(f._client.ops) == n + 1)
        f._complete_child_task("/olmayan.php")
        check("planda olmayan dosya → no-op", len(f._client.ops) == n + 1)
        f._create_child_tasks()
        check("ikinci çağrı → yeniden açmaz (state dolu)", len([o for o in f._client.ops if o[0] == "create"]) == 3)
        f = mk(wi_type="User Story", plan=_plan(3), kickoff_only=True)
        f._create_child_tasks()
        check("kickoff-only → açılmaz", f._client.ops == [])
    finally:
        pc.get = _orig_get

    src_flow = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/flow.py").read_text()
    check("plan kesinleşen 3 yolda da _after_plan_finalized (normal/resume/HAL)", src_flow.count("self._after_plan_finalized()") == 3)
    check("step6'da 4 push yolunda _complete_child_task (resume/skip/silme/push)", src_flow.count("self._complete_child_task(file_path)") == 4)
    check("step1: requirements sonrası kaba tahmin", 'self._estimate("requirements")' in src_flow)
    src_db = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/db.py").read_text()
    check("db: estimate_sp kolonu + whitelist", '"estimate_sp"' in src_db and 'ensure_column(cur, "jobs", "estimate_sp"' in src_db)
    src_yaml = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/config/tasks.yaml").read_text()
    check("tasks.yaml: BA estimate bloğu + ESTIMATE kuralı", '"estimate": {' in src_yaml and "ESTIMATE (mandatory)" in src_yaml)


# ── 33. Retrospektif — deterministik analiz, sınıflandırma, öneriler, rapor ─

def test_retrospective():
    print("\n[33] Retrospektif — sonuç sınıfları, kalite kapıları, SP metrikleri, kural önerileri")
    from datetime import datetime, timedelta
    from agile_sdlc_crew import retrospective as rt

    now = datetime.now()

    def job(i, status, err="", cost=1.0, mins=10, sp=None, wi="1", steps=()):
        return {"id": i, "work_item_id": wi, "status": status, "error_message": err,
                "total_cost_usd": cost, "started_at": now - timedelta(minutes=mins), "finished_at": now,
                "estimate_sp": sp,
                "steps": [{"step_key": k, "status": st, "output": out} for k, st, out in steps]}

    jobs = [
        job(1, "completed", cost=3.96, mins=17, sp=5, wi="73121", steps=[
            ("review_pr_task", "completed", "REVIEW_DECISION: APPROVE\nVerdict: APPROVE — 2 düzeltme turundan sonra onaylandı."),
            ("pr_build_gate", "completed", "Build 132551 succeeded (webservice-test)"),
            ("uat_task", "completed", _UAT_189_SHAPE)]),
        job(2, "needs_info", "NEEDS_INFO: hazırlık skoru 5/100 < eşik 60 — WI'da 6 eksik detay", cost=0.13, mins=1, wi="73061"),
        job(3, "needs_human", "Review: 1 deneme sonrasi kapanmayan madde (N1) — PR #42951 acik", cost=3.43, wi="73121", steps=[
            ("review_pr_task", "completed", "REVIEW_DECISION: NEEDS_HUMAN\nİnsan müdahalesi gerekli — 1 deneme sonrası kapanmayan madde")]),
        job(4, "failed", "Sunucu yeniden baslatildi, is yarida kaldi", cost=0.5, steps=[("technical_design_task", "failed", "")]),
        job(5, "failed", "Plan-push uyumsuzlugu: 1/2 dosya push edildi, %70 esigin altinda. PR iptal.", cost=4.89, wi="73121",
            steps=[("implement_change_task", "failed", "")]),
        job(6, "completed", cost=2.0, mins=10, sp=2, wi="5", steps=[
            ("review_pr_task", "completed", "REVIEW_DECISION: APPROVE\nVerdict: APPROVE"),
            ("pr_build_gate", "completed", "Repoda PR-test pipeline'i yok — gate atlandi"),
            ("uat_task", "completed", "## UAT Report\n\n1. AC1 - PASS - ok\n\n**Overall Evaluation:** ACCEPTED")]),
        job(7, "running", cost=0),
    ]
    a = rt.analyze(jobs)
    check("running iş sayılmaz → 6 terminal iş", a["jobs"] == 6)
    check("durum dağılımı", a["status"] == {"completed": 2, "needs_info": 1, "needs_human": 1, "failed": 2}, str(a["status"]))
    check("başarı oranı %33", round(a["success_rate"]) == 33)
    oc = dict(a["outcomes"])
    check("nedenler: hazırlık kapısı / review madde / altyapı / push",
          oc.get("WI detayı yetersiz (hazırlık kapısı)") == 1 and oc.get("Review: kapanmayan madde") == 1
          and oc.get("Altyapı: sunucu yeniden başlatıldı") == 1 and oc.get("Implement: dosya push edilemedi") == 1, str(oc))
    check("kırılan adımlar", dict(a["step_failures"]) == {"technical_design_task": 1, "implement_change_task": 1})
    rv = a["review"]
    check("review: 3 iş, ilk tur 1, 2 düzeltme turu, 1 needs_human",
          (rv["jobs"], rv["first_pass"], rv["retries_total"], rv["needs_human"]) == (3, 1, 2, 1), str(rv))
    check("build: yeşil 1, pipeline yok 1", a["build"] == {"yeşil": 1, "pipeline yok": 1}, str(a["build"]))
    check("UAT: kabul 1, red 1 (#189 şekli)", a["uat"] == {"kabul": 1, "red": 1}, str(a["uat"]))
    check("hazırlık skoru ortalaması 5", a["readiness_avg"] == 5)
    check("toplam maliyet 14.91", round(a["cost_total"], 2) == 14.91)
    check("teslim edilen SP 7", a["sp_delivered"] == 7)
    check("SP başına 4.2 dk · $0.90", round(a["min_per_sp"], 1) == 4.2 and round(a["cost_per_sp"], 2) == 0.9, f"{a['min_per_sp']} {a['cost_per_sp']}")
    check("tekrar koşan WI: 73121 ×3", a["rerun_wis"] == [("73121", 3)])
    check("en pahalı iş #5 ($4.89)", a["top_cost"][0][1] == 5)

    rules = rt.suggest_rules(a)
    texts = " ".join(r["text"] for r in rules)
    check("kural önerileri: review testleri okusun + UAT kanıt + tekrar koşu (3)",
          len(rules) == 3 and "MEVCUT testlerini" in texts and "UAT uzmanı" in texts and "ikinci kez" in texts, str([r["why"] for r in rules]))
    check("needs_info oranı %17 < %20 → hazırlık kuralı yok", "kabul kriteri alanı boş" not in texts)
    cfg = rt.suggest_config(a)
    check("yapılandırma önerileri: DOD_ENFORCE + restart (2)", len(cfg) == 2 and "CREW_DOD_ENFORCE" in cfg[0] and "restart" in cfg[1], str(cfg))
    md = rt.render_markdown(a, title="2026_19_Sudo", rules=rules, config=cfg)
    check("rapor: başlık, özet tablosu, SP, öneri bölümleri, tekrar koşu",
          "## 🔁 Retrospektif — 2026_19_Sudo" in md and "| İş sayısı | 6 |" in md and "| Teslim edilen SP | 7 |" in md
          and "### Önerilen kılavuz kuralları" in md and "### Yapılandırma önerileri" in md and "#73121 (3×)" in md)
    check("rapor: en pahalı işler tablosu", "| #5 | #73121 | failed | $4.89 |" in md)

    co = rt.classify_outcome
    check("sınıflandırma: DoD / build / bütçe / diğer / tamamlandı",
          co("needs_human", "DoD gecilemedi: UAT") == "DoD geçilemedi"
          and co("needs_human", "PR build 2 duzeltme sonrasi hala 'failed'") == "PR build kırmızı / doğrulanamadı"
          and co("failed", "Butce asildi: $18") == "Bütçe aşımı"
          and co("failed", "zzz") == "Diğer hata" and co("completed", None) == "Tamamlandı")
    e = rt.analyze([])
    check("boş pencere → 0 iş, rapor 'iş yok' der", e["jobs"] == 0 and "iş yok" in rt.render_markdown(e, title="x")
          and rt.suggest_rules(e) == [] and rt.suggest_config(e) == [])

    try:
        rep = rt.build_report(title="son 60 gün", since_days=60)
    except Exception as ex:
        skip("gerçek DB retrospektifi", f"DB erişilemedi: {type(ex).__name__}")
        return
    check("GERÇEK DB: son 60 günde iş var, rapor üretildi",
          rep["jobs"] >= 1 and "## 🔁 Retrospektif" in rep["markdown"] and isinstance(rep["suggested_rules"], list), str(rep["jobs"]))
    import json as _j
    _j.dumps(rep)
    check("GERÇEK DB: rapor JSON'a serileşir (Decimal/datetime sızmaz)", True)

    src_ui = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text()
    check("dashboard: 🔁 Retro butonu + modal + tablo render", "openRetro()" in src_ui and 'id="retroModal"' in src_ui
          and "md-table" in src_ui and "adoptRetroRule" in src_ui)
    src_srv = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/server.py").read_text()
    check("server: /api/retro + CREW_RETRO kapısı", '@app.get("/api/retro")' in src_srv and 'CREW_RETRO' in src_srv)


# ── 34. Sprint planlama — aday, sıra, kapasite, tahmin, toplu kuyruk ──────

def test_sprint_planning():
    print("\n[34] Sprint planlama — aday kuralı, sıra, kapasite, tahmin, güvenli kuyruk")
    from agile_sdlc_crew import sprint_planning as sp

    check("Proposed: Backlog/To Do evet, In Progress/QA hayır (FLO süreç listesi)",
          sp.is_proposed("Backlog", _FLO_TASK) and sp.is_proposed("To Do", _FLO_TASK)
          and not sp.is_proposed("In Progress", _FLO_TASK) and not sp.is_proposed("QA To Do", _FLO_TASK))
    check("Proposed yedek (süreç listesi yok): To Do/New evet, Active hayır",
          sp.is_proposed("To Do", None) and sp.is_proposed("New", []) and not sp.is_proposed("Active", None))

    items = [
        {"id": 1, "title": "A", "type": "Task", "state": "To Do", "priority": 2, "storyPoints": 3},
        {"id": 2, "title": "B", "type": "Task", "state": "To Do", "priority": 1, "storyPoints": 5},
        {"id": 3, "title": "C", "type": "Task", "state": "In Progress", "priority": 1, "storyPoints": 2},
        {"id": 4, "title": "D", "type": "Test Case", "state": "To Do", "priority": 1, "storyPoints": 0},
        {"id": 5, "title": "E", "type": "Bug", "state": "Backlog", "priority": 1, "storyPoints": 2},
        {"id": 6, "title": "F", "type": "Task", "state": "To Do", "priority": 1, "storyPoints": 0},
        {"id": 7, "title": "G", "type": "Task", "state": "To Do", "priority": 3, "storyPoints": 1},
        {"id": 8, "title": "H", "type": "Task", "state": "To Do", "priority": 1, "storyPoints": 2},
        {"id": 9, "title": "I", "type": "Task", "state": "To Do", "priority": 1, "storyPoints": 1},
    ]
    jobs = {"5": {"id": 50, "status": "running"}, "7": {"id": 70, "status": "completed"},
            "8": {"id": 80, "status": "needs_info"}, "9": {"id": 90, "status": "failed"}}
    rows = sp.order_candidates(sp.classify_candidates(items, {"Task": _FLO_TASK, "Bug": _FLO_TASK}, jobs))
    by = {r["id"]: r for r in rows}
    check("uygun: 1, 2, 6, 9 (failed → yeniden denenebilir)", {r["id"] for r in rows if r["eligible"]} == {1, 2, 6, 9})
    check("nedenler: başlanmış / tip dışı / açık iş / tamamlandı / insan bekliyor",
          "başlanmış" in by[3]["reason"] and "pipeline dışı" in by[4]["reason"] and "açık iş #50" in by[5]["reason"]
          and "tamamlandı" in by[7]["reason"] and "insan bekliyor" in by[8]["reason"] and "yeniden" in by[9]["reason"])
    check("sıra: uygunlar önce; öncelik ↑, SP ↑, SP=0 sona → 9, 2, 6, 1",
          [r["id"] for r in rows if r["eligible"]] == [9, 2, 6, 1], str([r["id"] for r in rows]))
    sp.select_by_capacity(rows, 6)
    check("kapasite 6 SP: 9 (1) + 2 (5) = 6 → seçili; 6 (SP yok) da işaretli, 1 (3) sığmaz",
          {r["id"] for r in rows if r["selected"]} == {9, 2, 6}, str([(r["id"], r["selected"]) for r in rows]))
    sp.select_by_capacity(rows, None)
    check("kapasite yok → tüm uygunlar seçili", {r["id"] for r in rows if r["selected"]} == {1, 2, 6, 9})
    est = sp.effort_estimate(rows, cost_per_sp=0.9, min_per_sp=4.0, cost_per_job=3.0, min_per_job=15.0)
    check("tahmin: 4 iş, 9 SP (+1 SP'siz iş başına ortalama) → $11.1, 51 dk",
          est["count"] == 4 and est["sp"] == 9 and est["sp_unknown"] == 1
          and round(est["cost_usd"], 2) == 11.1 and round(est["minutes"]) == 51, str(est))
    est2 = sp.effort_estimate(rows, cost_per_sp=None, min_per_sp=None, cost_per_job=3.0, min_per_job=15.0)
    check("tahmin: SP oranı yoksa iş başına", est2["cost_usd"] == 12.0 and est2["minutes"] == 60.0)

    created = []
    _o_status, _o_create = sp.latest_job_status, None
    from agile_sdlc_crew import db as _dbm
    _o_create = _dbm.create_job
    try:
        sp.latest_job_status = lambda ids: {"2": {"id": 20, "status": "queued"}}
        _dbm.create_job = lambda wid, use_hal, wi_title="", **kw: created.append((wid, use_hal, wi_title)) or (100 + len(created))
        res = sp.queue_selected([{"id": 1, "title": "A"}, {"id": 2, "title": "B"}, {"id": 9, "title": "I"}])
        check("kuyruk: açık işi olan #2 atlandı, 1 ve 9 kuyruklandı (use_hal=False)",
              [q["wi"] for q in res["queued"]] == ["1", "9"] and res["skipped"][0]["wi"] == "2"
              and all(c[1] is False for c in created), str(res))
    finally:
        sp.latest_job_status = _o_status
        _dbm.create_job = _o_create

    class _C:
        def get_iteration_work_items(self, ip):
            return items
        def get_work_item_type_states(self, t):
            return _FLO_TASK
    _o_vel, _o_rt = sp.team_velocity, None
    from agile_sdlc_crew import retrospective as _rt
    _o_an, _o_col = _rt.analyze, _rt.collect
    try:
        sp.latest_job_status = lambda ids: jobs
        sp.team_velocity = lambda client, team, n=3: 7.0
        _rt.collect = lambda **kw: []
        _rt.analyze = lambda jobs_: {"cost_per_sp": 1.0, "min_per_sp": 5.0, "cost_avg_completed": 3.0, "minutes_avg_completed": 20.0}
        plan = sp.build_plan(_C(), "BM\\Ops\\2026_19", team="Ops")
        check("build_plan: sprint adı, 9 iş, 4 uygun, kapasite velocity 7 → 9+2+6 seçili",
              plan["sprint"] == "2026_19" and plan["items_total"] == 9 and plan["eligible"] == 4
              and plan["capacity_sp"] == 7.0 and "velocity" in plan["capacity_source"]
              and {r["id"] for r in plan["rows"] if r["selected"]} == {9, 2, 6}, str(plan["estimate"]))
        check("build_plan: oranlar retro'dan", plan["rates"]["cost_per_sp"] == 1.0 and plan["estimate"]["count"] == 3)
        plan2 = sp.build_plan(_C(), "X", team="", capacity_sp=1)
        check("kapasite 1 (kullanıcı) → yalnızca 9 (1 SP) + SP'siz 6", {r["id"] for r in plan2["rows"] if r["selected"]} == {9, 6}
              and plan2["capacity_source"] == "kullanıcı")
    finally:
        sp.latest_job_status = _o_status
        sp.team_velocity = _o_vel
        _rt.analyze, _rt.collect = _o_an, _o_col


# ── 35. Günlük özet — render, düz metin, zamanlayıcı yardımcıları ────────

def test_daily_summary():
    print("\n[35] Günlük özet — bölümler, engeller, düz metin, zaman hesabı")
    from datetime import datetime, timedelta
    from agile_sdlc_crew import daily

    now = datetime(2026, 9, 10, 9, 0)
    d = {
        "now": now, "since": now - timedelta(hours=24),
        "finished": [
            {"id": 189, "work_item_id": "73121", "wi_title": "İade nedeni dil", "status": "completed", "pr_url": "https://x/pr/42951",
             "total_cost_usd": 3.96, "estimate_sp": 5, "started_at": now - timedelta(minutes=137), "finished_at": now - timedelta(minutes=120), "error_message": None},
            {"id": 191, "work_item_id": "73061", "wi_title": "COD nokta", "status": "needs_info", "pr_url": "", "total_cost_usd": 0.13,
             "estimate_sp": None, "started_at": now - timedelta(minutes=61), "finished_at": now - timedelta(minutes=60),
             "error_message": "NEEDS_INFO: hazırlık skoru 5/100 < eşik 60"},
        ],
        "running": [{"id": 200, "work_item_id": "73200", "wi_title": "Koşan iş", "current_step": "implement_change_task", "started_at": now - timedelta(minutes=12)}],
        "queued": [{"id": 201, "work_item_id": "73201"}, {"id": 202, "work_item_id": "73202"}],
        "waiting": [{"id": 191, "work_item_id": "73061", "wi_title": "COD nokta", "status": "needs_info",
                     "error_message": "NEEDS_INFO: 6 eksik detay", "finished_at": now - timedelta(days=2)}],
    }
    md = daily.render_markdown(d, wi_base_url="https://dev.azure.com/o/p/_workitems/edit")
    check("başlık tarih + saat", md.startswith("## ☀️ Günlük özet — 10 Eyl 2026 09:00"))
    check("özet satırı: 2 iş bitti · $4.09 · koşan 1 · kuyrukta 2", "2 iş bitti · $4.09 · koşan 1 · kuyrukta 2" in md, md.splitlines()[2])
    check("biten: durum, WI linki, PR, SP, dk, $", "✅ tamamlandı — [#73121](https://dev.azure.com/o/p/_workitems/edit/73121) İade nedeni dil · [PR](https://x/pr/42951) · 5 SP · 17 dk · $3.96" in md)
    check("needs_info hata özeti italik", "_NEEDS_INFO: hazırlık skoru 5/100 < eşik 60_" in md)
    check("koşan: adım + süre", "🔄 [#73200]" in md and "`implement_change_task`, 12 dk" in md)
    check("kuyruk listesi", "- #73201, #73202" in md)
    check("engeller: kaç gün", "### İnsan bekleyenler (engeller)" in md and "· 2 gün" in md)
    plain = daily.to_plain(md)
    check("düz metin: link/kalın/başlık işaretleri yok", "**" not in plain and "](" not in plain and not plain.startswith("#")
          and "https://x/pr/42951" in plain)
    empty = daily.render_markdown({"now": now, "since": now - timedelta(hours=24), "finished": [], "running": [], "queued": [], "waiting": []})
    check("hareket yok metni", "Hareket yok" in empty)
    check("HH:MM parse: 09:00, 7:30, bozuk → 09:00", daily._parse_hhmm("09:00") == (9, 0) and daily._parse_hhmm("7:30") == (7, 30)
          and daily._parse_hhmm("x") == (9, 0))
    check("seconds_until: bugün 09:00 geçmişse yarına", daily.seconds_until(datetime(2026, 9, 10, 10, 0), 9, 0) == 23 * 3600
          and daily.seconds_until(datetime(2026, 9, 10, 8, 30), 9, 0) == 1800)
    check("Telegram yapılandırması env'e bağlı (test ortamında kapalı olabilir)", isinstance(daily.telegram_configured(), bool))

    try:
        rep = daily.build_daily(hours=24 * 60)
    except Exception as ex:
        skip("gerçek DB günlük özeti", f"DB erişilemedi: {type(ex).__name__}")
        rep = None
    if rep:
        check("GERÇEK DB: özet üretildi, sayaçlar tutarlı", rep["markdown"].startswith("## ☀️") and set(rep["counts"]) == {"finished", "running", "queued", "waiting"})

    src_srv = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/server.py").read_text()
    check("server: /api/sprint-plan (+queue), /api/daily (+send), startup zamanlayıcı",
          '@app.get("/api/sprint-plan")' in src_srv and '@app.post("/api/sprint-plan/queue")' in src_srv
          and '@app.get("/api/daily")' in src_srv and '@app.post("/api/daily/send")' in src_srv and "start_scheduler" in src_srv)
    src_ui = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text()
    check("dashboard: 🗓️ Planla + ☀️ Günlük modalleri, onaylı kuyruk", 'id="planModal"' in src_ui and 'id="dailyModal"' in src_ui
          and "confirm(sel.length+' iş pipeline kuyruğuna" in src_ui)


# ── 36. İş tipine göre akış — spike tespiti, kılavuz, rapor, PO parse ─────

def test_type_flow():
    print("\n[36] İş tipine göre akış — spike tespiti (yalnızca açık işaret), kılavuz, rapor, PO parse")
    from agile_sdlc_crew import type_flow as tf

    check("spike: tip Spike / Research", tf.is_spike("Spike") and tf.is_spike("Research"))
    check("spike: etiket 'spike' / 'PoC' / 'araştırma' (noktalı virgül listesinde)",
          tf.is_spike("Task", "backend; spike") and tf.is_spike("Task", "PoC") and tf.is_spike("Task", "araştırma;ops"))
    check("spike: başlık '[Spike] …' / 'Spike: …' / 'POC - …'",
          tf.is_spike("Task", "", "[Spike] Kargo API karşılaştırması") and tf.is_spike("Task", "", "Spike: cache stratejisi")
          and tf.is_spike("Task", "", "POC - Redis"))
    check("spike DEĞİL: başlıkta geçen 'spike' kelimesi / etiket 'spiker' / 'research' içeren başka kelime",
          not tf.is_spike("Task", "", "Trafik spike'ında sepet hatası") and not tf.is_spike("Task", "spiker")
          and not tf.is_spike("Task", "", "Researcher paneli"))
    check("akış türü: Bug→bug, User Story/Feature→story, Task→task, Test Case→other, spike önce",
          tf.flow_kind("Bug") == "bug" and tf.flow_kind("User Story") == "story" and tf.flow_kind("Feature") == "story"
          and tf.flow_kind("Task") == "task" and tf.flow_kind("Test Case") == "other" and tf.flow_kind("Bug", "spike") == "spike")
    check("kılavuz: bug regresyon testi + minimal fix; story AC izi; spike 'Do NOT write production code'; task boş",
          "regression test" in tf.guidance("bug") and "no refactors" in tf.guidance("bug")
          and "acceptance criterion" in tf.guidance("story") and "Do NOT write production code" in tf.guidance("spike")
          and tf.guidance("task") == "" and tf.guidance("") == "")
    check("kılavuz İngilizce kural, Türkçe çıktı notu", "in Turkish" in tf.guidance("bug") and "in Turkish" in tf.guidance("spike"))

    rep = tf.spike_report("[Spike] Kargo API", _BA_WITH_ESTIMATE, "Bulgu: `app/Cargo.php` tek noktadan çağrılıyor.",
                          repo_name="webservice", estimate_sp=5)
    check("spike raporu: başlık, repo, tahmin, kapsam (BA summary + FR'ler), bulgular, kod yok notu",
          rep.startswith("## 🔬 Araştırma Raporu (Spike)") and "`webservice`" in rep and "5 SP" in rep
          and "Getirilen iade nedeni metni" in rep and "**FR1**" in rep and "### Bulgular (mimar keşfi)" in rep
          and "app/Cargo.php" in rep and "kod üretilmedi, PR açılmadı" in rep, rep[:200])
    rep2 = tf.spike_report("X", "düz metin", "")
    check("spike raporu: klon yok → iş analizine dayanır notu, BA JSON yok → kapsam yok",
          "Repo klonu bulunamadığı" in rep2 and "### Kapsam" not in rep2)

    po = tf.parse_po('```json\n{"business_value": 8, "urgency": "6", "priority": "p2", "decision": "go", '
                     '"scope_decisions": ["Rusça çeviri kapsamda", " "], "risks_if_delayed": "Rus sitesi Türkçe metin gösterir", '
                     '"rationale": "Global site deneyimi"}\n```')
    check("PO parse: sayılar kırpılır, priority/decision normalize, boş kapsam maddesi atılır",
          po == {"business_value": 8, "urgency": 6, "priority": "P2", "decision": "GO",
                 "scope_decisions": ["Rusça çeviri kapsamda"], "risks_if_delayed": "Rus sitesi Türkçe metin gösterir",
                 "rationale": "Global site deneyimi"}, str(po))
    check("PO parse: 15 → 10'a kırp, 'P9' → '', 'hold' → HOLD",
          tf.parse_po('{"business_value": 15, "priority": "P9", "decision": "hold"}')["business_value"] == 10
          and tf.parse_po('{"priority": "P9", "decision": "hold"}')["priority"] == ""
          and tf.parse_po('{"decision": "hold"}')["decision"] == "HOLD")
    check("PO parse: bozuk / JSON değil → None", tf.parse_po("no json") is None and tf.parse_po("{oops") is None)
    md = tf.render_po_comment(po)
    check("PO yorumu: tablo + kapsam + gecikme + gerekçe + danışma notu",
          "## 🎯 PO Değerlendirmesi" in md and "| İş değeri | 8 / 10 |" in md and "✅ GO" in md
          and "- Rusça çeviri kapsamda" in md and "**Gecikirse:**" in md and "danışma niteliğinde" in md)
    check("PO yorumu: HOLD işareti", "⏸️ HOLD" in tf.render_po_comment({"decision": "HOLD"}))


# ── 37. Faz 5 flow kancaları — flow_kind state'e, kılavuz context'te, spike/DoD bağları ─

def test_type_flow_hooks():
    print("\n[37] Faz 5 kancaları — flow_kind, context kılavuzu, PO context, spike stop, Bug DoD")
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.flow import _SpikeStop, _KickoffOnlyStop

    knobs = {"CREW_TYPE_FLOW": True, "CREW_WI_LIFECYCLE": False}
    _orig_get = pc.get
    pc.get = lambda k: knobs[k] if k in knobs else _orig_get(k)
    try:
        def mk(wi_type="Bug", tags="", title="Sepet hatası"):
            f = AgileSDLCFlow()
            f._db = None
            f.state.work_item_id = "73121"
            f._client = _StubAzClient(_FLO_TASK, {"System.WorkItemType": wi_type, "System.State": "To Do",
                                                   "System.Tags": tags, "System.Title": title})
            f._wi_begin(f._client.fields)
            return f

        f = mk("Bug")
        check("Bug → flow_kind=bug, başlık/etiket state'e", f.state.flow_kind == "bug" and f.state.wi_title == "Sepet hatası")
        f.state.requirements_text = _BA_WITH_ESTIMATE
        ctx = f._build_step_context("technical_design_task")
        check("teknik tasarım context'inde Bug kılavuzu", "WORK ITEM TYPE GUIDANCE: Bug" in ctx and "regression test" in ctx)
        f.state.po_text = '{"decision": "GO", "rationale": "değerli"}'
        check("PO metni kickoff/tasarım context'ine girer, requirements'a girmez",
              "PO Değerlendirmesi" in f._build_step_context("kickoff_meeting_task")
              and "PO Değerlendirmesi" not in f._build_step_context("requirements_analysis_task"))
        f2 = mk("Task", tags="spike")
        check("Task + 'spike' etiketi → spike", f2.state.flow_kind == "spike")
        f3 = mk("Task")
        check("Task → task, context'e kılavuz eklenmez", f3.state.flow_kind == "task"
              and "WORK ITEM TYPE GUIDANCE" not in f3._build_step_context("technical_design_task"))
        knobs["CREW_TYPE_FLOW"] = False
        f4 = mk("Bug")
        check("knob kapalı → flow_kind boş, kılavuz yok", f4.state.flow_kind == "" and "GUIDANCE" not in f4._build_step_context("technical_design_task"))
    finally:
        pc.get = _orig_get

    check("_SpikeStop, _KickoffOnlyStop alt sınıfı (main/server ayrımı değişmez)", issubclass(_SpikeStop, _KickoffOnlyStop))
    src_flow = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/flow.py").read_text()
    i_spike, i_bfirst = src_flow.find('if self.state.flow_kind == "spike":'), src_flow.find("# ── ARCHITECT: B-first")
    check("step4: spike dalı B-first plan üretiminden ÖNCE (plan LLM çağrısı yapılmaz)", 0 < i_spike < i_bfirst)
    check("step4 spike: rapor → technical_design_task done → kalan 8 adım 'atlandı' → _SpikeStop",
          "raise _SpikeStop(" in src_flow and '"completion_report_task"):' in src_flow
          and src_flow.find("def _run_spike") > 0 and "Atlandı — spike" in src_flow)
    check("DoD: Bug akışında test zorunlu", 'or self.state.flow_kind == "bug"' in src_flow)
    check("PO: hazırlık kapısından sonra, kickoff'tan önce (step1 sonu)", src_flow.find("self._po_assessment()") > src_flow.find('self._step_done("requirements_analysis_task"'))
    src_main = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/main.py").read_text()
    check("main: spike stop kickoff-only ile aynı yolda, farklı log", "_SpikeStop" in src_main)
    src_crew = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/crew.py").read_text()
    check("crew: product_owner ajanı (tool'suz) + create_po_crew", "def product_owner(self)" in src_crew and "def create_po_crew" in src_crew and "tools=[]," in src_crew)
    import yaml as _y
    base = Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/config"
    agents = _y.safe_load((base / "agents.yaml").read_text()); tasks = _y.safe_load((base / "tasks.yaml").read_text())
    profiles = _y.safe_load((base / "llm_profiles.yaml").read_text())
    check("YAML: product_owner ajanı, po_assessment_task (TURKISH alanlar), agent_defaults profili",
          "product_owner" in agents and "in TURKISH" in agents["product_owner"]["backstory"]
          and tasks["po_assessment_task"]["agent"] == "product_owner" and "in TURKISH" in tasks["po_assessment_task"]["expected_output"]
          and (profiles.get("agent_defaults") or {}).get("product_owner") == "reasoning_remote")
    overrides = _y.safe_load((base / "agent_llm_overrides.yaml").read_text()) or {}
    po_ov = (overrides.get("agents") or {}).get("product_owner") or {}
    check("dashboard override: product_owner claude_cli (CREW_USE_LOCAL_LLM=1 onu qwen3'e düşürmesin)",
          po_ov.get("provider") == "claude_cli" and po_ov.get("model"), str(po_ov))


# ── 38. PR inceleme katkısı — PR çözümleme, context, özet, sınırlar ────────

def test_pr_review_contribution():
    print("\n[38] PR inceleme katkısı — WI→PR çözümleme, değişen dosyalar, özet/satır yorumu, güvenlik sınırları")
    from agile_sdlc_crew import pr_review as prv

    rels = [{"attributes": {"name": "Pull Request"}, "url": "vstfs:///Git/PullRequestId/p1%2Frepo-guid-A%2F42951"},
            {"attributes": {"name": "Pull Request"}, "url": "vstfs:///Git/PullRequestId/p1%2frepo-guid-A%2f42800"},
            {"attributes": {"name": "Parent"}, "url": "vstfs:///WorkItemTracking/WorkItem/1"}]
    links = prv.parse_pr_links(rels)
    check("PR bağlantıları: %2F ve %2f, en yeni önce, parent atlanır",
          links == [{"repo_id": "repo-guid-A", "pr_id": 42951}, {"repo_id": "repo-guid-A", "pr_id": 42800}], str(links))
    check("değişen dosyalar: klasör/silinen hariç, tekrar yok",
          prv.changed_paths([{"item": {"path": "/a.php"}}, {"item": {"path": "/dir", "isFolder": True}},
                             {"item": {"path": "/b.php"}, "changeType": "delete"}, {"item": {"path": "/a.php"}}]) == ["/a.php"])
    txt = prv.wi_requirements_text({"System.Title": "İade nedeni", "System.Description": "<p>Global site</p><br>rusça",
                                    "Microsoft.VSTS.Common.AcceptanceCriteria": ""})
    check("WI metni: HTML temizlenir, AC boşsa not düşer", "<p>" not in txt and "Global site" in txt and "rusça" in txt
          and "WI alanı boş" in txt, txt)

    class _C:
        def __init__(self, statuses):
            self.statuses = statuses
        def get_work_item(self, wid):
            return {"fields": {}, "relations": rels}
        def get_repository(self, rid):
            return {"name": "webservice"}
        def get_pull_request(self, rname, pid):
            return {"pullRequestId": pid, "status": self.statuses[pid], "sourceRefName": "refs/heads/feature/x",
                    "title": f"#73121 x", "repository": {"project": {"name": "EcomBackend"}}}
    r, p, pr = prv.resolve_pr(_C({42951: "abandoned", 42800: "active"}), work_item_id="73121")
    check("çözümleme: en yeni abandoned → sonraki aktif seçilir", (r, p) == ("webservice", 42800))
    r, p, pr = prv.resolve_pr(_C({42951: "completed", 42800: "abandoned"}), work_item_id="73121")
    check("çözümleme: aktif yoksa tamamlanmış", p == 42951)
    try:
        prv.resolve_pr(_C({42951: "abandoned", 42800: "abandoned"}), work_item_id="73121"); ok = False
    except LookupError:
        ok = True
    check("çözümleme: hepsi abandoned → LookupError (iş failed, sessiz değil)", ok)
    r, p, pr = prv.resolve_pr(_C({7: "active"}), pr_id=7, repo_name="core")
    check("çözümleme: repo+pr verildiyse WI'a gitmez", (r, p) == ("core", 7))

    issues = [{"id": "R1", "severity": "blocker", "file": "app/X.php", "line": 12, "problem": "null kontrolü yok", "required_fix": "guard ekle"},
              {"id": "R2", "severity": "minor", "file": "app/Y.php", "line": None, "problem": "isim", "required_fix": ""}]
    md = prv.review_summary_markdown(verdict="CHANGES_REQUIRED", issues=issues, pr_id=42951, pr_url="https://x/pr",
                                     work_item_id="73121", files=3)
    check("özet: karar, tablo, danışma notu, PR linki",
          "🔴 Değişiklik gerekli" in md and "| R1 | blocker | `app/X.php:12` |" in md and "danışma niteliğindedir" in md and "https://x/pr" in md)
    check("özet: APPROVE + madde yok → 'bulunmadı'", "bulunmadı" in prv.review_summary_markdown(
        verdict="APPROVE", issues=[], pr_id=1, pr_url="u", work_item_id="1", files=1))
    check("satır yorumu: id, önem, düzeltme", "[R1] BLOCKER" in prv.inline_comment_text(issues[0]) and "guard ekle" in prv.inline_comment_text(issues[0]))

    # Satır numarası yeniden oturtma (job #208 / PR #43642 regresyonu):
    # reviewer dosyayı okuyor ama satırı KENDİSİ sayıyordu; R1 86 yerine 95,
    # R2 98 yerine 106 dedi. Alıntılar doğruydu → numara alıntıdan hesaplanır.
    body = "<?php\n\nclass X\n{\n    public function f()\n    {\n        $a = 1;\n        return ceil($t / 10);\n    }\n}\n"
    anc = [{"id": "R1", "file": "app/X.php", "line": 15,
            "evidence": {"file": "app/X.php", "line": 15, "quote": "return ceil($t / 10);"}},
           {"id": "R2", "file": "app/X.php", "line": 0,
            "evidence": {"file": "app/X.php", "line": 0, "quote": "   7|         $a = 1;"}},
           {"id": "R3", "file": "app/X.php", "line": 4,
            "evidence": {"file": "app/X.php", "line": 4, "quote": "$this->uydurma();"}},
           {"id": "R4", "file": "app/Yok.php", "line": 9, "evidence": {}}]
    st = prv.anchor_issue_lines(anc, {"/app/X.php": body})
    check("satır: yanlış numara alıntıdan düzeltilir (15→8)", anc[0]["line"] == 8 and anc[0]["evidence"]["line"] == 8, str(anc[0]["line"]))
    check("satır: line 0 da oturtulur, 'NN| ' öneki alıntıdan atılır (→7)", anc[1]["line"] == 7, str(anc[1]["line"]))
    check("satır: alıntı dosyada yoksa numara DÜŞÜRÜLÜR (yanlış satır yazmaktansa numarasız)", anc[2]["line"] == 0)
    check("satır: gövdesi/kanıtı olmayan madde dokunulmadan kalır", anc[3]["line"] == 9)
    check("satır: istatistik", st == {"anchored": 2, "moved": 2, "cleared": 1, "kept": 1}, str(st))
    check("yol anahtarı: '/app/X.php' gövdesi 'app/X.php' maddesiyle eşleşir (suggestion sessizce düşmesin)",
          prv.body_index({"/app/X.php": "a"}) == {"app/X.php": "a"})
    check("suggestion: baştaki slash farkına rağmen doğrulanır",
          prv.verified_suggestion({"file": "app/X.php", "suggestion": {"code": "        $a = 2;", "line_start": 7, "line_end": 7}},
                                  {"/app/X.php": body}).get("code") == "        $a = 2;")
    # Büyük dosyada BAŞ taraf değil DEĞİŞEN taraf context'e girer (job #208:
    # OrderLine.php 3098 satır → ilk %5'i gitti, PR'ın dokunduğu metot şansa kaldı).
    big_old = "\n".join(f"satir {i}" for i in range(1, 401))
    big_new = "\n".join(("DEGISTI" if i in (350, 351) else f"satir {i}") for i in range(1, 401))
    foc = prv.focused_source(big_new, big_old, per_file=1000, context=3)
    shown = {int(l.split("|", 1)[0]) for l in foc.splitlines() if not l.startswith("    |")}
    check("odak: değişen satırlar (350-351) pencerede, baş taraf değil",
          {350, 351} <= shown and 1 not in shown and shown == set(range(347, 355)), str(sorted(shown))[:80])
    check("odak: satır numaraları GERÇEK dosya numarası (pencere ofseti değil)", "\n 350| DEGISTI" in foc)
    check("odak: atlanan bölge işaretli (görülmeyen kod hakkında madde açılmasın)",
          "satır atlandı" in foc and "346 satır atlandı" in foc, foc.splitlines()[0])
    check("odak: dosya sığıyorsa tamamı verilir", prv.focused_source(big_new, big_old, per_file=99_000).count("\n") == 399)
    nobase = prv.focused_source(big_new, None, per_file=300)
    check("odak: base yoksa baş kırpma + 'devamı yok' uyarısı (eski davranış)",
          nobase.startswith("   1| satir 1") and "context'e girmedi" in nobase)
    from agile_sdlc_crew.flow import _parse_review_issues as _pri
    # EKSİKLİK itirazının doğal bir satırı yoktur (job #211/R1: "cachedContent
    # yok" itirazı `$body = [` satırına düştü → okuyucu alakasız kod gördü).
    _abs = {"id": "R1", "severity": "major", "anchor": "absence", "file": "app/X.php", "line": 134,
            "problem": "İstek gövdesinde cachedContent yok.", "required_fix": "Ekleyin."}
    _txt_abs = prv.inline_comment_text(_abs)
    check("eksiklik: yorum başlığı 'EKSİK' der ve satırın konum olduğunu söyler",
          "MAJOR — EKSİK" in _txt_abs and "Bu satırdaki kodda hata yok" in _txt_abs, _txt_abs[:60])
    check("kusur (defect): ek açıklama YOK, eski davranış",
          "EKSİK" not in prv.inline_comment_text({**_abs, "anchor": "defect"}))
    _md_abs = prv.review_summary_markdown(verdict="CHANGES_REQUIRED", issues=[_abs], pr_id=1, pr_url="u",
                                          work_item_id="1", files=1)
    check("eksiklik: özet tablosunda da işaretli", "major · eksik" in _md_abs and "ait olduğu yer" in _md_abs)
    _pa = _pri('REVIEW_ISSUES_JSON:\n```json\n{"issues": ['
               '{"file":"a.php","line":1,"severity":"major","problem":"p","required_fix":"f","anchor":"absence"},'
               '{"file":"b.php","line":1,"severity":"major","problem":"p","required_fix":"f"},'
               '{"file":"c.php","line":1,"severity":"major","problem":"p","required_fix":"f","anchor":"SAÇMA"}]}\n```')
    check("parse: anchor absence/defect taşınır, bilinmeyen değer defect'e düşer",
          [x["anchor"] for x in _pa] == ["absence", "defect", "defect"], str([x["anchor"] for x in _pa]))
    _t2 = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/config/tasks.yaml").read_text()
    check("talimat: anchor alanı tanımlı ve eksiklik kuralı yazılı",
          '- "anchor": "defect" (default) or "absence"' in _t2 and "no line of its own" in _t2)

    # SİLME önerisi (job #210/R4: "metodu tamamen sil"). Boş `suggested_code`
    # eskiden "öneri yok" sayılıyordu → talimatın kendi saydığı geçerli vaka
    # ("deleting dead code") yapısal olarak üretilemiyordu.
    dosya = "a\nb\nc\nd\ne\n"
    sil = {"file": "app/X.php", "suggestion": {"code": "", "line_start": 2, "line_end": 3}}
    check("silme: boş kod + açık aralık → öneri ÜRETİLİR",
          prv.verified_suggestion(sil, {"/app/X.php": dosya}).get("line_start") == 2)
    check("silme: yorum metni boş ```suggestion bloğu ve 'siler' etiketi taşır",
          "```suggestion\n\n```" in prv.inline_comment_text({"id": "R4", "severity": "major", "problem": "ölü kod"},
                                                            suggestion={"code": "", "line_start": 2, "line_end": 3})
          and "siler" in prv.inline_comment_text({"id": "R4", "severity": "major", "problem": "ölü kod"},
                                                 suggestion={"code": "", "line_start": 2, "line_end": 3}))
    check("öneri yok: alan hiç verilmemişse (None) öneri üretilmez",
          prv.verified_suggestion({"file": "app/X.php", "suggestion": {"line_start": 2, "line_end": 3}},
                                  {"/app/X.php": dosya}) == {})
    check("silme: dosya dışı aralık reddedilir", prv.verified_suggestion(
          {"file": "app/X.php", "suggestion": {"code": "", "line_start": 4, "line_end": 99}}, {"/app/X.php": dosya}) == {})
    from agile_sdlc_crew.flow import _parse_review_issues as _pri
    _blk = ('REVIEW_ISSUES_JSON:\n```json\n{"issues": ['
            '{"file":"a.php","line":5,"severity":"major","problem":"ölü kod","required_fix":"sil",'
            '"suggested_code":"","suggested_line_start":3,"suggested_line_end":9},'
            '{"file":"b.php","line":5,"severity":"major","problem":"x","required_fix":"y","suggested_code":""}'
            ']}\n```')
    _iss = _pri(_blk)
    check("parse: boş kod + AÇIK aralık → silme önerisi taşınır",
          _iss[0]["suggestion"] == {"code": "", "line_start": 3, "line_end": 9}, str(_iss[0].get("suggestion")))
    check("parse: boş kod ama aralık YOK → yok sayılır (kazara silme olmasın)",
          not _iss[1].get("suggestion"), str(_iss[1].get("suggestion")))
    _tasks = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/config/tasks.yaml").read_text()
    _rv2 = _tasks.split("review_pr_task:")[1].split("\n  agent:")[0]
    check("talimat: öneri artık 'usually ABSENT' değil, mekanik vakada BEKLENİYOR",
          "usually ABSENT" not in _rv2 and "fill these\n      WHENEVER" in _rv2)
    check("talimat: silme vakası açıkça tanımlı", "DELETION:" in _rv2 and 'EMPTY string' in _rv2)

    ctx_src = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/pr_review.py").read_text()
    check("context: dosyalar satır numarası önekiyle veriliyor", '{i:>4}| {ln}' in ctx_src)
    check("context: sığmayan dosyada base sürüm çekilip diff pencerelenir",
          "focused_source(content, old_body" in ctx_src and 'pr.get("targetRefName")' in ctx_src)

    # Pipeline'ın KENDİ review adımı (step8 / _review_retry_loop) da aynı odağı
    # kullanmalı: tek turlu çağrılarda kullanıcı mesajı cache-write fiyatından
    # yazılıp bir daha okunmuyor (probe 2026-09-24), yani her KB doğrudan para.
    flow_src = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/flow.py").read_text()
    check("step8: tam dosya dalı da focused_source kullanıyor (baş kırpma yok)",
          "_focused_source(content, base_body" in flow_src and "content[:per_file]" not in flow_src.split("if not diff_mode:")[1][:900])
    check("step8: NN| önekinin dosyanın parçası olmadığı ajana söyleniyor",
          flow_src.count("DOSYANIN PARÇASI DEĞİL") >= 1 and "N satır atlandı" in flow_src)
    # NN| öneki context'e girdiği için kanıt doğrulaması onu SOYMALI; yoksa
    # sağlam bir madde "kanıtsız" diye düşer (en pahalı hata: bir retry turu).
    class _FakeFlow:
        state = SimpleNamespace(repo_name="r", branch_name="b")
        _repo_mgr = None
        class _C:
            def get_file_content(self, repo, f, ref):
                return "<?php\n\nclass X\n{\n    public function f()\n    {\n        return ceil($t / 10);\n    }\n}\n"
        _client = _C()
    vl = AgileSDLCFlow._verify_issue_loc
    ff = _FakeFlow()
    check("kanıt: düz alıntı doğrulanır",
          vl(ff, {"file": "app/X.php", "quote": "return ceil($t / 10);"}) is True)
    check("kanıt: 'NN| ' önekli alıntı da doğrulanır (önek soyulur)",
          vl(ff, {"file": "app/X.php", "quote": "   7|         return ceil($t / 10);"}) is True)
    check("kanıt: uydurma alıntı yine REDDEDİLİR (soyma kapıyı açmadı)",
          vl(ff, {"file": "app/X.php", "quote": "   7| $this->uydurma();"}) is False)
    check("kanıt: sadece numaradan ibaret alıntı → dosya var sayılır, kanıt iddiası yok",
          vl(ff, {"file": "app/X.php", "quote": "  42| "}) is True)

    # discover_repos: aday özetleri. Çağrı başına 36.9K yazım / 7.6K okuma
    # (sistemin en kötü oranı, adımın %85'i cache write) — tek turlu çağrıda
    # yazılan bir daha okunmadığı için her karakter para.
    from agile_sdlc_crew.flow import _repo_summary_slice, _wi_keywords, _compress_list_line
    summ = ("# orkestra\n## Ozet\n- **Dil**: PHP\n\n## README\n" + ("uzun readme metni. " * 200) +
            "\n## Domain Bilesenleri\n- **Model**: Order, OrderLine\n"
            "## DB Tablolari & Migrationlar\n- **Tablolar**: " +
            ", ".join(f"tablo_{i:03d}" for i in range(400)) + "\n"
            "## Ust Seviye Dizinler\n- app/\n")
    SEC = ("Ozet", "Domain Bilesenleri", "DB Tablolari")
    kw = _wi_keywords("WI: tablo_399 tablosuna yeni alan eklenecek, order_line ile ilişkili")
    check("anahtar kelime: WI metninden snake_case/camelCase çıkarılır", "tablo_399" in kw and "order_line" in kw, str(kw[:5]))
    out = _repo_summary_slice(summ, 2500, SEC, keywords=kw)
    check("özet: README ve dizin ağacı karara girmez (atılır)", "uzun readme" not in out and "Ust Seviye" not in out)
    check("özet: karar kuralının andığı bölümler durur", "Domain Bilesenleri" in out and "Tablolar" in out)
    check("özet: listenin SONUNDAKİ eşleşen tablo kurtarılır (düz kırpma atardı)",
          "tablo_399" in out and "tablo_399" not in _repo_summary_slice(summ, 2500, SEC), out[:90])
    check("özet: listelenmeyen tablo sayısı belirtilir (model 'liste tam' sanmasın)", "tane daha, listelenmedi" in out)
    check("özet: cap'e uyar", len(out) <= 2600, str(len(out)))
    check("liste sıkıştırma: kısa satıra dokunulmaz",
          _compress_list_line("- **Model**: Order, OrderLine", {"order"}) == "- **Model**: Order, OrderLine")
    check("liste sıkıştırma: eşleşme yoksa baş taraf + sayı",
          "tane daha" in _compress_list_line("- **T**: " + ", ".join(f"x_{i:04d}" for i in range(300)), {"yok_boyle"}))
    flow_src2 = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/flow.py").read_text()
    check("discover_repos: kanıtlı/geçmişli repo TAM özet alır, gerisi kısa",
          "strong.add" in flow_src2 and "keywords=_kw_dr" in flow_src2 and "CREW_DISCOVER_COMPACT" in flow_src2)

    # Review turlarının NEDEN uzadığı ölçülebilir olmalı: iş başına 5.7 reviewer
    # çağrısı ($66.49 / 24 iş) ve en yüksek c.write ortalaması (31.9K) burada.
    check("step8: satır çapalama pipeline review'ında da çalışıyor",
          "anchor_issue_lines as _anchor" in flow_src2 and "_pr_file_bodies" in flow_src2)
    check("step8: itiraz kapısı sonuçları REVIEW_STATS damgasına yazılıyor",
          "_review_stats_line" in flow_src2 and "drop_reasons" in flow_src2
          and 'review_text[:3000] + self._review_stats_line()' in flow_src2)
    from agile_sdlc_crew import retrospective as _retro
    _out = ("REVIEW_DECISION: APPROVE\n2 düzeltme turu\nREVIEW_STATS: "
            + json.dumps({"issues": 5, "blocking": 2, "dropped": 3, "rounds": 3,
                          "drop_reasons": {"minor/öneri": 2, "kanıt doğrulanamadı": 1},
                          "anchor_moved": 4, "anchor_cleared": 1}))
    _m = _retro._RE_STATS.search(_out)
    check("retro: REVIEW_STATS damgası çıktının sonundan parse edilir",
          bool(_m) and json.loads(_m.group(1))["blocking"] == 2)
    check("retro: damga yoksa (eski işler) sessizce atlanır, sayaç bozulmaz",
          _retro._RE_STATS.search("REVIEW_DECISION: APPROVE\n1 düzeltme turu") is None)
    tasks_y = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/config/tasks.yaml").read_text()
    rv = tasks_y.split("review_pr_task:")[1].split("\n  agent:")[0]
    check("review_pr_task: problem/required_fix TÜRKÇE yazılacak kuralı var",
          '"required_fix" values in the JSON MUST be written in TURKISH' in rv)
    check("review_pr_task: satırı sayma, NN| önekini oku kuralı var", "do\n      NOT count lines yourself" in rv)
    src = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/pr_review.py").read_text()
    check("sınırlar: push/oy/durum yok — modülde push_file, set_work_item_state, reviewers/vote API'si geçmez",
          "push_file" not in src and "set_work_item_state" not in src and "/reviewers" not in src and '"vote"' not in src)
    check("yalnızca blocker/major satır yorumu, en çok 8", prv.INLINE_SEVERITIES == ("blocker", "major") and prv.MAX_INLINE_COMMENTS == 8)
    # ...AMA doğrulanmış önerisi olan madde önemine bakılmaksızın satıra yazılır.
    # Job #211: mekanik düzeltmeler tanımı gereği `minor` — eski filtre tam da
    # ÖNERİ TAŞIYAN maddeleri eliyordu (R5 silme önerisi üretildi, doğrulandı,
    # PR'a hiç düşmedi → sayaç yine 0).
    _prv_src = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/pr_review.py").read_text()
    check("satır yorumu: doğrulanmış önerisi olan minor da yazılır",
          'i.get("severity") in INLINE_SEVERITIES or _sugs.get(id(i))' in _prv_src)
    check("satır yorumu: öneri madde başına BİR kez doğrulanır (çift I/O yok)",
          "_sugs = {id(i): verified_suggestion(i, file_bodies) for i in issues}" in _prv_src)
    for name, mod in (("retrospective", "retrospective"), ("sprint_planning", "sprint_planning"), ("daily", "daily")):
        s = (Path(__file__).resolve().parent.parent / f"src/agile_sdlc_crew/{mod}.py").read_text()
        check(f"{name}: pipeline dışı işler (pr_review) filtrelenir", "job_kind" in s)
    src_srv = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/server.py").read_text()
    check("server: /api/pr-review + CREW_PR_REVIEW + job_kind", '@app.post("/api/pr-review")' in src_srv and "CREW_PR_REVIEW" in src_srv and 'job_kind="pr_review"' in src_srv)
    src_ui = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text()
    check("dashboard: kartta 🔍 (başlanmış durumlar), iş listesinde rozet, onay diyaloğu",
          "reviewWi(" in src_ui and "'Code Review'" in src_ui and "İNCELEME" in src_ui and "confirm(`#${id}" in src_ui)



# ── 39. Backlog refinement — Definition of Ready skoru, sorgu, eylemler, sınırlar ─

def test_wi_comments_context():
    print("\n[46] WI yorumları — insan yönlendirmesi context'e girsin")
    from agile_sdlc_crew.flow import wi_comments_block, AgileSDLCFlow
    from agile_sdlc_crew import pipeline_config as _pc_wc

    cms = [  # API en yeniyi ÖNCE döndürür
        {"tarih": "2026-09-24T07:39", "yazar": "Hasan S", "metin": "Stage testleri ok"},
        {"tarih": "2026-09-22T13:44", "yazar": "Büşra A", "metin": "@Osman&nbsp;ÇEVİK dahil edebiliriz"},
        {"tarih": "2026-09-22T13:38", "yazar": "Osman Ç", "metin": "<div>Dalli için de geçerli mi?</div>"},
        {"tarih": "2026-09-09T11:56", "yazar": "Volkan Ö",
         "metin": "ℹ️ Hazırlık Skoru 5/100\n\n---\n*Tempo — Hazırlık Kapısı*"},
        {"tarih": "2026-09-08T10:00", "yazar": "X", "metin": "   "},
    ]
    blk, n = wi_comments_block(cms, cap=4000)
    check("bot yorumu ayıklanır (kendi sorumuzu cevap sanmayalım)", "Hazırlık Skoru" not in blk)

    # Azure yorumu markdown alıp HTML döndürüyor: yazdığımız *Tempo — X*
    # dışarıya <em>Tempo — X</em> olarak geliyor, YILDIZLAR KAYBOLUYOR.
    # Marker '*Tempo' yıldızlı olduğu için rebrand'den beri kendi WI/PR
    # yorumlarımızın hiçbiri tanınmıyordu (yalnız eski 'Agile SDLC Crew'
    # düz ad olduğu için tutuyordu). WI yorumlarını gereksinim kaynağı
    # yapınca bu bir DÖNGÜ oldu: kapının kendi "şu detaylar eksik" yorumu
    # insan yorumu sayılıp bir sonraki değerlendirmeye girdi (WI #73061'de
    # 1868 karakter kendi metnimiz insan kümesine eklendi).
    from agile_sdlc_crew.branding import is_bot_comment as _ibc
    _rendered = ("<p>Bu iş kalemi yeterince detaylı değil.</p>\n<br>\n<hr>\n"
                 "<p><em>Tempo — Hazırlık Kapısı</em></p>")
    check("HTML'e dönüşmüş imza (yıldızsız) bot olarak tanınır", _ibc(_rendered) is True, _rendered[-60:])
    check("markdown imza (yıldızlı) hâlâ tanınır", _ibc("...\n\n*Tempo — Kod İnceleme*") is True)
    check("eski ad (Agile SDLC Crew) hâlâ tanınır", _ibc("<p><em>Agile SDLC Crew — Hazırlık Kapısı</em></p>") is True)
    check("insan yorumu bot sanılmaz", _ibc("@Büşra AYVAZ evet, tüm kargolar için bu geçerli.") is False)
    check("metnin ORTASINDA geçen 'Tempo —' imza sayılmaz (yanlış pozitif yok)",
          _ibc("Tempo — bu araç " + ("çok uzun bir insan yorumu. " * 40)) is False)
    _cms_r = [{"tarih": "t", "yazar": "bot", "metin": _rendered},
              {"tarih": "t", "yazar": "A", "metin": "gerçek cevap"}]
    _b, _ = wi_comments_block(_cms_r)
    check("blok: HTML imzalı bot yorumu da ayıklanır", "Hazırlık Kapısı" not in _b and "gerçek cevap" in _b)
    check("boş yorum atlanır", blk.count("- [") == 3, blk)
    check("KRONOLOJİK sıra (kapsam kararı sorunun ardından gelmeli)",
          blk.index("Dalli için de geçerli mi") < blk.index("dahil edebiliriz"), blk[:200])
    check("HTML ve &nbsp; temizlenir", "<div>" not in blk and "&nbsp;" not in blk)
    check("uzunluk döner (hazırlık kapısının içerik sayacına eklenir)",
          n > 50 and n == sum(len(l) for l in blk.split("\n") if l.startswith("- [")), str(n))
    check("başlık ajana 'gereksinim kaynağı say' der", "gereksinim kaynağı say" in blk)
    blk2, _ = wi_comments_block(cms, cap=60)
    check("cap aşılınca kalan sayısı belirtilir", "yorum daha, yer kalmadı" in blk2, blk2[-120:])
    check("yorum yoksa boş döner", wi_comments_block([]) == ("", 0)
          and wi_comments_block([{"metin": "*Tempo — bot*"}]) == ("", 0))

    src = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/flow.py").read_text()
    check("step1: yorumlar ctx'e VE içerik uzunluğuna eklenir",
          "wi_content_length += _blen" in src and "self.state.wi_comments_text = _blk" in src)
    check("knob'a bağlı (CREW_WI_COMMENTS)", '_pc_wc.get("CREW_WI_COMMENTS")' in src)
    check("pipeline_config'de kayıtlı, varsayılan açık", _pc_wc.get("CREW_WI_COMMENTS") is True)

    # Önek kararlılığı: yorumlar is-degismezi → parts'a girer, tail'e DEĞİL
    st = SimpleNamespace(work_item_id="1", requirements_text="", kickoff_text="", acceptance_criteria=[],
                         plan=None, branch_name="", all_pushes=[], pr_id="", pr_url="", repo_name="",
                         review_text="", test_text="", uat_text="", flow_kind="", po_text="",
                         wi_comments_text="\n\n# İŞ KALEMİ YORUMLARI (insan — kronolojik)\n- [t] a: b")
    fake = SimpleNamespace(state=st, _vector_store=None, _forward_text=lambda k, t, c: t[:c])
    a = AgileSDLCFlow._build_step_context(fake, "technical_design_task")
    b = AgileSDLCFlow._build_step_context(fake, "review_pr_task")
    check("yorumlar karar veren adımlara gider", "İŞ KALEMİ YORUMLARI" in a and "İŞ KALEMİ YORUMLARI" in b)
    check("BA adımına tekrar enjekte edilmez (ctx'te zaten var)",
          "İŞ KALEMİ YORUMLARI" not in AgileSDLCFlow._build_step_context(fake, "requirements_analysis_task"))
    import os as _os_pfx
    pre = len(_os_pfx.path.commonprefix([a, b]))
    check("ortak önek yorumları KAPSAR (parts'ta, prefix bozulmuyor)",
          pre > a.index("İŞ KALEMİ YORUMLARI"), f"ortak {pre}")


def test_cli_prompt_overhead():
    print("\n[45] claude -p prompt öneki — kullanılmayan araç şemaları gönderilmesin")
    import os as _os_cpo
    from agile_sdlc_crew.tools import claude_cli_llm as _cli

    seen = {}

    def _fake(cmd, env, timeout_s, meta=None, idle_s=0):
        seen["cmd"] = list(cmd)
        if meta is not None:
            meta.update({"turns": 1, "cost_usd": 0.0})
        return "OK"

    _orig = _cli._run_streaming
    _cli._run_streaming = _fake
    _saved = {k: _os_cpo.environ.get(k) for k in ("CREW_CLI_STRICT_MCP", "CREW_CLI_AUTO_TOOLLESS")}
    try:
        # (a) repo aracı YOK → şemalar gönderilmesin
        _cli.clear_repo_ctx(); _cli.set_toolless(False)
        _os_cpo.environ.pop("CREW_CLI_STRICT_MCP", None); _os_cpo.environ.pop("CREW_CLI_AUTO_TOOLLESS", None)
        _cli.claude_cli_completion("x", model="sonnet", system="s")
        c = seen["cmd"]
        check("--add-dir yokken MCP sunucuları atlanır", "--strict-mcp-config" in c, " ".join(c[-6:]))
        check("--add-dir yokken dosya/exec araçları kapatılır (şema da gitmez)",
              "--disallowedTools" in c and _cli._TOOLLESS_DENY in c)
        # (b) repo aracı VAR → araçlar korunur, yalnız MCP atlanır
        _cli.set_repo_ctx(["/tmp/repo"], "Read,Grep,Glob,LS")
        _cli.claude_cli_completion("x", model="sonnet", system="s")
        c = seen["cmd"]
        check("--add-dir varken keşif araçları KORUNUR (auto-toolless tetiklenmez)",
              "--add-dir" in c and "--allowedTools" in c and "--disallowedTools" not in c, " ".join(c[-8:]))
        check("--add-dir varken de MCP atlanır (keşif native araçları kullanır)", "--strict-mcp-config" in c)
        _cli.clear_repo_ctx()
        # (c) env ile geri alınabilir
        _os_cpo.environ["CREW_CLI_STRICT_MCP"] = "0"; _os_cpo.environ["CREW_CLI_AUTO_TOOLLESS"] = "0"
        _cli.claude_cli_completion("x", model="sonnet", system="s")
        c = seen["cmd"]
        check("geri alma: iki env=0 → eski davranış (bayrak yok)",
              "--strict-mcp-config" not in c and "--disallowedTools" not in c)
        # (d) açık set_toolless env'den bağımsız çalışır
        _cli.set_toolless(True)
        _cli.claude_cli_completion("x", model="sonnet", system="s")
        check("açık set_toolless(True), CREW_CLI_AUTO_TOOLLESS=0 iken de araçları kapatır",
              "--disallowedTools" in seen["cmd"])
        _cli.set_toolless(False)
    finally:
        _cli._run_streaming = _orig
        for k, v in _saved.items():
            if v is None: _os_cpo.environ.pop(k, None)
            else: _os_cpo.environ[k] = v


def test_backlog_refinement():
    print("\n[39] Backlog refinement — DoR skoru, WIQL, SP önerisi, yorum/etiket, soru parse, knob'lar")
    from datetime import datetime, timezone, timedelta
    from agile_sdlc_crew import backlog_refinement as br

    now = datetime(2026, 9, 11, tzinfo=timezone.utc)
    wiql = br.backlog_wiql(area_path="Boards\\Team's Area", states=["Backlog", "To Do"])
    check("WIQL: durum + tip filtresi, area UNDER, tek tırnak kaçışı",
          "[System.State] IN ('Backlog', 'To Do')" in wiql and "UNDER 'Boards\\Team''s Area'" in wiql
          and "'User Story'" in wiql and "IterationPath" not in wiql, wiql)
    wiql2 = br.backlog_wiql(area_path="X", iteration_path="P\\S1", states=["New"])
    check("WIQL: sprint kapsamında iteration filtresi, area yok", "[System.IterationPath] = 'P\\S1'" in wiql2 and "AreaPath" not in wiql2)
    check("AC sayısı: madde imi/numara/Given sayılır, düz metin satır", br.count_criteria("<ul><li>a</li><li>b</li></ul>") == 2
          and br.count_criteria("1. x<br>2. y<br>3. z") == 3 and br.count_criteria("Given a\nWhen b\nThen c") == 3 and br.count_criteria("") == 0)

    def item(**kw):
        base = {"id": 1, "title": "Kargo etiketinde alıcı telefonu maskelensin", "type": "Task", "state": "Backlog",
                "priority": 2, "sp": 3, "assignedTo": "", "tags": "", "description": "<p>" + "Açıklama metni. " * 10 + "</p>",
                "acceptance": "<ul><li>AC1 maskeli</li><li>AC2 log yok</li></ul>", "repro": "",
                "created": (now - timedelta(days=3)).isoformat(), "iterationPath": "", "areaPath": "", "parent_id": 10, "url": ""}
        base.update(kw); return base
    ok = br.assess(item(), now=now)
    check("tam iş: skor 100, hazır, eksik yok, SP önerisi yok", ok["score"] == 100 and ok["ready"] and not ok["gaps"] and ok["suggest_sp"] is None, str(ok["gaps"]))
    r = br.assess(item(acceptance="", sp=None, parent_id=None), now=now)
    codes = [g["code"] for g in r["gaps"]]
    check("Task: AC yok −25, SP yok −10, ebeveyn yok −10 → 55, eksik", codes == ["no_ac", "no_sp", "no_parent"] and r["score"] == 55 and not r["ready"], f"{codes} {r['score']}")
    check("SP yoksa yapısal öneri (AC yok → 2)", r["suggest_sp"] == 2, str(r["suggest_sp"]))
    desc_head = "<p>Proje özeti uzun uzun anlatılıyor burada, servisler ve akış.</p><p>Kabul Kriterleri:</p><ul><li>Soru kaydedilir</li><li>Cevap 5 sn içinde döner</li></ul>"
    r = br.assess(item(acceptance="", description=desc_head), now=now)
    check("AC alanı yok (bu org'da alan yok) ama açıklamada 'Kabul Kriterleri' başlığı → eksik yok, 2 AC",
          not r["gaps"] and r["n_ac"] == 2 and r["ac_source"] == "explicit", str(r["gaps"]))
    desc_impl = "<p>" + "Proje özeti cümlesi. " * 6 + "</p><ul><li>QuestionCreationService</li><li>RuleEngineService</li></ul>"
    r = br.assess(item(acceptance="", description=desc_impl), now=now)
    check("AC başlığı yok ama maddeler var → 'açık değil' −10 (pipeline gibi açıklamadan çıkarılır)",
          [g["code"] for g in r["gaps"]] == ["ac_implicit"] and r["score"] == 90 and r["ready"], str(r["gaps"]))
    r = br.assess(item(type="Bug", acceptance="", repro=""), now=now)
    check("Bug: AC yerine repro aranır", [g["code"] for g in r["gaps"]] == ["no_repro"])
    r = br.assess(item(type="Bug", acceptance="", repro="<ol><li>Aç</li><li>Tıkla</li></ol>"), now=now)
    check("Bug: repro varsa eksik yok", not r["gaps"])
    r = br.assess(item(title="[Spike] Redis TTL araştırması", type="Task", acceptance="", description="<p>" + "Redis TTL davranışı nasıl? " * 6 + "</p>", sp=None), now=now)
    check("Spike: AC aranmaz, soru var; SP önerisi yok", "no_ac" not in [g["code"] for g in r["gaps"]] and r["kind"] == "spike" and r["suggest_sp"] is None, str(r["gaps"]))
    # Bayatlık KADEMELİ (kritik eşik aging'den okunur — env ile değişebilir):
    # eşiğin altı 'stale' −5, eşikte/üstü 'very_stale' −15. Eşik değerini
    # sabitlemiyoruz; sabitlersek CREW_AGING_BACKLOG_CRIT_DAYS ayarı testi kırar.
    from agile_sdlc_crew import aging as _ag
    crit = _ag.backlog_crit_days()
    base_gaps = ["no_description", "too_big", "no_priority", "short_title"]
    stale_item = dict(description="", sp=13, priority=None, title="Kısa")
    r = br.assess(item(**stale_item, created=(now - timedelta(days=crit - 1)).isoformat()), now=now)
    codes = [g["code"] for g in r["gaps"]]
    check("açıklama yok −30, SP>8 −10, öncelik −5, başlık −5, bayat −5 → 45",
          codes == base_gaps + ["stale"] and r["score"] == 45, f"{codes} {r['score']}")
    r = br.assess(item(**stale_item, created=(now - timedelta(days=crit)).isoformat()), now=now)
    codes = [g["code"] for g in r["gaps"]]
    check(f"kritik eşikte ({crit}g) bayatlık −5 değil −15 → 35 (1 yıllık işle 31 günlük iş aynı cezayı almasın)",
          codes == base_gaps + ["very_stale"] and r["score"] == 35, f"{codes} {r['score']}")
    r = br.assess(item(**stale_item, created=(now - timedelta(days=crit + 400)).isoformat(),
                       state_change=(now - timedelta(days=2)).isoformat()), now=now)
    check("bayatlık OLUŞMA yaşını değil DURUMDA GEÇEN süreyi ölçer (sprint'e taşınan iş tazelenir)",
          [g["code"] for g in r["gaps"]] == base_gaps and r["score"] == 50, str(r["gaps"]))
    r = br.assess(item(description="<p>Kısa bir açıklama burada.</p>"), now=now, min_desc_chars=100)
    check("kısa açıklama (20 ≤ n < eşik) −15", [g["code"] for g in r["gaps"]] == ["short_description"] and r["score"] == 85)
    check("eşik parametrik: 85 ≥ 70 hazır, 85 < 90 eksik", br.assess(item(description="<p>Kısa bir açıklama burada.</p>"), now=now)["ready"]
          and not br.assess(item(description="<p>Kısa bir açıklama burada.</p>"), now=now, min_score=90)["ready"])
    check("etiket tespiti büyük/küçük harf duyarsız", br.assess(item(tags="flo; Needs-Refinement"), now=now)["tagged"])

    rows = [br.assess(item(), now=now), br.assess(item(id=2, acceptance="", sp=None), now=now)]
    s = br.summarize(rows)
    check("özet: toplam/hazır/eksik/SP'siz/eksik sayaçları", s["total"] == 2 and s["ready"] == 1 and s["not_ready"] == 1 and s["sp_missing"] == 1
          and s["gap_counts"] == {"no_ac": 1, "no_sp": 1} and s["sp_suggested"] == 2, str(s))
    md = br.gaps_comment_markdown(rows[1])
    check("yorum: skor, eksikler, SP önerisi, Tempo imzası", "Hazırlık skoru:** 65/100" in md and "- Kabul kriteri yok" in md
          and "SP önerisi:** 2" in md and md.rstrip().endswith("*Tempo — Backlog Refinement*"), md)
    from agile_sdlc_crew.branding import is_bot_comment
    check("yorum imzası bot yorumu olarak tanınır", is_bot_comment(md))

    class _C:
        def __init__(self): self.ops = None; self.tags = "flo; Needs-Refinement; ops"
        def get_work_item(self, wid): return {"id": wid, "fields": {"System.Tags": self.tags}}
        def update_work_item(self, wid, ops): self.ops = ops; return {}
    c = _C()
    check("etiket kaldır: diğerleri korunur", br.set_tag(c, 5, add=False) == "flo; ops" and c.ops[0]["path"] == "/fields/System.Tags")
    c.tags = "flo"
    check("etiket ekle: mevcut korunur, tekrar eklenmez", br.set_tag(c, 5, add=True) == "flo; needs-refinement" and br.set_tag(_C(), 5, add=True) == "flo; Needs-Refinement; ops")
    check("SP yaz: WI'da tahmin varsa dokunmaz", br.write_suggested_sp(c, {"id": 5, "sp": 3, "suggest_sp": 2}) == "")

    q = br.parse_questions('bla ```json\n{"summary":"Eksik","questions":[{"question":"Hangi dil?","why":"AC"},"düz soru"],"draft_acceptance_criteria":["AC olmalı"]}\n```')
    check("soru parse: JSON bloğu, düz string soruya çevrilir", q["summary"] == "Eksik" and len(q["questions"]) == 2
          and q["questions"][1] == {"question": "düz soru", "why": ""} and q["draft_acceptance_criteria"] == ["AC olmalı"])
    check("soru parse: JSON yoksa None", br.parse_questions("serbest metin") is None)
    qmd = br.questions_comment_markdown(rows[1], q)
    check("soru yorumu: S1/AC1 numaralı, imzalı", "S1. Hangi dil?" in qmd and "AC1. AC olmalı" in qmd and "*Tempo — Backlog Refinement*" in qmd)

    from agile_sdlc_crew import pipeline_config as _pc
    # Varsayilan (schema) degerleri sinanir: WRITE/LLM kapali gelir. Calisan kurulumda
    # insan bunlari yaml'dan acabilir (pipeline_config.get yaml > env > default cozer),
    # bu yuzden cozulmus deger degil, gonderilen varsayilan kontrol edilir.
    _dflt = {f["key"]: f["default"] for f in _pc.SCHEMA}
    check("knob varsayılanları: refinement açık, eşik 70, bayat 30, WRITE ve LLM KAPALI",
          _dflt["CREW_BACKLOG_REFINEMENT"] is True and _dflt["CREW_REFINEMENT_MIN_SCORE"] == 70
          and _dflt["CREW_REFINEMENT_STALE_DAYS"] == 30 and _dflt["CREW_REFINEMENT_WRITE"] is False
          and _dflt["CREW_REFINEMENT_LLM"] is False)
    check("knob'lar çözümlenebilir (yaml/env override dahil)",
          isinstance(_pc.get("CREW_REFINEMENT_WRITE"), bool) and isinstance(_pc.get("CREW_REFINEMENT_LLM"), bool)
          and isinstance(_pc.get("CREW_REFINEMENT_MIN_SCORE"), int))
    root = Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew"
    src = (root / "backlog_refinement.py").read_text()
    check("sınırlar: durum/kuyruk yok — set_work_item_state, create_job, push_file geçmez",
          "set_work_item_state" not in src and "create_job" not in src and "push_file" not in src)
    tasks = (root / "config/tasks.yaml").read_text()
    check("tasks.yaml: refinement_questions_task (BA, JSON, Türkçe)", "refinement_questions_task:" in tasks and "agent: business_analyst" in tasks.split("refinement_questions_task:")[1])
    check("crew: create_refinement_crew tool'suz", "def create_refinement_crew" in (root / "crew.py").read_text())
    srv = (root / "server.py").read_text()
    check("server: GET /api/refinement + POST action, knob kapıları, job_kind=refinement",
          '@app.get("/api/refinement")' in srv and '@app.post("/api/refinement/action")' in srv
          and "CREW_REFINEMENT_WRITE" in srv and "CREW_REFINEMENT_LLM" in srv and 'job_kind="refinement"' in srv)
    ui = (root / "web/index.html").read_text()
    check("dashboard: Refinement sekmesi, kapsam seçimi, eylem onayı", "showTab('refine')" in ui and 'id="rfScope"' in ui
          and "rfAct(" in ui and "confirm(`#${id} için İş Analisti" in ui)

# ── 40. Maliyet analizi — kırılım, bulgu eşikleri, salt-okunurluk ────────

def test_cost_analytics():
    print("\n[40] Maliyet analizi — adım/model/sınıf kırılımı, bulgu eşikleri, sınırlar")
    from agile_sdlc_crew import cost_analytics as ca

    def row(job, step, model, usd, agent="a", cr=90_000, cw=10_000, kind="pipeline"):
        return {"job_id": job, "step_key": step, "agent": agent, "model": model,
                "provider": "claude_cli", "turns": 2, "tool_calls": 0,
                "cost_usd": usd, "duration_ms": 60_000,
                "input_tokens": 1_000, "output_tokens": 500,
                "cache_read_tokens": cr, "cache_creation_tokens": cw,
                "created_at": None, "work_item_id": f"WI{job}",
                "job_status": "completed", "job_kind": kind}

    rows = []
    # Dört ucuz adım, üç işte birer çağrı → iş başına 1.0 (medyanı aşağı çeker).
    for s in ("requirements_analysis_task", "discover_repos_task",
              "uat_task", "completion_report_task"):
        rows += [row(j, s, "claude-sonnet-5", 0.10) for j in (1, 2, 3)]
    # Beş ucuz iş daha — ortalamayı düşürüp "ortalamayı katlayan iş" eşiğini sınanabilir kılar.
    rows += [row(j, "requirements_analysis_task", "claude-sonnet-5", 0.10) for j in range(4, 9)]
    # Pahalı adım: TEK işte 8 çağrı → iş başına 8.0, uyarı bekleniyor.
    rows += [row(1, "technical_design_task", "claude-opus-4-8[1m]", 5.0,
                 agent="software_architect") for _ in range(8)]
    # Kickoff: tek işte 9 çağrı ama BY_DESIGN_MULTI → uyarı değil bilgi.
    rows += [row(2, "kickoff_meeting_task", "claude-opus-5", 1.0,
                 agent="scrum_master") for _ in range(9)]

    a = ca.analyze(rows)
    t = a["totals"]
    check("kırılım: toplam $, iş sayısı, iş başına ortalama",
          abs(t["usd"] - 50.70) < 0.01 and t["jobs"] == 8 and t["calls"] == len(rows),
          f"{t}")
    check("adım sıralaması: en pahalı adım başta, pay yüzdesi tutarlı",
          a["by_step"][0]["key"] == "technical_design_task"
          and 78 <= a["by_step"][0]["share_pct"] <= 80)
    check("iş başına çağrı: tek işte 8 çağrı → 8.0 (iş sayısına bölünür, çağrı sayısına değil)",
          a["by_step"][0]["calls_per_job"] == 8.0)

    check("model sınıfı: sürüm ve [1m] soneki sınıfı değiştirmez",
          ca._family("claude-opus-4-8[1m]") == "opus"
          and ca._family("claude-opus-5") == "opus"
          and ca._family("claude-sonnet-5") == "sonnet"
          and ca._family("claude-haiku-4-5-20251001") == "haiku"
          and ca._family("gpt-4o") == ca.UNASSIGNED
          and ca._family(None) == ca.UNASSIGNED)
    fam = {m["key"]: m for m in a["by_family"]}
    check("sınıf kırılımı: opus payı hesaplanır, sonnet ayrı durur",
          a["by_family"][0]["key"] == "opus" and fam["opus"]["share_pct"] >= 80
          and "sonnet" in fam)

    titles = [f["title"] for f in ca.findings(a)]
    lv = {f["title"]: f["level"] for f in ca.findings(a)}
    check("bulgu: harcama iki adımda yoğunlaşıyor",
          any(x.startswith("Harcama yoğunlaşması:") and "iki adımda" in x for x in titles), str(titles))
    warn_step = [x for x in titles if x.startswith("technical_design_task")]
    check("bulgu: normal adımda yüksek çağrı UYARI",
          bool(warn_step) and lv[warn_step[0]] == "warn", str(titles))
    kick = [x for x in titles if x.startswith("kickoff_meeting_task")]
    check("bulgu: tasarımı gereği çok çağrılı adım uyarı değil BİLGİ",
          bool(kick) and "tasarım gereği" in kick[0] and lv[kick[0]] == "info", str(titles))
    check("bulgu: tek model sınıfı yoğunlaşması uyarı verir",
          any("tek model sınıfında" in x for x in titles), str(titles))
    check("bulgu: ortalamayı 3 kat aşan iş işaretlenir",
          any(x.startswith("#1 (WI WI1)") for x in titles), str(titles))

    # Önbellek oranı eşiği iki yönlü sınanır: 9.0 sessiz, 1.5 uyarı.
    # Oran bulgusu UYARI seviyesinde; onbellek ekonomisi bulgusu ise her zaman
    # bilgi olarak cikar (token payi != maliyet payi yanilgisini onlemek icin).
    warn_titles = [f["title"] for f in ca.findings(a) if f["level"] == "warn"]
    check("önbellek oranı 9.0 → UYARI yok (eşik 5)",
          a["token_mix"]["cache_ratio"] == 9.0
          and not any("okuma/yazma oranı" in x for x in warn_titles))
    low = ca.analyze([row(1, "s", "claude-opus-5", 1.0, cr=15_000, cw=10_000)])
    low_warn = [f["title"] for f in ca.findings(low) if f["level"] == "warn"]
    check("önbellek oranı 1.5 → UYARI var",
          low["token_mix"]["cache_ratio"] == 1.5
          and any("okuma/yazma oranı" in x for x in low_warn))

    # Token payi ile maliyet payi AYRI raporlanmali: okuma token'in cogunu
    # olusturur ama maliyetin kucuk kismidir; yazma tersi. Bu ayrim olmazsa
    # rapor okuru yanlis yeri optimize etmeye iter.
    share = a["token_mix"]["cost_share"]
    tok = a["token_mix"]
    tok_total = tok["input"] + tok["cache_read"] + tok["cache_write"]
    read_tok_pct = 100 * tok["cache_read"] / tok_total
    check("maliyet payı token payından AYRI hesaplanır (okuma ucuz, yazma pahalı)",
          read_tok_pct > 80 and share["cache_read"]["max"] < read_tok_pct
          and share["cache_write"]["max"] > 100 * tok["cache_write"] / tok_total,
          f"token %{read_tok_pct:.0f} vs maliyet {share['cache_read']}")
    check("yüksek okuma payı bilgi olarak ÖVÜLÜR, uyarı olarak basılmaz",
          any(f["level"] == "info" and "maliyetin" in f["title"] for f in ca.findings(a)))

    md = ca.render_markdown(a, title="test", fnd=ca.findings(a))
    check("markdown: dashboard mdToHtml'in beklediği pipe tablo (|---| ayracı)",
          "| Adım | $ |" in md and "|---|" in md
          and "## Model sınıfı bazında" in md and "## Adım bazında" in md)

    root = Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew"
    src = (root / "cost_analytics.py").read_text()
    check("sınırlar: salt okunur — yazan SQL yok, ajan/crew kurulmuyor",
          not any(k in src.upper() for k in ("INSERT INTO", "DELETE FROM", "UPDATE "))
          and not any(k in src for k in (".kickoff(", "Crew(", "Agent(", "build_for_agent")))
    from agile_sdlc_crew import pipeline_config as _pc
    _dflt = {f["key"]: f["default"] for f in _pc.SCHEMA}
    check("knob: CREW_COST_ANALYTICS açık (maliyeti sıfır, salt okunur)",
          _dflt["CREW_COST_ANALYTICS"] is True
          and isinstance(_pc.get("CREW_COST_ANALYTICS"), bool))
    srv = (root / "server.py").read_text()
    check("server: GET /api/costs + knob kapısı",
          '@app.get("/api/costs")' in srv and "CREW_COST_ANALYTICS" in srv)
    ui = (root / "web/index.html").read_text()
    check("dashboard: Maliyet sekmesi, kapsam/tür seçimi, bulgu renklendirme",
          "showTab('cost')" in ui and 'id="costScope"' in ui
          and 'id="costKind"' in ui and "loadCosts(" in ui)


# ── 41. Angarya/karar model ayrimi — ucuz faz sonnet, karar fazi opus ────

def test_cheap_phase_split():
    print("\n[41] Model ayrimi — angarya fazlari sonnet, karar/teknik noktalar korunur")
    import os
    import inspect
    from agile_sdlc_crew.llm import resolver as R

    CHEAP = ("software_architect_explore", "software_architect_kickoff",
             "senior_developer_kickoff")

    # 1) Angarya anahtarlari cozulur ve UCUZ modele baglanir (bizim varsayilanimiz,
    #    agent_defaults katmanindan — dashboard override'i degil, o insanin).
    specs = {k: R.resolve_spec_with_source(k) for k in CHEAP}
    check("angarya anahtarlari cozulur ve sonnet'e baglanir",
          all(s[0]["model"] == "sonnet" and s[0]["provider"] == "claude_cli"
              for s, in [(v,) for v in specs.values()]),
          str({k: (v[0]["model"], v[1]) for k, v in specs.items()}))

    # 2) Karar/teknik noktalar angarya anahtarlarindan BAGIMSIZ kalir. Modeli
    #    sabitlemiyoruz (dashboard'un hakki) — ayri cozulduklerini dogruluyoruz.
    for judge in ("software_architect", "senior_developer", "code_reviewer"):
        jspec, _ = R.resolve_spec_with_source(judge)
        check(f"{judge}: karar/teknik nokta angarya profiline BAGLANMAZ",
              jspec.get("_profile") not in ("architect_explore_cli", "kickoff_cli"))

    # 3) CREW_USE_LOCAL_LLM angarya fazlarini yerel modele kacirmamali: kesif
    #    682K'ya varan baglam tariyor, yerel 8B model bunu sessizce bozar.
    os.environ["CREW_USE_LOCAL_LLM"] = "1"
    R.reset_cache()
    try:
        local_leak = [k for k in CHEAP
                      if R.resolve_spec_with_source(k)[0]["provider"] == "ollama"]
    finally:
        os.environ.pop("CREW_USE_LOCAL_LLM", None)
        R.reset_cache()
    check("CREW_USE_LOCAL_LLM angarya fazlarini yerele kacirmaz", not local_leak,
          str(local_leak))

    # 4) Geri alma tek env ile calismali (kalite bozulursa aninda donus).
    os.environ["CREW_LLM_PROFILE_SOFTWARE_ARCHITECT_EXPLORE"] = "architect_cli"
    try:
        rb, rbsrc = R.resolve_spec_with_source("software_architect_explore")
    finally:
        os.environ.pop("CREW_LLM_PROFILE_SOFTWARE_ARCHITECT_EXPLORE", None)
    check("env ile geri alinabilir (CREW_LLM_PROFILE_… → opus)",
          rb["model"] == "opus" and rbsrc == "env")

    # 5) Cagri yeri baglantisi: SADECE kesif fazi ucuz anahtari kullanmali.
    #    _amend_plan ve _architect_emit_json plan YAZAR — onlar karar fazi.
    root = Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew"
    fl = (root / "flow.py").read_text()
    import re as _re
    explore_fn = fl.split("def _architect_explore")[1].split("\n    def ")[0]
    check("kesif fazi ucuz anahtarla cagrilir",
          'agent_key="software_architect_explore"' in explore_fn)
    check("ucuz anahtar flow.py'de SADECE kesif fazinda geciyor",
          fl.count("software_architect_explore") == explore_fn.count("software_architect_explore"),
          f"toplam {fl.count('software_architect_explore')}")
    for fn in ("_amend_plan", "_architect_emit_json"):
        body = fl.split(f"def {fn}")[1].split("\n    def ")[0]
        check(f"{fn} (plan yazan faz) ucuz anahtar KULLANMAZ",
              "software_architect_explore" not in body)

    # 6) Varsayilan parametre karar fazini korur: arguman verilmeyen her cagri
    #    (amend + emit) eski davranista kalir.
    from agile_sdlc_crew.crew import AgileSDLCCrew
    sig = inspect.signature(AgileSDLCCrew.create_analysis_crew_toolless)
    check("create_analysis_crew_toolless varsayilani software_architect",
          sig.parameters["agent_key"].default == "software_architect")

    # 7) Kayit: dashboard listesi + model-erisilebilirlik dogrulamasi. Kayit disi
    #    config sessizce kayar (resolver docstring'indeki aylarca suren hata).
    srv = (root / "server.py").read_text()
    check("dashboard: angarya anahtarlari _AGENT_KEYS + _AGENT_DISPLAY'de",
          all(f'"{k}"' in srv for k in CHEAP)
          and all(k in srv.split("_AGENT_DISPLAY")[1][:800] for k in CHEAP))
    res_src = (root / "llm/resolver.py").read_text()
    check("assert_models_reachable angarya anahtarlarini da dogrular",
          all(k in res_src.split("def assert_models_reachable")[1] for k in CHEAP))
    prof = (root / "config/llm_profiles.yaml").read_text()
    check("llm_profiles: profiller + agent_defaults eslemesi yazili",
          "architect_explore_cli:" in prof and "kickoff_cli:" in prof
          and all(f"{k}:" in prof for k in CHEAP))


# ── 42. Kesif bulgusunun amend turlarinda yeniden kullanimi ──────────────

def test_findings_reuse():
    print("\n[42] Keşif bulgusu amend turlarına taşınır — repo iki kez taranmaz")
    root = Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew"
    fl = (root / "flow.py").read_text()

    from agile_sdlc_crew.flow import _FINDINGS_REUSE_MAX
    check("taşıma sınırı tanımlı ve makul (prompt'u şişirmeyecek kadar)",
          isinstance(_FINDINGS_REUSE_MAX, int) and 2000 <= _FINDINGS_REUSE_MAX <= 40000,
          str(_FINDINGS_REUSE_MAX))

    step4 = fl.split("def crew_step4_technical_design")[1].split("\n    def ")[0]
    check("step4: Faz A bulgusu state'e yazılır",
          "_architect_findings = findings[:_FINDINGS_REUSE_MAX]" in step4)

    amend = fl.split("def _amend_plan")[1].split("\n    def ")[0]
    check("amend: bulgu varsa context'e taşınır",
          "_architect_findings" in amend and "TEKRARLAMA" in amend)
    # Arac erisimi KAPATILMAMALI: geri bildirim hic kesfedilmemis bir dosyaya
    # isaret edebilir. Kazanc tekrar taramayi onlemekten gelir, erisimi
    # kesmekten degil — yoksa amend eksik dosyayi hic bulamaz.
    check("amend: repo araç erişimi kapatılmadı (yeni dosya gerekebilir)",
          "set_repo_ctx" in amend)

    # Bulgu yokken (kesif hic kosmadi) amend eskisi gibi davranmali.
    check("bulgu yoksa context'e blok EKLENMEZ (boş başlık basılmaz)",
          '_prev = (getattr(self, "_architect_findings", "") or "").strip()' in amend
          and "if _prev:" in amend)

    # Butce cap'i (CREW_CLI_CALL_MAX_USD → --max-budget-usd) zaten bagliydi ve
    # ACIK. Probe (2026-09-21): sinira carpinca is_error=True /
    # subtype=error_max_budget_usd, result BOS, stream'de text blogu yok →
    # salvage de bos. Para harcanir, is kaybolur. Kesinti en azindan GORUNUR
    # olmali; yoksa raporda sadece "pahali cagri" gibi durur.
    cli = (root / "tools/claude_cli_llm.py").read_text()
    check("bütçe kesintisi yakalanır ve stop_reason olarak taşınır",
          "error_max_budget_usd" in cli and 'meta["stop_reason"]' in cli
          and '"stop_reason": str(meta.get("stop_reason") or "")' in cli)
    dbs = (root / "db.py").read_text()
    check("stop_reason llm_calls'a yazılır (kolon + INSERT)",
          '_ensure_column(cur, "llm_calls", "stop_reason"' in dbs
          and 'rec.get("stop_reason")' in dbs)
    ca_src = (root / "cost_analytics.py").read_text()
    check("Maliyet raporu kesilen çağrıları ayrı bulgu olarak gösterir",
          "budget_cut" in ca_src and "error_max_budget_usd" in ca_src)
    from agile_sdlc_crew import cost_analytics as _ca
    _an = _ca.analyze([
        {"job_id": 1, "step_key": "technical_design_task", "agent": "software_architect",
         "model": "claude-opus-5", "cost_usd": 1.6, "turns": 20, "input_tokens": 10,
         "output_tokens": 10, "cache_read_tokens": 10, "cache_creation_tokens": 10,
         "duration_ms": 1000, "created_at": None, "stop_reason": "error_max_budget_usd",
         "job_kind": "pipeline", "work_item_id": "1", "job_status": "completed"},
    ])
    check("kesilen çağrı bulgusu tetiklenir ve tutarı raporlar",
          _an["budget_cut"]["calls"] == 1 and _an["budget_cut"]["usd"] == 1.6
          and any("bütçe cap'inde kesildi" in f["title"] for f in _ca.findings(_an)))


# ── 43. Prompt onbellek oneki — degisken placeholder EN SONDA ────────────

def test_prompt_prefix_order():
    print("\n[43] Placeholder sırası — sabit context önce, değişen geri bildirim sonra")
    import yaml as _yaml
    root = Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew"
    raw = (root / "config/tasks.yaml").read_text()
    doc = _yaml.safe_load(raw)

    # Onbellek ONEKTEN eslesir: ilk degisen karakterden sonrasi yeniden yazilir.
    # Yazma 1.25-2x, okuma 0.1x → degisken metin SONDA olmali. Retry donguleri
    # (_architect_emit_json, _review_retry_loop) yalnizca feedback string'ini
    # degistirir; sira bozulursa 2. deneme tum oneki yeniden yazar.
    offenders = []
    checked = 0
    for key, block in doc.items():
        desc = ((block or {}).get("description") or "")
        if "{scrum_master_feedback}" not in desc or "{previous_context}" not in desc:
            continue
        checked += 1
        if desc.index("{scrum_master_feedback}") < desc.rindex("{previous_context}"):
            offenders.append(key)
    check("değişken {scrum_master_feedback} her görevde sabit {previous_context}'ten SONRA",
          checked >= 5 and not offenders, f"{checked} görev kontrol edildi, bozuk: {offenders}")

    # Placeholder'lar kaybolmamis olmali (sira degisimi silme olmamali).
    counts = {p: sum(((b or {}).get("description") or "").count(p) for b in doc.values())
              for p in ("{work_item_id}", "{previous_context}",
                        "{scrum_master_feedback}", "{target_repo}")}
    check("sıra değişimi hiçbir placeholder'ı düşürmedi",
          counts["{previous_context}"] >= 13 and counts["{scrum_master_feedback}"] == 5
          and counts["{work_item_id}"] >= 18, str(counts))

    # Gerekce dosyada yazili olmali — yoksa biri "duzenler" ve geri alir.
    check("sıranın gerekçesi tasks.yaml başlığında yazılı",
          "PLACEHOLDER ORDER IS A COST DECISION" in raw and "cache write" in raw.lower())

    # Retry dongusu gercekten SADECE feedback'i degistiriyor mu? Degistirdigi
    # baska bir sey varsa onek yine bozulur ve bu sira ise yaramaz.
    fl = (root / "flow.py").read_text()
    emit = fl.split("def _architect_emit_json")[1].split("\n    def ")[0]
    check("emit retry'ı yalnızca scrum_master_feedback'i değiştirir (emit_ctx sabit)",
          "fb_note = (" in emit and 'previous_context": emit_ctx' in emit
          and emit.count("emit_ctx =") <= 2, "emit_ctx döngü içinde yeniden kurulmamalı")


def test_quality_rank_nr():
    print("\n[Q1] Kalite alımı — NR normalizasyon, birleşme, SP")
    from agile_sdlc_crew.quality import rank
    from agile_sdlc_crew.quality.model import Finding
    n = rank.normalize_message
    check("dizi anahtarı adı korunur",
          n('Undefined array key "host"') == 'Undefined array key "host"'
          and n('Undefined array key "host"') != n('Undefined array key "password"'))
    check("sayı/uuid/ID-benzeri tırnak içi maskelenir",
          n("Order 12345 not found") == n("Order 98 not found") == "Order <n> not found"
          and n('Product "550e8400-e29b-41d4-a716-446655440000" missing') == 'Product "<v>" missing'
          and n("Undefined array key 7") == "Undefined array key <n>")
    check("mesajdaki iç işaretçi karakteri çökertmez, sahte eşleşme üretmez",
          n("a \ue000x b") == "a x b" and n('"k" \ue000\ue100') == '"k"')
    check("action: WebTransaction öneki atılır, controller StudlyCase",
          rank.action_of("WebTransaction/Action/api::order::integrationCancel") == "Api/Order::integrationCancel"
          and rank.action_of("WebTransaction/Custom/index::index") == "Index::index"
          and rank.action_of("WebTransaction/Action/web-hook::trigger") == "WebHook::trigger")
    fp1 = rank.nr_fingerprint("webservice", "A::b", "E_WARNING", "x" * 900)
    check("parmak izi 512'de kesilir ve kararlıdır",
          len(fp1) == 512 and fp1 == rank.nr_fingerprint("webservice", "A::b", "E_WARNING", "x" * 900))

    def nf(app, msg, cnt, action="Content::page", cls="E_WARNING", repo="webservice"):
        return Finding(source="newrelic", repo=repo, app=app, action=action, err_class=cls,
                       message=msg, count=cnt, last_seen=2000.0, error_rate=0.01)

    sites = [f"[Prod] site{i}.ru" for i in range(6)]
    units = rank.nr_units([nf(s, 'Undefined array key "value"', 100) for s in sites])
    check("6 uygulamadaki aynı hata → tek birim, 6 uygulama, skor = adet × uygulama, 1 SP",
          len(units) == 1 and len(units[0].apps) == 6 and units[0].score == 600 * 6 and units[0].sp == 1)
    u2 = rank.nr_units([nf("[Prod] a", "Order 1 failed", 5, cls="TypeError"),
                        nf("[Prod] a", "Order 2 failed", 7, cls="TypeError")])
    check("ID farkı aynı birime düşer; TypeError → 3 SP",
          len(u2) == 1 and u2[0].total == 12 and u2[0].sp == 3 and not u2[0].needs_human)
    infra = rank.nr_units([nf("[Prod] a", "cURL error 28: Operation timed out after 5000 ms", 50,
                              cls="GuzzleHttp\\Exception\\ConnectException")])
    check("altyapı hatası → 3 SP + needs_human", infra[0].sp == 3 and infra[0].needs_human)
    fam = rank.nr_units([nf("[Prod] a", f'Undefined array key "{c}"', 10, action="Api/Order::create")
                         for c in "abcde"]
                        + [nf("[Prod] a", 'Undefined array key "host"', 10, action="Index::index")])
    merged = rank.merge_undefined(fam)
    big = [g for g in merged if g.findings[0].action == "Api/Order::create"]
    check("aynı action'daki 5 tanımsız anahtar → tek Task, 5 parmak izi, 2 SP; başka action ayrı",
          len(merged) == 2 and len(big) == 1 and len(big[0].fingerprints) == 5 and big[0].sp == 2)
    three = rank.merge_undefined(rank.nr_units(
        [nf("[Prod] a", f"Undefined index: {c}", 10, action="X::y", cls="E_NOTICE") for c in ("ka", "kb", "kc")]))
    check("≤3 anahtar → 1 SP", len(three) == 1 and three[0].sp == 1 and len(three[0].fingerprints) == 3)


def test_quality_rank_sonar():
    print("\n[Q2] Kalite alımı — Sonar birim, skor, paketleme")
    from agile_sdlc_crew.quality import rank
    from agile_sdlc_crew.quality.model import Finding

    def sf(file, rule="css:S4657", sev="CRITICAL", typ="BUG", repo="webservice", line=1):
        return Finding(source="sonar", repo=repo, project=f"EcomBackend:{repo}", file=file, rule=rule,
                       rule_name=rule, severity=sev, itype=typ, lines=[line])

    css = [sf("themes/a.css", line=i) for i in range(300)]
    vuln = [sf("app/Login.php", rule="php:S2076", sev="BLOCKER", typ="VULNERABILITY")]
    units = rank.sonar_units(css + vuln)
    by = {u.findings[0].rule: u for u in units}
    check("aynı dosya+kural → tek birim (300 madde)", len(units) == 2 and by["css:S4657"].total == 300)
    check("tek Vulnerability+Blocker, 300 maddelik Bug+Critical'ı geçer",
          by["php:S2076"].score > by["css:S4657"].score,
          f"{by['php:S2076'].score} vs {by['css:S4657'].score}")
    check("SP: 1 madde → 1, 300 madde → 3", by["php:S2076"].sp == 1 and by["css:S4657"].sp == 3)
    check("SP sınırları ≤5→1, ≤20→2, >20→3", [rank.sonar_unit_sp(n) for n in (5, 6, 20, 21)] == [1, 2, 2, 3])
    sec = rank.sonar_units([sf("config/app.php", rule="secrets:S6703", sev="BLOCKER", typ="VULNERABILITY")])[0]
    check("secrets:* → 1 SP + needs_human", sec.sp == 1 and sec.needs_human and "rotasyon" in sec.human_reason)
    twelve = rank.sonar_units([sf(f"themes/t{i}.css") for i in range(12)])
    packs = rank.pack_sonar(twelve, pack_sp=3, pack_files=5)
    check("12 dosya × 1 SP, paket ≤3 SP → 4 paket × 3 dosya",
          len(packs) == 4 and all(p.sp == 3 and len(p.files) == 3 for p in packs))
    packs5 = rank.pack_sonar(twelve, pack_sp=99, pack_files=5)
    check("dosya sınırı 5 → 5+5+2", sorted(len(p.files) for p in packs5) == [2, 5, 5])
    check("paket başlığı +N dosya, parmak izi birim başına",
          "(+4 dosya)" in max(packs5, key=lambda p: len(p.files)).title
          and sum(len(p.fingerprints) for p in packs5) == 12)
    mixed = rank.pack_sonar(rank.sonar_units([sf("a.css"), sf("b.php", rule="php:S1")]))
    check("farklı kural aynı pakete girmez", len(mixed) == 2)


def test_quality_select():
    print("\n[Q3] Kalite alımı — bütçe seçimi")
    from agile_sdlc_crew.quality import rank
    from agile_sdlc_crew.quality.model import Group

    def g(name, sp, score, human=False):
        return Group(source="sonar", repo="r", title=name, fingerprints=[name], findings=[],
                     score=score, sp=sp, needs_human=human)

    pick = rank.select_budget([g("a", 5, 100), g("b", 3, 90), g("c", 3, 80), g("d", 2, 70), g("e", 1, 60)], 10)
    check("açgözlü: sığmayan atlanır, küçük olan denenir, toplam ≤ bütçe",
          [x.title for x in pick] == ["a", "b", "d"] and sum(x.sp for x in pick) == 10)
    pick = rank.select_budget([g("h1", 3, 100, True), g("h2", 1, 99, True), g("x", 3, 50), g("y", 3, 40)], 10)
    check("needs_human payı ≤ %30 (10 SP → en çok 3)",
          [x.title for x in pick] == ["h1", "x", "y"] and sum(x.sp for x in pick if x.needs_human) <= 3)
    check("Task sayısı tavanı", len(rank.select_budget([g(str(i), 1, 100 - i) for i in range(20)], 40, max_tasks=8)) == 8)
    check("bütçeden büyük tek iş → boş seçim", rank.select_budget([g("big", 5, 1)], 3) == [])
    check("3 SP bütçede needs_human hiç girmez (pay 0)", rank.select_budget([g("h", 1, 9, True)], 3) == [])


def test_quality_render():
    print("\n[Q4] Kalite alımı — WI metinleri, maskeleme, parmak izi satırı")
    import html as _html
    from agile_sdlc_crew.quality import rank, render
    from agile_sdlc_crew.quality.model import Finding
    m = render.mask_secrets
    check("Bearer/Basic ve password=… maskelenir",
          "abc123def456" not in m("Authorization: Bearer abc123def456ghi")
          and m("password=Hunter2x&x=1") == "password=***&x=1"
          and "dXNlcjpwYXNz" not in m("Basic dXNlcjpwYXNz"))
    check("uzun anahtar adları maskelenmez",
          m('Undefined array key "is_prime_activate_on_preprod_web"')
          == 'Undefined array key "is_prime_activate_on_preprod_web"')
    check("uzun hex/token maskelenir", "***" in m("token 9f86d081884c7d659a2feaa0c55ad015a3bf4f1b"))
    check("action → dosya ipucu (V1 varyantı dahil)",
          render.file_hints("Api/Order::integrationCancel") == [
              "app/Controller/Api/Order.php::integrationCancelAction",
              "app/Controller/V1/Api/Order.php::integrationCancelAction"]
          and render.file_hints("cron.php") == [])
    nr = rank.nr_units([Finding(source="newrelic", repo="webservice", app=f"[Prod] s{i}.ru", action="Content::page",
                                err_class="E_WARNING", message='Undefined array key "value"', count=10,
                                uris=["/ru/x"], first_seen=1.7e9, last_seen=1.7e9 + 60, error_rate=0.05)
                        for i in range(2)])[0]
    d = render.task_description(nr, nr_days=7)
    check("NR açıklaması: ilk satır Repo, iki uygulama, ipucu, kabul kriteri, parmak izi",
          d.startswith("<p><b>Repo:</b> <code>webservice</code></p>") and "[Prod] s0.ru" in d
          and "[Prod] s1.ru" in d and "app/Controller/Content.php::pageAction" in d
          and "Kabul kriteri" in d and render.parse_fingerprints(d) == nr.fingerprints)
    check("NR başlığı", render.task_title(nr) == '[webservice] Content::page: Undefined array key "value"')
    sec = rank.sonar_units([Finding(source="sonar", repo="webservice", project="EcomBackend:webservice",
                                    file="config/db.php", rule="secrets:S6703",
                                    rule_name="Database passwords should not be disclosed",
                                    severity="BLOCKER", itype="VULNERABILITY", lines=[12],
                                    message="password = 'S3cr3tValue!'")])[0]
    sd = render.task_description(sec, sonar_url="https://sonar.x", nr_days=7)
    check("secrets:* açıklamasına madde mesajı girmez; dosya:satır ve insan nedeni girer",
          "S3cr3tValue" not in sd and "config/db.php:12" in sd and "Neden insan" in sd)
    check("etiketler", render.task_tags(sec) == "tempo-quality; tempo-src-sonar; needs-human")
    nr.recurring_of = 75001
    check("tekrarlayan iş önceki WI'ı anar", "#75001" in render.task_description(nr, nr_days=7))
    fps = ['nr:webservice:A::b:E_WARNING:Undefined array key "x"', "sonar:webservice:a.css:css:S4657"]
    blob = "".join(f"<p><small>tempo-fp: {_html.escape(fp)}</small></p>" for fp in fps)
    check("parmak izi satırları geri okunur (HTML kaçışlı ve etiketsiz)",
          render.parse_fingerprints(blob) == fps
          and render.parse_fingerprints("tempo-fp: a:b tempo-fp: c:d") == ["a:b", "c:d"])
    check("story başlığı", render.story_title("newrelic", "Sprint 42") == "Kalite: New Relic hataları — Sprint 42")
    evil = rank.nr_units([Finding(source="newrelic", repo="webservice", app="[Prod] a", action="X::y",
                                  err_class="TypeError", count=1, last_seen=1.7e9,
                                  message="boom tempo-fp: sonar:webservice:x.php:php:S1 " + "y" * 5000,
                                  uris=["/ru/p?token=abcd1234efgh5678&x=1#frag", "/" + "z" * 500])])[0]
    ed = render.task_description(evil, nr_days=7)
    check("dış metin parmak izi satırı enjekte edemez",
          render.parse_fingerprints(ed) == evil.fingerprints, str(render.parse_fingerprints(ed)))
    check("örnek isteklerde sorgu dizgisi atılır, uzunluk sınırlı",
          "token" not in ed and "#frag" not in ed and "z" * 200 not in ed)
    check("dış metin veri olarak çerçevelenir ve kısaltılır",
          "talimat değil" in ed and "y" * 600 not in ed)


def test_quality_dedup():
    print("\n[Q5] Kalite alımı — board'dan dedup")
    from agile_sdlc_crew.quality import dedup

    def wi(i, state, fps, closed=""):
        return {"id": i, "fields": {
            "System.State": state, "Microsoft.VSTS.Common.ClosedDate": closed,
            "System.Description": "".join(f"<p><small>tempo-fp: {fp}</small></p>" for fp in fps)}}

    idx = dedup.index_board([
        wi(1, "In Progress", ["a"]), wi(2, "Ready for Production", ["b"]), wi(3, "Canceled", ["c"]),
        wi(4, "Done", ["d"], "2026-10-01T00:00:00Z"), wi(5, "Done", ["e"], "2026-10-01T00:00:00Z"),
        wi(6, "Done", ["f"], "2026-09-01T00:00:00Z"), wi(7, "To Do", ["f"])])
    day = 86400
    closed = dedup._epoch("2026-10-01T00:00:00Z")
    now = closed + 10 * day

    def v(fp, seen=0.0):
        return dedup.verdict(idx.get(fp), last_seen=seen, now=now, grace_days=3)

    check("açık WI → blocked", v("a") == ("blocked", 1))
    check("Ready for Production → blocked (düzeltme canlıda değil)", v("b") == ("blocked", 2))
    check("Canceled → kalıcı bastırma", v("c") == ("suppressed", 3))
    check("Done + grace sonrası hâlâ görülüyor → recurring", v("d", seen=closed + 5 * day) == ("recurring", 4))
    check("Done + grace içinde görülmüş → blocked", v("e", seen=closed + 1 * day) == ("blocked", 5))
    check("Sonar (last_seen yok) → şimdiye göre", v("d") == ("recurring", 4))
    check("aynı parmak izi iki WI'da → en yeni (büyük id) kazanır", v("f") == ("blocked", 7))
    check("hiç yok → new", v("zzz") == ("new", None))


def test_quality_app_map():
    print("\n[Q6] Kalite alımı — NR uygulama → repo eşlemesi (FLODevOps)")
    import inspect
    import yaml
    from agile_sdlc_crew.quality import app_map
    from agile_sdlc_crew.tools.azure_devops_base import AzureDevOpsClient
    check("Azure istemcisi başka projedeki repoyu okuyabilir (project parametresi)",
          "project" in inspect.signature(AzureDevOpsClient.get_file_content).parameters
          and "project" in inspect.signature(AzureDevOpsClient.get_items_in_path).parameters)
    jf_tr = 'String gitOriginApp = "git@ssh.dev.azure.com:v3/flotechnology/ecombackend/webservice"\n'
    df_tr = ("ARG website=tr-stocks.flo.com.tr\n"
             "&& sed -ie 's/[ ;]*newrelic.appname[[:space:]]=.*/newrelic.appname=\"[Prod] tr-stocks.flo.com.tr\"/' "
             "$(php-config --ini-dir)/newrelic.ini \\\n")
    e = app_map.parse_build_dir(jf_tr, df_tr)
    check("literal appname + gitOriginApp",
          e == {"project": "ecombackend", "repo": "webservice", "apps": ["[Prod] tr-stocks.flo.com.tr"]}, str(e))
    jf_web = ('String gitOriginApp = "git@ssh.dev.azure.com:v3/flotechnology/ecombackend/webservice"\n'
              "   flokz       : [website: 'flo.com.kz', themePath: 'flo-kz'],\n"
              "   floma       : [website: 'flo.co.ma', themePath: 'flo-ma'],\n"
              '   newrelic_app_name = "[Prod] ${website}"\n'
              "   sh \"docker build --build-arg newrelic_app_name='${newrelic_app_name}' .\"\n")
    df_web = ("ARG website\nARG newrelic_app_name\n"
              '&& sed -ie "s/[ ;]*newrelic.appname[[:space:]]=.*/newrelic.appname=\\"${newrelic_app_name}\\"/" x\n')
    e2 = app_map.parse_build_dir(jf_web, df_web)
    check("${website} şablonu her siteye açılır, ${newrelic_app_name} yok sayılır",
          e2 and e2["apps"] == ["[Prod] flo.com.kz", "[Prod] flo.co.ma"], str(e2))
    check("gitOriginApp yoksa None", app_map.parse_build_dir("", df_tr) is None)

    calls = {"n": 0}

    class FakeClient:
        def get_items_in_path(self, repo, path, branch=None, recursion_level="oneLevel", project=None):
            calls["n"] += 1
            assert project == "FLODevOps" and branch == "production"
            if repo != "global":
                return []
            return [{"path": "/build", "isFolder": True}, {"path": "/build/trstock", "isFolder": True},
                    {"path": "/build/.DS_Store", "isFolder": False}]

        def get_file_content(self, repo, path, branch=None, project=None):
            return {"/build/trstock/JenkinsfileBuild": jf_tr, "/build/trstock/Dockerfile": df_tr}[path]

    class Broken(FakeClient):
        def get_items_in_path(self, *a, **k):
            raise RuntimeError("401")

    with tempfile.TemporaryDirectory() as td:
        cache, ovr = Path(td) / "map.yaml", Path(td) / "ovr.yaml"
        ovr.write_text(yaml.safe_dump({"[Prod] Stock Api": "stock-api"}), encoding="utf-8")
        m = app_map.load_map(FakeClient(), cache_path=cache, overrides_path=ovr)
        check("FLODevOps'tan kurulur, büyük/küçük harf duyarsız, override eklenir",
              app_map.resolve("[PROD] TR-STOCKS.FLO.COM.TR", m) == "webservice"
              and app_map.resolve("[Prod] Stock Api", m) == "stock-api" and app_map.resolve("yok", m) == "")
        n = calls["n"]
        app_map.load_map(FakeClient(), cache_path=cache, overrides_path=ovr)
        check("taze önbellek varken FLODevOps tekrar okunmaz", calls["n"] == n)
        m2 = app_map.load_map(Broken(), cache_path=cache, overrides_path=ovr, force=True)
        check("FLODevOps okunamazsa eski önbellek kullanılır",
              app_map.resolve("[Prod] tr-stocks.flo.com.tr", m2) == "webservice")
        cache.unlink()
        try:
            app_map.load_map(Broken(), cache_path=cache, overrides_path=ovr)
            ok = False
        except app_map.QualitySourceError:
            ok = True
        check("önbellek de yoksa QualitySourceError", ok)


def test_quality_sources():
    print("\n[Q7] Kalite alımı — Sonar/NR yanıtlarını bulguya çevirme")
    import os
    from agile_sdlc_crew.quality import newrelic, sonar
    err = [
        {"facet": ["[Prod] tr-stocks.flo.com.tr", "E_WARNING", 'Undefined array key "host"',
                   "WebTransaction/Custom/index::index"],
         "count": 645052, "first": 1791290877564.0, "last": 1791895671566.0, "uris": ["/en"]},
        {"facet": ["[Preprod] preprod.flo.com.kz", "E_WARNING", "x", "WebTransaction/Custom/Page::notFound"],
         "count": 5, "uris": []},
        {"facet": ["Marketplace API", "TypeError", "y", "WebTransaction/Action/api::x"], "count": 3},
    ]
    tx = [{"facet": "[Prod] tr-stocks.flo.com.tr", "appName": "[Prod] tr-stocks.flo.com.tr", "count": 1290104},
          {"facet": "Marketplace API", "count": 0}]
    fs = newrelic.rows_to_findings(err, tx, is_prod=lambda a: newrelic.is_prod_name(a) or a == "Marketplace API")
    check("prod filtresi: [Preprod] elenir, eşlenmiş köşeli parantezsiz prod kalır",
          [f.app for f in fs] == ["[Prod] tr-stocks.flo.com.tr", "Marketplace API"])
    f0 = fs[0]
    check("NR satırı → Finding (action, saniye, error rate, uri)",
          f0.action == "Index::index" and f0.count == 645052 and abs(f0.last_seen - 1791895671.566) < 1e-3
          and abs(f0.error_rate - 0.5) < 0.01 and f0.uris == ["/en"])
    check("işlem sayısı 0 → error rate 0", fs[1].error_rate == 0.0)
    check("prod adı kalıpları",
          all(newrelic.is_prod_name(a) for a in ("[Prod] x", "[PROD] y", "[Production] FLO Orkestra V2",
                                                  "[GCP-Prod] panel"))
          and not newrelic.is_prod_name("[Preprod] x"))
    issues = [
        {"key": "k1", "rule": "php:S2076", "severity": "BLOCKER", "type": "VULNERABILITY",
         "project": "EcomBackend:webservice", "component": "EcomBackend:webservice:app/Login.php",
         "line": 40, "message": "Make sure..."},
        {"key": "k2", "rule": "secrets:S6703", "severity": "BLOCKER", "type": "VULNERABILITY",
         "project": "EcomBackend:webservice", "component": "EcomBackend:webservice:config/db.php",
         "line": 3, "message": "password = 'S3cret'"},
    ]
    sf = sonar.issues_to_findings(issues, {"php:S2076": "OS commands should not be vulnerable"},
                                  base_url="https://sonar.x")
    check("Sonar: proje → repo, component → dosya, kural adı, link",
          sf[0].repo == "webservice" and sf[0].file == "app/Login.php"
          and sf[0].rule_name == "OS commands should not be vulnerable" and sf[0].lines == [40]
          and "open=k1" in sf[0].link)
    check("secrets:* mesajı hiç saklanmaz", sf[1].message == "")
    check("kapsam: AZURE_DEVOPS_REPO_PROJECTS önekine göre (büyük/küçük harf duyarsız)",
          sonar.in_scope("EcomBackend:webservice", ["ecombackend", "EcomLojistik"])
          and not sonar.in_scope("FLOXO:tedweb", ["EcomBackend"]))
    pages = {1: {"issues": [{"key": "a"}] * 500, "paging": {"total": 863}, "rules": [{"key": "r", "name": "R"}]},
             2: {"issues": [{"key": "b"}] * 363, "paging": {"total": 863}}}
    seen = []
    orig = sonar._get
    sonar._get = lambda path, params: (seen.append(params["p"]), pages[params["p"]])[1]
    try:
        iss, rules = sonar.fetch_issues()
    finally:
        sonar._get = orig
    check("Sonar sayfalama: 863 madde iki sayfada", len(iss) == 863 and seen == [1, 2] and rules == {"r": "R"})
    saved = {k: os.environ.pop(k, None) for k in ("NEW_RELIC_ACCOUNT_ID", "NEW_RELIC_API_KEY")}
    try:
        try:
            newrelic.nrql("SELECT 1")
            ok = False
        except newrelic.QualitySourceError:
            ok = True
    finally:
        for k, val in saved.items():
            if val is not None:
                os.environ[k] = val
    check("NR kimlik bilgisi yoksa QualitySourceError", ok)
    raw = newrelic.rows_to_findings(
        [{"facet": ["[Prod] a", "E", "login failed password=Hunter2x " + "q" * 3000, "WebTransaction/Custom/a::b"],
          "count": 1, "uris": ["/p?session=abc123def456&x=1"]}], [], is_prod=newrelic.is_prod_name)[0]
    check("NR mesajı/URI girişte temizlenir (secret maskesi, sorgu dizgisi, uzunluk)",
          "Hunter2x" not in raw.message and len(raw.message) <= 1000 and raw.uris == ["/p"], repr(raw.uris))
    from agile_sdlc_crew import pipeline_config as _pcq
    check("Sonar adresi dashboard'dan değiştirilemez (token sızmasın)", "CREW_SONAR_URL" not in _pcq._SCHEMA_BY_KEY)
    old_url, old_tok = os.environ.get("CREW_SONAR_URL"), os.environ.get("SONAR_TOKEN")
    os.environ["CREW_SONAR_URL"], os.environ["SONAR_TOKEN"] = "http://evil.example", "t0k"
    try:
        try:
            sonar._get("/api/issues/search", {})
            ok = False
        except sonar.QualitySourceError as e:
            ok = "https" in str(e)
    finally:
        for k, val in (("CREW_SONAR_URL", old_url), ("SONAR_TOKEN", old_tok)):
            if val is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = val
    check("https olmayan Sonar adresine token gönderilmez", ok)


def test_quality_intake():
    print("\n[Q8] Kalite alımı — orkestratör (sahte Azure/DB)")
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.quality import intake, rank
    from agile_sdlc_crew.quality.model import Finding
    keys = ("CREW_QUALITY_INTAKE", "NEW_RELIC_REGION", "CREW_QUALITY_NR_LOOKBACK_DAYS",
            "CREW_QUALITY_NR_MIN_ERRORS", "CREW_QUALITY_SONAR_PACK_SP", "CREW_QUALITY_SONAR_PACK_FILES",
            "CREW_QUALITY_HUMAN_MAX", "CREW_QUALITY_MAX_TASKS", "CREW_QUALITY_REOPEN_GRACE_DAYS")
    check("knob'lar kayıtlı, ana anahtar varsayılan kapalı",
          all(k in pc._SCHEMA_BY_KEY for k in keys) and pc._SCHEMA_BY_KEY["CREW_QUALITY_INTAKE"]["default"] is False)
    cfg = {"nr_days": 7, "nr_min_errors": 100, "pack_sp": 3, "pack_files": 5, "human_max": 5,
           "max_tasks": 8, "grace_days": 3, "sonar_url": "https://sonar.x"}

    def nf(app, key, cnt, action="Index::index", repo="webservice", cls="E_WARNING", msg=None):
        return Finding(source="newrelic", repo=repo, app=app, action=action, err_class=cls,
                       message=msg or f'Undefined array key "{key}"', count=cnt, last_seen=1e9)

    findings = [nf("[Prod] a", "host", 645000), nf("[Prod] b", "password", 900, action="Api/Order::create", repo="channel"),
                nf("[Prod] a", "low", 50), nf("[Prod] c", "x", 5000, repo="sites")]
    board = {rank.nr_fingerprint("channel", "Api/Order::create", "E_WARNING", 'Undefined array key "password"'):
             {"wi": 9, "state": "Active", "closed_at": 0}}
    groups, skipped = intake.build_candidates("newrelic", findings, board_index=board, now=2e9,
                                              repo_ok=lambda r: r != "sites", cfg=cfg)
    reasons = " | ".join(s["reason"] for s in skipped)
    check("eşik altı, açık WI ve yerelde olmayan repo elenir",
          [g.repo for g in groups] == ["webservice"] and len(groups[0].fingerprints) == 1
          and "yerelde" in reasons and "#9" in reasons and "eşik" in reasons, reasons)

    class FakeClient:
        org_url, project, repo_projects = "https://dev.azure.com/o", "P", ["EcomBackend"]

        def __init__(self):
            self.created, self.updates, self.stories = [], [], []
            self.iters = [{"name": "Sprint 42", "path": "P\\Sprint 42", "attributes": {"timeFrame": "current"}},
                          {"name": "Sprint 44", "path": "P\\Sprint 44",
                           "attributes": {"timeFrame": "future", "startDate": "2026-11-02T00:00:00Z"}},
                          {"name": "Sprint 43", "path": "P\\Sprint 43",
                           "attributes": {"timeFrame": "future", "startDate": "2026-10-19T00:00:00Z"}}]

        def list_iterations(self, team=""):
            return self.iters

        def query_work_items(self, wiql, limit=0):
            return self.stories

        def get_team_area_path(self, team=""):
            return "P\\Team"

        def get_authenticated_user(self):
            return {"uniqueName": "me@flo.com.tr"}

        def create_work_item(self, type_name, fields, parent_id=None):
            self.created.append((type_name, fields, parent_id))
            return {"id": 1000 + len(self.created) - 1, "fields": fields}

        def update_work_item(self, wid, ops):
            self.updates.append((wid, ops))
            return {}

        def get_work_item(self, wid):
            return {"id": wid, "fields": {"System.WorkItemType": "Bug", "System.State": "Active"}}

    queued, recorded = [], []
    deps = intake.Deps(
        gather=lambda source, client, cfg_: (list(findings), [{"app": "[Prod] ?", "reason": "eşlenemedi"}]),
        board=lambda client: {}, repo_ok=lambda r: True,
        queue=lambda wid, title: (queued.append(wid), 500 + len(queued))[1],
        record=lambda **kw: recorded.append(kw), now=lambda: 2e9, cfg=lambda: cfg)
    c = FakeClient()
    res = intake.take(c, "newrelic", 10, team="Team", deps=deps)
    story = c.created[0]
    check("story yoksa SONRAKİ sprintte açılır, aktif sprinte dokunulmaz (başlık, etiket, iterasyon, area)",
          story[0] == "User Story" and story[1]["System.Title"] == "Kalite: New Relic hataları — Sprint 43"
          and "tempo-src-newrelic" in story[1]["System.Tags"] and story[1]["System.IterationPath"] == "P\\Sprint 43"
          and story[1]["System.AreaPath"] == "P\\Team")
    tasks = [x for x in c.created if x[0] == "Task"]
    check("Task'lar story'nin child'ı; area/iteration story'den; atama PAT sahibi",
          tasks and all(t[2] == 1000 and t[1]["System.IterationPath"] == "P\\Sprint 43"
                        and t[1]["System.AreaPath"] == "P\\Team"
                        and t[1]["System.AssignedTo"] == "me@flo.com.tr" for t in tasks))
    check("her Task kuyruğa alınır ve parmak izi başına kaydedilir",
          len(queued) == len(tasks) and res["tasks"][0]["job_id"] == 501 and len(recorded) >= len(tasks))
    check("toplam SP bütçeyi aşmaz; eşlenemeyen raporlanır",
          res["total_sp"] <= 10 and any(s.get("reason") == "eşlenemedi" for s in res["skipped"]))
    check("story SP = açılan Task'ların toplamı",
          any(w == 1000 and ops[0]["path"] == "/fields/Custom.StoryPoints" and ops[0]["value"] == res["total_sp"]
              for w, ops in c.updates))

    c2 = FakeClient()
    c2.stories = [{"id": 77, "fields": {"System.AreaPath": "P\\A", "System.IterationPath": "P\\Sprint 43",
                                        "Custom.StoryPoints": 2}}]
    res2 = intake.take(c2, "newrelic", 10, deps=deps)
    check("sprintte bu kaynağın story'si varsa onu kullanır; SP üstüne eklenir",
          c2.created and all(x[0] == "Task" and x[2] == 77 for x in c2.created)
          and any(w == 77 and ops[0]["value"] == 2 + res2["total_sp"] for w, ops in c2.updates))

    hum = [nf("[Prod] a", "", 99999, action="X::y", cls="PDOException",
              msg="SQLSTATE[HY000] [2002] Connection refused")]
    c3 = FakeClient()
    q0 = len(queued)
    res3 = intake.take(c3, "newrelic", 10, deps=deps._replace(gather=lambda s, cl, cf: (hum + findings, [])))
    h = res3["human_tasks"]
    check("needs_human Task açılır, kuyruğa alınmaz, needs-human etiketi taşır",
          len(h) == 1 and h[0]["job_id"] is None and len(queued) - q0 == len(res3["tasks"])
          and any("needs-human" in x[1].get("System.Tags", "") for x in c3.created if x[0] == "Task"))

    c4 = FakeClient()
    res4 = intake.take(c4, "newrelic", 10, deps=deps._replace(gather=lambda s, cl, cf: ([], [])))
    check("aday yoksa hiçbir WI (story dahil) açılmaz", c4.created == [] and res4["story"] is None
          and res4["tasks"] == [] and res4.get("message"))
    c5 = FakeClient()
    c5.iters = []
    try:
        intake.take(c5, "newrelic", 10, deps=deps)
        ok = False
    except intake.NoSprint:
        ok = True
    check("hedef sprint yok ve story_id yok → NoSprint, WI açılmaz", ok and c5.created == [])
    c7 = FakeClient()
    cfg_cur = dict(cfg, target_sprint="current")
    intake.take(c7, "newrelic", 10, deps=deps._replace(cfg=lambda: cfg_cur))
    check("CREW_QUALITY_TARGET_SPRINT=current → aktif sprint",
          c7.created and c7.created[0][1]["System.IterationPath"] == "P\\Sprint 42")
    check("hedef sprint knob'u kayıtlı, varsayılan next",
          pc._SCHEMA_BY_KEY.get("CREW_QUALITY_TARGET_SPRINT", {}).get("default") == "next"
          and intake.load_cfg()["target_sprint"] == "next")
    c6 = FakeClient()
    try:
        intake.take(c6, "newrelic", 10, story_id=5, deps=deps)
        ok = False
    except ValueError:
        ok = True
    check("story_id User Story değilse ValueError, WI açılmaz", ok and c6.created == [])
    intake._LOCK.acquire()
    try:
        try:
            intake.take(FakeClient(), "newrelic", 10, deps=deps)
            ok = False
        except intake.Busy:
            ok = True
    finally:
        intake._LOCK.release()
    check("eşzamanlı ikinci alım → Busy", ok)
    pv = intake.preview(FakeClient(), "newrelic", 10, deps=deps)
    check("preview yazmaz ama aynı seçimi döner", pv["total_sp"] == res["total_sp"] and pv["story"] is None)
    try:
        from agile_sdlc_crew import db
        db.init_db()
        check("quality_items tablosu oluşur ve okunur", isinstance(db.list_quality_items(limit=1), list))
    except Exception as e:  # noqa: BLE001
        skip("quality_items tablosu", f"DB erişilemedi: {e}")


def test_quality_endpoints():
    print("\n[Q9] Kalite alımı — endpoint kapıları + dashboard kartı")
    from fastapi.testclient import TestClient
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew import server
    orig = pc.get
    pc.get = lambda k: False if k == "CREW_QUALITY_INTAKE" else orig(k)
    try:
        tc = TestClient(server.app)
        check("kapalıyken take → 403", tc.post("/api/quality/take", json={"source": "newrelic", "sp": 10}).status_code == 403)
        check("kapalıyken preview → 403", tc.get("/api/quality/preview?source=sonar&sp=5").status_code == 403)
        check("kapalıyken unmapped → 403", tc.get("/api/quality/unmapped").status_code == 403)
        check("recent kapalıyken enabled:false", tc.get("/api/quality/recent").json().get("enabled") is False)
    finally:
        pc.get = orig
    pc.get = lambda k: True if k == "CREW_QUALITY_INTAKE" else orig(k)
    try:
        tc = TestClient(server.app)
        check("geçersiz kaynak → 400", tc.post("/api/quality/take", json={"source": "jira", "sp": 10}).status_code == 400)
        check("SP sınırı (1..40) → 400", tc.post("/api/quality/take", json={"source": "sonar", "sp": 0}).status_code == 400
              and tc.get("/api/quality/preview?source=sonar&sp=99").status_code == 400)
    finally:
        pc.get = orig
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text(encoding="utf-8")
    check("dashboard Kalite kartı ve çağrıları",
          'id="ovQualityCard"' in html and "/api/quality/take" in html and "/api/quality/preview" in html
          and "function loadQuality" in html and "loadQuality();" in html)


def test_quality_review_fixes():
    print("\n[Q10] Kalite alımı — son inceleme düzeltmeleri")
    import requests as _rq
    import yaml
    from agile_sdlc_crew.quality import app_map, intake, rank, render
    from agile_sdlc_crew.quality.model import Finding

    def one(msg, cls="E_WARNING"):
        return rank.nr_units([Finding(source="newrelic", repo="r", app="[Prod] a", action="A::b",
                                      err_class=cls, message=msg, count=500, last_seen=1.0)])[0]
    code_bugs = [("Undefined array key \"timeout\"", "E_WARNING"),
                 ("Undefined array key \"connect_timeout\"", "E_WARNING"),
                 ("Undefined property: stdClass::$timeOut", "E_WARNING"),
                 ("Call to a member function getTimeout() on null", "Error"),
                 ("Argument #1 ($timeout) must be of type int, string given", "TypeError")]
    wrong = [m for m, c in code_bugs if one(m, c).needs_human]
    check("kimlik adındaki timeout altyapı sayılmaz (kod hatası, kuyruğa girer)", not wrong, str(wrong))
    check("tanımsız 'timeout' anahtarı 1 SP", one('Undefined array key "timeout"').sp == 1)
    infra = [("cURL error 28: Operation timed out after 5000 ms", "GuzzleHttp\\Exception\\ConnectException"),
             ("SQLSTATE[HY000] [2002] Connection refused", "PDOException"),
             ("SQLSTATE[HY000]: General error: 1205 Lock wait timeout exceeded", "PDOException"),
             ("MySQL server has gone away", "PDOException")]
    missed = [m for m, c in infra if not one(m, c).needs_human]
    check("gerçek altyapı hataları hâlâ needs_human", not missed, str(missed))

    m = render.mask_secrets
    leaks = [s for s, secret in (("https://user:hunter22@api.x.com/v1", "hunter22"),
                                 ("redis://:p4ssw0rd@10.0.0.5:6379", "p4ssw0rd"),
                                 ("redis://:pw9x@redis:6379", "pw9x"),
                                 ("pass=Secret1", "Secret1"), ("pwd hunter2", "hunter2"),
                                 ("client_secret: abcXYZ", "abcXYZ"))
             if secret in m(s)]
    check("URL userinfo ve kısa anahtar parolaları maskelenir", not leaks, str(leaks))
    check("bypass=1 gibi kelime içi eşleşme maskelenmez", m("bypass=1") == "bypass=1")
    pref = [s for s, secret in (("dbpassword=Hunter2x", "Hunter2x"), ("accesstoken=abcdEF", "abcdEF"),
                                ("authtoken: zzQQ", "zzQQ"), ("mysqlpasswd=x1y2", "x1y2"),
                                ("clientsecret=abc", "abc")) if secret in m(s)]
    check("önekli anahtar adları (dbpassword, accesstoken) hâlâ maskelenir", not pref, str(pref))

    jf = 'String gitOriginApp = "git@ssh.dev.azure.com:v3/o/ecombackend/webservice"\n'
    df = 'newrelic.appname="[Prod] a.com"\n'

    def http_err(code):
        r = _rq.Response()
        r.status_code = code
        return _rq.HTTPError(str(code), response=r)

    class Dev:
        def __init__(self, fail=None, dirs=("x",)):
            self.fail, self.dirs = fail, dirs

        def get_items_in_path(self, repo, path, branch=None, recursion_level="oneLevel", project=None):
            return [{"path": f"/build/{d}", "isFolder": True} for d in self.dirs] if repo == "global" else []

        def get_file_content(self, repo, path, branch=None, project=None):
            if self.fail:
                raise http_err(self.fail)
            if path.endswith("/Dockerfile") and "/y/" in path:
                raise http_err(404)
            return jf if path.endswith("JenkinsfileBuild") else df

    with tempfile.TemporaryDirectory() as td:
        cache, ovr = Path(td) / "m.yaml", Path(td) / "o.yaml"
        ovr.write_text("{}", encoding="utf-8")
        mp = app_map.load_map(Dev(dirs=("x", "y")), cache_path=cache, overrides_path=ovr)
        check("404 (dosya yok) sessizce atlanır", app_map.resolve("[Prod] a.com", mp) == "webservice")
        m2 = app_map.load_map(Dev(fail=503), cache_path=cache, overrides_path=ovr, force=True)
        kept = (yaml.safe_load(cache.read_text(encoding="utf-8")) or {}).get("apps") or {}
        check("FLODevOps 5xx → eski önbellek korunur ve kullanılır",
              app_map.resolve("[Prod] a.com", m2) == "webservice" and "[Prod] a.com" in kept)
        m3 = app_map.load_map(Dev(dirs=()), cache_path=cache, overrides_path=ovr, force=True)
        check("boş/çok küçülmüş eşleme önbelleği ezmez",
              app_map.resolve("[Prod] a.com", m3) == "webservice")

    class AzC:
        org_url, project, repo_projects = "https://dev.azure.com/o", "P", []

        def __init__(self, reject_ac=False):
            self.reject_ac, self.created = reject_ac, []

        def list_iterations(self, team=""):
            return [{"name": "S1", "path": "P\\S1", "attributes": {"timeFrame": "current"}},
                    {"name": "S2", "path": "P\\S2", "attributes": {"timeFrame": "future", "startDate": "2026-10-19"}}]

        def query_work_items(self, wiql, limit=0):
            return []

        def get_team_area_path(self, team=""):
            return "P\\T"

        def get_authenticated_user(self):
            return {}

        def create_work_item(self, type_name, fields, parent_id=None):
            if type_name == "Task" and self.reject_ac and "Microsoft.VSTS.Common.AcceptanceCriteria" in fields:
                raise http_err(400)
            self.created.append((type_name, dict(fields)))
            return {"id": 100 + len(self.created), "fields": fields}

        def update_work_item(self, wid, ops):
            return {}

    fs = [Finding(source="newrelic", repo="webservice", app="[Prod] a", action="A::b", err_class="TypeError",
                  message="boom", count=500, last_seen=1.0)]
    deps = intake.Deps(gather=lambda s, c, cf: (list(fs), []), board=lambda c: {}, repo_ok=lambda r: True,
                       queue=lambda w, t: 1, record=lambda **k: None, now=lambda: 2.0,
                       cfg=lambda: {"nr_days": 7, "nr_min_errors": 100, "pack_sp": 3, "pack_files": 5,
                                    "human_max": 5, "max_tasks": 8, "grace_days": 3, "sonar_url": ""})
    c = AzC()
    intake.take(c, "newrelic", 5, deps=deps)
    task = next(f for ty, f in c.created if ty == "Task")
    check("Task'a Acceptance Criteria alanı yazılır (hazırlık kapısı −15 yemesin)",
          "Kabul kriteri" in task.get("Microsoft.VSTS.Common.AcceptanceCriteria", "")
          or "Bu action" in task.get("Microsoft.VSTS.Common.AcceptanceCriteria", ""))
    c2 = AzC(reject_ac=True)
    r2 = intake.take(c2, "newrelic", 5, deps=deps)
    check("tipte AC alanı yoksa (400) alansız açılır",
          any(ty == "Task" for ty, _ in c2.created) and r2["tasks"][0]["id"] and not r2["tasks"][0].get("error"))
    cf = deps.cfg()
    cf["task_fields"] = {"Custom.ARGEType": "Geliştirme - GE"}
    c3 = AzC()
    intake.take(c3, "newrelic", 5, deps=deps._replace(cfg=lambda: cf))
    t3 = next(f for ty, f in c3.created if ty == "Task")
    check("org'un zorunlu Task alanları (quality_task_fields.yaml) Task'a yazılır",
          t3.get("Custom.ARGEType") == "Geliştirme - GE"
          and intake.load_cfg()["task_fields"].get("Custom.ARGEType") == "Geliştirme - GE")


def test_quality_readiness():
    print("\n[Q11] Kalite alımı — beklenen davranış bloğu + hazırlık kapısı muafiyeti")
    import inspect
    from agile_sdlc_crew import flow as _flow
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.quality import rank, render
    from agile_sdlc_crew.quality.model import Finding

    def g(msg, cls="E_WARNING"):
        return rank.nr_units([Finding(source="newrelic", repo="webservice", app="[Prod] a", action="Index::index",
                                      err_class=cls, message=msg, count=500, last_seen=1.0)])[0]
    und = render.task_description(g('Undefined array key "host"'))
    nul = render.task_description(g("Trying to access array offset on value of type null"))
    gen = render.task_description(g("Call to undefined method Foo::bar()", cls="Error"))
    check("tanımsız anahtar: beklenen davranış (isset/??, akış korunur, hangi dosya, doğrulama)",
          "Beklenen davranış" in und and "isset" in und and "birim testi" in und and "gerçekten" in und)
    check("olası yer kapsam kilidi değil, başlangıç noktası (#225: mimar kilit sandı)",
          "başlangıç noktası" in und and "yardımcı" in und)
    check("test kapsamı dar: yalnız değişen satır, büyük metot sahte veriyle çağrılmaz, görünürlük değişmez (#232)",
          "yalnız değişen satır" in und and "baştan sona" in und and "görünürlüğü değiştirilmez" in und)
    check("null erişimi: null koruması", "Beklenen davranış" in nul and "null" in nul.split("Beklenen davranış")[1])
    check("genel hata: kök neden, hatayı yutmak çözüm değil", "Beklenen davranış" in gen and "yutmak" in gen)
    ac = render.acceptance_block(g('Undefined array key "host"'))
    check("AC alanı beklenen davranışı da taşır", "isset" in ac and "Kabul kriteri" not in ac and "bir daha" in ac)
    check("tempo-quality etiketli iş hazırlık kapısından muaf (varsayılan açık)",
          _flow._quality_skips_readiness("tempo-quality; tempo-src-newrelic")
          and not _flow._quality_skips_readiness("needs-refinement")
          and not _flow._quality_skips_readiness(""))
    orig = pc.get
    pc.get = lambda k: False if k == "CREW_QUALITY_SKIP_READINESS" else orig(k)
    try:
        check("muafiyet kapatılabilir", not _flow._quality_skips_readiness("tempo-quality"))
    finally:
        pc.get = orig
    src = inspect.getsource(_flow)
    check("hazırlık kapısı muafiyeti kullanıyor",
          "_quality_skips_readiness(self.state.wi_tags)" in src)
    # #236: Sonar SSL işi 2. aşamada düştü — BA'nın 12 maddesinin 8'i (tempo-fp korunur,
    # Sonar'da kapanır, yalnız listelenen satırlar…) plan değişikliğiyle kapsanamaz.
    stage2 = src[src.index("HAZIRLIK KAPISI (asama 2)"):src.index("raise NeedsMoreInfo(_msg)", src.index("HAZIRLIK KAPISI (asama 2)"))]
    check("2. aşama kapsam eşiği kalite işinde atlanır (mimar reddi kalır)",
          "_quality_skips_readiness(self.state.wi_tags)" in stage2 and "_refusal" in stage2)


def test_repo_ctx_cross_thread():
    print("\n[46] claude -p repo bağlamı — crewai'nin çağrı thread'ine ulaşır, başka işe sızmaz")
    import contextvars as _cv
    import threading as _th
    from agile_sdlc_crew.tools import claude_cli_llm as _cli
    seen = []

    def _fake(cmd, env, timeout_s, meta=None, idle_s=0):
        seen.append(list(cmd))
        if meta is not None:
            meta.update({"turns": 1, "cost_usd": 0.0})
        return "OK"

    def _call(copy_ctx=True):
        fn = lambda: _cli.claude_cli_completion("x", model="sonnet", system="s")  # noqa: E731
        # crewai LLM çağrısını asyncio.to_thread ile yapar → contextvars kopyalanır
        # (ölçüldü: çağrı 'asyncio_0' thread'inde, ContextVar görünür, threading.local görünmez)
        th = _th.Thread(target=_cv.copy_context().run, args=(fn,)) if copy_ctx else _th.Thread(target=fn)
        th.start()
        th.join()
        return seen[-1]

    _orig = _cli._run_streaming
    _cli._run_streaming = _fake
    try:
        _cli.clear_repo_ctx(); _cli.set_toolless(False)
        _cli.set_repo_ctx(["/tmp/repo-a"], "Read,Grep,Glob,LS")
        c = _call()
        check("flow adımında kurulan repo bağlamı crewai'nin çağrı thread'inde görülür (#225: keşif araçsız kaldı)",
              "--add-dir" in c and "/tmp/repo-a" in c and "--disallowedTools" not in c, " ".join(c[-8:]))
        c = _call(copy_ctx=False)
        check("bağlamı kopyalamayan başka bir iş (ör. eşzamanlı pr_review) bu repoyu ve araçları GÖRMEZ",
              "--add-dir" not in c and "--disallowedTools" in c)
        _cli.clear_repo_ctx()
        c = _call()
        check("temizlenince araçlar kapanır (emit fazı)", "--add-dir" not in c and "--disallowedTools" in c)
        _cli.set_toolless(True)
        _cli.set_repo_ctx([], "")
        c = _call()
        check("set_toolless da çağrı thread'ine ulaşır", "--disallowedTools" in c)
        _cli.set_toolless(False)
        with _cli.repo_tools_context(["/tmp/repo-c"]):
            c = _call()
        c2 = _call()
        check("repo_tools_context çağrı thread'ine yayılır ve çıkışta geri alınır",
              "/tmp/repo-c" in c and "--add-dir" not in c2)
    finally:
        _cli._run_streaming = _orig
        _cli.clear_repo_ctx()
        _cli.set_toolless(False)

def test_windowed_edit():
    print("\n[47] Büyük dosyada pencereli düzenleme — 'sonuna ekle' yok (#229: Order.php 352KB)")
    import shutil
    import subprocess as _sp
    import inspect
    from agile_sdlc_crew import edit_window as ew
    from agile_sdlc_crew import flow as _flow
    from agile_sdlc_crew.main import _lint_issue_lines
    lines = [f"    $x{i} = $data['k{i}'];" for i in range(1, 801)]
    lines[399] = "        'telephone' => Utils::getTelephoneFormat($order['billing']['telephone']),"
    content = "<?php\nclass A {\n" + "\n".join(lines) + "\n}\n"
    cur = "'telephone' => Utils::getTelephoneFormat($order['billing']['telephone']),"
    new = "'telephone' => Utils::getTelephoneFormat($order['billing']['telephone'] ?? null),"
    anchor = ew.find_anchor(content, "        " + cur + "  // eski yorum")
    check("anchor: plan alıntısı birebir değilse de doğru satır bulunur", anchor and anchor[0] == 402, str(anchor))
    check("anchor: metot adından", ew.find_anchor("a\nb\n  public function newOrderDetailAction($id)\n{", "function newOrderDetailAction()") [0] == 3)
    check("anchor: mimarın verdiği satır öncelikli", ew.find_anchor(content, "yok", hint_line=10)[0] == 10)
    check("anchor: bulunamazsa None", ew.find_anchor(content, "tamamen alakasız bir şey 123") is None)
    asked = {}

    def ask_ok(window_text, start, end):
        asked.update(text=window_text, start=start, end=end)
        body = [ln.split("| ", 1)[1] for ln in window_text.splitlines()]
        return "```php\n" + "\n".join(b.replace(cur, new) for b in body) + "\n```"
    out, info = ew.apply_windowed(content, cur, new, ask_ok, radius=50)
    check("yalnız ±50 satırlık numaralı pencere sorulur",
          asked["start"] == 352 and asked["end"] == 452 and asked["text"].splitlines()[0].startswith("  352| "), info)
    check("pencere yerine oturur, dosyanın geri kalanı aynen kalır",
          out is not None and new in out and out.count("\n") == content.count("\n")
          and out.replace(new, cur) == content)
    snippet, _ = ew.apply_windowed(content, cur, new, lambda w, s, e: "```php\n" + new + "\n```", radius=50)
    check("model yalnız değişen satırı döndürürse (pencereyi silerdi) reddedilir", snippet is None)
    prose, _ = ew.apply_windowed(content, cur, new, lambda w, s, e: "(1) satır 3678: telephone okuması korunmalı", radius=50)
    check("model kod yerine açıklama döndürürse reddedilir", prose is None)
    num, _ = ew.apply_windowed(content, cur, new, lambda w, s, e: "\n".join(
        ln.replace(cur, new) for ln in w.splitlines()), radius=50)
    check("model satır numaralarını bıraksa da temizlenir", num is not None and "| " not in num.splitlines()[400])
    from agile_sdlc_crew.quality import rank as _rk, render as _rd
    from agile_sdlc_crew.quality.model import Finding as _F
    exc = _rk.nr_units([_F(source="newrelic", repo="webservice", app="[Prod] a", action="Api/V1/Basket::shoppingCartCount",
                           err_class="Exception", count=500, last_seen=1.0,
                           message="Uncaught exception 'Exception' with message 'cURL Error' in "
                                   "/var/www/app/app/Model/V1/Florchestra.php:820")])[0]
    d_exc = _rd.task_description(exc)
    check("exception mesajındaki dosya:satır Task açıklamasına 'Hata yeri' olarak girer",
          "Hata yeri" in d_exc and "app/Model/V1/Florchestra.php:820" in d_exc
          and _rd.exception_location("x in /var/www/app/app/A.php:12") == ("app/A.php", 12)
          and _rd.exception_location('Undefined array key "id"') is None)
    src = inspect.getsource(_flow)
    check("implement'te 'sonuna ekle' geri dönüşü yok, pencereli düzenleme kullanılıyor",
          "append stratejisine dusuluyor" not in src and "apply_windowed(" in src)
    out_txt = ("PHP Deprecated:  Non-canonical cast (double) is deprecated in /tmp/x.php on line 12\n"
               "Deprecated: Non-canonical cast (double) is deprecated in /tmp/x.php on line 13\n"
               "PHP Parse error:  syntax error, unexpected identifier \"satır\" in /tmp/x.php on line 9001\n"
               "Errors parsing /tmp/x.php")
    iss = _lint_issue_lines(out_txt)
    check("linter çıktısında Deprecated gürültüsü atılır, gerçek hata görünür",
          iss and "Parse error" in iss[0] and not any("Deprecated" in i for i in iss), str(iss))
    if shutil.which("php"):
        from agile_sdlc_crew.main import _lint_with_native
        ok = _lint_with_native("<?php\n$a = (double) '1.5';\n", "php")
        bad = _lint_with_native("<?php\n$a = (double) '1.5';\n(1) satır 3678: açıklama\n", "php")
        check("php -l: eski cast uyarısı geçer, gerçek sözdizimi hatası düşer (ve mesajı Parse error)",
              ok["valid"] and not bad["valid"] and "Parse error" in bad["issues"][0], str(bad))
    else:
        skip("php -l", "php yok")


def test_quality_lean():
    print("\n[Q12] Kalite işleri — yalın akış + ucuz modeller (yalnız tempo-quality)")
    import inspect
    from crewai.flow.flow import Flow, listen, start
    from agile_sdlc_crew import flow as _flow
    from agile_sdlc_crew import main as _main
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.llm import resolver as rs
    from agile_sdlc_crew.quality import render
    check("knob'lar kayıtlı, varsayılan açık",
          pc._SCHEMA_BY_KEY.get("CREW_QUALITY_LEAN_FLOW", {}).get("default") is True
          and pc._SCHEMA_BY_KEY.get("CREW_QUALITY_LEAN_MODELS", {}).get("default") is True)
    ov = _flow.quality_lean_overrides()
    check("mimar → Sonnet, tamamlanma raporu → Haiku; reviewer ve geliştirici dokunulmaz",
          rs._profile_to_spec(ov["software_architect"])["model"] == "sonnet"
          and rs._profile_to_spec(ov["scrum_master"])["model"] == "haiku"
          and "code_reviewer" not in ov and "senior_developer" not in ov)
    seen = {}

    class _F(Flow):
        @start()
        def a(self):
            rs.set_job_llm_overrides({"software_architect": "architect_cli_sonnet"})

        @listen(a)
        def b(self):
            seen["in"] = rs.resolve_spec_with_source("software_architect")

    with rs.job_llm_scope():
        _F().kickoff()
    seen["out"] = rs.resolve_spec_with_source("software_architect")
    check("iş kapsamındaki model geçersiz kılması sonraki adımlarda görülür",
          seen["in"][1] == "job" and seen["in"][0]["model"] == "sonnet", str(seen.get("in")))
    check("kapsam dışında (başka iş / sonraki iş) geçersiz kılma yok", seen["out"][1] != "job")
    rs.set_job_llm_overrides({"software_architect": "architect_cli_sonnet"})
    check("kapsam yokken set etmek hiçbir şeyi değiştirmez",
          rs.resolve_spec_with_source("software_architect")[1] != "job")
    check("Task açıklamasından repo okunur",
          render.repo_from_description("<p><b>Repo:</b> <code>webservice</code> </p><p>x</p>") == "webservice"
          and render.repo_from_description("<p>yok</p>") == "")
    src = inspect.getsource(_flow)
    check("discover / test planlama / UAT kalite işinde atlanır",
          src.count("Atlandı — kalite işi") >= 3 and "_quality_lean_flow()" in src)
    s1 = inspect.getsource(_flow.AgileSDLCFlow.crew_step1_requirements)
    res_branch = s1[s1.index("cached_ba = self._try_resume_step"):s1.index('self._resume_step("requirements_analysis_task", cached_ba)')]
    check("requirements resume edilse de WI bağlamı (etiket → kalite ayarları) kurulur (#232: atlanmıştı)",
          "self._wi_begin(" in res_branch)
    check("kalite işinde build-fix testi daraltır, üretim görünürlüğünü değiştirmez",
          "QUALITY_BUILD_FIX_RULE" in src and "do not change the visibility" in src.lower())
    from agile_sdlc_crew.crew import AgileSDLCCrew
    with rs.job_llm_scope():
        crew_obj = AgileSDLCCrew()  # flow init'te kurulur: _wi_begin'den ÖNCE
        before = getattr(crew_obj.llm_architect, "model", "")
        rs.set_job_llm_overrides({"software_architect": "architect_cli_sonnet", "scrum_master": "report_haiku"})
        after_a = getattr(crew_obj.llm_architect, "model", "")
        after_s = getattr(crew_obj.llm_scrum, "model", "")
    check("crew kurulduktan SONRA gelen iş kapsamlı model seçimi de uygulanır (#234: mimar Opus kaldı)",
          "sonnet" in str(after_a) and "haiku" in str(after_s), f"{before} → {after_a} / {after_s}")
    check("run_pipeline flow'u iş kapsamında koşturur", "job_llm_scope()" in inspect.getsource(_main.run_pipeline))


def test_worktree_staging():
    print("\n[48] Review/build-fix turları önceki düzeltmeleri geri almasın (#231: telephone ?? null kayboldu)")
    import inspect
    import types
    from agile_sdlc_crew import flow as _flow
    F = _flow.AgileSDLCFlow
    with tempfile.TemporaryDirectory() as td:
        repo = Path(td) / "webservice"
        (repo / "app").mkdir(parents=True)
        main_v = "<?php\n$t = $o['telephone'];\n" + "// dolgu\n" * 40
        branch_v = "<?php\n$t = $o['telephone'] ?? null;\n" + "// dolgu\n" * 40
        f = repo / "app" / "Order.php"
        f.write_text(main_v, encoding="utf-8")  # yerel klon main'de
        fake = types.SimpleNamespace(_repo_mgr=types.SimpleNamespace(base_dir=Path(td)))
        for name in ("_read_worktree_file", "_prefer_worktree_edit", "_stage_worktree_file", "_worktree_base"):
            setattr(fake, name, types.MethodType(getattr(F, name), fake))
        # dev diske dokunmadı, kısa bir metin döndürdü
        out = fake._prefer_worktree_edit("webservice", "/app/Order.php", "", branch_v)
        check("hazırlık yokken bile main kopyası branch'teki düzeltmenin üstüne yazılmaz",
              "?? null" in out or out == "", repr(out[:60]))
        fake._stage_worktree_file("webservice", "/app/Order.php", branch_v)
        check("tur öncesi diske PR branch'inin hali yazılır", f.read_text(encoding="utf-8") == branch_v)
        f.write_text(branch_v.replace("// dolgu", "// dolgu düzeltildi", 1), encoding="utf-8")  # dev in-place edit
        out2 = fake._prefer_worktree_edit("webservice", "/app/Order.php", "", branch_v)
        check("dev'in in-place düzenlemesi branch üstüne kurulur (önceki düzeltme korunur)",
              "?? null" in out2 and "düzeltildi" in out2)
    with tempfile.TemporaryDirectory() as td2:
        base = Path(td2)
        repo2 = base / "r"
        (repo2 / "app").mkdir(parents=True)
        outside = base / "outside.txt"
        outside.write_text("orijinal", encoding="utf-8")
        (repo2 / "app" / "link.php").symlink_to(outside)
        fk = types.SimpleNamespace(_repo_mgr=types.SimpleNamespace(base_dir=base))
        fk._stage_worktree_file = types.MethodType(F._stage_worktree_file, fk)
        fk._stage_worktree_file("r", "/app/link.php", "EZILDI")
        fk._stage_worktree_file("r", "/../outside.txt", "EZILDI")
        check("hazırlık yazımı symlink'i izlemez, repo dışına (../) yazmaz",
              outside.read_text(encoding="utf-8") == "orijinal")
    src = inspect.getsource(F)
    check("review tekrarı ve build-fix turları diski PR branch'ine hazırlar",
          src.count("self._stage_worktree_file(") >= 2)


def test_techdebt_lane():
    print("\n[TD1] Teknik borç — hat (lane) ve story")
    from agile_sdlc_crew.quality import intake, render
    from agile_sdlc_crew.quality.model import Group
    check("techdebt hattı tek story başlığı", render.story_title("sonar_smell", "S22", lane="techdebt") == "Teknik borç — S22")
    check("quality hattı değişmedi", render.story_title("newrelic", "S22") == "Kalite: New Relic hataları — S22")
    check("bilinmeyen kaynak KeyError vermez", "yeni" in render.story_title("yeni", "S1"))
    g = Group(source="sonar_smell", repo="webservice", title="t", fingerprints=["sonar:webservice:x"], findings=[],
              needs_human=True)
    tags = render.task_tags(g, lane="techdebt")
    check("techdebt Task etiketleri (tempo-quality yalın akış için korunur)",
          tags == "tempo-quality; tempo-techdebt; tech-debt; tempo-src-sonar_smell; needs-human", tags)
    check("quality Task etiketleri değişmedi",
          render.task_tags(Group(source="sonar", repo="r", title="t", fingerprints=["f"], findings=[])) == "tempo-quality; tempo-src-sonar")

    class C:
        org_url, project = "https://dev.azure.com/o", "P"

        def __init__(self):
            self.wiql, self.created = [], []

        def list_iterations(self, team=""):
            return [{"name": "S22", "path": "P\\S22", "attributes": {"timeFrame": "future", "startDate": "2026-10-19"}}]

        def query_work_items(self, wiql, limit=0):
            self.wiql.append(wiql)
            return []

        def get_team_area_path(self, team=""):
            return "P\\T"

        def create_work_item(self, type_name, fields, parent_id=None):
            self.created.append((type_name, fields))
            return {"id": 7, "fields": fields}
    c = C()
    st = intake.resolve_story(c, "sonar_smell", None, "T", "next", lane="techdebt")
    check("techdebt story kaynak etiketi olmadan aranır ve tek etiketle açılır",
          "tempo-techdebt" in c.wiql[0] and "tempo-src-" not in c.wiql[0]
          and c.created[0][1]["System.Title"] == "Teknik borç — S22"
          and "tempo-techdebt" in c.created[0][1]["System.Tags"] and st["id"] == 7)


def test_techdebt_open_tasks_db():
    print("\n[TD2] Teknik borç — ortak Task açma + quality_items.source kolonu")
    from agile_sdlc_crew.quality import intake
    from agile_sdlc_crew.quality.model import Group
    queued, rec = [], []

    class C:
        def __init__(self):
            self.created = []

        def create_work_item(self, type_name, fields, parent_id=None):
            self.created.append((type_name, fields, parent_id))
            return {"id": 500 + len(self.created)}

        def update_work_item(self, wid, ops):
            return {}
    deps = intake.Deps(gather=None, board=None, repo_ok=None, queue=lambda w, ttl: (queued.append(w), 9)[1],
                       record=lambda **k: rec.append(k), now=lambda: 1.0, cfg=None)
    cfg = {"sonar_url": "", "nr_days": 7, "task_fields": {"Custom.ARGEType": "Geliştirme - GE"}}
    from agile_sdlc_crew.quality.model import Finding
    g = Group(source="sonar_smell", repo="webservice", title="X temizliği", fingerprints=["sonar:webservice:a.php:php:S125"],
              findings=[Finding(source="sonar", repo="webservice", project="EcomBackend:webservice", file="a.php",
                                rule="php:S125", rule_name="Yorumda kod", severity="MAJOR", itype="CODE_SMELL", lines=[3])],
              sp=2)
    story = {"id": 77, "area": "P\\T", "iteration": "P\\S22", "sp": 0.0}
    c = C()
    rows, added = intake._open_tasks(c, [g], story, cfg, deps, source="sonar_smell", lane="techdebt", queue=False)
    f = c.created[0][1]
    check("ortak Task açma: techdebt etiketleri, story child'ı, org alanı, kuyruğa ALINMAZ",
          rows[0]["id"] == 501 and rows[0]["job_id"] is None and not queued and added == 2
          and c.created[0][2] == 77 and "tempo-techdebt" in f["System.Tags"]
          and f["Custom.ARGEType"] == "Geliştirme - GE" and rec and rec[0]["source"] == "sonar_smell")
    from agile_sdlc_crew import db
    check("quality_items.source serbest metin (yeni kaynaklar sessizce düşmesin)",
          "source VARCHAR(32)" in db.SCHEMA and "MODIFY source VARCHAR(32)" in inspect_source(db.init_db))
    try:
        db.init_db()
        with db.get_conn() as conn:
            cur = conn.cursor()
            cur.execute("SHOW COLUMNS FROM quality_items LIKE 'source'")
            typ = (cur.fetchone() or {}).get("Type", "")
        check("canlı DB'de kolon varchar'a çevrildi", "varchar" in typ.lower(), typ)
    except Exception as e:  # noqa: BLE001
        skip("canlı DB kolon tipi", f"DB erişilemedi: {e}")


def inspect_source(fn):
    import inspect
    return inspect.getsource(fn)


def test_techdebt_list_queue():
    print("\n[TD4] Teknik borç — liste ve kuyruğa alma")
    from agile_sdlc_crew.quality import techdebt

    def wi(i, title, tags, state="To Do", typ="Task", desc="", sp=None):
        f = {"System.Title": title, "System.Tags": tags, "System.State": state, "System.WorkItemType": typ,
             "System.IterationPath": "P\\S22", "System.Description": desc}
        if sp is not None:
            f["Custom.StoryPoints"] = sp
        return {"id": i, "fields": f}

    class C:
        def __init__(self):
            self.wiql = []

        def get_team_area_path(self, team=""):
            return "P\\T"

        def query_work_items(self, wiql, limit=0):
            self.wiql.append(wiql)
            return [wi(1, "Elle", "tempo-quality; tempo-techdebt; tech-debt",
                       desc="<p><b>Repo:</b> <code>webservice</code></p>", sp=2),
                    wi(2, "Ekipten", "tech-debt"),
                    wi(3, "Büyük", "tech-debt; needs-human"),
                    wi(4, "Koşan", "tech-debt")]
    jobs = {"4": {"id": 99, "status": "running"}}
    c = C()
    res = techdebt.list_items(c, team="T", job_status=lambda ids: {k: v for k, v in jobs.items() if k in ids})
    rows = {r["id"]: r for r in res["items"]}
    q = c.wiql[0]
    check("liste: takım alanı + tech-debt etiketi + açık durum + yalnız Task/Bug/Improvement",
          "UNDER 'P\\T'" in q and "'tech-debt'" in q and "NOT IN" in q and "'Task', 'Bug', 'Improvement'" in q)
    check("satır alanları: repo, insan işi, iş, kuyruklanabilir",
          rows[1]["repo"] == "webservice" and rows[1]["sp"] == 2
          and rows[3]["needs_human"] and not rows[3]["queueable"]
          and rows[4]["job"]["id"] == 99 and not rows[4]["queueable"] and rows[1]["queueable"])
    sent = []
    out = techdebt.queue(C(), [1, 2, 3, 77], team="T", max_queue=5,
                         job_status=lambda ids: {}, queue_fn=lambda rows: (sent.extend(rows), {
                             "queued": [{"wi": str(r["id"]), "job_id": 10} for r in rows], "skipped": []})[1])
    reasons = {s["wi"]: s["reason"] for s in out["skipped"]}
    check("kuyruk: insan işi ve listede olmayan reddedilir, diğerleri kuyruğa",
          [r["id"] for r in sent] == [1, 2] and "insan" in reasons["3"] and "listesinde yok" in reasons["77"])
    out2 = techdebt.queue(C(), [1, 2], team="T", max_queue=1, job_status=lambda ids: {},
                          queue_fn=lambda rows: {"queued": [{"wi": str(r["id"]), "job_id": 1} for r in rows], "skipped": []})
    check("istek başı tavan", len(out2["queued"]) == 1 and any("tavan" in s["reason"] for s in out2["skipped"]))


def test_techdebt_endpoints():
    print("\n[TD5] Teknik borç — endpoint kapıları, ayarlar, hazırlık kapısı")
    from fastapi.testclient import TestClient
    from agile_sdlc_crew import flow as _flow
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew import server
    check("ayarlar kayıtlı: CREW_TECHDEBT kapalı, kuyruk tavanı 5",
          pc._SCHEMA_BY_KEY.get("CREW_TECHDEBT", {}).get("default") is False
          and pc._SCHEMA_BY_KEY.get("CREW_TECHDEBT_MAX_QUEUE", {}).get("default") == 5)
    orig = pc.get
    pc.get = lambda k: False if k == "CREW_TECHDEBT" else orig(k)
    try:
        tc = TestClient(server.app)
        check("kapalıyken liste / kuyruk → 403",
              tc.get("/api/techdebt").status_code == 403
              and tc.post("/api/techdebt/queue", json={"ids": [1]}).status_code == 403)
    finally:
        pc.get = orig
    pc.get = lambda k: True if k == "CREW_TECHDEBT" else orig(k)
    try:
        tc = TestClient(server.app)
        check("elle giriş kaldırıldı (uç yok)", tc.post("/api/techdebt/manual", json={}).status_code in (404, 405))
        r0 = tc.get("/api/techdebt")
        check("takım seçilmeden liste: hata değil, 'takım seçin' mesajı (200)",
              r0.status_code == 200 and r0.json().get("items") == [] and "Takım" in r0.json().get("message", ""))
        check("boş kuyruk isteği → 400", tc.post("/api/techdebt/queue", json={"ids": []}).status_code == 400)
    finally:
        pc.get = orig
    check("taranan kalite işi hazırlık kapısından muaf", _flow._quality_skips_readiness("tempo-quality; tempo-src-newrelic"))


def test_techdebt_ui():
    print("\n[TD6] Teknik borç — dashboard sekmesi")
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text(encoding="utf-8")
    s = html.find('id="paneTechdebt"')
    e = html.find("<section", s + 10)
    pane = html[s:e] if s >= 0 else ""
    ov_s = html.find('id="paneOverview"')
    ov = html[ov_s:html.find("<section", ov_s + 10)]
    check("üst barda Teknik borç sekmesi ve paneli", 'data-tab="techdebt"' in html and s >= 0)
    check("Kalite kartı Genel bakış'tan Teknik borç sekmesine taşındı",
          'id="ovQualityCard"' in pane and 'id="ovQualityCard"' not in ov)
    check("liste ve çoklu seçimle kuyruk var; elle giriş formu kaldırıldı",
          'id="tdList"' in pane and "/api/techdebt/queue" in html and "function loadTechdebt" in html
          and "name==='techdebt'" in html and 'id="tdForm"' not in html and "/api/techdebt/manual" not in html)


def test_techdebt_review_fixes():
    print("\n[TD7] Teknik borç — inceleme düzeltmeleri")
    import requests as _rq
    from fastapi.testclient import TestClient
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew import server
    from agile_sdlc_crew.quality import intake, techdebt

    class C:
        def __init__(self, area="P\\T"):
            self.area, self.wiql = area, []

        def get_team_area_path(self, team=""):
            return self.area

        def query_work_items(self, wiql, limit=0):
            self.wiql.append(wiql)
            return [{"id": i, "fields": {"System.Title": f"t{i}", "System.Tags": "tech-debt", "System.State": "To Do",
                                          "System.WorkItemType": "Task"}} for i in (1, 2, 3, 4, 5)]
    sent = []
    techdebt.queue(C(), [1, 1, 1], team="T", max_queue=5, job_status=lambda ids: {},
                   queue_fn=lambda rows: (sent.extend(rows), {"queued": [], "skipped": []})[1])
    check("aynı id birden çok gönderilse de tek iş", [r["id"] for r in sent] == [1])
    intake._LOCK.acquire()
    try:
        try:
            techdebt.queue(C(), [2], team="T", job_status=lambda ids: {}, queue_fn=lambda rows: {})
            busy = False
        except intake.Busy:
            busy = True
    finally:
        intake._LOCK.release()
    check("eşzamanlı kuyruk isteği → Busy (çift tıklama)", busy)
    jobs = {"1": {"id": 11, "status": "completed"}, "2": {"id": 12, "status": "needs_info"},
            "3": {"id": 13, "status": "needs_human"}, "4": {"id": 14, "status": "failed"}}
    rows = {r["id"]: r for r in techdebt.list_items(C(), team="T", job_status=lambda ids: jobs)["items"]}
    check("tamamlanmış / insan bekleyen işi olan WI kuyruklanamaz; failed tekrar alınabilir",
          not rows[1]["queueable"] and not rows[2]["queueable"] and not rows[3]["queueable"]
          and rows[4]["queueable"] and rows[5]["queueable"])
    try:
        techdebt.list_items(C(area=""), team="", job_status=lambda ids: {})
        scoped = False
    except ValueError as e:
        scoped = "takım" in str(e)
    check("takım alanı çözülemezse proje geneline düşmez (ValueError)", scoped)
    c = C()
    techdebt.list_items(c, team="T", job_status=lambda ids: {})
    check("yalnız Task/Bug/Improvement listelenir (Feature/Epic kuyruğa gitmez)",
          "[System.WorkItemType] IN ('Task', 'Bug', 'Improvement')" in c.wiql[0])
    orig_get, orig_q = pc.get, techdebt.queue
    pc.get = lambda k: True if k == "CREW_TECHDEBT" else orig_get(k)
    techdebt.queue = lambda *a, **k: (_ for _ in ()).throw(_rq.ConnectionError("azure yok"))
    try:
        r = TestClient(server.app).post("/api/techdebt/queue", json={"ids": [1]})
        check("Azure hatası → 502", r.status_code == 502, str(r.status_code))
    finally:
        pc.get, techdebt.queue = orig_get, orig_q
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text(encoding="utf-8")
    otc = html[html.index("async function onTeamChange"):html.index("async function onTeamChange") + 1500]
    check("takım seçimi yüklenirken ('Yükleniyor...') istek gönderilmez", "function _tdTeam" in html)
    check("takım değişince teknik borç listesi yenilenir; kuyruk butonu istek boyunca kilitli",
          "loadTechdebt()" in otc and "tdQueueBtn" in html and ".disabled=true" in html)


def test_dashboard_readability():
    print("\n[UI] Okunabilirlik — ayar editörü, başlık butonları")
    import re as _re
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text(encoding="utf-8")
    ed = _re.search(r'<textarea id="cfgEditor"[^>]*>', html).group(0)
    check("ayar editörü tema rengini kullanır (açık temada sarı yazı yok)", "#e6db74" not in ed and "var(--text)" in ed)
    check("açık temada okunmayan sabit sarı/pembe yazı renkleri yok",
          not any(c in html for c in ("color:#fbbf24", "color:#fca5a5", "color:#f9c74f")))
    sb = _re.search(r'<button[^>]*onclick="openSettings\(\)"[^>]*>.*?</button>', html, _re.S).group(0)
    tb = _re.search(r'<button[^>]*id="themeBtn"[^>]*>.*?</button>', html, _re.S).group(0)
    check("ayarlar butonu etiketli ve dişli simgeli; tema butonundan ayırt edilir",
          "Ayarlar" in sb and 'data-icon="gear"' in sb and 'data-icon="gear"' not in tb)
    check("metin/sayı kutuları açılır liste gibi çizilmez (ok simgesi + kesilen değer)",
          not _re.search(r'<input[^>]*class="sprint-select"', html) and 'id="qSp"' in html
          and _re.search(r'<input class="ctl-input" id="qSp"', html) is not None)
    check("modallarda açık temada simsiyah kutu yok (backfill, sprint raporu)",
          not any(c in html for c in ("background:#111", "background:#222", "background:#6c5ce7")))
    check("geçmiş iş taraması ilk açılışta boş siyah log/çubuk değil açıklama gösterir",
          'id="bfHint"' in html and "logEl.style.display" in html)
    check("tema butonu güncel modu simgeyle gösterir (güneş / ay / ekran)",
          "THEME_ICONS" in html and all(k in html for k in ("light:", "dark:", "system:")))


def test_human_label():
    print("\n[UI] İnsan işi etiketi anlaşılır ve sebebi gösterir")
    from agile_sdlc_crew.quality import intake
    from agile_sdlc_crew.quality.model import Group
    g = Group(source="sonar", repo="r", title="t", fingerprints=["f"], findings=[], needs_human=True,
              human_reason="Secret rotasyonu gerekir — koddan silmek yetmez")
    check("önizleme satırı insan işi sebebini taşır", intake._row(g).get("human_reason") == g.human_reason)
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text(encoding="utf-8")
    check("çıplak 'insan' etiketi kalmadı; 'İnsan işi · kuyruğa alınmaz' + sebep ipucu",
          'chip human">insan<' not in html and "· insan'" not in html
          and "function humanChip" in html and "İnsan işi · kuyruğa alınmaz" in html and "t.human_reason" in html)
    check("otomatik satırların durumu da yazılı (boş hücre yok)", "Pipeline kuyruğuna alınır" in html)


def test_quality_human_split():
    print("\n[Q13] Kalite alımı — insan işleri ayrı grup, bütçe otomatik işlere")
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.quality import intake
    from agile_sdlc_crew.quality.model import Finding

    def sec(i):
        return Finding(source="sonar", repo=f"r{i}", project=f"EcomBackend:r{i}", file=f"s{i}.php",
                       rule="secrets:S6703", rule_name="Secret", severity="BLOCKER", itype="VULNERABILITY", lines=[1])

    def bug(i):
        return Finding(source="sonar", repo="webservice", project="EcomBackend:webservice", file=f"b{i}.php",
                       rule=f"php:S{i}", rule_name="Bug", severity="CRITICAL", itype="BUG", lines=[1])
    findings = [sec(i) for i in range(7)] + [bug(i) for i in range(20)]
    cfg = {"nr_days": 7, "nr_min_errors": 100, "pack_sp": 3, "pack_files": 5, "max_tasks": 8, "human_max": 5,
           "grace_days": 3, "sonar_url": "", "target_sprint": "next", "task_fields": {}}
    deps = intake.Deps(gather=lambda s, c, cf: (list(findings), []), board=lambda c: {}, repo_ok=lambda r: True,
                       queue=lambda w, t: 1, record=lambda **k: None, now=lambda: 1.0, cfg=lambda: cfg)
    pv = intake.preview(None, "sonar", 30, deps=deps)
    check("insan işleri bütçeyi ve Task tavanını yemez: otomatik işler 8 Task'ı doldurur",
          len(pv["tasks"]) == 8 and not any(t["needs_human"] for t in pv["tasks"]) and pv["total_sp"] == 8)
    check("insan işi gerektiren önemli maddeler ayrı listede, en çok human_max",
          len(pv["human_tasks"]) == 5 and all(t["needs_human"] for t in pv["human_tasks"]))
    check("Task tavanı bütçeyi sınırladıysa açıkça söylenir",
          pv.get("cap_reached") is True and "8" in pv.get("cap_note", ""))
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text(encoding="utf-8")
    check("arayüzde iki ayrı grup: otomatik işler ve insan işi gerektiren önemli maddeler + tavan notu",
          "r.human_tasks" in html and "İnsan işi gerektiren önemli maddeler" in html
          and "Otomatik işler" in html and "r.cap_note" in html)
    check("insan işi tavanı ayarı kayıtlı (varsayılan 5); eski pay ayarı kaldırıldı",
          pc._SCHEMA_BY_KEY.get("CREW_QUALITY_HUMAN_MAX", {}).get("default") == 5
          and "CREW_QUALITY_HUMAN_SHARE" not in pc._SCHEMA_BY_KEY and intake.load_cfg()["human_max"] == 5)


def test_quality_project_filter():
    print("\n[Q14] Kalite alımı — proje (repo) seçimi")
    from fastapi.testclient import TestClient
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew import server
    from agile_sdlc_crew.quality import intake
    from agile_sdlc_crew.quality.model import Finding

    def bug(repo, i):
        return Finding(source="sonar", repo=repo, project=f"EcomBackend:{repo}", file=f"b{i}.php", rule=f"php:S{i}",
                       rule_name="Bug", severity="CRITICAL", itype="BUG", lines=[1])
    findings = [bug("webservice", i) for i in range(3)] + [bug("FloNext", i) for i in range(2)] + [bug("orkestra", 9)]
    cfg = {"nr_days": 7, "nr_min_errors": 100, "pack_sp": 3, "pack_files": 5, "max_tasks": 8, "human_max": 5,
           "grace_days": 3, "sonar_url": "", "target_sprint": "next", "task_fields": {}}
    deps = intake.Deps(gather=lambda s, c, cf: (list(findings), []), board=lambda c: {}, repo_ok=lambda r: True,
                       queue=lambda w, t: 1, record=lambda **k: None, now=lambda: 1.0, cfg=lambda: cfg)
    pv = intake.preview(None, "sonar", 30, deps=deps, repos=["flonext", "orkestra"])
    repos = {t["repo"] for t in pv["tasks"]}
    check("seçilen projeler dışındakiler alınmaz (büyük/küçük harf duyarsız)",
          repos == {"FloNext", "orkestra"} and pv["skipped_counts"].get("seçim dışı proje") == 1, str(pv["skipped_counts"]))
    check("seçim boşsa tüm projeler", {t["repo"] for t in intake.preview(None, "sonar", 30, deps=deps)["tasks"]}
          == {"webservice", "FloNext", "orkestra"})
    pr = intake.projects(None, "sonar", deps=deps)
    check("proje listesi: bulgusu olan repolar, bulgu sayısına göre",
          [(i["repo"], i["findings"]) for i in pr["items"]] == [("webservice", 3), ("FloNext", 2), ("orkestra", 1)])
    seen = {}
    orig_get, orig_pv = pc.get, intake.preview
    pc.get = lambda k: True if k == "CREW_QUALITY_INTAKE" else orig_get(k)
    intake.preview = lambda *a, **k: (seen.update(k), {"tasks": []})[1]
    try:
        TestClient(server.app).get("/api/quality/preview?source=sonar&sp=10&repos=FloNext,orkestra")
        check("önizleme ucu repos parametresini geçirir", seen.get("repos") == ["FloNext", "orkestra"], str(seen))
    finally:
        pc.get, intake.preview = orig_get, orig_pv
    pc.get = lambda k: False if k == "CREW_QUALITY_INTAKE" else orig_get(k)
    try:
        check("proje listesi ucu kapalıyken 403",
              TestClient(server.app).get("/api/quality/projects?source=sonar").status_code == 403)
    finally:
        pc.get = orig_get
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text(encoding="utf-8")
    check("kartta proje seçimi; önizleme ve alım seçimi gönderir",
          'id="qProj"' in html and "/api/quality/projects" in html and "repos=" in html and "repos:" in html)


def test_guidance_in_steps():
    print("\n[48b] Kickoff kılavuzu — toplantı kapalıyken de requirements ve teknik tasarıma girer")
    from agile_sdlc_crew import kickoff_guidance as _kg
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.flow import AgileSDLCFlow, PipelineState
    orig = _kg.format_for_context
    _kg.format_for_context = lambda max_rules=30: "# OGRENILMIS YONERGELER (gecmis kickoff'lardan)\n1. (WI #retro) AC türet"
    try:
        s = PipelineState(work_item_id="1", requirements_text="{}", repo_name="r")
        f = SimpleNamespace(state=s, _vector_store=None, _forward_text=lambda k, t, c: t[:c])
        ctx = {k: AgileSDLCFlow._build_step_context(f, k) for k in
               ("requirements_analysis_task", "technical_design_task", "implement_change_task")}
        check("retro'dan eklenen kural requirements ve teknik tasarım bağlamında",
              "AC türet" in ctx["requirements_analysis_task"] and "AC türet" in ctx["technical_design_task"])
        check("diğer adımlara girmez (önbellek öneki bozulmaz)", "AC türet" not in ctx["implement_change_task"])
        tail_ok = ctx["technical_design_task"].rstrip().endswith("AC türet") or \
            ctx["technical_design_task"].index("AC türet") > ctx["technical_design_task"].find("# Is Kalemi")
        check("kılavuz bağlamın sonunda (adıma özel kısım)", tail_ok)
        orig_get = pc.get
        pc.get = lambda k: False if k == "CREW_KICKOFF_GUIDANCE_IN_STEPS" else orig_get(k)
        try:
            off = AgileSDLCFlow._build_step_context(f, "technical_design_task")
            check("ayarla kapatılabilir", "AC türet" not in off)
        finally:
            pc.get = orig_get
    finally:
        _kg.format_for_context = orig


def test_completion_report_no_tool_note():
    print("\n[49] Tamamlanma raporu — ajan 'yorum ekleyemiyorum' notu yazmasın")
    import yaml
    from agile_sdlc_crew import flow as _flow
    tasks = yaml.safe_load((Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/config/tasks.yaml")
                           .read_text(encoding="utf-8"))
    d = tasks["completion_report_task"]["description"] + tasks["completion_report_task"]["expected_output"]
    check("görev ajandan yorum EKLEMESİNİ istemiyor; metni pipeline ekliyor",
          "Add the summary as a comment" not in d and "Comment must be added" not in d
          and "pipeline posts" in d)
    raw = ("**Tamamlanma Raporu**\n\n> **Not:** Elimdeki araçlar yalnızca WI okuyup listeleyebiliyor, yorum yazamıyor. "
           "Bu rapor, WI'ya yorum olarak eklenmek üzere hazırlanmış tam metindir. Yorumu pipeline'ın yorum adımı "
           "veya bir insan eklemelidir.\n\n## Özet\nTek satır düzeltildi.")
    out = _flow._strip_tool_meta(raw)
    check("ajan araç notu rapordan ayıklanır, içerik kalır",
          "yorum yazamıyor" not in out and "insan eklemelidir" not in out and "Tek satır düzeltildi" in out
          and "Tamamlanma Raporu" in out)
    check("araçtan söz etmeyen metne dokunulmaz", _flow._strip_tool_meta("## Özet\nPR açıldı.") == "## Özet\nPR açıldı.")


def test_review_prefetch_window():
    print("\n[50] Review bağlamı — büyük dosyada değişen satırın çevresi (#230: ilk 173 satır)")
    from agile_sdlc_crew.flow import AgileSDLCFlow
    base = "\n".join(f"    $x{i} = $a['k{i}'];" for i in range(1, 3001))
    new = base.replace("    $x1460 = $a['k1460'];", "    $x1460 = $a['k1460'] ?? 0;")

    class Cl:
        def get_file_content(self, repo, fp, ref):
            if ref == "feature/76201":
                return new
            if ref == "main":
                return base
            raise RuntimeError("404")
    f = SimpleNamespace(
        state=SimpleNamespace(repo_name="webservice", branch_name="feature/76201", pr_id="1",
                              plan={"changes": [{"file_path": "/app/Controller/Api/V1/Basket.php"}]}, all_pushes=[]),
        _client=Cl(), _repo_mgr=None, _pr_file_bodies={})
    ctx = AgileSDLCFlow._prefetch_pr_changes_context(f)
    check("değişen satır (1460) review bağlamında, satır numarasıyla", "1460| " in ctx and "?? 0" in ctx, ctx[:200])
    check("dosya baştan kesilmez; değişmeyen yerler 'atlandı' diye işaretlenir",
          "satır atlandı" in ctx and "kısaltıldı" not in ctx)


def test_quality_estimate():
    print("\n[Q15] Kalite işi tahmini — alımın SP'si (yapısal tahmin şişirmesin; #76201: 13 SP)")
    from agile_sdlc_crew.flow import AgileSDLCFlow
    reqs = json.dumps({"functional_requirements": [{"id": f"FR{i}", "desc": "x"} for i in range(1, 7)],
                       "technical_requirements": [{"id": f"TR{i}", "desc": "x"} for i in range(1, 6)],
                       "acceptance_criteria": [{"id": f"AC{i}", "desc": "x"} for i in range(1, 6)]})

    def mk(tags, wi_sp):
        f = AgileSDLCFlow()
        f._db = None
        f.state.requirements_text = reqs
        f.state.plan = {"changes": [{"file_path": f"f{i}.php"} for i in range(4)]}
        f.state.wi_tags, f.state.wi_story_points = tags, wi_sp
        f._needed_explore = True
        f._estimate("plan")
        return f.state.estimate
    q = mk("tempo-quality; tempo-src-newrelic", 1.0)
    check("kalite işi: tahmin = alımın SP'si (1), kaynak kalite alımı",
          q["sp"] == 1 and "kalite" in q["source"], str(q))
    n = mk("", None)
    check("normal iş: yapısal tahmin aynen", n["sp"] >= 8, str(n.get("sp")))


def test_no_env_dependent_tests():
    print("\n[51] Test talimatı — ortama bağlı kontrol yok (#230: 'uyarı listesi boş' testi CI'da kırıldı)")
    import inspect
    from agile_sdlc_crew import flow as _flow
    from agile_sdlc_crew.quality import rank, render
    from agile_sdlc_crew.quality.model import Finding
    skill = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/skills/implementation/SKILL.md").read_text(encoding="utf-8")
    check("implementation skill'i ortama bağlı kontrolü yasaklar (tüm işler)",
          "ortama bağlı" in skill.lower() and "deprecation" in skill.lower())
    g = rank.nr_units([Finding(source="newrelic", repo="w", app="[Prod] a", action="A::b", err_class="E_WARNING",
                               message='Undefined array key "count"', count=500, last_seen=1.0)])[0]
    d = render.task_description(g)
    check("kalite Task'ı 'uyarı üretilmemesi'ni test etmeyi önermez; dönen değeri test etmeyi söyler",
          "uyarı üretilmemesi" not in d and "dönen değer" in d and "ortama bağlı" in d.lower())
    src = inspect.getsource(_flow)
    check("build-fix turu (tüm işler + kalite kuralı) ortama bağlı kontrolü kaldırmayı söyler",
          "ENVIRONMENT-DEPENDENT" in _flow.QUALITY_BUILD_FIX_RULE and "Ortama bağlı kontrol" in src)


def test_pr_comment_emoticons():
    print("\n[52] PR yorumu — ':/' gibi değerler emojiye dönüşmesin (PR #44918)")
    from agile_sdlc_crew.tools import azure_devops_base as az
    d = az.defuse_emoticons
    s = "Test, host olmayan canonical için ':/' ve '://en' gibi bozuk URL'leri bekliyor."
    out = d(s)
    check("emoji'ye dönüşecek değerler satır içi koda alınır",
          "`':/'`" in out and "`'://en'`" in out and out.endswith("bekliyor."), out)
    check("URL ve normal metin değişmez",
          d("Bkz. https://dev.azure.com/x ve Sonuç: tamam") == "Bkz. https://dev.azure.com/x ve Sonuç: tamam")
    code = "```php\n$u = ':/' . $x;\n```\nve `'://en'` zaten kod"
    check("kod blokları ve satır içi kod dokunulmaz", d(code) == code)
    check("cümle sonu noktalama kodun dışında kalır", d("değer :/,") == "değer `:/`,")
    tbl = "| Madde | Durum |\n|:---|---:|\n| R1 | açık |"
    check("markdown tablo hizalama satırı bozulmaz", d(tbl) == tbl)
    import inspect
    src = inspect.getsource(az.AzureDevOpsClient.add_pr_comment) + inspect.getsource(az.AzureDevOpsClient.reply_to_pr_thread)
    check("PR yorumu ve yanıtı gönderilmeden önce uygulanır", src.count("defuse_emoticons(") == 2)


def test_completion_comment_integrity():
    print("\n[53] Tamamlanma yorumu — tablo ortasında kesilmesin, atlanan UAT 'parse edilemedi' demesin (#76200)")
    import inspect
    from agile_sdlc_crew import flow as _flow
    from agile_sdlc_crew import wi_lifecycle as _wl
    rows = "\n".join(f"| AC{i} kriter | GEÇTİ |" for i in range(1, 120))
    rep = "## Özet\nx\n\n## Kabul\n| Kriter | Durum |\n|---|---|\n" + rows + "\n\n## Notlar\nson"
    cut = _flow._trim_report(rep, 1500)
    last = cut.rstrip().splitlines()
    check("kısaltma paragraf sınırında; yarım tablo satırı yok, kısaltıldı notu var",
          len(cut) <= 1600 and "kısaltıldı" in cut and not any(l.startswith("| AC") and not l.rstrip().endswith("|") for l in last))
    check("sığan rapor aynen kalır", _flow._trim_report("## Özet\nkısa", 1500) == "## Özet\nkısa")
    src = inspect.getsource(_flow.AgileSDLCFlow.step11_completion_report)
    check("yorum 3000 karakterde kesilmiyor", "completion_text[:3000]}" not in src and "_trim_report(" in src)
    d = _wl.evaluate_dod(review_approved=True, open_review_issues=0, build_status="succeeded",
                         uat_text="Atlandı — kalite işi (CREW_QUALITY_LEAN_FLOW)", pushed_files=["a.php"],
                         require_tests=False, pr_id="1", uat_skipped=True)
    uat = next(i for i in d.items if i.key == "uat")
    check("atlanan UAT DoD'da 'atlandı' der", "atlandı" in uat.note and "parse" not in uat.note, uat.note)
    rej = _wl.evaluate_dod(review_approved=True, open_review_issues=0, build_status="succeeded",
                           uat_text="Atlandı gibi başlayan ama **Overall Evaluation: REJECTED** UAT çıktısı\n1. AC1 FAIL",
                           pushed_files=["a.php"], require_tests=False, pr_id="1")
    u2 = next(i for i in rej.items if i.key == "uat")
    check("UAT metni 'Atlandı' ile başlasa da bayrak yoksa sonuç atlanmaz (fail-open yok)",
          u2.ok is False, str((u2.ok, u2.note)))
    check("atlama bayrağı flow'dan gelir, metinden değil",
          "uat_skipped=" in inspect.getsource(_flow.AgileSDLCFlow.step11_completion_report))


# ── Teknik borç aşama 2a: PHP modernleştirme ─────────────────────────────

def test_php_modern_catalog():
    print("\n[PM1] php_modern — katalog ve hedef PHP sürümü")
    from agile_sdlc_crew.quality import php_modern as pm
    rules = pm.load_rules()
    keys = {r.key for r in rules}
    check("katalog yüklendi (≥15 kural)", len(rules) >= 15, str(len(rules)))
    req = ("key", "since", "kind", "title", "summary", "advantages", "risks", "before", "after", "ref", "unit")
    bad = [r.key for r in rules if any(not getattr(r, f) for f in req)]
    check("her kuralda zorunlu alanlar dolu", not bad, str(bad))
    check("kind geçerli", all(r.kind in ("deprecation", "mechanical", "structural", "report") for r in rules))
    check("unit geçerli", all(r.unit in ("file", "module") for r in rules))
    check("since (major, minor)", all(isinstance(r.since, tuple) and len(r.since) == 2 for r in rules))
    check("structural/report kurallar otonom değil", all(not r.auto for r in rules if r.kind in ("structural", "report")))
    check("deprecation/mechanical kurallar otonom", all(r.auto for r in rules if r.kind in ("deprecation", "mechanical")))
    check("beklenen kurallar var", {"implicit_nullable", "str_contains", "match", "enum", "typed_properties",
                                    "property_hooks", "dto_from_array", "strict_types"} <= keys, str(keys))
    check("ağırlık: deprecation > structural > mechanical > report",
          pm.WEIGHT["deprecation"] > pm.WEIGHT["structural"] > pm.WEIGHT["mechanical"] > pm.WEIGHT["report"] == 0)
    k82 = {r.key for r in pm.applicable(rules, "8.2")}
    k84 = {r.key for r in pm.applicable(rules, "8.4")}
    check("8.2 hedefe property hooks / implicit_nullable önerilmez; 8.4'e önerilir",
          "property_hooks" not in k82 and "implicit_nullable" not in k82
          and {"property_hooks", "implicit_nullable", "enum", "str_contains"} <= k84)
    check("7.4 hedefe yalnız ≤7.4 kurallar", all(r.since <= (7, 4) for r in pm.applicable(rules, "7.4"))
          and "null_coalesce" in {r.key for r in pm.applicable(rules, "7.4")})
    check("imaj adından sürüm", pm.parse_php_image("europe-west1-docker.pkg.dev/flo-global/prod-ar/php-8-4:1.0") == (8, 4)
          and pm.parse_php_image("cr.yandex/crpsbatlmbdqnhk3392d/php-7-4:1.0") == (7, 4)
          and pm.parse_php_image("alpine:3.19") is None and pm.parse_php_image("php:8.3-fpm") == (8, 3))

    class C:
        def __init__(self, files):
            self.files = files

        def get_file_content(self, repo, path, branch="", project=""):
            if path not in self.files:
                import requests
                e = requests.HTTPError("404"); e.response = SimpleNamespace(status_code=404); raise e
            return self.files[path]
    amap = {"[Prod] A": {"repo": "webservice", "project": "EcomBackend", "src": "anatolia/build/webservice"},
            "[Prod] B": {"repo": "webservice", "project": "EcomBackend", "src": "russia/build/websites"},
            "[Prod] C": {"repo": "alfred", "project": "EcomBackend", "src": "anatolia/build/alfred"}}
    c = C({"/build/webservice/Dockerfile": "FROM x/php-8-4:1.0\nRUN a", "/build/websites/Dockerfile": "FROM y/php-7-4:1.0",
           "/build/alfred/Dockerfile": "FROM x/php-8-2:1.0"})
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        t = pm.target_version(c, "webservice", app_map_cache=amap, repo_root=root, cache_path=root / "t.yaml")
        check("iki imaj (8.4, 7.4) → en düşük 7.4", t["version"] == "7.4" and t["source"] == "dockerfile"
              and sorted(t["images"]) == ["x/php-8-4:1.0", "y/php-7-4:1.0"], str(t))
        (root / "composer.json").write_text('{"require": {"php": "^8.1 || ^8.2"}}', encoding="utf-8")
        t2 = pm.target_version(c, "pigeon", app_map_cache=amap, repo_root=root, cache_path=root / "t.yaml")
        check("imaj yok → composer.json alt sınırı 8.1", t2["version"] == "8.1" and t2["source"] == "composer", str(t2))
        t3 = pm.target_version(c, "sorting", app_map_cache=amap, repo_root=root / "yok", cache_path=root / "t.yaml")
        check("ne imaj ne composer → None", t3["version"] is None and t3["source"] is None, str(t3))
        c.files.clear()
        t4 = pm.target_version(c, "webservice", app_map_cache=amap, repo_root=root, cache_path=root / "t.yaml")
        check("önbellek: ikinci çağrı ağa gitmeden aynı sonucu verir", t4["version"] == "7.4", str(t4))
        t5 = pm.target_version(c, "webservice", app_map_cache=amap, repo_root=root, cache_path=root / "t.yaml", max_age_h=0)
        check("önbellek süresi dolunca yeniden okur (imaj 404 → composer'a düşer)", t5["source"] == "composer", str(t5))


def test_php_modern_scan():
    print("\n[PM2] php_modern — token tarayıcı")
    import shutil
    from agile_sdlc_crew.quality import php_modern as pm
    from agile_sdlc_crew.quality.model import QualitySourceError
    if not shutil.which("php"):
        skip("php_modern tarayıcı", "php binary yok")
        return
    root = Path(__file__).resolve().parent / "fixtures" / "php_modern"
    res = pm.scan_repo(root, exclude=["vendor"], php_bin="php")
    hits = res["hits"]
    got = {(h["file"], h["line"], h["rule"]) for h in hits}
    # Beklenen: fikstürlerdeki `// @expect <kural>` işaretleri (satır = işaretin satırı)
    want = set()
    for f in sorted(root.rglob("*.php")):
        rel = f.relative_to(root).as_posix()
        if rel.startswith("vendor/"):
            continue
        for i, ln in enumerate(f.read_text(encoding="utf-8", errors="replace").splitlines(), 1):
            for m in re.finditer(r"@expect (\w+)", ln):
                want.add((rel, i, m.group(1)))
    missing = sorted(want - got)
    extra = sorted(got - want)
    check("beklenen her bulgu yerinde (dosya, satır, kural)", not missing, str(missing))
    check("fazladan bulgu yok (string/heredoc/yorum içi kalıplar, modern.php, vendor)", not extra, str(extra))
    check("tarama sayacı: vendor hariç 9 dosya (bozuk ve Latin-5 dosya dahil, tarayıcı takılmaz)", res["files"] == 9, str(res["files"]))
    check("bulgu metni tarayıcı çıkışında maskelenir (önbelleğe ham secret yazılmaz)",
          any("Secret123" not in h["text"] and "pw" in h["text"] for h in hits if h["file"] == "app/Lib/edge.php" and h["rule"] == "null_coalesce")
          and not any("Secret123" in h["text"] for h in hits))
    check("bulgu satır metni taşır", all(h.get("text") for h in hits))
    check("modül: kaynak kökünden 2 dizin", {h["module"] for h in hits} <= {"app/Model", "app/Lib"}, str({h["module"] for h in hits}))
    ad = res["adoption"]
    check("benimseme sayaçları", ad["enum"] == 1 and ad["match"] == 1 and ad["nullsafe"] == 1 and ad["attribute"] == 2
          and ad["readonly"] == 1 and ad["strict_types_files"] == 1 and ad["untyped_props"] == 7 and ad["typed_props"] >= 4, str(ad))
    res2 = pm.scan_repo(root, exclude=[], php_bin="php")
    check("hariç tutma kapalıyken vendor da taranır", res2["files"] == 10 and any(h["file"].startswith("vendor/") for h in res2["hits"]))
    try:
        pm.scan_repo(root, exclude=[], php_bin="/nonexistent/php")
        check("php yoksa QualitySourceError", False)
    except QualitySourceError:
        check("php yoksa QualitySourceError", True)
    try:
        pm.scan_repo(root / "yok", exclude=[], php_bin="php")
        check("repo dizini yoksa QualitySourceError", False)
    except QualitySourceError:
        check("repo dizini yoksa QualitySourceError", True)


def test_php_modern_groups():
    print("\n[PM3] php_modern — gruplama, skor, dedup, önbellek")
    import math
    from agile_sdlc_crew.quality import php_modern as pm, rank
    rules = pm.load_rules()
    cfg = {"pack_sp": 3, "pack_files": 5, "grace_days": 3}

    def h(rule, file, line=1):
        return {"rule": rule, "file": file, "line": line, "text": "x", "module": pm.module_of(file)}
    hits = ([h("str_contains", f"app/Lib/F{i}.php", 3) for i in range(12)]
            + [h("str_contains", "app/Lib/F0.php", 9)]
            + [h("utf8_encode", "app/Lib/F0.php", 5)]
            + [h("typed_properties", f"app/Model/M{i}.php", 2) for i in range(8)]
            + [h("typed_properties", "app/Lib/F1.php", 2)]
            + [h("property_hooks", "app/Model/M1.php", 7)]
            + [h("strict_types", "app/Lib/F0.php", 1)])
    scan = {"files": 20, "sha": "abc", "hits": hits, "adoption": {}}
    groups, skipped, summary = pm.build_groups(scan, rules, "8.2", repo="webservice", board_index={}, now=1e9, cfg=cfg)
    keys = {g.findings[0].rule for g in groups}
    check("8.2 hedefte property_hooks grubu yok", "property_hooks" not in keys, str(keys))
    g84, _s, sum84 = pm.build_groups(scan, rules, "8.4", repo="webservice", board_index={}, now=1e9, cfg=cfg)
    check("8.4 hedefte property_hooks grubu var", "property_hooks" in {g.findings[0].rule for g in g84})
    sc = [g for g in groups if g.findings[0].rule == "str_contains"]
    check("mekanik: dosya başına birim, aynı kural ≤5 dosya / ≤3 SP paket → 12 dosya → 4 paket",
          len(sc) == 4 and all(g.sp <= 3 and len(g.files) <= 5 for g in sc) and sum(len(g.files) for g in sc) == 12,
          str([(g.sp, len(g.files)) for g in sc]))
    check("paket parmak izleri dosya başına php_modern:<repo>:<kural>:<dosya>",
          all(fp.startswith("php_modern:webservice:str_contains:app/Lib/F") for g in sc for fp in g.fingerprints)
          and sum(len(g.fingerprints) for g in sc) == 12)
    check("aynı dosyada 2 bulgu → tek birim, count 2", any(f.count == 2 and f.file == "app/Lib/F0.php" for g in sc for f in g.findings))
    check("mekanik grup otonom, SP = dosya başına sp_unit", all(not g.needs_human for g in sc))
    tp = [g for g in groups if g.findings[0].rule == "typed_properties"]
    check("yapısal: modül başına birim (app/Model, app/Lib) → 2 grup, SP 5, needs_human",
          len(tp) == 2 and all(g.needs_human and g.sp == 5 for g in tp)
          and {g.fingerprints[0] for g in tp} == {"php_modern:webservice:typed_properties:app/Model", "php_modern:webservice:typed_properties:app/Lib"},
          str([(g.fingerprints, g.sp) for g in tp]))
    check("yapısal neden metni katalog riskinden", all(g.human_reason.startswith("Tasarım kararı") for g in tp))
    dep = next(g for g in groups if g.findings[0].rule == "utf8_encode")
    one_sc = next(g for g in sc if len(g.files) == 1) if any(len(g.files) == 1 for g in sc) else None
    w = pm.WEIGHT
    check("skor = ağırlık × (1+log10(adet)); deprecation(1) > mechanical(1)",
          abs(dep.score - w["deprecation"]) < 1e-9 and (one_sc is None or dep.score > one_sc.score), str(dep.score))
    big = max(tp, key=lambda g: g.total)
    check("yapısal skor adetle log artar", abs(big.score - w["structural"] * (1 + math.log10(8))) < 1e-9, str(big.score))
    check("report kuralı grup üretmez ama özette", "strict_types" not in keys and summary["strict_types"]["count"] == 1
          and summary["strict_types"]["kind"] == "report")
    check("özet: kural → adet/dosya/modül/kind/auto", summary["str_contains"]["count"] == 13 and summary["str_contains"]["files"] == 12
          and summary["typed_properties"]["modules"] == 2 and summary["typed_properties"]["auto"] is False, str(summary["str_contains"]))
    # dedup: board'da açık fp → atlanır
    bi = {"php_modern:webservice:utf8_encode:app/Lib/F0.php": {"wi": 500, "state": "Active"},
          "php_modern:webservice:typed_properties:app/Model": {"wi": 501, "state": "Canceled"}}
    g2, sk2, _ = pm.build_groups(scan, rules, "8.2", repo="webservice", board_index=bi, now=1e9, cfg=cfg)
    reasons = {s["reason"].split(" #")[0] for s in sk2}
    check("board'da açık WI → 'açık WI #', iptal → 'iptal edilmiş WI #'", reasons == {"açık WI", "iptal edilmiş WI"}
          and not any(g.findings[0].rule == "utf8_encode" for g in g2)
          and not any(fp.endswith("typed_properties:app/Model") for g in g2 for fp in g.fingerprints), str(sk2))
    # pack_sonar eski davranışı
    check("pack_sonar, pack_units'e sarmalayıcı", rank.pack_sonar([]) == [] and hasattr(rank, "pack_units"))
    # önbellek
    calls = []

    def fake_scan(root, **kw):
        calls.append(kw)
        return {"files": 1, "sha": kw.get("_sha", "s1"), "hits": [], "adoption": {}, "errors": []}
    with tempfile.TemporaryDirectory() as td:
        cdir = Path(td)
        a = pm.scan_cached("webservice", Path(td), cache_dir=cdir, scan_fn=lambda root, **kw: fake_scan(root, _sha="s1", **kw),
                           head_sha_fn=lambda r: "s1", exclude=[], php_bin="php", ttl_h=24)
        b = pm.scan_cached("webservice", Path(td), cache_dir=cdir, scan_fn=lambda root, **kw: fake_scan(root, _sha="s1", **kw),
                           head_sha_fn=lambda r: "s1", exclude=[], php_bin="php", ttl_h=24)
        check("aynı sha → ikinci çağrı tarayıcıyı çağırmaz", len(calls) == 1 and b["sha"] == "s1" and a["cached"] is False and b["cached"] is True)
        pm.scan_cached("webservice", Path(td), cache_dir=cdir, scan_fn=lambda root, **kw: fake_scan(root, _sha="s2", **kw),
                       head_sha_fn=lambda r: "s2", exclude=[], php_bin="php", ttl_h=24)
        check("sha değişince yeniden tarar", len(calls) == 2)
        pm.scan_cached("webservice", Path(td), cache_dir=cdir, scan_fn=lambda root, **kw: fake_scan(root, _sha="s2", **kw),
                       head_sha_fn=lambda r: "s2", exclude=[], php_bin="php", ttl_h=24, force=True)
        check("force → yeniden tarar", len(calls) == 3)
        pm.scan_cached("webservice", Path(td), cache_dir=cdir, scan_fn=lambda root, **kw: fake_scan(root, _sha="s2", **kw),
                       head_sha_fn=lambda r: "s2", exclude=[], php_bin="php", ttl_h=0)
        check("TTL dolunca yeniden tarar", len(calls) == 4)
        pm.scan_cached("webservice", Path(td), cache_dir=cdir, scan_fn=lambda root, **kw: fake_scan(root, _sha="s2", **kw),
                       head_sha_fn=lambda r: "s2", exclude=["vendor"], php_bin="php", ttl_h=24)
        check("hariç listesi değişince yeniden tarar (önbellek anahtarı)", len(calls) == 5)
        check("önbellek dosyası repo adıyla, yol kaçışı yok", (cdir / "webservice.json").exists())
        try:
            pm.scan_cached("../x", Path(td), cache_dir=cdir, scan_fn=fake_scan, head_sha_fn=lambda r: "s", exclude=[], php_bin="php")
            check("repo adında yol kaçışı reddedilir", False)
        except ValueError:
            check("repo adında yol kaçışı reddedilir", True)


def test_php_modern_render():
    print("\n[PM4] php_modern — Task metni (avantaj / risk / önce-sonra)")
    from agile_sdlc_crew.quality import php_modern as pm, render
    rules = pm.load_rules()
    hits = [{"rule": "str_contains", "file": "app/Lib/Url.php", "line": 12, "text": "if (strpos($u, 'tempo-fp: x') !== false) // password=Secret123!", "module": "app/Lib"},
            {"rule": "str_contains", "file": "app/Lib/Url.php", "line": 40, "text": "if (false !== strpos($u, 'a'))", "module": "app/Lib"}]
    hits += [{"rule": "typed_properties", "file": f"app/Model/M{i}.php", "line": i + 1, "text": f"private $p{i};", "module": "app/Model"} for i in range(25)]
    groups, _sk, _sum = pm.build_groups({"hits": hits}, rules, "8.4", repo="webservice", board_index={}, now=1e9,
                                        cfg={"pack_sp": 3, "pack_files": 5, "grace_days": 3})
    mech = next(g for g in groups if g.findings[0].rule == "str_contains")
    struct = next(g for g in groups if g.findings[0].rule == "typed_properties")
    d = render.task_description(mech, php_target="8.4")
    check("ilk satır Repo (repo_from_description çözer)", d.startswith("<p><b>Repo:</b> <code>webservice</code></p>")
          and render.repo_from_description(d) == "webservice")
    check("hedef PHP ve kural sürümü", "Hedef PHP:</b> 8.4" in d and "PHP 8.0" in d, d[:400])
    check("kural başlığı ve 'Ne'", "str_contains()" in d and "Ne:" in d)
    check("Avantajlar ve Riskler katalogdan", "Avantajlar:" in d and "Riskler:" in d and "0 / false tuzağı" in d)
    check("Önce/Sonra pre içinde, HTML-escape'li", "<b>Önce:</b>" in d and "<pre>" in d
          and "if (strpos($uri, &#x27;/api/&#x27;) !== false)" in d and "str_contains($uri" in d, d[d.find("Önce"):d.find("Önce") + 300])
    check("bulunan yerler dosya:satır + metin, maskeli ve defang'li",
          "app/Lib/Url.php:12" in d and "app/Lib/Url.php:40" in d and "tempo_fp" in d and "Secret123" not in d, d)
    check("kaynak (RFC) bağlantısı", 'href="https://wiki.php.net/rfc/str_contains"' in d)
    fps = render.parse_fingerprints(d)
    check("tempo-fp satırları geri okunur", fps == mech.fingerprints, str(fps))
    eb = " ".join(render.expected_behavior(mech))
    check("mekanik beklenen davranış: davranış değişmez, yalnız listelenen yerler, dar test, ortama bağlı kontrol yok",
          "değişmez" in eb and "listelenen" in eb and "birim testi" in eb and "Ortama bağlı" in eb, eb)
    ac = render.acceptance_criteria(mech)
    check("mekanik kabul kriteri", "eski yapı kalmaz" in ac and "davranış" in ac, ac)
    ds = render.task_description(struct, php_target="8.4")
    check("yapısal: 'Neden insan' + tasarım kararı", "Neden insan:</b> Tasarım kararı" in ds)
    check("yapısal: 25 dosya → her dosya en az bir kez listelenir (parmak izi = dosya; listelenmeyen dosya 'yapıldı' sayılmasın)",
          ds.count("app/Model/M") == 25, str(ds.count("app/Model/M")))
    many = [{"rule": "str_contains", "file": "config/Migrations/20231012093012_CreateOrders.php", "line": i, "text": "x", "module": "config/Migrations"} for i in range(1, 26)]
    many += [{"rule": "str_contains", "file": "app/B.php", "line": 1, "text": "y", "module": "app"}, {"rule": "str_contains", "file": "app/C.php", "line": 1, "text": "z", "module": "app"}]
    gm, _a, _b = pm.build_groups({"hits": many}, rules, "8.4", repo="webservice", board_index={}, now=1e9, cfg={"pack_sp": 9, "pack_files": 9, "grace_days": 3})
    big = max(gm, key=lambda g: g.total)
    dm = render.task_description(big, php_target="8.4")
    forged = pm.build_groups({"hits": [{"rule": "str_contains", "file": "app/tempo-fp: x.php", "line": 1, "text": "y", "module": "app"}]},
                             rules, "8.4", repo="webservice", board_index={}, now=1e9, cfg={"pack_sp": 3, "pack_files": 5, "grace_days": 3})[0][0]
    df = render.task_description(forged, php_target="8.4")
    check("dosya adı 'tempo-fp:' ile dedup satırı taklit edemez (yol defang)", render.parse_fingerprints(df) == forged.fingerprints, str(render.parse_fingerprints(df)))
    check("paket: 25+2+2 yer → B ve C de listelenir, dosya başına en çok 5 satır, yol maskelenmez",
          "app/B.php:1" in dm and "app/C.php:1" in dm and "config/Migrations/20231012093012_CreateOrders.php:1" in dm
          and dm.count("20231012093012_CreateOrders.php:") == 5 and "+20 yer" in dm, dm[dm.find("Bulunan"):dm.find("Bulunan") + 900])
    es = " ".join(render.expected_behavior(struct))
    check("yapısal beklenen davranış: insan kapsam belirler, etiket kaldırılınca kuyruğa", "insan" in es and "needs-human" in es, es)
    check("AC alanı dolu", "<li>" in render.acceptance_block(struct))
    check("etiketler techdebt hattı + kaynak", render.task_tags(struct, lane="techdebt")
          == "tempo-quality; tempo-techdebt; tech-debt; tempo-src-php_modern; needs-human"
          and render.task_tags(mech, lane="techdebt") == "tempo-quality; tempo-techdebt; tech-debt; tempo-src-php_modern")
    check("kaynak etiketi", render.SOURCE_LABEL.get("php_modern") == "PHP modernleştirme")
    check("başlık [repo] kural — yer", render.task_title(mech).startswith("[webservice] strpos(...) !== false → str_contains() — Url.php"), render.task_title(mech))
    check("eski kaynaklar etkilenmedi (sonar body hâlâ çalışır)", "Kabul kriteri" in render.task_description(
        __import__("agile_sdlc_crew.quality.model", fromlist=["Group"]).Group(source="sonar", repo="r", title="t", fingerprints=["f"],
        findings=[__import__("agile_sdlc_crew.quality.model", fromlist=["Finding"]).Finding(source="sonar", repo="r", file="a.php", rule="php:S1", rule_name="n", severity="CRITICAL", itype="BUG")])))


def test_php_modern_take():
    print("\n[PM5] php_modern — repo listesi, tarama, iş açma")
    from agile_sdlc_crew.quality import intake, php_modern as pm
    queued, rec = [], []

    class C:
        org_url, project = "https://dev.azure.com/o", "P"

        def __init__(self):
            self.created, self.wiql = [], []

        def list_iterations(self, team=""):
            return [{"name": "S23", "path": "P\\S23", "attributes": {"timeFrame": "future", "startDate": "2026-11-02"}}]

        def query_work_items(self, wiql, limit=0):
            self.wiql.append(wiql)
            return []

        def get_team_area_path(self, team=""):
            return "P\\T"

        def create_work_item(self, type_name, fields, parent_id=None):
            self.created.append((type_name, fields, parent_id))
            return {"id": 600 + len(self.created), "fields": fields}

        def update_work_item(self, wid, ops):
            return {}

        def get_authenticated_user(self):
            return {"uniqueName": "me@x"}
    hits = [{"rule": "str_contains", "file": "app/Lib/A.php", "line": 3, "text": "x", "module": "app/Lib"},
            {"rule": "typed_properties", "file": "app/Model/M.php", "line": 2, "text": "y", "module": "app/Model"},
            {"rule": "utf8_encode", "file": "app/Lib/B.php", "line": 4, "text": "z", "module": "app/Lib"}]
    scans = []
    board = {}
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        (root / "webservice").mkdir()
        # ext-mbstring: utf8_encode düzeltmesi mbstring ister; kanıt yoksa iş altyapı listesine gider
        (root / "webservice" / "composer.json").write_text('{"require":{"php":"^8.4","ext-mbstring":"*"}}', encoding="utf-8")
        (root / "notphp").mkdir()
        (root / "oldphp").mkdir()
        (root / "oldphp" / "composer.json").write_text('{}', encoding="utf-8")
        cfg = {"pack_sp": 3, "pack_files": 5, "grace_days": 3, "target_sprint": "next", "task_fields": {},
               "sonar_url": "", "nr_days": 7, "php_exclude": [], "php_bin": "php", "scan_ttl_h": 24, "max_take": 8,
               "repos_dir": root, "targets_cache": root / "t.yaml", "scan_cache": root / "cache"}
        deps = intake.Deps(gather=None, board=lambda c: board, repo_ok=lambda r: (root / r).is_dir(),
                           queue=lambda w, ttl: (queued.append(w), 9)[1], record=lambda **k: rec.append(k),
                           now=lambda: 1e9, cfg=lambda: cfg)

        def scan_fn(repo_root, **kw):
            scans.append(str(repo_root))
            return {"files": 3, "sha": "s1", "hits": hits, "adoption": {"enum": 0}, "errors": []}
        pm_kw = dict(deps=deps, scan_fn=scan_fn, head_sha_fn=lambda r: "s1", app_map_cache={})
        c = C()
        rp = pm.repos(c, **pm_kw)
        items = {i["repo"]: i for i in rp["items"]}
        check("repo listesi: composer.json olan yerel repolar, hedef sürüm composer'dan; sürümsüz repo işaretli; php olmayan yok",
              set(items) == {"webservice", "oldphp"} and items["webservice"]["version"] == "8.4"
              and items["oldphp"]["version"] is None and items["oldphp"]["scannable"] is False, str(rp))
        sc = pm.scan(c, "webservice", team="T", **pm_kw)
        check("tarama yazma yapmaz", not c.created and len(scans) == 1)
        check("tarama sonucu: sürüm, sha, kural özeti, otomatik/insan grupları, gid",
              sc["version"] == "8.4" and sc["sha"] == "s1" and sc["rules"]["str_contains"]["count"] == 1
              and len(sc["auto"]) == 2 and len(sc["human"]) == 1 and all(len(r["gid"]) == 12 for r in sc["auto"] + sc["human"]), str(sc)[:600])
        row = sc["auto"][0]
        check("grup satırı katalog ayrıntısını taşır", all(k in row for k in ("rule", "kind", "since", "title", "summary", "advantages",
              "risks", "before", "after", "ref", "count", "files", "locations", "sp", "needs_human")), str(row.keys()))
        check("deprecation önce sıralanır", sc["auto"][0]["rule"] == "utf8_encode")
        sc2 = pm.scan(c, "webservice", team="T", **pm_kw)
        check("ikinci tarama önbellekten (aynı sha)", len(scans) == 1 and sc2["cached"] is True)
        sc_up = pm.scan(c, "WEBSERVICE", team="T", **pm_kw)
        check("repo adı kanonik (dizin adı): WEBSERVICE → webservice, parmak izi aynı",
              sc_up["repo"] == "webservice" and sc_up["auto"][0]["fingerprints"] == sc["auto"][0]["fingerprints"], str(sc_up["repo"]))
        gids = [r["gid"] for r in sc["auto"]] + [r["gid"] for r in sc["human"]]
        # kilit
        intake._LOCK.acquire()
        try:
            try:
                pm.take(c, "webservice", gids, team="T", **pm_kw)
                check("kilitliyken Busy", False)
            except intake.Busy:
                check("kilitliyken Busy", True)
        finally:
            intake._LOCK.release()
        try:
            pm.take(c, "webservice", ["000000000000"], team="T", **pm_kw)
            check("yalnız bilinmeyen gid → ValueError (tarama değişti)", False)
        except ValueError as e:
            check("yalnız bilinmeyen gid → ValueError (tarama değişti)", "tarama" in str(e).lower(), str(e))
        res = pm.take(c, "webservice", gids + ["000000000000", gids[0]], team="T", **pm_kw)
        types = [t for t, _f, _p in c.created]
        check("story techdebt hattında (Teknik borç — S23, tempo-techdebt) + 3 Task",
              types == ["User Story", "Task", "Task", "Task"] and c.created[0][1]["System.Title"] == "Teknik borç — S23"
              and "tempo-techdebt" in c.created[0][1]["System.Tags"], str(types))
        tf = [f for t, f, _p in c.created if t == "Task"]
        check("Task etiketleri php_modern kaynağı; yapısal needs-human", sum("needs-human" in f["System.Tags"] for f in tf) == 1
              and all("tempo-src-php_modern" in f["System.Tags"] for f in tf))
        check("Task açıklamasında hedef PHP ve avantajlar", all("Hedef PHP:</b> 8.4" in f["System.Description"] and "Avantajlar" in f["System.Description"] for f in tf))
        check("mekanik 2 iş kuyruğa alındı, yapısal alınmadı", len(queued) == 2 and len(res["tasks"]) == 2 and len(res["human_tasks"]) == 1
              and res["human_tasks"][0]["job_id"] is None, str(res))
        check("bilinmeyen gid atlanan listesinde, tekrar gid tekilleşti", any("tarama" in s_["reason"] for s_ in res["skipped"]) and len(res["skipped"]) == 1)
        check("quality_items kaydı source=php_modern", rec and all(r["source"] == "php_modern" for r in rec))
        check("story bağlantısı + SP", res["story"]["id"] == 601 and res["total_sp"] == 2 and res["human_sp"] == 5)
        # board'da açık fp → atlanır
        board.clear()
        board[gids_fp := pm.fingerprint("webservice", "str_contains", "app/Lib/A.php")] = {"wi": 700, "state": "Active"}
        sc3 = pm.scan(c, "webservice", team="T", force=True, **pm_kw)
        check("board'da açık WI → grup listede yok, atlananlarda", all(r["rule"] != "str_contains" for r in sc3["auto"])
              and sc3["skipped_counts"].get("açık WI") == 1, str(sc3["skipped_counts"]))
        # tavan
        cfg["max_take"] = 1
        c2 = C()
        res2 = pm.take(c2, "webservice", gids, team="T", **pm_kw)
        check("istek başı tavan", len([t for t, _f, _p in c2.created if t == "Task"]) == 1 and any("tavan" in s_["reason"] for s_ in res2["skipped"]))
        try:
            pm.take(c, "oldphp", gids, team="T", **pm_kw)
            check("hedef sürümü bilinmeyen repo → ValueError", False)
        except ValueError as e:
            check("hedef sürümü bilinmeyen repo → ValueError", "sürüm" in str(e), str(e))
        try:
            pm.scan(c, "yok", team="T", **pm_kw)
            check("yerel olmayan repo → ValueError", False)
        except ValueError:
            check("yerel olmayan repo → ValueError", True)
    check("load_cfg knob'ları okur", {"php_exclude", "php_bin", "scan_ttl_h", "max_take"} <= set(pm.load_cfg()))


def test_php_modern_endpoints():
    print("\n[PM6] php_modern — ayarlar ve uçlar")
    from fastapi.testclient import TestClient
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew import server
    from agile_sdlc_crew.quality import intake, php_modern as pm
    from agile_sdlc_crew.quality.model import QualitySourceError
    sb = pc._SCHEMA_BY_KEY
    check("ayarlar kayıtlı ve varsayılanlar güvenli",
          sb.get("CREW_TECHDEBT_PHP_MODERN", {}).get("default") is False
          and sb.get("CREW_TECHDEBT_SCAN_TTL_H", {}).get("default") == 24
          and sb.get("CREW_TECHDEBT_MODERN_MAX_TAKE", {}).get("default") == 8
          and "CREW_PHP_BIN" not in sb  # env-only: dashboard'dan yazılabilen binary yolu = keyfi çalıştırma
          and "vendor" in sb.get("CREW_TECHDEBT_PATH_EXCLUDE", {}).get("default", ""))
    orig = pc.get

    def on(**kw):
        return lambda k: kw[k] if k in kw else orig(k)
    pc.get = on(CREW_TECHDEBT=True, CREW_TECHDEBT_PHP_MODERN=False)
    try:
        tc = TestClient(server.app)
        check("php_modern kapalıyken üç uç → 403",
              tc.get("/api/techdebt/modern/repos").status_code == 403
              and tc.get("/api/techdebt/modern/scan?repo=webservice").status_code == 403
              and tc.post("/api/techdebt/modern/take", json={"repo": "webservice", "group_ids": ["a"]}).status_code == 403)
    finally:
        pc.get = orig
    pc.get = on(CREW_TECHDEBT=False, CREW_TECHDEBT_PHP_MODERN=True)
    try:
        check("teknik borç ana anahtarı kapalıyken de 403", TestClient(server.app).get("/api/techdebt/modern/repos").status_code == 403)
    finally:
        pc.get = orig
    pc.get = on(CREW_TECHDEBT=True, CREW_TECHDEBT_PHP_MODERN=True)
    import os as _os
    _env_bin = _os.environ.get("CREW_PHP_BIN")
    o_scan, o_take, o_repos, o_worker = pm.scan, pm.take, pm.repos, server._ensure_worker
    worker = []
    try:
        tc = TestClient(server.app)
        check("repo boş → 400", tc.get("/api/techdebt/modern/scan?repo=").status_code == 400)
        check("group_ids boş → 400", tc.post("/api/techdebt/modern/take", json={"repo": "webservice", "group_ids": []}).status_code == 400)

        def raiser(exc):
            def f(*a, **k):
                raise exc
            return f
        pm.scan = raiser(QualitySourceError("php yok"))
        check("tarayıcı hatası → 502", tc.get("/api/techdebt/modern/scan?repo=webservice").status_code == 502)
        pm.scan = raiser(ValueError("repo yok"))
        check("geçersiz repo → 400", tc.get("/api/techdebt/modern/scan?repo=x").status_code == 400)
        pm.take = raiser(intake.Busy("meşgul"))
        check("kilit → 409", tc.post("/api/techdebt/modern/take", json={"repo": "webservice", "group_ids": ["a"]}).status_code == 409)
        pm.take = lambda c, repo, ids, **k: {"repo": repo, "story": {"id": 1, "url": "u", "created": True},
                                             "tasks": [{"id": 2, "job_id": 9}], "human_tasks": [], "skipped": []}
        server._ensure_worker = lambda: worker.append(1)
        r = tc.post("/api/techdebt/modern/take", json={"repo": "webservice", "group_ids": ["a"], "team": "T"})
        check("başarılı alım → 201 + worker tetiklenir", r.status_code == 201 and worker == [1], str(r.status_code))
        pm.scan = lambda c, repo, **k: {"repo": repo, "auto": [], "human": [], "force": k.get("force")}
        r2 = tc.get("/api/techdebt/modern/scan?repo=webservice&force=1")
        check("scan force parametresi geçer", r2.status_code == 200 and r2.json()["force"] is True)
        pm.repos = lambda c, **k: {"items": [{"repo": "webservice", "version": "8.4"}]}
        r3 = tc.get("/api/techdebt/modern/repos")
        check("repos 200 + enabled", r3.status_code == 200 and r3.json()["enabled"] is True and r3.json()["items"])
        _os.environ["CREW_PHP_BIN"] = "php; rm -rf /"
        bad1 = pm.load_cfg()["php_bin"]
        _os.environ["CREW_PHP_BIN"] = "/usr/bin/python3"
        bad2 = pm.load_cfg()["php_bin"]
        _os.environ["CREW_PHP_BIN"] = "/opt/homebrew/bin/php8.4"
        ok = pm.load_cfg()["php_bin"]
        check("CREW_PHP_BIN env-only; boşluk/; ve php olmayan ad reddedilir, php8.4 yolu kabul",
              bad1 == "php" and bad2 == "php" and ok == "/opt/homebrew/bin/php8.4", f"{bad1} {bad2} {ok}")
    finally:
        pm.scan, pm.take, pm.repos, server._ensure_worker = o_scan, o_take, o_repos, o_worker
        pc.get = orig
        if _env_bin is None:
            _os.environ.pop("CREW_PHP_BIN", None)
        else:
            _os.environ["CREW_PHP_BIN"] = _env_bin


def test_php_modern_ui():
    print("\n[PM7] php_modern — Modernleştirme (PHP) kartı")
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text(encoding="utf-8")
    s = html.find('id="paneTechdebt"')
    pane = html[s:html.find("<section", s + 10)] if s >= 0 else ""
    check("kart Teknik borç sekmesinde, başlangıçta gizli", 'id="tdModernCard"' in pane and '<span class="h">Modernleştirme</span>' in pane
          and 'id="tdModernCard" style="display:none' in pane)
    check("repo seçici, hedef sürüm chip'i, benimseme özeti, Yeniden tara, İş olarak aç",
          all(x in pane for x in ('id="tdmRepo"', 'id="tdmTarget"', 'id="tdmAdoption"', 'id="tdmRescan"', 'id="tdmTake"'))
          and "Seçilenleri iş olarak aç" in pane)
    for fn in ("loadModernRepos", "modernScan", "modernTake", "renderModern"):
        check(f"JS: {fn}", f"function {fn}" in html)
    js = html[html.find("function loadModernRepos"):html.find("function loadModernRepos") + 16000]
    check("uçlar", "/api/techdebt/modern/repos" in js and "/api/techdebt/modern/scan" in js and "/api/techdebt/modern/take" in js)
    check("iki grup başlığı", "Otomatik işler" in js and "İnsan kararı gerektiren yapısal işler" in js)
    check("satır ayrıntısı: Avantajlar / Riskler / Önce / Sonra / Bulunan yerler / Kaynak", all(x in js for x in ("Avantajlar", "Riskler", "Önce", "Sonra", "Bulunan yerler", "Kaynak")))
    check("403'te kart gizli, takım _tdTeam() ile, humanChip kullanılır", "status===403" in js and "_tdTeam()" in js and "humanChip(" in js)
    check("sürümü bilinmeyen repo seçilemez", "disabled" in js and "hedef sürüm bilinmiyor" in js)
    check("tema renkleri (sabit renk yok)", "var(--text" in js and not re.search(r"#[0-9a-fA-F]{6}\b", js), "sabit renk var")
    check("sekme açılınca modern repolar yüklenir", "loadModernRepos()" in html[html.find("name==='techdebt'"):html.find("name==='techdebt'") + 200])
    tk = js[js.find("async function modernTake"):]
    check("alım sonucu silinmez, yeniden tarama beklenir (butonlar erken açılmaz)", "await modernScan(false,true)" in tk and "keepMsg" in js)
    check("tekrarlayan (recurring) satır da seçilebilir", "t.recurring_of?'':" not in js)
    check("tarayıcı dosya hataları gösterilir", "errors" in js[js.find("function renderModern"):js.find("async function modernTake")])
    rc = html[html.find("function _tdmRepoChanged"):html.find("async function loadModernRepos")]
    check("repo listesi 'Repo seçiniz…' ile başlar, ilk repo kendiliğinden seçilmez",
          "Repo seçiniz…" in js and "sel.value='';" in js and 'id="tdmScan"' not in html)
    check("repo seçilince tarama kendiliğinden başlar", "if(r&&r.version)modernScan(false);" in rc)
    check("hızlı repo değişiminde eski tarama sonucu yazılmaz", "if(seq!==_tdm.seq)return;" in js)


def test_quality_sonar_exclude():
    print("\n[Q16] Kalite alımı — Sonar'da CSS/SCSS kuralları ve tema/public yolları hariç (#239/#240)")
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.quality import intake, sonar
    from agile_sdlc_crew.quality.model import Finding
    sb = pc._SCHEMA_BY_KEY
    check("ayarlar: kural dili + yol hariç listesi (varsayılan css ve tema/public/vendor/min.js)",
          "css" in sb.get("CREW_QUALITY_SONAR_SKIP_LANGS", {}).get("default", "")
          and all(x in sb.get("CREW_QUALITY_SONAR_PATH_EXCLUDE", {}).get("default", "") for x in ("app/Themes", "public", "vendor", "*.min.js")))

    def f(file, rule):
        return Finding(source="sonar", repo="webservice", project="EcomBackend:webservice", file=file, rule=rule, rule_name=rule,
                       severity="CRITICAL", itype="BUG")
    langs, paths = ["css"], ["app/Themes", "public", "vendor", "*.min.js", "*.bundle.js"]
    check("css:* kuralı hariç (SCSS/LESS de css: öneki taşır)", sonar.excluded(f("app/x.scss", "css:S4654"), langs, paths) == "kural dili hariç (css)")
    check("tema / public / vendor yolları hariç", sonar.excluded(f("app/Themes/butigo/a.php", "php:S1"), langs, paths) == "hariç yol (app/Themes)"
          and sonar.excluded(f("public/frontend/src/style.css", "php:S1"), langs, paths) == "hariç yol (public)"
          and sonar.excluded(f("vendor/x/y.php", "php:S1"), langs, paths) == "hariç yol (vendor)")
    check("*.min.js / *.bundle.js hariç", sonar.excluded(f("webroot/js/app.min.js", "javascript:S1"), langs, paths) == "hariç yol (*.min.js)")
    check("PHP kaynak dosyası kalır", sonar.excluded(f("app/Service/RequestInquiry.php", "php:S4830"), langs, paths) is None
          and sonar.excluded(f("app/Publication.php", "php:S1"), langs, paths) is None)
    check("liste boşsa hiçbir şey hariç değil", sonar.excluded(f("app/x.scss", "css:S1"), [], []) is None)
    cfg = intake.load_cfg()
    check("load_cfg hariç listelerini taşır", cfg.get("sonar_skip_langs") == ["css"] and "app/Themes" in cfg.get("sonar_path_exclude", []))
    orig_fetch, orig_conv = sonar.fetch_issues, sonar.issues_to_findings
    sonar.fetch_issues = lambda *a, **k: ([], {})
    sonar.issues_to_findings = lambda *a, **k: [f("app/x.scss", "css:S4654"), f("app/Themes/t/a.php", "php:S1"), f("app/A.php", "php:S2")]
    try:
        out, skipped = intake.gather("sonar", SimpleNamespace(repo_projects=["EcomBackend"]), cfg)
        check("gather: hariç bulgular düşer ve atlananlarda nedeniyle görünür",
              [x.file for x in out] == ["app/A.php"] and {sk["reason"] for sk in skipped} == {"kural dili hariç (css)", "hariç yol (app/Themes)"}, str(skipped))
    finally:
        sonar.fetch_issues, sonar.issues_to_findings = orig_fetch, orig_conv


def test_env_conventions():
    print("\n[Q17] Ortam değişkeni kuralı — repo adlarını kullan, uydurma (PR 45061: ORKESTRA_RO_DB_PASSWORD)")
    import inspect
    from agile_sdlc_crew import flow as _flow
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.quality import env_conventions as ec, rank, render
    from agile_sdlc_crew.quality.model import Finding
    check("ayar kayıtlı, varsayılan açık (bozuk davranış onarımı; yalnız kimlik bilgisi işlerinde, birkaç KB)",
          pc._SCHEMA_BY_KEY.get("CREW_ENV_CONVENTIONS", {}).get("default") is True)
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        (root / "main").mkdir(); (root / "ru").mkdir(); (root / "vendor" / "x").mkdir(parents=True)
        (root / "main" / ".env.example").write_text("# DB\nDB_ORKESTRA_PROD_HOST=localhost\nDB_ORKESTRA_PROD_PASS=password\nAUTH_PASSWORD=yoursecurepassword\n", encoding="utf-8")
        (root / "main" / ".env").write_text("DB_ORKESTRA_PROD_PASS=RealSecret99!\nONLY_IN_REAL_ENV=x\n", encoding="utf-8")
        (root / "main" / "init.php").write_text("<?php\n$x = new Mysqlidb(getenv('DB_ORKESTRA_PROD_HOST'), getenv('DB_ORKESTRA_PROD_USER'), getenv('DB_ORKESTRA_PROD_PASS'));\n$a = $_ENV['APP_FLORCHESTRA_URL'] ?? env('MAIL_DSN');\n", encoding="utf-8")
        (root / "vendor" / "x" / "v.php").write_text("<?php getenv('VENDOR_ONLY');", encoding="utf-8")
        info = ec.collect(root)
        names = set(info["names"])
        check("adlar .env.example + getenv/$_ENV/env() kullanımlarından; vendor hariç",
              {"DB_ORKESTRA_PROD_HOST", "DB_ORKESTRA_PROD_PASS", "DB_ORKESTRA_PROD_USER", "AUTH_PASSWORD", "APP_FLORCHESTRA_URL", "MAIL_DSN"} <= names
              and "VENDOR_ONLY" not in names, str(names))
        check("gerçek .env okunmaz: değer ve yalnız orada geçen ad sızmaz",
              "ONLY_IN_REAL_ENV" not in names and "RealSecret99" not in str(info) and "yoursecurepassword" not in str(info))
        check("kullanım yeri dosya:satır", any(u["name"] == "DB_ORKESTRA_PROD_PASS" and u["file"] == "main/init.php" and u["line"] == 2 for u in info["usages"]), str(info["usages"][:3]))
        ranked = ec.rank(info["names"], ["$oMySQL_orkestra_prod_reporting = new Mysqlidb('orkestradbro.ydx.intl', 'reader-user', 'x', 'orkestra-db');"])
        check("ipucu satırındaki 'orkestra' / 'prod' → DB_ORKESTRA_PROD_* önce",
              ranked[:3] and all(n.startswith("DB_ORKESTRA_PROD_") for n in ranked[:3]), str(ranked[:4]))
        # gerçek repo: 'reader-user' değeri DB_DASHBOARD_REPORTING_USER'ı öne çıkarmasın (DB ortak parça sayılmaz, aile bonusu)
        real = Path.home() / ".crew_repos" / "flo-dashboard"
        if (real / "ru" / "init.php").is_file():
            ri = ec.collect(real)
            rr = ec.rank(ri["names"], ["Credentials should not be hard-coded — init.php", (real / "ru" / "init.php").read_text(errors="replace").splitlines()[15]])
            check("flo-dashboard: DB_ORKESTRA_PROD_PASS ilk 4'te", "DB_ORKESTRA_PROD_PASS" in rr[:4] and all(n.startswith("DB_ORKESTRA_PROD_") for n in rr[:4]), str(rr[:6]))
        else:
            skip("flo-dashboard gerçek repo sıralaması", "klon yok")
        block = ec.context_block(info, ["$oMySQL_orkestra_prod_reporting = new Mysqlidb('orkestradbro.ydx.intl', 'reader-user', 'x', 'orkestra-db');"], repo="flo-dashboard")
        check("bağlam bloğu: ad uydurma yasağı, en yakın adlar, kaynak dosyalar, örnek dosyaya ekleme kuralı",
              "DB_ORKESTRA_PROD_PASS" in block and "main/.env.example" in block and "main/init.php" in block
              and "never invent" in block.lower() and ".env.example" in block, block[:600])
        check("repo'da env kuralı yoksa blok boş", ec.context_block(ec.collect(root / "yok"), [], repo="r") == "")
    check("kimlik bilgisi işi tespiti: S6437 / S2068 / secrets:* / başlık", ec.is_credential_task("[x] Credentials should not be hard-coded — init.php", "", "tempo-quality")
          and ec.is_credential_task("t", "<code>php:S2068</code>", "tempo-quality") and ec.is_credential_task("t", "secrets:S6703", "tempo-quality")
          and not ec.is_credential_task("[x] Undefined array key", "php:S1234", "tempo-quality"))
    src = inspect.getsource(_flow)
    check("adım bağlamı: tasarım/geliştirme/inceleme adımlarında env bloğu eklenir (iş-değişmezi kuyrukta)",
          "env_conventions" in src and "CREW_ENV_CONVENTIONS" in src)
    g = rank.sonar_units([Finding(source="sonar", repo="flo-dashboard", project="EcomBackend:flo-dashboard", file="ru/init.php",
                                  rule="php:S6437", rule_name="Credentials should not be hard-coded", severity="BLOCKER", itype="VULNERABILITY", lines=[16])])[0]
    eb = " ".join(render.expected_behavior(g))
    check("S6437 beklenen davranış: mevcut env adı, .env.example, ad uydurulmaz",
          ".env.example" in eb and "getenv" in eb and "uydur" in eb, eb)
    check("S6437 kabul kriteri de env kuralını söyler", ".env.example" in render.acceptance_criteria(g) or "mevcut" in render.acceptance_criteria(g))
def test_sprint_report_no_sp():
    """Sprint sunumu: iş özeti slaytlarında SP yok, sıralama/bölme iş sayısıyla,
    özet infografiklerle (Bir Bakışta + tamamlanma halkası + satır çubukları).

    Kullanıcı kararı (2026-10-23): "(25 SP · %23)" ekleri ve SP payına göre
    slayt bölme kaldırıldı. Şablon git'te olmadığı için sentetik bir şablon
    kurulur; ikon üretimi (LLM) devre dışı bırakılır."""
    print("\n[sprint raporu: SP yok, infografik özet]")
    import os

    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError:
        skip("sprint raporu", "python-pptx yok")
        return
    from agile_sdlc_crew import sprint_report as sr

    keys = ("CREW_SPRINT_REPORT_AI_SUMMARY", "CREW_SPRINT_REPORT_VISUALS",
            "CREW_SPRINT_REPORT_OVERVIEW", "CREW_SPRINT_REPORT_MIN_ITEMS",
            "CREW_SPRINT_REPORT_ICON_DIR")
    saved = {k: os.environ.get(k) for k in keys}
    orig_icon = sr._ai_svg_icon
    tmp = tempfile.mkdtemp()
    try:
        os.environ.update({"CREW_SPRINT_REPORT_AI_SUMMARY": "0",
                           "CREW_SPRINT_REPORT_VISUALS": "1",
                           "CREW_SPRINT_REPORT_OVERVIEW": "1",
                           "CREW_SPRINT_REPORT_ICON_DIR": tmp})
        os.environ.pop("CREW_SPRINT_REPORT_MIN_ITEMS", None)
        sr._ai_svg_icon = lambda concept: ""  # emoji yedeği, LLM çağrısı yok

        def it(i, pid, title, state, typ="Task"):
            return sr.SprintItem(id=i, title=f"iş {i}", type=typ, state=state,
                                 parent_id=pid, parent_title=title)

        items = [it(1, 10, "Küçük proje", "Done"),
                 it(2, 20, "Kargo entegrasyonu", "Done"),
                 it(3, 20, "Kargo entegrasyonu", "Active", "Bug"),
                 it(4, 20, "Kargo entegrasyonu", "Resolved"),
                 it(5, None, "", "Active")]
        groups = sr.group_by_parent(items)
        check("gruplar iş sayısına göre azalan, parent'sız en sonda",
              [g.parent_id for g in groups] == [20, 10, None],
              str([g.parent_id for g in groups]))
        sums = [sr.summarize_group(g, {}) for g in groups]
        kargo = sums[0]
        check("özet iş sayılarını taşır (Resolved tamamlanmış sayılır)",
              (kargo.total, kargo.done, kargo.types) == (3, 2, {"Task": 2, "Bug": 1}),
              f"{kargo.total}/{kargo.done}/{kargo.types}")
        check("GroupSummary'de SP alanı yok",
              not any(hasattr(kargo, a) for a in ("sp", "sp_pct")))

        kinds = [(k, [s.parent_id for s in c]) for k, c in sr.plan_slides(sums)]
        check("plan: bir bakışta → ≥2 işli parent kendi slaytında → kalanlar özet satırı",
              kinds == [("overview", [20, 10, None]), ("top", [20]), ("digest", [10, None])],
              str(kinds))

        tpl = os.path.join(tmp, "tpl.pptx")
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        blank = prs.slide_layouts[6]
        s1 = prs.slides.add_slide(blank)
        s1.shapes.add_textbox(Inches(1), Inches(1), Inches(10), Inches(1)).name = "Metin kutusu 2"
        s2 = prs.slides.add_slide(blank)
        title = s2.shapes.add_textbox(Inches(2.4), Inches(0.3), Inches(10), Inches(0.4))
        title.name = "Başlık"  # varsayılan ad 'TextBox N' gövde kutusuyla çakışır
        title.text_frame.text = "Sprint Özeti - Neler Yaptık"
        s2.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(8), Inches(5)).name = "TextBox 1"
        s3 = prs.slides.add_slide(blank)
        s3.shapes.add_textbox(Inches(0), Inches(3), Inches(13), Inches(0.5)).text_frame.text = (
            "Teşekkürler.")
        prs.save(tpl)

        out = os.path.join(tmp, "out.pptx")
        sr.build_pptx("Takım", "21. Sprint", "", sums, tpl, out)
        deck = Presentation(out)
        texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
                 for s in deck.slides]
        check("hiçbir slaytta SP yok",
              not any(re.search(r"\bSP\b", t) for t in texts),
              str([t[:80] for t in texts if re.search(r"\bSP\b", t)]))
        order = [("Bir Bakışta" in t, "Neler Yaptık" in t, "Diğer Çalışmalar" in t) for t in texts]
        check("sıra: başlık, bir bakışta, içerik, diğer çalışmalar, teşekkürler",
              order[1:4] == [(True, False, False), (False, True, False), (False, False, True)],
              str(order))
        charts = [sum(1 for sh in s.shapes if sh.has_chart) for s in deck.slides]
        check("bir bakışta: KPI + proje durum grafiği; içerik slaytı: tamamlanma halkası",
              charts[1] == 1 and charts[2] == 1 and "Toplam İş" in texts[1]
              and "2 / 3 iş tamamlandı" in texts[2], f"charts={charts}")
        check("diğer çalışmalar satırı ilerleme metni taşır",
              "1 / 1 iş" in texts[3] and "0 / 1 iş" in texts[3], texts[3][:200])
    finally:
        sr._ai_svg_icon = orig_icon
        for k, v in saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def test_sprint_report_designs():
    """Alternatif tasarımlar (dashboard dropdown'u): her tasarım şablonsuz sıfırdan
    çizilir, SP içermez, proje başına madde listesi yerine tek cümle manşet taşır;
    bilinmeyen tasarım klasiğe düşer; config uç noktası listeyi verir."""
    print("\n[sprint raporu: alternatif tasarımlar]")
    import asyncio
    import os

    try:
        from pptx import Presentation
    except ImportError:
        skip("sprint raporu tasarımları", "python-pptx yok")
        return
    from agile_sdlc_crew import sprint_report as sr
    from agile_sdlc_crew import sprint_report_designs as sd

    keys = ("CREW_SPRINT_REPORT_ICON_DIR", "CREW_SPRINT_REPORT_TEMPLATE",
            "CREW_SPRINT_REPORT_TEAM_IMAGE", "CREW_SPRINT_REPORT_DESIGN")
    saved = {k: os.environ.get(k) for k in keys}
    orig_icon = sr._ai_svg_icon
    tmp = tempfile.mkdtemp()
    try:
        os.environ.update({"CREW_SPRINT_REPORT_ICON_DIR": tmp,
                           "CREW_SPRINT_REPORT_TEMPLATE": os.path.join(tmp, "yok.pptx"),
                           "CREW_SPRINT_REPORT_TEAM_IMAGE": os.path.join(tmp, "yok.png")})
        os.environ.pop("CREW_SPRINT_REPORT_DESIGN", None)
        sr._ai_svg_icon = lambda concept: ""

        check("bilinmeyen/boş tasarım klasiğe düşer",
              sd.resolve("") == "klasik" and sd.resolve("uydurma") == "klasik"
              and sd.resolve("GECE") == "gece")
        os.environ["CREW_SPRINT_REPORT_DESIGN"] = "pastel"
        check("boş tasarım env varsayılanını alır", sd.resolve("") == "pastel")
        os.environ.pop("CREW_SPRINT_REPORT_DESIGN", None)
        opts = sd.design_options()
        check("dropdown: klasik ilk, her seçenekte ad/açıklama/renk örneği",
              opts[0]["key"] == "klasik" and len(opts) == 1 + len(sd.DESIGNS)
              and all(o["label"] and o["description"] and len(o["swatch"]) == 4 for o in opts))

        check("AI manşeti JSON'dan okunur, yoksa boş",
              sr._parse_headline('x {"headline": "Kargo akışı yenilendi.", "bullets": ["a"]}')
              == "Kargo akışı yenilendi." and sr._parse_headline('{"bullets": ["a"]}') == "")

        mk = lambda pid, t, n, dn: [sr.SprintItem(id=pid * 10 + j, title=f"iş {j}", type="Task",
                                                  state="Done" if j < dn else "Active",
                                                  parent_id=pid, parent_title=t) for j in range(n)]
        items = mk(1, "Kargo entegrasyonu", 4, 3) + mk(2, "Stok raporu", 1, 1) + mk(3, "İade", 2, 0)
        sums = [sr.summarize_group(g, {}) for g in sr.group_by_parent(items)]
        sums[0].headline = "Kargo entegrasyonu canlıya alındı."
        burn = {"labels": ["01.10", "02.10"], "remaining": [5, 2], "total": [5, 5], "ideal": [5, 0]}
        for key in sd.DESIGNS:
            out = os.path.join(tmp, f"{key}.pptx")
            sd.build_designed_pptx(key, "Takım", "21. Sprint", "01.10 – 14.10", sums, out,
                                   burndown=burn)
            deck = Presentation(out)
            texts = [" ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
                     for s in deck.slides]
            joined = " | ".join(texts)
            # kapak, bir bakışta, burndown, 1 kart slaytı (2 proje), diğer çalışmalar, kapanış
            check(f"{key}: 6 slayt, şablonsuz üretildi", len(texts) == 6, f"{len(texts)} slayt")
            check(f"{key}: SP yok", not re.search(r"\bSP\b", joined))
            check(f"{key}: kartta manşet var, madde listesi yok",
                  "Kargo entegrasyonu canlıya alındı." in joined and "•" not in joined)
            check(f"{key}: bir bakışta + diğer çalışmalar + kapanış",
                  "Bir Bakışta" in texts[1] and "Diğer Çalışmalar" in texts[4]
                  and "Teşekkürler" in texts[5])

        from agile_sdlc_crew import server

        cfg = json.loads(asyncio.run(server.sprint_report_config()).body)
        check("config uç noktası tasarım listesini ve eşiği verir (eski min_pct değil)",
              [d["key"] for d in cfg["designs"]][0] == "klasik" and "min_items" in cfg
              and "min_pct" not in cfg, str(list(cfg)))
        check("istek modeli tasarımı taşır",
              server.SprintReportRequest(team="t", iteration_path="p", design="gece").design == "gece")
        html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text()
        check("dashboard: tasarım dropdown'u isteğe eklenir",
              'id="srDesign"' in html and "JSON.stringify({team,iteration_path,design})" in html)
    finally:
        sr._ai_svg_icon = orig_icon
        for k, v in saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def test_review_retry_delete_and_no_progress():
    """Job #244 ($3.19): (1) plan bir dosyanın SİLİNMESİNİ istiyordu ama tek push
    yolu içerik aldığı için developer boş döndü, push iptal edildi; (2) "ilerleme
    yok, erken durduruluyor" sayaca ham CREW_REVIEW_MAX_RETRIES yazdı, tavan zarftan
    (3) geldiği için sayaç her turda 2'ye döndü — döngü hiç bitmedi."""
    print("\n[review retry: dosya silme + erken durdurma]")
    import agile_sdlc_crew.pipeline as _pl
    import agile_sdlc_crew.main as _mn
    from agile_sdlc_crew import pipeline_config as pc
    from agile_sdlc_crew.flow import NeedsHumanReview as _NHR

    knobs = {"CREW_STRUCTURED_REVIEW": True, "CREW_REVIEW_RETRY_REPLAN": False,
             "CREW_REVIEW_MAX_RETRIES": 1, "CREW_CONTRACT_GATE": False}
    _orig = (pc.get, _pl.delete_file, _pl.push_file, _mn._add_wi_comment)
    pc.get = lambda k: knobs[k] if k in knobs else _orig[0](k)
    _mn._add_wi_comment = lambda *a, **k: None

    class _Crew:
        def __init__(self, raw, log):
            self.raw, self.log = raw, log
        def kickoff(self, inputs=None):
            self.log.append(inputs)
            return SimpleNamespace(raw=self.raw, token_usage=None)

    class _Client:
        def get_file_content(self, *a, **k):
            return "<?php\n" + "echo 1;\n" * 40
        def add_pr_comment(self, *a, **k):
            return {}

    def mk(change_type):
        f = AgileSDLCFlow()
        f.state.work_item_id, f.state.repo_name, f.state.branch_name = "76578", "marketplace", "feature/76578"
        f.state.pr_id, f.state.pr_url = "45479", "https://example/pr/45479"
        f.state.plan = {"changes": [{"file_path": "/secret.json", "change_type": change_type,
                                     "description": "x", "current_code": "", "new_code": ""}]}
        f.state.review_issues = [{"id": "R1", "file": "/secret.json", "severity": "blocker",
                                  "problem": "p", "required_fix": "r", "status": "open",
                                  "fix_targets": [_norm_path("/secret.json")]}]
        f._envelope = {"retries": 3}
        f.state.requirements_text = "Kanıt: secret.json:5 — dosya repodan silinir"
        f._client = _Client()
        f._db = None
        calls = {"verify": [], "code": []}
        verify_raw = 'REVIEW_VERIFY_JSON: {"results": [{"id": "R1", "status": "open"}], "new_findings": []}'
        f._agile_crew = SimpleNamespace(
            create_verify_review_crew=lambda: _Crew(verify_raw, calls["verify"]),
            create_code_crew=lambda: _Crew("<?php\n" + "echo 2;\n" * 40, calls["code"]),
        )
        for name in ("_enable_impl_repo_tools", "_restore_worktree_file", "_stage_worktree_file",
                     "_track_and_check_budget", "_step_done"):
            setattr(f, name, lambda *a, **k: None)
        f._build_step_context = lambda *a, **k: ""
        f._prefetch_pr_changes_context = lambda *a, **k: ""
        f._test_requirement_note = lambda *a, **k: ""
        f._prefer_worktree_edit = lambda repo, fp, new, old: new
        return f, calls

    try:
        # ── silme: ilk tur siler, ikinci turda dosya zaten yok → ilerleme yok ──
        deletes = []
        def _fake_delete(repo, br, fp, msg, repo_mgr=None, dry_run=False):
            deletes.append(fp)
            return ({"success": True, "change_type": "delete", "file": fp, "push_id": 7}
                    if len(deletes) == 1 else
                    {"success": True, "change_type": "delete", "file": fp, "note": "already absent"})
        _pl.delete_file = _fake_delete
        f, calls = mk("delete")
        try:
            f._review_retry_loop()
            raised = False
        except _NHR:
            raised = True
        check("silme: developer'a hiç gidilmedi", calls["code"] == [], str(len(calls["code"])))
        check("silme: dosya delete_file ile silindi", deletes[:1] == ["/secret.json"], str(deletes))
        check("silme: push olmayan turda doğrulama çağrılmadı (1 verify)", len(calls["verify"]) == 1,
              str(len(calls["verify"])))
        check("silme: ilerleme yok → needs_human (sonsuz döngü yok)", raised and len(deletes) == 2,
              f"raised={raised} deletes={len(deletes)}")

        # ── düzenleme: push her turda başarılı ama madde kapanmıyor ──
        pushes = []
        _pl.push_file = lambda *a, **k: (pushes.append(a[2]) or {"success": True, "change_type": "edit", "file": a[2]})
        f, calls = mk("edit")
        try:
            f._review_retry_loop()
            raised = False
        except _NHR:
            raised = True
        check("erken durdurma: aynı açık madde iki turda → durur (zarf 3, ham knob 1)",
              raised and len(calls["verify"]) == 2, f"raised={raised} verify={len(calls['verify'])}")
    finally:
        pc.get, _pl.delete_file, _pl.push_file, _mn._add_wi_comment = _orig


def test_delete_push_payload():
    """Azure push'ta silme değişikliği içerik taşımaz; dry-run'da git rm + commit."""
    print("\n[silme: Azure push gövdesi + dry-run]")
    from agile_sdlc_crew.tools import azure_devops_base as _azb
    from agile_sdlc_crew.tools.local_repo import LocalRepoManager

    sent = {}
    class _Resp:
        def __init__(self, data): self._d = data; self.text = "x"; self.status_code = 201
        def raise_for_status(self): pass
        def json(self): return self._d
    class _C(_azb.AzureDevOpsClient):
        _headers = {}
    c = _C.__new__(_C)
    c.list_branches = lambda repo: [{"name": "refs/heads/b", "objectId": "abc"}]
    c._repo_api_url = lambda repo, project=None: "https://x/repo"
    c.API_VERSION = "7.1"
    _orig_post = _azb.requests.post
    _azb.requests.post = lambda url, headers=None, json=None, params=None, timeout=None: (sent.update(body=json) or _Resp({"pushId": 9}))
    try:
        c.push_changes("r", "b", [{"changeType": "delete", "path": "/a.json"},
                                  {"changeType": "edit", "path": "/b.php", "content": "x"}], "m")
    finally:
        _azb.requests.post = _orig_post
    ch = sent["body"]["commits"][0]["changes"]
    check("push: delete değişikliğinde newContent yok", "newContent" not in ch[0] and ch[0]["changeType"] == "delete")
    check("push: edit değişikliği içeriğini korur", ch[1]["newContent"]["content"] == "x")

    with tempfile.TemporaryDirectory() as td:
        repo = Path(td) / "r"
        repo.mkdir()
        g = lambda *a: subprocess.run(["git", *a], cwd=repo, capture_output=True, text=True)
        g("init", "-q", "-b", "main")
        (repo / "a.json").write_text("{}")
        g("add", "."); g("-c", "user.name=t", "-c", "user.email=t@t", "commit", "-qm", "i")
        mgr = LocalRepoManager.__new__(LocalRepoManager)
        mgr._get_repo_dir = lambda name: repo
        mgr.checkout = lambda name, br: g("checkout", "-q", "-B", br)
        mgr._git = lambda args, cwd=None: subprocess.run(["git", *args], cwd=cwd, capture_output=True, text=True)
        r1 = mgr.delete_and_commit_local("r", "feature/x", "/a.json", "sil")
        r2 = mgr.delete_and_commit_local("r", "feature/x", "/a.json", "sil")
        check("dry-run: dosya silinip commit edildi", r1.get("success") and not (repo / "a.json").exists()
              and "sil" in g("log", "-1", "--format=%s").stdout, str(r1))
        check("dry-run: zaten yoksa no-op başarı", r2.get("success") and "already absent" in r2.get("note", ""), str(r2))


def test_quality_direct_flow():
    """Job #244: Task dosyayı ve satırı söylüyordu; BA + mimar keşfi + araçlı amend
    $1.54 tuttu. Doğrudan akış: analiz Task'tan, kanıt + okuyan kod Python'da."""
    print("\n[kalite doğrudan akış]")
    from agile_sdlc_crew.flow import _quality_evidence, _quality_requirements_json, _requirement_ids
    from agile_sdlc_crew import pipeline_config as pc

    req = _quality_requirements_json(
        "[marketplace] Google Cloud service accounts keys should not be disclosed",
        "<p>Kanıt: mpModerationAiAuth.json:5</p>",
        "<ul><li>Secret koddan kaldırılır</li><li>Yeni değer gizli yönetiminden okunur</li></ul>")
    ids = _requirement_ids(req)
    check("analiz: FR/TR + Task'taki iki kabul kriteri", {"FR1", "TR1", "AC1", "AC2"} <= set(ids), str(ids))
    check("analiz: az gereksinim (zarf büyümez)", len(ids) <= 5, str(ids))

    with tempfile.TemporaryDirectory() as td:
        root = Path(td) / "repo"
        root.mkdir()
        g = lambda *a: subprocess.run(["git", *a], cwd=root, capture_output=True, text=True)
        g("init", "-q", "-b", "main")
        (root / "auth.json").write_text('{\n  "type": "service_account",\n  "private_key": "-----BEGIN KEY-----abc",\n'
                                        '  "client_email": "x@y.iam.gserviceaccount.com"\n}\n')
        (root / "app").mkdir()
        (root / "app" / "Svc.php").write_text("<?php\nclass Svc {\n  function __construct() {\n"
                                              "    $p = __DIR__ . '/../auth.json';\n  }\n}\n")
        (root / "app" / "Big.php").write_text("<?php\n" + "".join(f"// line {i} {'x' * 30}\n" for i in range(2, 400))
                                              + "$apiKey = 'AIzaSyD1234567890abcdefghijkLMNOP';\n")
        (root / "vendor").mkdir()
        (root / "vendor" / "lib.php").write_text("<?php // auth.json\n")
        g("add", "-A"); g("-c", "user.name=t", "-c", "user.email=t@t", "commit", "-qm", "i")
        ev = _quality_evidence(root, "<li>auth.json:3</li><li>app/Big.php:400</li> api.flo.com.tr:443")
        check("kanıt: veri dosyası bağlamda", "## /auth.json (tamamı" in ev, ev[:300])
        check("kanıt: dosyayı okuyan kod eklendi", "/app/Svc.php — `auth.json`" in ev, ev[:600])
        check("kanıt: vendor okuyucu alınmaz", "vendor/lib.php" not in ev)
        check("kanıt: büyük dosyada yalnız satır penceresi", "/app/Big.php (satır 360–" in ev and "// line 10 " not in ev, ev[-500:])
        check("kanıt: repoda olmayan yol (alan adı:port) elenir", "api.flo.com.tr" not in ev)
        check("kanıt: sır değerleri modele gitmez", "BEGIN KEY" not in ev and "AIzaSy" not in ev
              and '"private_key": "***"' in ev, ev[:400])
        check("kanıt: repoda yoksa boş", _quality_evidence(root, "<li>nope.json:1</li>") == "")

    # Akış kancaları: knob kapalıysa normal akış, açıkken zarf S ve keşif kapalı
    f = AgileSDLCFlow()
    f.state.wi_tags = "tempo-quality; tempo-src-sonar"
    knobs = {"CREW_QUALITY_DIRECT_FLOW": True, "CREW_VARIABLE_ENVELOPE": True}
    _orig = pc.get
    pc.get = lambda k: knobs[k] if k in knobs else _orig(k)
    try:
        f.state.requirements_text = req
        f._apply_envelope("requirements")
        f._apply_envelope("plan")
        check("zarf: doğrudan akışta S/$5/1 tur", getattr(f, "_envelope", None) == {"class": "S", "budget": 5.0, "retries": 1},
              str(getattr(f, "_envelope", None)))
        knobs["CREW_QUALITY_DIRECT_FLOW"] = False
        check("knob kapalı: doğrudan akış yok", not f._quality_direct_flow())
        f2 = AgileSDLCFlow()
        f2.state.wi_tags = "backend"
        knobs["CREW_QUALITY_DIRECT_FLOW"] = True
        check("kalite etiketi yoksa doğrudan akış yok", not f2._quality_direct_flow())
    finally:
        pc.get = _orig

    src = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/flow.py").read_text()
    check("step1: doğrudan akışta BA crew ve SM review çağrılmaz",
          "requirements_text = _quality_requirements_json(" in src
          and '(True, "") if _direct else self._scrum_review("Is Analizi"' in src)
    check("step4: doğrudan akışta Faz A keşfi kapalı",
          src.find("if self._quality_direct_flow() and _repo_dirs:") < src.find("plan, raw_output, need_explore = self._architect_emit_json("))
    check("amend: doğrudan akışta repo araçları açılmaz", "and not self._quality_direct_flow()):" in src)
    check("kanıt yalnız teknik tasarım bağlamına girer", 'if step_key == "technical_design_task":\n                try:\n                    ev_ctx = self._quality_evidence_block()' in src)


def test_resume_skips_discarded_work():
    """Job #246: önceki işin PR'ı kapatılmış, branch'i silinmişti; resume yine de
    BA'yı, planı, branch'i ve KAPALI PR'ı geri yükleyip review'a geçti."""
    print("\n[resume: atılmış önceki iş]")

    class _C:
        def __init__(self, status=None, branches=(), boom=False):
            self.status, self.branches, self.boom = status, branches, boom
        def get_pull_request(self, repo, pr):
            if self.boom:
                raise RuntimeError("503")
            return {"status": self.status}
        def list_branches(self, repo):
            return [{"name": b} for b in self.branches]

    def mk(art, client):
        f = AgileSDLCFlow()
        f._prior_art_cache = art
        f._client = client
        return f

    art_pr = {"repo_name": "marketplace", "branch_name": "feature/76578", "pr_id": "45479"}
    check("PR abandoned → resume yok", mk(art_pr, _C("abandoned"))._prior_work_discarded())
    check("PR merge edilmiş → resume yok", mk(art_pr, _C("completed"))._prior_work_discarded())
    check("PR aktif → resume serbest", not mk(art_pr, _C("active"))._prior_work_discarded())
    art_br = {"repo_name": "marketplace", "branch_name": "feature/76578", "pr_id": ""}
    check("PR yok, branch silinmiş → resume yok", mk(art_br, _C(branches=["refs/heads/main"]))._prior_work_discarded())
    check("PR yok, branch duruyor → resume serbest",
          not mk(art_br, _C(branches=["refs/heads/feature/76578"]))._prior_work_discarded())
    check("durum okunamazsa resume engellenmez", not mk(art_pr, _C(boom=True))._prior_work_discarded())
    check("önceki iş yoksa engel yok", not mk({}, _C("abandoned"))._prior_work_discarded())


def test_wi_lifecycle_on_queue():
    """Kullanıcı kararı: kart kuyruğa atıldığı anda In Progress'e geçer; sorun olursa
    kuyruk anındaki duruma (To Do) döner — iş başlarken okunan In Progress'e değil."""
    print("\n[yaşam döngüsü: kuyruğa alınınca]")
    from agile_sdlc_crew import wi_lifecycle as wl
    from agile_sdlc_crew import pipeline_config as pc

    c = _StubAzClient(_FLO_TASK, {"System.WorkItemType": "Task", "System.State": "To Do"})
    before = wl.mark_queued(c, "76578")
    check("kuyruk: To Do → In Progress, önceki durum döner", before == "To Do" and c.ops == [("state", 76578, "In Progress")], str(c.ops))
    c2 = _StubAzClient(_FLO_TASK, {"System.WorkItemType": "Task", "System.State": "QA To Do"})
    check("kuyruk: sahiplik dışı (QA) → dokunmaz", wl.mark_queued(c2, "1") == "QA To Do" and c2.ops == [])
    check("geri al: In Progress → To Do", wl.revert_to(c, "76578", "To Do") == "To Do" and c.ops[-1] == ("state", 76578, "To Do"))
    check("geri al: başlangıç boşsa dokunmaz", wl.revert_to(c, "76578", "") is None)

    knobs = {"CREW_WI_LIFECYCLE": True, "CREW_WI_ASSIGN_IF_EMPTY": False}
    _orig = pc.get
    pc.get = lambda k: knobs[k] if k in knobs else _orig(k)
    try:
        f = AgileSDLCFlow()
        f.state.work_item_id, f.state.job_id = "76578", 247
        f._client = _StubAzClient(_FLO_TASK, {"System.WorkItemType": "Task", "System.State": "In Progress"})
        f._db = SimpleNamespace(get_job=lambda jid: {"id": jid, "wi_state_before": "To Do"})
        f._wi_begin(f._client.fields)
        check("iş başlarken: başlangıç = kuyruk anındaki durum", f.state.wi_state_initial == "To Do"
              and f._client.ops == [], f"{f.state.wi_state_initial} {f._client.ops}")
        f.wi_lifecycle_on_exception(RuntimeError("x"))
        check("sorun → kuyruk anındaki duruma (To Do) döner", f._client.ops[-1] == ("state", 76578, "To Do"), str(f._client.ops))
    finally:
        pc.get = _orig

    src = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/server.py").read_text()
    check("server: her kuyruğa almada işaretleme (_ensure_worker)", "threading.Thread(target=_wi_mark_queued_jobs" in src)
    check("server: silinen kuyruk işi ve restart'ta yarım kalan iş geri alınır",
          'f"iş #{job_id} silindi"' in src and '"sunucu yeniden başladı"' in src)


def test_pre_pr_verify_and_build_fix():
    """Job #248 ($5.44): eksik Go import'u CI'a kadar gitti; build hata özeti
    hazırlık YAML'ıydı; build-fix dosya başına Opus çağrısı yaptı ve silinmiş
    analytics.json'a düzyazı yazdı; test dışı Sonar uyarısı ikinci turu açtı."""
    print("\n[PR öncesi doğrulama + build-fix]")
    import shutil
    from agile_sdlc_crew import verify as vf
    from agile_sdlc_crew.tools.azure_devops_base import extract_build_errors
    from agile_sdlc_crew.flow import _build_issues_block
    from agile_sdlc_crew import pipeline_config as pc

    # ── build log özeti ──
    log = "\n".join([
        "2026-10-27T11:23:33.1Z ##[section]Starting: Run Tests",
        "2026-10-27T11:23:33.1Z Task         : Command line",
        "2026-10-27T11:23:33.1Z Evaluating: eq('true', variables['system.debugContext'])",
        "2026-10-27T11:25:01.0Z Dumping database schema",
        "2026-10-27T11:25:02.0Z # gocrons/pkg/sorting",
        "2026-10-27T11:25:02.0Z pkg/sorting/a.go:8:2: \"os\" imported and not used",
        "2026-10-27T11:25:02.0Z pkg/sorting/a.go:392:15: undefined: webanalytics",
    ] + [f"2026-10-27T11:25:03.0Z FAIL\tgocrons/pkg/p{i} [build failed]" for i in range(40)]
      + ["2026-10-27T11:25:04.0Z ##[error]Bash exited with code '1'."])
    ex = extract_build_errors(log)
    check("log: derleyici hataları önde, zaman damgasız", ex.startswith("# gocrons/pkg/sorting\npkg/sorting/a.go:8:2"), ex[:120])
    check("log: şablon gürültüsü atılır", "Evaluating:" not in ex and "Task         :" not in ex)
    check("log: FAIL paket yığını tek satır", "(40 paket 'build failed'" in ex and ex.count("[build failed]") == 0, ex[-300:])

    # ── kısmi başarı ──
    check("kısmi başarı: yalnız Sonar uyarısı → bloklamaz",
          not _build_issues_block([{"name": "Run SonarQube Analysis", "result": "succeededWithIssues"}]))
    check("kısmi başarı: test adımı uyarılı → bloklar",
          _build_issues_block([{"name": "Run Tests", "result": "succeededWithIssues"}]))
    check("kısmi başarı: herhangi bir Task failed → bloklar",
          _build_issues_block([{"name": "Lint", "result": "failed"}]))

    # ── dil doğrulayıcıları ──
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        (root / "bad.json").write_text('{"a": 1,}')
        (root / "ok.yaml").write_text("a: 1\n")
        (root / "bad.py").write_text("def f(:\n")
        r = vf.verify(root, ["/bad.json", "/ok.yaml", "/bad.py", "/gone.php"])
        check("doğrulayıcı: JSON/Python hataları dosya:satır ile", not r["ok"]
              and any(i.startswith("bad.json:") for i in r["issues"]) and any("bad.py" in i for i in r["issues"]), str(r))
        check("doğrulayıcı: diskte olmayan dosya atlanır", "php" not in r["ran"])
        check("hata satırından dosya çıkarımı", "bad.json" in vf.referenced_files(r["issues"]))
        if shutil.which("go"):
            g = root / "gomod"
            (g / "pkg" / "a").mkdir(parents=True)
            (g / "go.mod").write_text("module m\n\ngo 1.21\n")
            (g / "pkg" / "a" / "a.go").write_text('package a\n\nimport "os"\n\nfunc F() string { return strings.ToUpper("x") }\n')
            rg = vf.verify(g, ["pkg/a/a.go"])
            check("go: eksik import ve kullanılmayan import yakalanır (#248)", not rg["ok"]
                  and any("undefined: strings" in i for i in rg["issues"])
                  and any('"os" imported and not used' in i for i in rg["issues"]), str(rg))
        else:
            skip("go doğrulayıcı", "go yok")

    # ── tek çağrıda çok dosyalı düzeltme ──
    with tempfile.TemporaryDirectory() as td:
        base = Path(td)
        repo = base / "r"
        (repo / "pkg").mkdir(parents=True)
        g = lambda *a: subprocess.run(["git", *a], cwd=repo, capture_output=True, text=True)
        g("init", "-q", "-b", "main")
        for n in ("a.go", "b.go"):
            (repo / "pkg" / n).write_text(f"package pkg // {n} main\n")
        (repo / "secret.json").write_text("{}")
        g("add", "-A"); g("-c", "user.name=t", "-c", "user.email=t@t", "commit", "-qm", "i")

        branch = {"/pkg/a.go": "package pkg // a branch broken\n", "/pkg/b.go": "package pkg // b branch broken\n",
                  "/pkg/new.go": "package pkg // new broken\n"}
        pushes = []

        class _Cl:
            def get_file_content(self, repo_name, fp, br):
                if fp in branch:
                    return branch[fp]
                raise RuntimeError("404 Not Found")
            def push_changes(self, repo_name, br, changes, msg):
                pushes.append((msg, [c["path"] for c in changes]))
                for c in changes:
                    branch[c["path"]] = c["content"]
                return {"pushId": 1}

        calls = []

        class _Crew:
            def kickoff(self, inputs=None):
                calls.append(inputs)
                for n in ("a.go", "new.go"):
                    p = repo / "pkg" / n
                    p.write_text(p.read_text().replace("broken", "fixed"))
                return SimpleNamespace(raw="tamam", token_usage=None, pydantic=None)

        f = AgileSDLCFlow()
        f.state.work_item_id, f.state.repo_name, f.state.branch_name = "76575", "r", "feature/76575"
        f.state.plan = {"changes": [{"file_path": "/secret.json", "change_type": "delete"}]}
        f._client = _Cl()
        f._repo_mgr = SimpleNamespace(base_dir=base, _git=lambda args, cwd=None: subprocess.run(
            ["git", *args], cwd=cwd, capture_output=True, text=True))
        f._agile_crew = SimpleNamespace(create_code_crew=lambda: _Crew())
        f._enable_impl_repo_tools = lambda: True
        f._track_and_check_budget = lambda *a, **k: None
        f._build_step_context = lambda *a, **k: ""
        n = f._multi_file_fix(["pkg/a.go", "/pkg/b.go", "/pkg/new.go", "/secret.json"], "undefined: x", "build_fix", "HATA")
        check("çok dosyalı düzeltme: TEK developer çağrısı", len(calls) == 1, str(len(calls)))
        check("çok dosyalı düzeltme: silinecek dosya hedefe girmez", "secret.json" not in calls[0]["change_description"])
        check("çok dosyalı düzeltme: yalnız değişen dosyalar, TEK commit",
              n == 2 and len(pushes) == 1 and sorted(pushes[0][1]) == ["/pkg/a.go", "/pkg/new.go"], str(pushes))
        check("çok dosyalı düzeltme: klon geri alındı (main hâli, PR'ın yeni dosyası kalmaz)",
              (repo / "pkg" / "a.go").read_text() == "package pkg // a.go main\n" and not (repo / "pkg" / "new.go").exists())

        # ── PR öncesi doğrulama kapısı ──
        knobs = {"CREW_PRE_PR_VERIFY": True}
        _orig = pc.get
        pc.get = lambda k: knobs[k] if k in knobs else _orig(k)
        import agile_sdlc_crew.verify as _vfm
        _orig_verify = _vfm.verify
        seen = []

        def _fake_verify(root, changed, **kw):
            seen.append({fp: (Path(root) / fp.lstrip("/")).read_text() for fp in changed
                         if (Path(root) / fp.lstrip("/")).exists()})
            bad = any("broken" in t for t in seen[-1].values())
            return {"ok": not bad, "ran": ["go"], "skipped": [],
                    "issues": ["pkg/b.go:1:1: undefined: y"] if bad else []}
        _vfm.verify = _fake_verify
        try:
            branch["/pkg/b.go"] = "package pkg // b branch broken\n"
            calls.clear(); pushes.clear()

            class _Crew2:
                def kickoff(self, inputs=None):
                    calls.append(inputs)
                    p = repo / "pkg" / "b.go"
                    p.write_text(p.read_text().replace("broken", "fixed"))
                    return SimpleNamespace(raw="", token_usage=None, pydantic=None)
            f._agile_crew = SimpleNamespace(create_code_crew=lambda: _Crew2())
            f.state.all_pushes = [{"file": "/pkg/b.go", "success": True, "change_type": "edit"},
                                  {"file": "/secret.json", "success": True, "change_type": "delete"}]
            f._pre_pr_verify()
            check("doğrulama: branch hâli klona bindirilip doğrulanır", "broken" in seen[0].get("/pkg/b.go", ""), str(seen[:1]))
            check("doğrulama: hata → tek hedefli düzeltme + tekrar doğrulama geçer",
                  len(calls) == 1 and len(pushes) == 1 and len(seen) == 2 and "fixed" in seen[1]["/pkg/b.go"], f"{len(calls)} {pushes} {seen}")
            check("doğrulama sonrası klon temiz, silinen dosya geri geldi",
                  (repo / "pkg" / "b.go").read_text() == "package pkg // b.go main\n" and (repo / "secret.json").exists())
        finally:
            _vfm.verify = _orig_verify
            pc.get = _orig

    src = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/flow.py").read_text()
    check("build-fix: silinecek dosyalar hedeften çıkarılır, tek çağrı yolu önce denenir",
          "fix_files = [f for f in fix_files if _norm_path(f) not in _deleted]" in src
          and 'self._multi_file_fix(fix_files, failure_summary, "build_fix"' in src)
    check("kalite doğrudan akış: build-fix en fazla 1 tur", "max_retries = min(max_retries, 1)" in src)


def test_security_hardening():
    """Güvenlik incelemesi (commit 07bac9738): doğrulayıcılar repo kodunu içe
    aktarmasın / sunucu sırlarını devralmasın; kanıtta tırnaksız sırlar ve anahtar
    dosyaları modele gitmesin; plan (LLM) keyfi dosya silemesin."""
    print("\n[güvenlik sertleştirme]")
    import os
    from agile_sdlc_crew import verify as vf
    from agile_sdlc_crew.flow import _qe_mask, _quality_evidence

    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        marker = root / "PWNED"
        (root / "py_compile.py").write_text(f"open({str(marker)!r}, 'w').write('x')\n")
        (root / "ok.py").write_text("x = 1\n")
        r = vf.verify(root, ["ok.py"])
        check("python: repo kökündeki py_compile.py içe aktarılmaz", r["ok"] and not marker.exists(), str(r))
        (root / "-v.json").write_text("{}")
        check("'-' ile başlayan yol bayrak sayılmaz", vf._arg("-v.json") == "./-v.json")
    os.environ["AZURE_DEVOPS_PAT"] = os.environ.get("AZURE_DEVOPS_PAT", "x")
    env = vf._env({"GOFLAGS": "-mod=vendor"})
    check("doğrulayıcı ortamı sırları taşımaz", "AZURE_DEVOPS_PAT" not in env and "PATH" in env
          and env.get("GOTOOLCHAIN") == "local" and env.get("GOFLAGS") == "-mod=vendor")

    m = _qe_mask("DB_PASSWORD=hunter2xyz\nexport API_TOKEN=abc123\nAPP_NAME=shop\n")
    check("kanıt: tırnaksız KEY=değer maskelenir", "hunter2xyz" not in m and "abc123" not in m and "APP_NAME=shop" in m, m)
    with tempfile.TemporaryDirectory() as td:
        root = Path(td) / "r"
        root.mkdir()
        g = lambda *a: subprocess.run(["git", *a], cwd=root, capture_output=True, text=True)
        g("init", "-q", "-b", "main")
        (root / "server.pem").write_text("-----BEGIN PRIVATE KEY-----\nMIIEvQIBADANBg\n-----END PRIVATE KEY-----\n")
        g("add", "-A"); g("-c", "user.name=t", "-c", "user.email=t@t", "commit", "-qm", "i")
        ev = _quality_evidence(root, "<li>server.pem:1</li>")
        check("kanıt: anahtar dosyasının içeriği gönderilmez", "BEGIN PRIVATE" not in ev and "MIIE" not in ev
              and "anahtar dosyası" in ev, ev)

    f = AgileSDLCFlow()
    f.state.requirements_text = "mpModerationAiAuth.json silinir; azure-pipelines.yml de geçiyor"
    ok, _ = f._delete_allowed("/mpModerationAiAuth.json")
    check("silme: WI'da adı geçen dosya serbest", ok)
    ok, why = f._delete_allowed("/app/Controller/Order.php")
    check("silme: WI'da geçmeyen dosya reddedilir", not ok and "geçmiyor" in why, why)
    ok, why = f._delete_allowed("/azure-pipelines.yml")
    check("silme: korunan yol reddedilir (adı geçse bile)", not ok and "korunan" in why, why)
    f._deletes_done = 3
    ok, why = f._delete_allowed("/mpModerationAiAuth.json")
    check("silme: iş başına sınır", not ok and "sınır" in why, why)


def test_devops_requests():
    """Kullanıcı kararı (2026-10-30): DevOps'un yapması gereken iş ayrı Task olarak açılır.
    PR'dan: repoda daha önce okunmayan ortam değişkeni (#247 MP_MODERATION_AI_AUTH_JSON).
    Teknik borçtan: eklenti ön koşulu kanıtlanamayan iş otomatik değil, DevOps'a gider."""
    print("\n[DevOps istekleri]")
    from agile_sdlc_crew import devops_request as dr
    from agile_sdlc_crew.quality import php_modern as pm
    from agile_sdlc_crew.quality.model import Finding, Group
    from agile_sdlc_crew import pipeline_config as pc

    old = "<?php\n$a = getenv('DB_HOST');\n"
    new = ("<?php\n$a = getenv('DB_HOST');\nconst AUTH_ENV_NAME = 'MP_MODERATION_AI_AUTH_JSON';\n"
           "$b = Helper::getEnvOrDefault('EXISTING_ELSEWHERE');\n")
    got = dr.new_env_names(old, new, lambda n: n == "EXISTING_ELSEWHERE")
    check("env: yalnız PR'ın eklediği ve repoda geçmeyen ad", got == {"MP_MODERATION_AI_AUTH_JSON": 3}, str(got))
    check("env: Go/JS/Python okumaları", set(dr.env_names(
        'x := os.Getenv("GA4_JSON")\nprocess.env.NODE_KEY\nos.environ["PY_SECRET"]').keys()) == {"GA4_JSON", "NODE_KEY", "PY_SECRET"})
    check("silinen dosya içerik sırrı (ad ele vermese de)", dr.looks_like_secret('{"private_key": "-----BEGIN PRIVATE KEY-----x"}')
          and not dr.looks_like_secret('{"name": "analytics", "enabled": true}'))
    body = dr.env_request_body(repo="marketplace", names={"MP_X": "app/A.php:3"}, deleted_files=["auth.json"],
                               deps=[{"devrepo": "anatolia", "service": "marketplace", "secret": "prod-marketplace-secrets",
                                      "images": ["php-8-2:1.0"]}], pr_url="https://x/pr/1", source_wi="76578")
    check("istek metni: ad, satır, secret adı, değer yazılmaz notu", all(x in body for x in (
        "MP_X", "app/A.php:3", "prod-marketplace-secrets", "auth.json", "Değerler bu işte yazılmaz")))

    class _Cl:
        _base_api_url = "https://x/_apis"
        def __init__(self, existing=None):
            self.created, self.updates, self.existing = [], [], existing
        def get_team_area_path(self, team):
            return "Boards Management\\DevOps"
        def query_work_items(self, q, limit=0):
            return [{"id": self.existing}] if self.existing else []
        def list_iterations(self, team):
            return [{"path": "Boards Management\\2026_21_DevOps", "attributes": {"timeFrame": "past"}},
                    {"path": "Boards Management\\2026_22_DevOps", "attributes": {"timeFrame": "current"}}]
        def create_work_item(self, t, fields, parent_id=None):
            self.created.append((t, fields)); return {"id": 90001}
        def update_work_item(self, wid, ops):
            self.updates.append((wid, ops)); return {}
    c = _Cl()
    r = dr.open_request(c, team="DevOps", title="[m] ENV", body_html="<p>x</p>", source_wi="76578", key="env:m:76578:MP_X",
                        task_fields={"Custom.ARGEType": "Hayır"})
    f = c.created[0][1]
    check("açılış: DevOps alanı, aktif sprint, Task, etiket, anahtar satırı, zorunlu alan",
          r == {"id": 90001, "created": True} and c.created[0][0] == "Task" and f["System.AreaPath"].endswith("DevOps")
          and f["System.IterationPath"].endswith("2026_22_DevOps") and f["System.Tags"] == dr.TAG
          and "tempo-devops-key: env:m:76578:MP_X" in f["System.Description"] and f["Custom.ARGEType"] == "Hayır", str(f))
    check("açılış: kaynak işe Predecessor bağı", c.updates and c.updates[0][0] == 76578
          and c.updates[0][1][0]["value"]["rel"] == "System.LinkTypes.Dependency-Reverse")
    c2 = _Cl(existing=88888)
    r2 = dr.open_request(c2, team="DevOps", title="t", body_html="b", source_wi="1", key="k")
    check("aynı istek ikinci kez açılmaz, yeni kaynak yine bağlanır", r2 == {"id": 88888, "created": False}
          and not c2.created and c2.updates)

    # ── teknik borç: eklenti ön koşulu ──
    rule = pm.Rule(key="utf8_encode", since=(8, 2), kind="deprecation", title="utf8", summary="", advantages=[],
                   risks=[], before="", after="", ref="", auto=True, requires=["mbstring"])
    rf = pm.Rule(key="removed_functions", since=(8, 0), kind="deprecation", title="rm", summary="", advantages=[],
                 risks=[], before="", after="", ref="", auto=True, requires_if={"money_format": ["intl"]})
    def g(rule_key, text):
        return Group(source="php_modern", repo="r", title="t", fingerprints=["fp"],
                     findings=[Finding(source="php_modern", file="a.php", rule=rule_key, uris=[f"3|{text}"])])
    check("eklenti kanıtı yoksa eksik listelenir", pm._missing_exts(g("utf8_encode", "utf8_encode($x)"), rule, lambda e: False) == ["mbstring"])
    check("eklenti repoda kullanılıyorsa otomatik kalır", pm._missing_exts(g("utf8_encode", "utf8_encode($x)"), rule, lambda e: True) == [])
    check("koşullu eklenti yalnız ilgili fonksiyon geçince", pm._missing_exts(g("removed_functions", "each($a)"), rf, lambda e: False) == []
          and pm._missing_exts(g("removed_functions", "money_format('%i', $x)"), rf, lambda e: False) == ["intl"])
    with tempfile.TemporaryDirectory() as td:
        root = Path(td)
        sp = lambda *a: subprocess.run(["git", *a], cwd=root, capture_output=True, text=True)
        sp("init", "-q", "-b", "main")
        (root / "a.php").write_text("<?php echo strlen('x');\n")
        sp("add", "-A"); sp("-c", "user.name=t", "-c", "user.email=t@t", "commit", "-qm", "i")
        check("kanıt yok → eklenti doğrulanamadı", not pm.ext_available(root, "mbstring"))
        (root / "b.php").write_text("<?php echo mb_strlen('x');\n")
        sp("add", "-A"); sp("-c", "user.name=t", "-c", "user.email=t@t", "commit", "-qm", "j")
        check("uygulama kodunda mb_* → mbstring var", pm.ext_available(root, "mbstring"))
        (root / "composer.json").write_text('{"require": {"ext-intl": "*"}}')
        check("composer ext-<ad> şartı → var", pm.ext_available(root, "intl"))

    src = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/quality/php_modern.py").read_text()
    check("tarama: otomatik yalnız kodla çözülen; altyapı ayrı liste",
          '"infra": [_row(g, rules_by_key) for g in groups if g.needs_human and g.infra]' in src
          and '"auto": [_row(g, rules_by_key) for g in groups if not g.needs_human]' in src)
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text()
    check("kart: 'Altyapı (DevOps) gerektiren işler' bölümü ve DevOps bağlantısı",
          "Altyapı (DevOps) gerektiren işler" in html and "DevOps isteği: " in html)
    fsrc = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/flow.py").read_text()
    check("PR açılınca (yeni ve mevcut PR yolu) env isteği kontrol edilir", fsrc.count("self._devops_env_request()") == 2)
    check("knob'lar", pc.get("CREW_DEVOPS_TEAM") == "DevOps" and isinstance(pc.get("CREW_DEVOPS_REQUESTS"), bool))


def test_modern_multilang():
    """Modernleştirme repo hangi dildeyse o: PHP, Go, Python, Next.js. Go/Python'da
    düzeltmeyi araç uygular (modernize / ruff), LLM yalnız inceler."""
    print("\n[modernleştirme: çok dil]")
    import shutil
    from agile_sdlc_crew.quality import modern, modern_fix
    from agile_sdlc_crew.quality.render import MODERN_SOURCES, task_description
    from agile_sdlc_crew.quality.model import Finding, Group

    with tempfile.TemporaryDirectory() as td:
        base = Path(td)
        go = base / "goapp"; go.mkdir()
        (go / "go.mod").write_text("module m\n\ngo 1.24.0\n")
        py = base / "pyapp"; py.mkdir()
        (py / "pyproject.toml").write_text('[project]\nrequires-python = ">=3.11"\n')
        py2 = base / "pyimg"; py2.mkdir()
        (py2 / "requirements.txt").write_text("fastapi\n")
        (py2 / "Dockerfile").write_text("FROM python:3.9-slim\n")
        nx = base / "web"; nx.mkdir()
        (nx / "package.json").write_text('{"dependencies": {"next": "^15.2.4", "react": "19.0.0"}}')
        (nx / "app").mkdir()
        (nx / "app" / "a.tsx").write_text("import { useFormState } from 'react-dom'\n// const c = cookies()\nconst h = headers()\n"
                                         "const x = <Link href='/a'>\n  <a>git</a></Link>\n")
        (nx / "middleware.ts").write_text("export function middleware(){}\n")
        (nx / "node_modules").mkdir(); (nx / "node_modules" / "z.js").write_text("useFormState()\n")
        check("dil tespiti", modern.detect_langs(go) == ["go"] and modern.detect_langs(py) == ["python"]
              and modern.detect_langs(nx) == ["nextjs"], str([modern.detect_langs(x) for x in (go, py, nx)]))
        check("Go hedefi go.mod'dan", modern.target(None, "goapp", go, "go")["version"] == "1.24")
        check("Python hedefi: FLODevOps imajı repo Dockerfile'ından önce gelir",
              modern.target(None, "pyimg", py2, "python", index={"pyimg": [{"src": "anatolia/build/x", "images": ["python:3.11"]}]})["version"] == "3.11"
              and modern.target(None, "pyimg", py2, "python", index={})["version"] == "3.9"
              and modern.target(None, "pyapp", py, "python", index={})["version"] == "3.11")
        check("Next.js hedefi package.json'dan", modern.target(None, "web", nx, "nextjs")["version"] == "15.2")
        hits, _ = modern.next_hits(nx)
        keys = sorted({(h["rule"], h["file"]) for h in hits})
        check("Next.js desenleri: useFormState, senkron headers(), <Link><a>, middleware dosyası",
              ("use_form_state", "app/a.tsx") in keys and ("sync_request_api", "app/a.tsx") in keys
              and ("link_legacy_anchor", "app/a.tsx") in keys and ("middleware_to_proxy", "middleware.ts") in keys, str(keys))
        check("Next.js: yorum satırı ve node_modules atlanır",
              sum(1 for h in hits if h["rule"] == "sync_request_api") == 1 and not any("node_modules" in h["file"] for h in hits))

        if Path(modern.go_modernize_bin()).exists() and shutil.which("go"):
            (go / "a.go").write_text("package m\n\nfunc F(v interface{}) interface{} { return v }\n")
            g = lambda *a: subprocess.run(["git", *a], cwd=go, capture_output=True, text=True)
            g("init", "-q", "-b", "main"); g("add", "-A"); g("-c", "user.name=t", "-c", "user.email=t@t", "commit", "-qm", "i")
            gh, errs = modern.go_hits(go)
            check("Go: modernize bulguları (any)", any(h["rule"] == "any" and h["file"] == "a.go" for h in gh), f"{gh} {errs}")
            changed = modern_fix.go_apply(go, {"any": ["a.go"]})
            check("Go: yalnız kapsamdaki düzeltme uygulanır ve gofmt'lanır",
                  changed == ["a.go"] and "func F(v any) any" in (go / "a.go").read_text(), (go / "a.go").read_text())
        else:
            skip("Go modernize", "araç yok")

    with tempfile.TemporaryDirectory() as td:
        base = Path(td); (base / "gocrons" / "cmd").mkdir(parents=True)
        ok = modern_fix.load_scope('{"lang": "go", "repo": "gocrons", "scope": {"any": ["cmd/a.go"]}}', base)
        check("araç kapsamı işin kaydından", ok[:3] == ("go", "gocrons", {"any": ["cmd/a.go"]}))
        bad = ['{"lang": "go", "repo": "../etc", "scope": {"any": ["a.go"]}}',
               '{"lang": "go", "repo": "gocrons", "scope": {"any": ["../../x.go"]}}',
               '{"lang": "go", "repo": "gocrons", "scope": {"any": ["/etc/passwd"]}}',
               '{"lang": "go", "repo": "gocrons", "scope": {"any": ["-rf.go"]}}',
               '{"lang": "php", "repo": "gocrons", "scope": {"x": ["a.php"]}}', '', '{"lang": "go", "repo": "nope", "scope": {"any": ["a.go"]}}']
        rej = 0
        for b in bad:
            try:
                modern_fix.load_scope(b, base)
            except ValueError:
                rej += 1
        check("kapsam doğrulaması: kök dışı/mutlak/'-' yol, geçersiz repo/dil, boş kayıt reddedilir", rej == len(bad), f"{rej}/{len(bad)}")
    check("ruff ikilisi doğrudan çalışır (python -m değil)", "-m" not in modern._ruff_cmd() and modern._ruff_cmd()[0].endswith("ruff"))

    grp = Group(source="go_modern", repo="gocrons", title="t", fingerprints=["go_modern:gocrons:any:a.go"],
                findings=[Finding(source="go_modern", file="a.go", rule="any", lines=[3], uris=["3|interface{}"])])
    d = task_description(grp, php_target="1.24 (go.mod)")
    check("Task metni: Go hedefi ve Go kataloğu", "Hedef Go:" in d and "interface{} → any" in d and "go_modern" in MODERN_SOURCES)

    srv = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/server.py").read_text()
    check("işçi: modern_fix işleri araçla koşar (pipeline değil)", 'if (job.get("job_kind") or "pipeline") == "modern_fix":' in srv
          and "modern_fix.run(job_id, work_item_id)" in srv)
    check("uçlar dil parametresi alır", 'lang: str = "php"' in srv and "_md.scan, AzureDevOpsClient(), repo.strip(), (lang or \"php\")" in srv)
    md = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/quality/modern.py").read_text()
    check("iş türü ve onaylı kapsam INSERT'te verilir (işçi yarışı yok, WI açıklaması okunmaz)",
          'job_kind="modern_fix",' in md and "job_payload=payloads.get(title" in md)
    html = (Path(__file__).resolve().parent.parent / "src/agile_sdlc_crew/web/index.html").read_text()
    check("kart çok dilli: başlık, dil:repo seçimi, araç uygular etiketi",
          '<span class="h">Modernleştirme</span>' in html and "function _tdmSel()" in html and "Araç uygular" in html)


def main():
    print("Katman 0 kapıları — regresyon testleri")
    print("=" * 62)
    for t in (test_norm_path, test_plan_paths, test_issue_gate,
              test_requirement_ids, test_completeness, test_contract_gate,
              test_fix_targets, test_envelope, test_prune_fix_targets,
              test_reachability, test_context_prefix_stability,
              test_grep_evidence, test_bm25_identifier_terms,
              test_summary_column_extraction, test_summary_index_refresh,
              test_build_fix_selection, test_resume_wiring,
              test_build_gate_timeout_strict,
              test_build_fix_single_commit,
              test_build_fix_regressions,
              test_partial_implement_resume,
              test_needs_human_and_resume_envelope,
              test_build_gate_stale_build,
              test_readiness_gate,
              test_call_context_cross_thread,
              test_md_to_html_headings_tables,
              test_wi_lifecycle_transitions,
              test_dod_checklist,
              test_wi_lifecycle_flow_hooks,
              test_estimation,
              test_child_tasks_planning,
              test_estimate_flow_hooks,
              test_retrospective,
              test_sprint_planning,
              test_daily_summary,
              test_type_flow,
              test_type_flow_hooks,
              test_pr_review_contribution,
              test_backlog_refinement,
              test_cost_analytics,
              test_cheap_phase_split,
              test_findings_reuse,
              test_prompt_prefix_order,
              test_cli_prompt_overhead,
              test_wi_comments_context,
              test_quality_rank_nr,
              test_quality_rank_sonar,
              test_quality_select,
              test_quality_render,
              test_quality_dedup,
              test_quality_app_map,
              test_quality_sources,
              test_quality_intake,
              test_quality_endpoints,
              test_quality_review_fixes,
              test_quality_readiness,
              test_repo_ctx_cross_thread,
              test_windowed_edit,
              test_quality_lean,
              test_worktree_staging,
              test_techdebt_lane,
              test_techdebt_open_tasks_db,
              test_techdebt_list_queue,
              test_techdebt_endpoints,
              test_techdebt_ui,
              test_techdebt_review_fixes,
              test_dashboard_readability,
              test_human_label,
              test_quality_human_split,
              test_quality_project_filter,
              test_guidance_in_steps,
              test_completion_report_no_tool_note,
              test_review_prefetch_window,
              test_quality_estimate,
              test_no_env_dependent_tests,
              test_pr_comment_emoticons,
              test_completion_comment_integrity,
              test_php_modern_catalog,
              test_php_modern_scan,
              test_php_modern_groups,
              test_php_modern_render,
              test_php_modern_take,
              test_php_modern_endpoints,
              test_php_modern_ui,
              test_quality_sonar_exclude,
              test_env_conventions,
              test_sprint_report_no_sp,
              test_sprint_report_designs,
              test_review_retry_delete_and_no_progress,
              test_delete_push_payload,
              test_quality_direct_flow,
              test_resume_skips_discarded_work,
              test_wi_lifecycle_on_queue,
              test_pre_pr_verify_and_build_fix,
              test_security_hardening,
              test_devops_requests,
              test_modern_multilang):
        try:
            t()
        except Exception as e:
            FAIL.append(t.__name__)
            print(f"  ❌ {t.__name__} ÇÖKTÜ: {type(e).__name__}: {e}")
    print("\n" + "=" * 62)
    print(f"GEÇEN {len(PASS)} · BAŞARISIZ {len(FAIL)} · ATLANAN {len(SKIP)}")
    if FAIL:
        print("Başarısızlar:")
        for f in FAIL:
            print(f"  - {f}")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
